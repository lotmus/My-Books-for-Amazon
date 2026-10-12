import re, io, os, datetime
from PIL import Image
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

B = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
P = os.path.join(B, "Step_Stories_Pics_Final")
TITLE1, TITLE2 = "Scally and the", "Sleepless Stream"
FULL = "Scally and the Sleepless Stream"
SERIES = "The Step Stories"
AUTHOR = "Anneliese Tauber"
STORY_FILE = "Book2_Story_v10.txt"
VERSION = re.search(r"_(v\d+)\.txt$", STORY_FILE).group(1)
STORY = os.path.join(B, "Step_Stories_Texts", STORY_FILE)
Q = "“"; R = "”"; A = "’"

raw = open(STORY, encoding="utf-8").read()
body = raw[raw.index("Chapter 1."):]
blocks = [b.strip("\n") for b in re.split(r"\n\s*\n", body) if b.strip()]

place = {
    'Hal had been timing the echo.': 'b65_echo_count.jpg',
    'And then Scally stood up.': 'b66_scally_stands.jpg',
    'She said it flatly, the way she said it when she meant it.': 'b67_hopsi_dont_know.jpg',
    'They crossed in single file, slowly.': 'b68_bog_crossing.jpg',
    'Hopsi had been watching, with her lips pressed together.': 'b69_hopsi_in_pot.jpg',
    '“The tree is the tallest thing for a mile,”': 'b70_hopsi_warns.jpg',
    '“One,” whispered Titch. “Two.”': 'b71_counting_storm.jpg',
    'Scally had been scratching his head.': 'b72_scally_pitches.jpg',
    'It was the longest minute of Hal’s life.': 'b73_run_high_ground.jpg',
    'They went home the long way.': 'b74_going_home.jpg',
    '“GUCCI!” screamed Titch.': 'b75_gucci_dragged.jpg',
    'There was a short, loud argument.': 'b76_the_argument.jpg',
    'That night every light in the valley went out early': 'b77_village_sleeps.jpg',
    'Hal put one in his palm and spun it on the step.': 'b78_egg_spin.jpg',
    'One grey afternoon at the end of October': 'b79_step_buns.jpg',
    '“Wade it,” said Scally.': 'b80_wade_it.jpg',
    'They launched the raft at noon.': 'b81_launch_raft.jpg',
    'In the morning Scally sat on the mill step': 'b82_painting_sign.jpg',
    'Scally, at the back of the crowd, had not spoken.': 'b49_box_scheme.jpg',
    'Before they left, Walt gave out his own rules': 'b50_walt_rules.jpg',
    'Hal wrote the plan on his slate': 'b51_slate_plan.jpg',
    'Scally began to laugh, and could not stop.': 'b53_laughing_blanket.jpg',
    'Walt took out his whistle.': 'b54_whistle_fog.jpg',
    'They ate Titch’s crusts': 'b55_supper_stew.jpg',
    'Before dawn Hal crept down': 'b56_hal_dawn.jpg',
    'In the morning they came to the Falls.': 'b57_the_falls.jpg',
    '“Listen,” shouted Hal.': 'b58_hello_echo.jpg',
    '“Everyone all right?” said Walt': 'b59_count_heads.jpg',
    'Titch had already done it.': 'b60_dumpling_pulls.jpg',
    '“I’m scared,” said Scally.': 'b61_scally_scared.jpg',
    '“Walter,” he said.': 'b62_mayor_hat.jpg',
    'Old Walt took on an apprentice.': 'b63_walt_apprentice.jpg',
    'In the morning, the miller fell asleep in the flour hopper': 'b64_black_buns.jpg',
    "Hal Millersen lay on his back": "b01_hum_night.jpg",
    "By ten o" + A + "clock a hand-painted sign": "b02_draught_sale.jpg",
    "He scooped up the tin.": "b03_pennies_back.jpg",
    "Hal went to the pond, to think.": "b04_reflection_up.jpg",
    "The Mayor stood on a barrel.": "b05_mayor_barrel.jpg",
    "It was a grey donkey, small and round": "b06_walt_arrives.jpg",
    Q + "You don" + A + "t shout at it," + R + " said Old Walt.": "b07_walt_stick.jpg",
    "They set out at dawn.": "b08_setting_off.jpg",
    "A bank had slipped into the stream.": "b09_green_pool.jpg",
    "Titch had a log in each hand.": "b10_apple_swap.jpg",
    "Hal found, in his surprise": "b11_float_test.jpg",
    "It took all morning. They cut six logs": "b12_raft_build.jpg",
    "The raft slid out.": "b13_raft_spin.jpg",
    "Four of them hauled": "b14_rescue_pool.jpg",
    "Walt sat down on a rock, in sections.": "b15_meander.jpg",
    "Hal knelt at the water. On the inside of the bend": "b16_pebbles.jpg",
    "They found him in clay": "b17_clay_rescue.jpg",
    "He built the fire himself": "b18_campfire.jpg",
    Q + "Scally," + R + " said Titch.": "b19_cheese_split.jpg",
    "They listened. And in the dark, under the stars": "b20_voice_in_water.jpg",
    "Nobody had an answer. In the silence, Hopsi got up.": "b21_three_stones.jpg",
    "They had to climb, one behind the other, on a rope.": "b22_falls_climb.jpg",
    "There was a pause. Then the cliff said": "b23_echo_trick.jpg",
    "The next came without any count at all.": "b24_oak_struck.jpg",
    "Scally was already beside him.": "b25_flood_rope.jpg",
    "They all got up, at last, on a ledge": "b26_flood_ledge.jpg",
    "It was the Source.": "b27_whirlpool.jpg",
    "He stepped up to the edge of the bank.": "b28_hat_lie.jpg",
    "Titch went after him.": "b29_gucci_saved.jpg",
    "Hal took it out. It was soggy": "b30_slate_data.jpg",
    "She was sitting very straight, on a stone": "b31_hopsi_i_think.jpg",
    "It sang. It sang four notes": "b32_fifth_note.jpg",
    "Just before it went, they all saw": "b33_mist_bows.jpg",
    "It was the quietest thing any of them had ever heard,": "b34_stream_sleeps.jpg",
    "They had been there since dawn.": "b35_homecoming.jpg",
    "He had Walt" + A + "s boots on.": "b36_walt_boots.jpg",
    "He got as far as the second E in PENNY.": "b37_sign_facedown.jpg",
    "It rocked once. Twice.": "b38_eggs_round.jpg",
    "Then, one by one, they cracked.": "b39_gosling_boots.jpg",
    Q + "Right," + R + " said Scally, and got up.": "b40_long_way.jpg",
    "The first drop fell, round and heavy": "b41_rain.jpg",
    '“Shelter!” said Scally.': 'b42_dry_seat.jpg',
    '“I know that voice,” he said.': 'b43_walt_story.jpg',
    'But Walt had turned to look at Dumpling': 'b44_scouting.jpg',
    '“Hal,” he said. “Come here.”': 'b45_brown_water.jpg',
    'Hal’s mother crouched in front of him': 'b46_mother_rules.jpg',
    'Walt sat down on the grass beside him,': 'b47_walt_sack.jpg',
    'Tabitha had said nothing all morning.': 'b48_counting_notes.jpg',
}

for _k, _v in place.items():
    _n = sum(1 for x in blocks if x.startswith(_k))
    if _n != 1 or not os.path.exists(os.path.join(P, _v)):
        print("ANCHOR PROBLEM", _n, _v, repr(_k[:50]))


def img_bytes(path, w=640, q=68):
    im = Image.open(path).convert("RGB")
    if im.width > w:
        im = im.resize((w, round(im.height * w / im.width)), Image.LANCZOS)
    b = io.BytesIO()
    im.save(b, "JPEG", quality=q, optimize=True)
    b.seek(0)
    return b


def make_cover():
    src = Image.open(os.path.join(B, "Cover_Front_1600x2560.jpg")).convert("RGB")
    out = os.path.join(P, "cover_1000x1600.jpg")
    src.resize((1000, 1600), Image.LANCZOS).save(out, quality=88)
    return out


COVER = make_cover()
d = Document()
sec = d.sections[0]
sec.page_width = Inches(6)
sec.page_height = Inches(9)
for a in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
    setattr(sec, a, Inches(0.8))
st = d.styles["Normal"]
st.font.name = "Georgia"
st.font.size = Pt(13)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Georgia")
st.paragraph_format.space_after = Pt(10)
st.paragraph_format.line_spacing = 1.3
NAVY = RGBColor(0x1F, 0x3A, 0x5F)
RUST = RGBColor(0xA8, 0x4B, 0x2A)
for n, sz in (("Heading 1", 24), ("Heading 2", 19)):
    h = d.styles[n]
    h.font.name = "Georgia"
    h.font.size = Pt(sz)
    h.font.bold = True
    h.font.color.rgb = NAVY
    h.element.rPr.rFonts.set(qn("w:eastAsia"), "Georgia")
    h.paragraph_format.keep_with_next = True
    h.paragraph_format.space_before = Pt(14)

bid = [0]


def bookmark(par, name):
    bid[0] += 1
    s = OxmlElement("w:bookmarkStart")
    s.set(qn("w:id"), str(bid[0]))
    s.set(qn("w:name"), name)
    e = OxmlElement("w:bookmarkEnd")
    e.set(qn("w:id"), str(bid[0]))
    par._p.insert(0, s)
    par._p.append(e)


def link(par, text, anchor):
    h = OxmlElement("w:hyperlink")
    h.set(qn("w:anchor"), anchor)
    r = OxmlElement("w:r")
    rp = OxmlElement("w:rPr")
    c = OxmlElement("w:color")
    c.set(qn("w:val"), "1F3A5F")
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rp.append(c)
    rp.append(u)
    r.append(rp)
    t = OxmlElement("w:t")
    t.text = text
    r.append(t)
    h.append(r)
    par._p.append(h)


def pagebreak():
    d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def centered(text, size=13, bold=False, italic=False, after=10, color=None):
    p = d.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(after)
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.bold = bold
    r.italic = italic
    if color:
        r.font.color.rgb = color
    return p


def pic(name, w_in=4.4):
    p = d.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(img_bytes(os.path.join(P, name)), width=Inches(w_in))


def heading(text, level, anchor=None, center=True):
    h = d.add_paragraph(text, style="Heading %d" % level)
    if center:
        h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if anchor:
        bookmark(h, anchor)
    return h


def toc_link():
    p = d.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    link(p, "↑ Contents", "toc")


def para(t, lead=None):
    p = d.add_paragraph()
    if lead:
        r = p.add_run(lead + " ")
        r.bold = True
    p.add_run(t)


def lead_bold(par, text):
    m = re.match(r"^(\d+\. [^.]+\.)(.*)$", text, re.S)
    if m:
        r = par.add_run(m.group(1))
        r.bold = True
        par.add_run(m.group(2))
    else:
        par.add_run(text)


chapters = [b.split("\n")[0] for b in blocks if b.startswith("Chapter ")]
TOC = [("Meet the Characters", "cast")] + [(c, "ch%d" % (i + 1)) for i, c in enumerate(chapters)] + [
    ("Hal" + A + "s Slate: What We Found Out", "slate"), ("For Parents and Teachers", "guide"),
    ("About the Author", "about"), ("Also by Anneliese Tauber", "alsoby")]

# cover
sec.top_margin = sec.bottom_margin = Inches(0.2)
sec.left_margin = sec.right_margin = Inches(0.2)
p = d.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(0)
p.add_run().add_picture(COVER, height=Inches(8.5))
sec2 = d.add_section()
for a in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
    setattr(sec2, a, Inches(0.8))

for _ in range(5):
    d.add_paragraph()
centered(TITLE1, 34, True, color=NAVY)
centered(TITLE2, 34, True, color=RUST)
centered("Book 2 of " + SERIES, 15, italic=True, after=30)
centered(AUTHOR, 18)
pagebreak()
heading("Contents", 1, "toc")
for t, a in TOC:
    p = d.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    link(p, t, a)
pagebreak()

heading("Meet the Characters", 1, "cast")
cast = [
    ("Hal (Halbert) Millersen", "8 and a bit. The miller" + A + "s son, with a slate for his notes. Looks, counts, tests, and writes WRONG when he is."),
    ("Scally, the wee scallywag", "8. A little devil with pointed ears and a leaf in his hair. He can never resist a scam, a lie or a shortcut. His true name is for special occasions."),
    ("Hopsi (Elfie Didlhop)", "9. An elf with long swept-back ears. She only says what is true, and this time she has to learn to say “I think”."),
    ("Titch Mitchel", "8. Red cap, quick fingers and a crust in each hand. The best knot-tier in the valley."),
    ("Tabitha Smoke", "8. Curly black hair in a bun. Watches carefully and notices what nobody else did."),
    ("Gucci", "A wolf cub, now the size of a cushion, who still thinks he is the size of a boot."),
    ("Fourpence", "A small brown bird who hatched from the last square egg. It sings four notes, exactly four."),
    ("Old Walt", "The stream-keeper, with a stiff leg, a hat made of eel-skin and fifty years of listening to water."),
    ("Dumpling", "Walt" + A + "s small round donkey, who walks as if she is paid by the hour."),
]
for n, t in cast:
    p = d.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(9)
    r = p.add_run(n)
    r.bold = True
    p.add_run("\n" + t)
centered("Ages 5–9. Little ones can follow the whole story from the pictures alone.", 12, italic=True)
pagebreak()

ci = 0
in_slate = False
first = True
for b in blocks:
    if b == "---":
        continue
    if b.startswith("Chapter "):
        ci += 1
        if not first:
            toc_link()
            pagebreak()
        first = False
        heading(b, 2, "ch%d" % ci)
        continue
    if b == "The End":
        centered("The End", 18, True).paragraph_format.space_before = Pt(18)
        toc_link()
        pagebreak()
        continue
    if b.startswith("Hal" + A + "s Slate: What We Found Out"):
        heading("Hal" + A + "s Slate: What We Found Out", 1, "slate")
        in_slate = True
        continue
    if len(b) < 60 and "\n" not in b and not re.search(r"[a-z]", b):
        centered(b, 16, True, after=12)
        continue
    if len(b.split("\n")) in (2, 3) and all(l.isupper() or not re.search(r"[a-z]", l) for l in b.split("\n")):
        p = d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        lines = b.split("\n")
        for i, l in enumerate(lines):
            r = p.add_run(l)
            r.bold = True
            if i < len(lines) - 1:
                r.add_break()
        continue
    p = d.add_paragraph()
    if in_slate:
        lead_bold(p, b.replace("\n", " "))
    else:
        p.add_run(b.replace("\n", " "))
    for k, v in place.items():
        if b.startswith(k):
            pic(v)
            break

toc_link()
pagebreak()
heading("For Parents and Teachers", 1, "guide")
para("This is a fantasy adventure to read aloud at ages 5–9. Children from about 5 can follow the whole story from the pictures alone (a “picture walk”: turn the pages and ask what is happening), and older children can read it by themselves. Every fact the elf Hopsi states is true; where she is not sure, she says “I think”, which is what a scientist says. The magic is marked as magic.")
heading("The habits of a scientist in the story", 2, center=False)
for lead, t in [
    (Q + "How do we know?" + R, "Look, count, test. Hal writes everything on his slate, and writes WRONG when he is."),
    ("Guess first, then test.", "What floats and what sinks? Hal guesses, then the pool tells him."),
    ("Say " + Q + "I think" + R + ".", "Hopsi learns that a scientist does not only know. She says what she thinks and how she would find out if she is wrong."),
    ("Data beats a hunch.", "On his slate Hal counts every shortcut on the journey and finds that none saved any time."),
    ("Honesty.", "Scally tries a scam, a lie or a cheat in every chapter, and the plain truth makes him think again."),
]:
    para(t, lead)
heading("What each chapter teaches", 2, center=False)
for lead, t in [
    ("The Hum:", "an honest question about a mystery, and why a trick to sell a cure does not work."),
    ("Old Walt:", "why water cannot be stopped, only moved, and why you ask where it is going. Family safety rules."),
    ("The Green Pool:", "what floats and what sinks, how a current carries a raft, and “reach or throw”."),
    ("The Wandering Stream:", "why streams bend, why pebbles are round, and why a shortcut across unknown ground is dangerous."),
    ("The Fire and the Voice:", "how to build and put out a campfire safely, and telling the truth about hunger."),
    ("The Count Between:", "echoes, how far a sound travels, counting between lightning and thunder, and where not to shelter in a storm."),
    ("The Brown Water:", "flash floods, high ground, and saying “I am scared”."),
    ("The Whirl:", "whirlpools, why you never go near one, and why a lie cannot be an ending."),
    ("The Long Way Round:", "reading evidence, and “I think”."),
    ("The Square Eggs:", "why an egg is shaped as it is, and not cheating when nobody is watching."),
    ("Rain:", "the water cycle, and choosing the long way yourself."),
]:
    para(t, lead)
heading("Staying safe: what this book models", 2, center=False)
for lead, tx in [
    ("Water.", "Never go in alone. Reach or throw: hold out a stick or throw a rope, and do not go in after a friend. Cold water takes your breath away. Stay away from the edge of a whirlpool and from rising streams."),
    ("Shortcuts and getting lost.", "Tell someone where you are going and when you will be back. Go in pairs. If you are lost, stand still, blow three blasts on a whistle and wait."),
    ("Fire.", "Build a campfire on bare stone with a ring of pebbles and a bucket of water. Put it out by drowning it, stirring it and drowning it again, until the ashes are cold to your hand."),
    ("Storms.", "Count the seconds between flash and bang. A house or a car is the safest place. If you are caught outside, stay away from tall trees, open water and metal, and crouch low in a dip. This is general advice; please check your local guidance."),
    ("Floods.", "Heavy rain upstream can make a stream rise very fast. Never camp in a dry gully, and go to high ground at once."),
    ("Strangers.", "The Soft Voice from Book 1 offers a shortcut, as a stranger might. Agree a family code word, as the children do."),
]:
    para(tx, lead)
heading("Talk about it", 2, center=False)
for q in [
    "Why did Scally try a trick in every chapter? What made him stop?",
    "Hal wrote down every shortcut and found none saved time. Do you know a time when the long way was better?",
    "Why does Hopsi say " + Q + "I think" + R + " instead of " + Q + "I know" + R + "? Is that braver or weaker?",
    "What floats and what sinks in your kitchen? Guess first, then test it with a grown-up.",
    "What would you do if you were lost, or if a stream started to rise?",
]:
    para(q)
heading("Words to grow", 2, center=False)
para("meander, current, pebble, echo, flash flood, whirlpool, float, evaporate, water cycle, data.")
heading("A note for adults", 2, center=False)
para("The science here is simplified for ages 5 to 9: a stream's bends come from faster water wearing the outside bank and slower water dropping sand on the inside; a thing floats if, for its size, it is lighter than the water it pushes aside; sound travels about a third of a kilometre per second, so a count of three between lightning and thunder is about a kilometre; whirlpools form where water drains toward a low outlet. The Mirror Mill and the humming stream are magic. Please have a teacher check the science and safety content before relying on it.")
toc_link()
pagebreak()
heading("About the Author", 1, "about")
para("Anneliese Tauber is a mother of four. When her children were small she read them a great many fairy tales, and she noticed that the best ones answered the questions children ask, and never needed to bend the truth. The Step Stories are for curious children and the grown-ups who read to them.")
centered("If your child enjoyed this book, a short, kind review on Amazon helps other curious families find it. Thank you.", 13, italic=True)
toc_link()
pagebreak()
heading("Also by Anneliese Tauber", 1, "alsoby")
centered("The Wee Scallywag and the Gentle Elf", 15, True, after=4)
centered("The Step Stories, Book 1. Where Scally first met Hal, Hopsi and a ghost in a mill.", 12, italic=True, after=14)
centered("Book 3 of The Step Stories is coming soon.", 12, italic=True, after=14)
centered("Bavarian Cooking Made Simple", 15, True, after=4)
centered("Classic dishes from a Bavarian kitchen, written so that any cook can enjoy a taste of Bavaria.", 12, italic=True)
toc_link()
centered("Typeface in the e-book edition: Comic Neue, © The Comic Neue Project Authors, used under the SIL Open Font License 1.1.", 10, italic=True)
d.core_properties.title = FULL
d.core_properties.author = AUTHOR
now = datetime.datetime.now(datetime.timezone.utc)
d.core_properties.created = now
d.core_properties.modified = now
d.core_properties.last_modified_by = AUTHOR
dst = os.environ.get("OUTDOCX") or os.path.join(B, "Scally_and_the_Sleepless_Stream_Book2_%s.docx" % VERSION)
d.save(dst)
print(dst, os.path.getsize(dst) // 1024, "KB")
