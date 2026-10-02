"""Build the plain-language Kindle docx from easy_book.md and the local pictures."""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor

root = Path(__file__).parent
src = (root / "easy_book.md").read_text(encoding="utf-8")
out = root / "Your_First_Book_That_Sells_Updated.docx"

d = Document()
s = d.sections[0]
s.page_width = Inches(7)
s.page_height = Inches(10)
for edge in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
    setattr(s, edge, Inches(0.7))
for name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2"):
    d.styles[name].font.name = "Calibri"
    d.styles[name].font.color.rgb = RGBColor(0, 0, 0)
d.styles["Normal"].font.size = Pt(12)
d.styles["Normal"].paragraph_format.space_after = Pt(8)
d.styles["Normal"].paragraph_format.line_spacing = 1.15
d.styles["Title"].font.size = Pt(32)
d.styles["Heading 1"].font.size = Pt(22)
d.styles["Heading 1"].paragraph_format.page_break_before = True
d.styles["Heading 2"].font.size = Pt(16)

d.add_paragraph("Your First Book\nThat Sells", "Title")
d.add_paragraph(
    "How to Self-Publish on Kindle, Attract Honest Reviews, and Build a Profitable Book Catalog",
    "Subtitle",
)
d.add_paragraph("A plain guide to getting paid for a book you wrote.")
d.add_paragraph("Updated 1 October 2026")


def add_pic(name, alt):
    p = d.add_paragraph()
    r = p.add_run()
    r.add_picture(str(root / name), width=Inches(5.4))
    r._r.xpath(".//wp:docPr")[0].set("descr", alt)


def bookmark(paragraph, name):
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bookmark.next_id))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bookmark.next_id))
    bookmark.next_id += 1
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


bookmark.next_id = 1


def link_to(paragraph, text, anchor):
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("w:anchor"), anchor)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0C2D5A")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rpr.append(color)
    rpr.append(underline)
    run.append(rpr)
    node = OxmlElement("w:t")
    node.text = text
    run.append(node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


from docx.oxml.ns import qn
from docx.oxml import OxmlElement

in_contents = False
for raw in src.splitlines():
    line = raw.rstrip()
    if not line.strip():
        continue
    if line.startswith("{{img:") and line.endswith("}}"):
        body = line[len("{{img:") : -2]
        name, alt = body.split("|", 1)
        add_pic(name.strip(), alt.strip())
        continue
    if line.startswith("# "):
        title = line[2:].strip()
        in_contents = title == "Contents"
        paragraph = d.add_paragraph(title, "Heading 1")
        slug = "h_" + "".join(ch.lower() if ch.isalnum() else "_" for ch in title)[:40]
        bookmark(paragraph, slug)
        continue
    if line.startswith("## "):
        d.add_paragraph(line[3:].strip(), "Heading 2")
        continue
    if in_contents:
        paragraph = d.add_paragraph()
        slug = "h_" + "".join(ch.lower() if ch.isalnum() else "_" for ch in line)[:40]
        link_to(paragraph, line.strip(), slug)
        continue
    d.add_paragraph(line.strip())

d.core_properties.title = "Your First Book That Sells"
d.core_properties.subject = "How to Self-Publish on Kindle, Attract Honest Reviews, and Build a Profitable Book Catalog"
d.core_properties.author = "Kevin Drew Peters"
d.save(out)
words = sum(len(p.text.split()) for p in d.paragraphs)
print("words", words)
