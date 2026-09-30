from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import shutil
root=Path(__file__).parent
source=root.parent/'The_First_Review_Problem.docx'
d=Document()
s=d.sections[0]
s.page_width=Inches(7);s.page_height=Inches(10)
s.top_margin=s.bottom_margin=s.left_margin=s.right_margin=Inches(.7)
for n in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
 d.styles[n].font.name='Calibri';d.styles[n].font.color.rgb=RGBColor(0,0,0)
d.styles['Normal'].font.size=Pt(11)
d.styles['Normal'].paragraph_format.space_after=Pt(7)
d.styles['Normal'].paragraph_format.line_spacing=1.08
d.styles['Title'].font.size=Pt(32)
d.styles['Heading 1'].font.size=Pt(21)
d.styles['Heading 1'].paragraph_format.page_break_before=True
d.styles['Heading 2'].font.size=Pt(14)
imgdir=Path(r'C:/Users/lomus/.codex/generated_images/01a0e497-3ff2-70f1-a67c-befd52efbb26')
def pic(src,name,alt):
 shutil.copy2(imgdir/src,root/name)
 p=d.add_paragraph();r=p.add_run();r.add_picture(str(root/name),width=Inches(5.6));r._r.xpath('.//wp:docPr')[0].set('descr',alt)
def add(txt):
 for line in txt.strip().splitlines():
  if line.startswith('# '):d.add_paragraph(line[2:],'Heading 1')
  elif line.startswith('## '):d.add_paragraph(line[3:],'Heading 2')
  elif line:d.add_paragraph(line)
d.add_paragraph('Your First Book\nThat Sells','Title')
d.add_paragraph('How to Self Publish on Kindle, Attract Honest Reviews, and Build a Profitable Book Catalog','Subtitle')
pic('exec-e01ba1ad-1041-49ca-a0ce-7e73dfd1c8e4.png','readers.png','An author and a diverse group of readers sharing an open book.')
d.add_paragraph('A practical guide to reader recruitment, pricing, promotion, and the next book.')
d.add_paragraph('Updated September 2026')
add('''# Start here
A finished manuscript is the beginning of a publishing business. The next task is to help the right readers discover it, understand what they will get, and decide whether it is worth their time and money. This guide gives you a process for doing that without buying praise or spending more on promotion than your books can reasonably earn.
The first-review problem is real: a new book has little evidence of reader satisfaction. But reviews are only one part of a purchase decision. The cover, description, sample, price, and book itself must make the same promise. A convincing review cannot repair an unreadable sample or a misleading description.
This guide is for independent authors preparing a first Kindle release or improving a quiet launch. Its focus is honest reviews, launch planning, pricing, and catalog economics. It includes a publication checklist, but does not replace manuscript editing or explain every feature of the KDP dashboard.
Use the chapters in order if you have not published. If your book is already live, begin with the product-page audit, then build a reader campaign. Keep a notebook or spreadsheet beside you. By the end, you should have a reader promise, an invitation, a launch calendar, a spending ceiling, and a way to evaluate results.
All numerical campaign examples are hypothetical teaching examples, not reported author earnings or promises of results. Official policy references were checked on September 27, 2026. Recheck the linked pages when you publish.
# 1 Define the reader and the promise
Before choosing a price or finding reviewers, describe who the book is for. “Anyone who likes reading” is too broad to guide a cover, invitation, or advertisement. Name a recognizable reading preference or a specific problem.
For fiction, identify the genre and central experience. For example: “A cozy mystery for readers who enjoy a small-town setting, a resourceful older heroine, and low on-page violence.” For nonfiction: “A practical guide for first-time independent authors who need a modest-budget plan for recruiting readers and testing Kindle pricing.”
Finish this sentence: “This book helps or entertains [specific reader] by providing [specific experience or result].” Then identify the evidence inside the manuscript that supports that promise. A guide promising a launch plan should contain a usable calendar. A mystery must resolve its central question.
Write a second sentence: “This book is unlikely to suit readers who want [different experience].” Use that distinction to improve your invitation. Poor matching wastes readers’ time and produces disappointment that better recruitment could have prevented.
## Study comparable books
Choose five to ten books serving a similar reader. Compare covers, descriptions, samples, length, ordinary price, and recurring reader complaints. Record the observation date and marketplace. A temporary discount is not an ordinary price.
Do not assume a famous author’s price is right for a debut. Reputation, a large catalog, and an established audience change the decision. Include reasonably comparable independent titles, not only category leaders.
Look for unmet expectations you can investigate: unclear instructions, missing examples, excessive repetition, weak endings, or misleading subtitles. Do not copy reviews into your book as if they were your own evidence.
## Your action
Write your reader promise, identify five comparable books, and name the most important expectation your manuscript must meet. If you cannot do that clearly, improve the offer before buying traffic.
# 2 Make the book worth recommending
Treat the manuscript, cover, description, and sample as one connected experience. The promise made before purchase should survive the opening chapter.
## Audit the product page
View the cover at thumbnail size. Can you read the title? Does the design suggest the right genre or subject? A beautiful image is not enough if the reader cannot tell what kind of book it represents.
Read the first two sentences of the description. Establish the story’s appeal or the nonfiction problem. Replace vague praise with specific stakes, benefits, or contents. Avoid income promises you cannot substantiate.
Open the sample on a phone-sized screen. Check headings, paragraph breaks, images, links, and the amount of front matter before useful content begins. Confirm that title, subtitle, author name, and series information agree across the manuscript, cover, and listing.
## Separate editing from advance reading
A beta reader helps improve a developing manuscript. An advance reader receives a near-final book and may independently choose to review it. These roles sometimes overlap, but the distinction helps you plan.
Do not use an advance-review campaign as a substitute for editing. If readers discover a serious factual mistake, broken ending, or unusable layout, pause and repair the book. A short delay can cost less than promoting a weak experience.
During development, ask specific questions: “Where did you lose interest?” or “Which instruction could you not follow?” Do not ask readers to keep negative reactions private while posting only positive ones. That steers the public evidence.
## Prepare a useful ending
Offer a relevant next step: the next available story, a final action checklist, or an optional author newsletter with a clear description. Include a brief, neutral invitation to leave an honest review.
Do not make a subscription, purchase, or review a condition for receiving the complete content promised by the book. A nonfiction sequel should solve a new problem, not withhold the essential final step of the current one.
## Your action
Ask an independent reader to inspect the cover, description, and opening together. Ask what they expect from the book. Compare that answer with your intended promise.
# 3 Invite reviews without creating obligations
An advance reader copy, or ARC, lets an interested person read a book before publication. Some recipients will finish, some will stop, and some will never post. Their opinion remains their own.
Amazon permits free or discounted book copies when a review is not required and the author does not attempt to influence it. Extra incentives such as gift cards invalidate this exception. Read the current Customer Reviews help page before running a campaign. [1]
Explain the book, reading window, and delivery method. State clearly that reviewing is optional and no particular rating is expected. A free copy is not payment for praise.
Do not buy reviews, arrange reciprocal reviews, reimburse purchases to manufacture a purchase badge, or promise future benefits for positive ratings. Avoid making close personal or financial connections the foundation of your review campaign. Consult the current Community Guidelines for eligibility and conflicts. [2]
Do not demand proof of posting or ask someone to revise an opinion you dislike. If a service promises a fixed number of favorable reviews, reject that promise. You may purchase delivery or discovery services; assess their actual practices before paying.
Encourage readers who review to identify that they received a free advance copy. Do not supply their review text. Publication details and an accurate link are helpful; the opinion itself must remain independent.
## When a review does not appear
Check the book, edition, marketplace, and whether reviewing is available. An ASIN or preorder page does not by itself guarantee that customers can submit reviews immediately. Display and moderation are controlled by the platform.
Do not create replacement accounts or encourage repeated submissions designed to evade a restriction. The reader can consult Amazon support. Continue focusing on a good reading experience and accurate links.
## Your action
Remove every sentence in your invitation that sounds like a debt, exchange, rating requirement, or threat of losing access. Keep the optional nature of reviewing unmistakable.
# 4 Find readers who want your book
The strongest prospect enjoys the kind of book you wrote and has time to read it. Enthusiasm for helping authors cannot compensate for a poor genre or subject match.
Start with people who asked to hear about your writing. Then consider relevant communities that allow advance-copy invitations, your own public channels, and a carefully chosen reader-matching service if needed. Follow community rules. Do not scrape addresses or add people to mailing lists without consent.
Include genre or subject, approximate length, relevant content expectations, format, publication date, and reading window. Give people enough information to decline before applying.
## Use a short signup form
Collect name, email, preferred format, and a simple question about fit. Include an acknowledgment that reviewing is optional. Offer any ongoing newsletter through a separate optional choice.
Collect only information needed to deliver the book and manage the campaign. Keep the roster private, honor removal requests, and do not share reader details with other authors.
## Plan with a range
Suppose you invite 40 suitable people, 24 accept, and 6 post. Sixty percent accepted; 25 percent of recipients posted. Those figures describe one hypothetical campaign, not a universal review rate.
For planning, test scenarios. At posting rates of 10, 20, and 30 percent, 30 recipients would produce 3, 6, or 9 reviews. These are arithmetic possibilities, not obligations owed by readers. Keep the campaign affordable at the low end.
Allow enough reading time for your actual book. Two to four weeks may be a starting point, but a long novel or busy audience may require more. Ask about availability rather than demanding a launch-day review.
## Your action
Choose two recruitment channels and tailor an invitation to each. Set a maximum number of copies you can support comfortably. Review the quality of responses before adding another channel.
# 5 Run a simple reader campaign
A useful campaign is easy to join, easy to leave, and easy to understand. Complicated requirements create work for everyone.
## Prepare delivery
Create a checked EPUB for comfortable ebook reading. Provide other formats only if you have inspected them. Use a clear file name with title and version date. Test the download as a recipient, including permissions and any expiry date.
If considering KDP Select, review its digital exclusivity terms before distributing files outside Amazon. Do not assume that a public full-book download is acceptable. Resolve uncertainty about a planned review-copy arrangement with KDP support before enrollment. [5]
## Use a few clear messages
Send confirmation, then the file. After publication, when reviewing is available, send the live page with a neutral invitation. If appropriate, send one gentle follow-up a week or two later, then stop. These are suggested contact points, not mandatory messages.
Honor opt-outs immediately. A reader reporting a file problem needs support, not another review request. Extend the reading window if your delivery failed.
Keep a modest tracker: name, consent date, recruitment source, file-sent date, support needed, follow-up date, and opt-out status. A voluntarily shared review link can help avoid reminders; do not turn the tracker into a rating-based reward system.
At the end, record what improved the process. Did applicants misunderstand the genre? Was the file difficult to open? Did a holiday make the schedule unrealistic? These observations improve the next campaign even when the review count disappoints you.
## Your action
Test the delivery link with one willing person before sending the campaign. Adapt the messages in Chapter 12 to your own book.
# 6 Price for earnings as well as sales
Price affects both the reader’s decision and your royalty. Lower prices may increase units without increasing income. Higher prices may increase earnings per copy while reducing units. Test the balance rather than treating price as a judgment of your manuscript’s worth.
As checked in September 2026, Amazon.com’s ordinary 70 percent ebook price band is $2.99 to $12.99. Price is only one eligibility condition; territory and other requirements matter. Consult the current requirements and the royalty estimate in your KDP account. [3]
## Understand the calculation
For a qualifying 70 percent sale, the basic calculation is 0.70 × (price minus applicable VAT minus delivery cost). Under the 35 percent option, delivery cost is not deducted. Amazon.com lists a delivery rate of $0.15 per megabyte based on Amazon’s determination of file size. [4]
The following examples assume no VAT, a 2 MB delivered file, ordinary sales, and eligibility for the stated option. Delivery cost is $0.30. Figures are rounded estimates, not settlement amounts.
$0.99 at 35 percent earns about $0.35.
$1.99 at 35 percent earns about $0.70.
$2.99 at 70 percent earns about $1.88.
$3.99 at 70 percent earns about $2.58.
$4.99 at 70 percent earns about $3.28.
$7.99 at 70 percent earns about $5.38.
At $2.99, the calculation is 0.70 × ($2.99 − $0.30) = $1.883. The delivery deduction belongs inside the parentheses. Check KDP’s actual estimate for the converted file, especially for an illustrated book.
## Choose and test a starting price
Compare books serving your reader. For a practical guide, usefulness and specificity matter alongside length. For fiction, consider genre expectations, length, reputation, and series position. A $2.99 or $3.99 starting point may suit some debuts, but neither is a universal rule.
Write down why you chose the price and what evidence would change your mind. Test one change at a time. Changing price, cover, description, and advertising together makes interpretation difficult.
Under the assumptions above, 30 copies at $2.99 produce $56.49 in royalties. Twenty-five copies at $3.99 produce about $64.58. Fewer sales produced more royalties, but comparable traffic and spending are still needed for a useful comparison.
A week with two sales is weak evidence about a price ceiling. Avoid daily reactions to ordinary fluctuations. Use a dated log and an observation window long enough to collect useful information.
## Keep print separate
For standard KDP paperback sales, the basic formula is (applicable royalty rate × list price) − printing cost. The printing cost is not multiplied by the royalty rate. Rates and costs vary by price, marketplace, and print specifications. [7]
A hypothetical $12 paperback at a 60 percent rate and a $4 print cost earns $3.20. This is an arithmetic example, not a quote for your book. Use KDP’s calculator for the actual edition.
# 7 Decide whether a promotion can pay
Every promotion needs a purpose, spending limit, and evaluation method. “Get the book moving” is too vague. “Test whether this audience generates enough royalties to justify another campaign” gives you a decision to make.
If you promote to learn, label that expenditure a research cost. Downloads or a better rank do not establish profit.
## Calculate break even first
At a $1.883 royalty per sale, a $60 placement requires about 32 additional sales to cover the placement alone. At an unrounded royalty of $0.3465, it requires about 174. Editing, cover design, and your time are not covered by this calculation.
Additional means beyond what you would reasonably expect without the campaign. If you normally sell ten books during a period and sell twenty during a promotion, attributing all twenty to the promotion exaggerates its effect.
## Understand the promotion you choose
A normal price cut and an eligible Kindle Countdown Deal differ. A qualifying Countdown Deal can retain the selected 70 percent rate below $2.99, with delivery costs still applying. Check current enrollment, marketplace, pricing-history, and timing requirements before scheduling. [6]
KDP Select is a 90-day ebook program with digital exclusivity obligations and access to Kindle Unlimited and promotional tools. Check the term and renewal settings. Enrollment is a distribution choice, not a guarantee of sales. [5]
Free downloads are not paid sales. Do not budget as though every downloader will read, review, or purchase a sequel. State what a free campaign is intended to accomplish and how you will evaluate it.
## Buy a service for a specific problem
Delivery services reduce file support work. Reader-matching services help suitable readers discover copies. Promotional newsletters introduce offers to subscribers. Advertising sells exposure or clicks. These purchases solve different problems.
Inspect current fees, audience fit, cancellation rules, and spending limits. Ask what is guaranteed. A delivery platform may promise infrastructure; it cannot promise reader enthusiasm.
## Your action
Record campaign cost, conservative royalty per additional sale, break-even units, and maximum acceptable loss. If the required result seems implausible, change the plan before paying.
# 8 Test advertising with a spending ceiling
Advertising introduces the book to readers and exposes weaknesses in the offer. Reviews can help shoppers assess a book, but no universal review count makes advertising profitable.
A carefully capped test may be useful even with few reviews if the cover, description, sample, and targeting are credible. Large spending without evidence is the problem. Start at a scale you can afford to lose.
## Work backward from earnings
If a click costs $0.30 and one in ten clicks becomes a sale, advertising costs $3 per sale. A $1.88 royalty cannot cover that cost on a standalone ebook.
A simplified break-even cost per click equals royalty per sale multiplied by click-to-sale conversion. At a $1.883 royalty and 10 percent conversion, the ceiling is about $0.19. At 5 percent, it is about $0.09. These are theoretical ceilings, not suggested bids.
Refunds, other costs, attribution delays, and uncertain later purchases change the result. Leave a margin below break even. Do not treat retail sales value in an ad dashboard as author earnings.
Track Kindle Unlimited income separately and avoid assigning every increase to an advertisement just because it occurred in the same week. Use consistent windows and acknowledge uncertain attribution.
## Diagnose the result
Impressions without clicks may indicate targeting or cover problems. Clicks without sales may indicate a mismatch between advertisement and listing, a weak sample, price resistance, or small-sample noise. Sales without profit may mean acquisition costs are too high.
Choose one plausible explanation and test one change. Do not answer every disappointing result by increasing the budget. Record your hypothesis, change, observation window, and next decision.
## Your action
Set a total test cap and review date before activating a campaign. Stop at the cap even if the result is inconclusive. Rethink the experiment instead of allowing unlimited spending.
# 9 Build a catalog readers want to continue
''')
pic('exec-0b3a3818-0c76-4d6d-bf94-47e2a3b36f46.png','catalog.png','An author calculates costs beside related books and a path for continuing readers.')
add('''A related catalog gives satisfied readers somewhere useful to go next. Fiction readers may want another story; nonfiction readers may need the next stage of a problem. The connection must make sense to them.
A second book does not automatically make the first profitable. Readers must continue, and additional royalties must exceed the cost of attracting them. Avoid using imagined future sales to justify losses today.
## Estimate read through conservatively
If 100 first-book purchasers produce 30 second-book purchases from that same group over an appropriate window, purchase read-through is 30 percent. Your dashboards may not identify that group cleanly. Monthly book-two units divided by monthly book-one units are a rough indicator, not precise customer-level tracking.
Suppose book one earns $0.3465, book two earns $2.583, and 30 percent continue. Expected royalties per first-book buyer are $0.3465 + (0.30 × $2.583), or about $1.12. That does not support a $2 acquisition cost.
If half of book-two buyers continue to a third book, the probability from book one is 0.30 × 0.50 = 0.15. At a $2.583 third-book royalty, that adds about $0.39, bringing the total to about $1.51. Do not count every sequel as though everyone buys the entire series.
## Make continuation easy
Provide the next book’s accurate title and working link when available. If it is not available, say so. Do not promise a release date you cannot reasonably meet. Finish the current book’s promise before asking the reader to buy something else.
## Your action
Write one sentence explaining why a satisfied reader would want your next book. Calculate a conservative two-book scenario. If the economics do not support a discount, let book one earn at its ordinary price.
# 10 Follow a six week launch plan
Use this calendar as a framework. Move dates to fit editing, readers, and platform processing. Readiness matters more than an arbitrary launch-day spike.
## Four weeks before publication
Complete editing, define the reader promise, audit the listing, and prepare the advance file. Decide initial price, distribution, and total budget. Begin recruitment. If major revision remains, move the launch before asking people to read.
## Three weeks before publication
Deliver copies, test support instructions, and correct delivery problems. Prepare accurate cover and publication details. Record confirmed announcements and promotional placements; follower counts are not guaranteed buyers.
## Two weeks before publication
Inspect the ebook on different screens. Check navigation, images, current pricing rules, and program terms. Prepare messages without sending premature review links. Allow time for platform processing and price changes.
## Publication week
Check the live listing, sample, price, and links. Tell interested subscribers the book is available. Send the live page to advance readers when appropriate. Handle problems before buying more traffic. Record spending, paid units, and estimated royalties.
## One week after publication
Send one gentle follow-up if appropriate. Review recurring feedback and delivery issues. Check advertising limits. Avoid changing several variables at once.
## Two weeks after publication
Write a short review of spending, earnings, reader fit, and lessons. Choose one next action: improve the listing, continue a modest test, pause promotion, or begin a related book.
# 11 Rescue a quiet launch
A slow start rarely identifies its own cause. Find the earliest stage that is not working rather than assuming price or review count explains everything.
## Nobody responds
Check audience fit, invitation clarity, and reading time. Improve the invitation before paying to reach more of a poorly matched audience.
## Readers cannot open the file
Correct permissions or formatting, send a tested link, and extend the window if necessary. A technical obstacle is not evidence that readers are ungrateful.
## Readers do not review
Some readers do not post. Check your optional invitation and link, then send at most the planned follow-up. A campaign can provide useful feedback without generating many public reviews.
## Several readers report the same weakness
Investigate repeated, specific problems. Test confusing instructions yourself. Repair a misleading description. Do not bargain for rating changes or argue with reviewers. Improve the product for future readers; existing reviews may remain after updates.
## Sales lose money
Separate royalties from retail sales value and subtract promotion spending. Check whether discounts replace full-price sales. Pause activities that exceed your loss limit. More units do not automatically solve negative margins.
## Your action
Choose a change addressing the earliest failing stage. Define an observable improvement and give the change a fair test before making another.
# 12 Copy and adapt these messages
Replace bracketed fields before sending. These are author messages, not scripts for readers’ reviews.
## Public invitation
I am inviting readers who enjoy [genre or subject] to receive a free advance copy of [Title]. It is [brief description and approximate length], planned for publication on [date].
If it sounds like a good fit and you have time during [reading window], request a copy here: [link]. An honest review after publication is welcome but entirely optional. No rating is requested. You are free to stop reading or decline to review.
## Confirmation
Thank you for your interest in [Title]. I plan to send the [format] copy on [date]. Publication is expected on [date]. If your availability changes, you can opt out using [method]. Receiving the book does not obligate you to review it.
## Delivery
Your advance copy of [Title] is ready: [download link]. To open it, [tested instructions]. If it does not work, reply with the problem and I will help.
I will send the live book page after publication. Any review is optional and your opinion is your own. Please do not redistribute the file.
## Publication
[Title] is available: [book page link]. If you have read it and would like to share an honest review, use the review option on that page. Reviewing is optional; no rating is requested. If you review, please mention that you received a free advance copy. Thank you for giving the book your time.
## One gentle follow up
A final note about [Title]: if you have finished and wish to leave an honest review, the page is [link]. If you already reviewed, thank you; no reply is needed. If you have not finished or prefer not to review, that is completely fine. I will not send further review reminders for this book.
## Request inside the book
Thank you for reading. If you would like to help other readers decide whether this book suits them, consider leaving an honest review on the store where you found it. Your review is optional, and your own experience is what matters.
# 13 Build your publishing workbook
Copy these prompts into a notebook or spreadsheet. Keep the original assumptions so you can compare them with actual results.
## Reader sheet
Book title and primary reader:
Reader promise:
Five comparable books and observation dates:
What the book delivers:
What it does not promise:
One product-page improvement before launch:
## Campaign sheet
Recruitment channels and copy limit:
Delivery date and reading window:
Publication target:
Consent and opt-out method:
Delivery test completed:
Maximum cost and follow-up date:
## Money sheet
List price and estimated royalty per copy:
Promotion cost and break-even additional units:
Maximum acceptable loss:
Actual royalties and promotion spending:
Net contribution after promotion:
Production costs still to recover:
Net contribution is what remains after the campaign’s direct spending. Do not call it full profit without considering editing, cover design, software, other production costs, and relevant taxes. Keep suitable business records.
## Decision log
Date and change:
Reason and what stayed the same:
Observation window:
Result and uncertainty:
Next action:
A useful entry might say: “Raised price from $2.99 to $3.99. No new promotion. Compared two similar four-week periods. Royalties rose, but traffic may have changed. Hold for another period.” This is more useful than “Higher prices work.”
# 14 Check the Kindle edition before release
A Word manuscript is an editable source, not proof that every Kindle screen will look correct. Preview the converted ebook before publishing. Readers can change text size, so favor a simple structure that reflows.
Use real headings and functioning navigation. Keep illustrations inline with alternative text. Avoid essential instructions based on fixed page numbers or color alone. A grayscale device should still convey every important point.
If a table becomes difficult on a small screen, use short entries instead. Do not embed essential prose inside images. External resources should supplement the book’s promised value, not replace it.
## Publication details
Confirm author or pen name, title, subtitle, rights information, description, cover, categories, keywords, territories, price, and optional program enrollment. Use accurate metadata. Prepare a separate store cover that remains readable as a thumbnail.
Check the converted file size and royalty estimate. Preserve image legibility while avoiding unnecessarily large files. Inspect pictures in grayscale and at small sizes.
KDP requires disclosure of AI-generated text and images when submitting a book, including generated content subsequently edited. AI-assisted editing is treated differently. Answer according to how the actual material was produced. [8]
## Final reader test
Ask someone to open the ebook, navigate to a chapter, follow an example, inspect an image, and return to the contents. Fix confusing steps. Finish with an editorial read for repetition, unexplained terms, and unsupported claims.
The book is ready for its next stage when it delivers its promise, readers can use it comfortably, and promotion has a spending limit. Sales remain uncertain. A clear process gives you a way to learn without handing control of the budget to hope.
# Official sources and updates
These sources support cited platform rules. Planning methods and hypothetical examples are editorial guidance, not Amazon requirements. Accessed September 27, 2026.
[1] Amazon KDP Customer Reviews
https://kdp.amazon.com/en_US/help/topic/G202101910
[2] Amazon Community Guidelines — follow the marketplace-specific Community Guidelines link on the Customer Reviews page above.
[3] Amazon KDP eBook List Price Requirements
https://kdp.amazon.com/en_US/help/topic/G200634560
[4] Amazon KDP Digital Book Pricing Page
https://kdp.amazon.com/en_US/help/topic/G200634500
[5] Amazon KDP Select
https://kdp.amazon.com/en_US/select
[6] Amazon KDP Kindle Countdown Deals
https://kdp.amazon.com/en_US/help/topic/G201293780
[7] Amazon KDP Paperback Royalty
https://kdp.amazon.com/en_US/help/topic/G201834330
[8] Amazon KDP Content Guidelines
https://kdp.amazon.com/en_US/help/topic/G200672390
[9] Amazon KDP Kindle Publishing Guidelines
https://kdp.amazon.com/en_US/help/topic/GU72M65VRFPH43L6
Independent guide. Not affiliated with or endorsed by Amazon. Check current official terms before acting.
''')
# Clickable contents with chapter bookmarks.
headings=[p for p in d.paragraphs if p.style.name=='Heading 1']
anchor=d.paragraphs[5]
toc=OxmlElement('w:p');anchor._p.addnext(toc)
from docx.text.paragraph import Paragraph
tp=Paragraph(toc,anchor._parent);tp.style=d.styles['Heading 1'];tp.add_run('Contents')
last=tp
for i,h in enumerate(headings):
 b=OxmlElement('w:bookmarkStart');b.set(qn('w:id'),str(i));b.set(qn('w:name'),'ch'+str(i));h._p.insert(0,b)
 e=OxmlElement('w:bookmarkEnd');e.set(qn('w:id'),str(i));h._p.append(e)
 p=OxmlElement('w:p');last._p.addnext(p);last=Paragraph(p,anchor._parent)
 link=OxmlElement('w:hyperlink');link.set(qn('w:anchor'),'ch'+str(i));r=OxmlElement('w:r');t=OxmlElement('w:t');t.text=h.text;r.append(t);link.append(r);p.append(link)
d.core_properties.title='Your First Book That Sells';d.core_properties.author=''
out=root/'Your_First_Book_That_Sells_Updated.docx';d.save(out)
print(str(out));print('Words',sum(len(p.text.split()) for p in d.paragraphs))
