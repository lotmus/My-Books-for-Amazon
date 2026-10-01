"""Insert chapter 15 into the Kindle manuscript, then audit it."""
from pathlib import Path
import shutil
from docx import Document
from docx.text.paragraph import Paragraph
from docx.oxml import OxmlElement

root = Path(r"C:\Users\lomus\OneDrive\My Books for Amazon\Your First Book That Sells")
src = root / "First_Review_Kindle_Edition" / "Your_First_Book_That_Sells_Updated.docx"
backup = root / "First_Review_Kindle_Edition" / "Your_First_Book_That_Sells_Updated.before-ch15.docx"
method_a = Path(r"C:\Users\lomus\OneDrive\Documents\kindle kdp\Make write generate a TEXTBOOK.txt")
method_b = Path(r"C:\Users\lomus\OneDrive\Documents\kindle kdp\Make write generate a  Popular science book.txt")

CHAPTER = [
    ("Heading 1", "15 Write one kind of book"),
    ("Normal", "Chapter 1 assumes a manuscript exists. If yours does not, stop the launch calendar. A price, a reviewer list, and an advertisement cannot repair a book that has no form. Write the book in one form, then come back."),
    ("Heading 2", "Choose textbook or popular explanation"),
    ("Normal", "A textbook and a popular explanation of the same subject are different products. Do not write both in one file and hope the reader sorts them out."),
    ("Normal", "A textbook is for a reader who must be able to do the next step. Each section needs a term that means one thing, a result stated so it can be used, and practice. Keep the order a student in that field already expects, unless a different order clearly teaches better. Rearranging for its own sake wastes the reader."),
    ("Normal", "A popular explanation is for a reader who can leave. What holds them is the voice, the picture you choose, and the path from what they believe now to what they can see at the end. Problem sets are not the product. A clever sentence that does not move the path is not the product either."),
    ("Normal", "Name the form in a sentence you can put on the title page. If you cannot, you do not yet know which reader you are writing for, and chapter 1 will not save you."),
    ("Heading 2", "Use one set of words"),
    ("Normal", "Write a single page of the words you will use, and the one meaning each word has in this book. The first time a word appears, define it in the sentence. After that, do not rename it because another book used a finer term."),
    ("Normal", "If you learned the subject from more than one teacher, they will have used different names for the same thing. Pick one name. A reader cannot check your facts while translating your vocabulary on every page."),
    ("Heading 2", "Set the path before you fill the pages"),
    ("Normal", "Write the path on one page before you draft sections. For a textbook, that page is the sequence of things the reader must be able to do. For a popular explanation, it is what the reader believes at the start and what they can see at the end. Fill the sections inside that page. Do not invent a new path every time a paragraph gets interesting."),
    ("Normal", "Another book's chapter order is not yours just because you remember it. If you can retell the book by listing someone else's chapters, you have not made a path. You have copied a table of contents."),
    ("Heading 2", "Make the examples yours"),
    ("Normal", "An example is part of the teaching. If you can point to the book where you first met it, it is still that book's example, even after you change a few words. Write a new one with your own numbers, your own case, or your own picture."),
    ("Normal", "In a textbook, put your careful effort on the few exercises that carry the course. Ordinary practice can follow a standard type, addition or a short proof or a worked substitution, but the numbers and the wording have to be written for this book. In a popular explanation, replace any comparison a well-read stranger would recognize as someone else's. A plain comparison you can explain is better than a famous one you borrowed."),
    ("Heading 2", "Check the facts, then the voice"),
    ("Normal", "A textbook fails on a wrong definition, a wrong unit, or a step that does not follow. Look those up. A popular explanation fails the same way, and it also fails when one section sounds like a lecture and the next sounds like a speech. Write four lines about how this book talks: how direct, how much humor, whether it says you. Draft first. Then make one pass for that sound. Do not stop on every paragraph of the first draft to make it lively."),
    ("Normal", "All numerical examples in this book are teaching examples, not reports of what you will earn. The same rule applies inside the manuscript you are writing. If a number is a measurement, a date, or a rule, name where a reader can check it. If you cannot name the source, cut the number."),
    ("Heading 2", "Your action"),
    ("Normal", "Write the form in one sentence, the word list on one page, and the path on one page. Mark every example you did not invent, and replace it before you ask anyone to read. Then begin chapter 1. Do not buy a cover, a review service, or an advertisement for a manuscript that still has no form."),
]

POINTER = "If the manuscript is not written yet, do that work in chapter 15 before you use the rest of this book."


def insert_before(paragraph, text, style):
    new_p = OxmlElement("w:p")
    paragraph._p.addprevious(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    new_para.style = style
    if text:
        new_para.add_run(text)
    return new_para


def insert_after(paragraph, text, style):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    new_para.style = style
    if text:
        new_para.add_run(text)
    return new_para


if not backup.exists():
    shutil.copy2(src, backup)
else:
    shutil.copy2(backup, src)

doc = Document(str(src))
pointer_host = None
contents_host = None
anchor = None
for para in doc.paragraphs:
    text = para.text.strip()
    style = para.style.name if para.style is not None else ""
    if text.startswith("A finished manuscript is the beginning of a publishing business."):
        pointer_host = para
    if style == "Normal" and text == "14 Check the Kindle edition before release":
        contents_host = para
    if style == "Heading 1" and text == "Official sources and updates":
        anchor = para
        break

if pointer_host is None or contents_host is None or anchor is None:
    raise SystemExit("missing pointer, contents line, or sources heading")

insert_after(pointer_host, POINTER, "Normal")
insert_after(contents_host, "15 Write one kind of book", "Normal")

cursor = anchor._p.getprevious()
if cursor is None:
    raise SystemExit("sources heading has no previous paragraph")
# The previous sibling may not be wrapped yet. Insert after a real paragraph object.
previous = None
for para in doc.paragraphs:
    if para._p is anchor._p:
        break
    previous = para
if previous is None:
    raise SystemExit("no paragraph before sources")
for style, text in CHAPTER:
    previous = insert_after(previous, text, style)

doc.save(str(src))

# Audit
doc = Document(str(src))
texts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
joined = "\n".join(texts)
print("POINTER", POINTER in joined)
print("CH15", "15 Write one kind of book" in joined)
print("CONTENTS", sum(1 for t in texts if t == "15 Write one kind of book"))
print("SOURCES_AFTER", joined.find("15 Write one kind of book") < joined.rfind("Official sources and updates"))

start = None
for i, t in enumerate(texts):
    if t == "15 Write one kind of book":
        start = i
end = texts.index("Official sources and updates", start)
chapter_bits = texts[start + 1 : end]
chapter = "\n".join(chapter_bits)
print("CHAPTER_WORDS", len(chapter.split()))

sources = (method_a.read_text(encoding="utf-8") + "\n" + method_b.read_text(encoding="utf-8")).lower()
needles = [
    "similarity audit",
    "commercially defensible",
    "nomenclature",
    "do not copy",
    "50 pages",
    "40–60",
    "40-60",
    "quintessential",
    "rhetorical",
    "time boxes",
    "diminishing returns",
    "legal review",
    "condense",
    "equalized",
    "ai does almost",
]
print("--- PHRASE FLAGS ---")
low = chapter.lower()
for n in needles:
    if n in low:
        print("FLAG", n)
print("--- DONE FLAGS ---")
