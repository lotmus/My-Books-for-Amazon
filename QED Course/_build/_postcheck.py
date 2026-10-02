# -*- coding: utf-8 -*-
from docx import Document
from docx.oxml.ns import qn

p = r"C:\Users\lomus\OneDrive\My Books for Amazon\Science Books\QED Course\Complete QED Course.docx"
d = Document(p)
text = "\n".join(x.text for x in d.paragraphs)
checks = [
    "Lesson 62 Vacuum Polarization",
    "Lesson 68 Running Electric Charge",
    "Lesson 75 The Anomalous Magnetic Moment",
    "e²_eff",
    "F₂(0)",
    "Bloch-Nordsieck",
    "Glossary of symbols",
    "sky is blue",
    "define the essential objects used in running",
    "d=4−2ε in dimensional regularization",
    "vertex factor −ieγμ",
    "+ieγ",
    "Rayleigh",
    "collinear fermion",
    "Course capstone",
]
open(r"C:\Users\lomus\OneDrive\My Books for Amazon\Science Books\QED Course\_build\_post.txt","w",encoding="utf-8").write(
    "\n".join(("YES " if c in text else "NO  ") + c for c in checks)
)
print("wrote post")
# heading sample around 62
h1=[]
for para in d.paragraphs:
    if para.style and para.style.name=="Heading 1" and para.text.startswith("Lesson 6"):
        h1.append(para.text)
open(r"C:\Users\lomus\OneDrive\My Books for Amazon\Science Books\QED Course\_build\_post.txt","a",encoding="utf-8").write("\n---\n"+"\n".join(h1[:40]))
print("h1", len(h1))
