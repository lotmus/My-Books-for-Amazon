# -*- coding: utf-8 -*-
import os as _os
HERE = _os.path.dirname(_os.path.abspath(__file__))
ROOT = _os.path.dirname(HERE)
from docx import Document
from collections import Counter

p = _os.path.join(ROOT, "Complete QED Course.docx")
d = Document(p)
paras = d.paragraphs
words = sum(len(x.text.split()) for x in paras)
styles = Counter(x.style.name for x in paras if x.style)
h1 = [x.text.strip() for x in paras if x.style and x.style.name == "Heading 1" and x.text.strip()]
out = []
out.append("words %d paras %d tables %d" % (words, len(paras), len(d.tables)))
out.append("h1 count %d" % len(h1))
out.append("styles %s" % styles.most_common(12))
sec = d.sections[0]
out.append("margins L%.3f R%.3f T%.3f B%.3f" % (
    sec.left_margin.inches, sec.right_margin.inches, sec.top_margin.inches, sec.bottom_margin.inches))
from docx.oxml.ns import qn
body = d.element.body
out.append("hyperlinks %d" % len(body.findall(".//" + qn("w:hyperlink"))))
xml = d.element.xml
out.append("drawings %d" % xml.count("w:drawing"))
out.append("oMath %d" % xml.count("m:oMath"))
out.append("--- H1 LESSONS ---")
out.append("\n".join([t for t in h1 if t.startswith("Lesson ") or t.startswith("Part ") or t.startswith("Glossary") or t.startswith("Course") or t.startswith("Consolidated") or t.startswith("Table")]))
text = "\n".join(x.text for x in paras)
needles = [
    "Update Field", "__QEDCourseTOC_END__", "sky is blue", "Rayleigh",
    "define the essential objects", "d=4−2ε", "Copyright", "ISBN",
    "Glossary of symbols", "Bloch-Nordsieck", "α/(2π)", "e²_eff",
    "FIG:", "collinear", "How to use this lesson",
]
out.append("--- NEEDLES ---")
for n in needles:
    out.append("%s %s" % ("YES" if n in text else "NO ", n))
# sample lesson 68 and 76 openings
for label, key in (("L68", "Lesson 68 "), ("L76", "Lesson 76 "), ("L86", "Lesson 86 "), ("GLOSS", "Glossary of symbols"), ("CAP", "Course capstone")):
    i = text.find(key)
    out.append("--- %s ---" % label)
    out.append(text[i:i+900] if i>=0 else "MISSING")
dest = _os.path.join(HERE, "_review2_facts.txt")
open(dest, "w", encoding="utf-8").write("\n".join(out))
print("ok", dest, "words", words)
