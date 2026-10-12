import re, io, os
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

B = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
P = os.path.join(B, "Step_Stories_Pics_Final")
SERIES = "The Step Stories"
TITLE1, TITLE2 = "The Wee Scallywag", "and the Gentle Elf"
AUTHOR = "Anneliese Tauber"

raw = open(os.path.join(B, "Step_Stories_Texts", "Book1_Story_v7.txt"), encoding="utf-8").read()
body = raw[raw.index("Chapter 1."):]
blocks = [b.strip("\n") for b in re.split(r"\n\s*\n", body) if b.strip()]

place = {
    '“Since the porridge,”': '02_wall_run.jpg',
    'They bowed. The scallywag bowed deeper.': '03_log_launch.jpg',
    'Scally opened the lid.': '05_ghost_trap.jpg',
    'They got out on the far bank': '06_toad_ford.jpg',
    'Under the willow, a little way off': '07_pike_gucci.jpg',
    'It was a thin, wavering howl': '08_howl_moon.jpg',
    'Hal borrowed the ghost’s lantern': '09_moon_model.jpg',
    'A tall figure stood at the end of the road.': '10_fog_stranger.jpg',
    'It was the same mill, and it was the wrong way round.': '12_mirror_mill.jpg',
    'Scally came out backwards': '13_pull_bag.jpg',
    'There was a soft, enormous whumph.': '19_flour_flash.jpg',
    'Out came a hundred years of boiled sweets': '14_stranger_flour.jpg',
    'There was a small, bright, interested sound': '18_fire_night.jpg',
    '“No,” said Tabitha. She held up one hand': '20_village_woman.jpg',
    'It was past midnight when the third day closed.': '15_stay_step.jpg',
    'Mrs Crumb was already counting sacks.': '16_parents_dawn.jpg',
    'The sun rose.': '17_ghost_sunrise.jpg',
    '“It’s on me,” said Sparefunk': '21_bird_on_head.jpg',
    '“They are all different,” said Hopsi at last.': '22_flowers_welcome.jpg',
    'Hal Millersen sat on the mill step': 'c01a_empty_yard.jpg',
    'Behind them, in the empty yard': 'c01c_boots_walk.jpg',
    'Next morning two messengers': 'c02a_messengers.jpg',
    'Hal looked at his slate.': 'c02b_lever_test.jpg',
    'A small girl with dark brown hair': 'c02c_storm_carriage.jpg',
    'Ten minutes later the sun came out': 'c02d_rainbow.jpg',
    '“Old Miller Millersen,” said the ghost.': 'c03a_ghost_tea.jpg',
    '“Real walls,': 'c03b_marble_hill.jpg',
    'He went in anyway.': 'c03c_mud_morning.jpg',
    '“Think of a recipe,': 'c04a_seedlings.jpg',
    'Hal guessed that the grasshoppers': 'c04b_counting_hoppers.jpg',
    'The toad took it with great care': 'c04c_toad_crust.jpg',
    '“A stick stays straight,': 'c05a_bent_stick.jpg',
    'Scally went straight across': 'c05b_puddle_race.jpg',
    'They caught the pike at dusk.': 'c05c_supper_bonfire.jpg',
    'Hal looked at the rectangular eggs': 'c06a_eggs_moon.jpg',
    'Afterwards he sat apart': 'c07a_apology.jpg',
    'She had set a crust down beside her': 'c08b_gucci_crust.jpg',
    '“She is right,” said Hopsi. “The Moon pulls': 'c08c_tides.jpg',
    '“Look,” said Tabitha. She had stopped.': 'c09a_willow_path.jpg',
    'Gucci would not go in.': 'c09b_gucci_refuses.jpg',
    '“It got through the earth,': 'c09c_light_coin.jpg',
    'Hal looked at it. He read it from the other end.': 'c09d_mirror_tag.jpg',
    'The stones glowed, a thin white light': 'c09e_stones_stop.jpg',
    '“Sparefunk,” said Hal.': 'c10a_name_reveal.jpg',
    'In the pond, very small and silver': 'c11a_reflection.jpg',
    'At dusk, when the village had gone home': 'c11b_boots_road.jpg',
    '“Pairs,” said Hal.': '11_culvert.jpg',
    '“He sent me.': 'd01_the_deal.jpg',
    'That evening Hal tried the wheel.': 'd02_jamming_the_wheel.jpg',
    '“A bag,” said the ghost.': 'd03_miller_falls.jpg',
    '“It’s mine,” said Scally.': 'd05_stamp_bonfire.jpg',
    '“I did it,”': 'd06_confession.jpg',
    '“No,” said Hal.': 'd07_hal_says_no.jpg',
    'Behind the mill, in the long grass, a rectangular egg': 'd08_egg_hatching.jpg',
    'Hal thought.': 'c01b_bargain_boots.jpg',
    '“Mind those,” he said.': '01_step_bargain.jpg',
    'Hal followed the sound across the yard.': '04_ghost_wheel.jpg',
    'Down by the bog, in the thin early light': 'd10_shadow.jpg',
    '“A soft voice sent me': 'd11_ghost_story_a.jpg',
    '“Boiled sweets,” said the ghost.': 'd12_ghost_story_b.jpg',
    '“You tell the truth when it’s funny,” said Hopsi quietly.': 'd09_bog_fall.jpg',
    'His father, the miller, had climbed down into the pit': 'd13_mother_marmalade.jpg',
    'On the second day they walked down to the village': 'd14_note_on_door.jpg',
    '“It was a fire,” said the miller, quite gently': 'd15_miller_scolds.jpg',
    'She told him. A fire does not stay the size it starts.': 'd16_what_if.jpg',
    '“There was a mill at Hollin Brook,': 'd17_hollin_brook.jpg',
}


def img_bytes(path, w=640, q=68):
    im = Image.open(path).convert("RGB")
    if im.width > w:
        im = im.resize((w, round(im.height * w / im.width)), Image.LANCZOS)
    b = io.BytesIO()
    im.save(b, "JPEG", quality=q, optimize=True)
    b.seek(0)
    return b


def make_cover():
    src = Image.open(os.path.join(B, "Cover_Front_1600x2560.jpg")).convert("RGB")
    cv = src.resize((1000, 1600), Image.LANCZOS)
    out = os.path.join(P, "cover_1000x1600.jpg")
    cv.save(out, quality=88)
    return out


# ANCHOR CHECK: every picture anchor must match exactly one story paragraph
_ps = [x for x in blocks]
for _k, _v in place.items():
    _n = sum(1 for x in _ps if x.startswith(_k))
    if _n != 1:
        print("ANCHOR PROBLEM", _n, _v, repr(_k[:50]))

COVER = make_cover()

d = Document()
sec = d.sections[0]
sec.page_width = Inches(6)
sec.page_height = Inches(9)
for a in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
    setattr(sec, a, Inches(0.8))
st = d.styles["Normal"]
st.font.name = "Georgia"
st.font.size = Pt(13)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Georgia")
st.paragraph_format.space_after = Pt(10)
st.paragraph_format.line_spacing = 1.3
NAVY = RGBColor(0x1F, 0x3A, 0x5F)
for n, sz in (("Heading 1", 24), ("Heading 2", 19)):
    h = d.styles[n]
    h.font.name = "Georgia"
    h.font.size = Pt(sz)
    h.font.bold = True
    h.font.color.rgb = NAVY
    h.element.rPr.rFonts.set(qn("w:eastAsia"), "Georgia")
    h.paragraph_format.keep_with_next = True
    h.paragraph_format.space_before = Pt(14)

bid = [0]


def bookmark(par, name):
    bid[0] += 1
    s = OxmlElement("w:bookmarkStart")
    s.set(qn("w:id"), str(bid[0]))
    s.set(qn("w:name"), name)
    e = OxmlElement("w:bookmarkEnd")
    e.set(qn("w:id"), str(bid[0]))
    par._p.insert(0, s)
    par._p.append(e)


def link(par, text, anchor):
    h = OxmlElement("w:hyperlink")
    h.set(qn("w:anchor"), anchor)
    r = OxmlElement("w:r")
    rp = OxmlElement("w:rPr")
    c = OxmlElement("w:color")
    c.set(qn("w:val"), "1F3A5F")
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rp.append(c)
    rp.append(u)
    r.append(rp)
    t = OxmlElement("w:t")
    t.text = text
    r.append(t)
    h.append(r)
    par._p.append(h)


def pagebreak():
    d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def centered(text, size=13, bold=False, italic=False, after=10, color=None):
    p = d.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(after)
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.bold = bold
    r.italic = italic
    if color:
        r.font.color.rgb = color
    return p


def pic(name, w_in=4.4):
    if name.startswith('22_'):
        w_in = 3.4
    p = d.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(img_bytes(os.path.join(P, name)), width=Inches(w_in))


def heading(text, level, anchor=None, center=True):
    h = d.add_paragraph(text, style="Heading %d" % level)
    if center:
        h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if anchor:
        bookmark(h, anchor)
    return h


def lead_bold(par, text):
    m = re.match(r"^(\d+\. [^.]+\.)(.*)$", text, re.S)
    if m:
        r = par.add_run(m.group(1))
        r.bold = True
        par.add_run(m.group(2))
    else:
        par.add_run(text)



def toc_link():
    p = d.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    link(p, "↑ Contents", "toc")


LINKS = [  # (item number, exact phrase in the story, first occurrence only)
    (1, "A lever"),
    (2, "A rainbow stood"),
    (3, "Think of a recipe"),
    (4, "Toads eat the ones they can see"),
    (5, "Light slows down"),
    (6, "The quickest"),
    (7, "called tunnelling"),
    (8, "back to front"),
    (9, "Sunlight is eight minutes old"),
    (10, "Make a model"),
    (11, "Two high tides and two low"),
    (12, "Wolves howl to call each other"),
    (13, "Fire needs three things"),
    (14, "That is the sign. That is enough."),
    (15, "Foxglove."),
    (16, "Reach or throw"),
    (17, "That counts"),
]
linked_done = set()


def _link_el(text, anchor, bm_id, bm_name):
    """A hyperlink wrapped in a bookmark so the explanation can link back to it."""
    s = OxmlElement("w:bookmarkStart"); s.set(qn("w:id"), str(bm_id)); s.set(qn("w:name"), bm_name)
    h = OxmlElement("w:hyperlink"); h.set(qn("w:anchor"), anchor)
    r = OxmlElement("w:r"); rp = OxmlElement("w:rPr")
    c = OxmlElement("w:color"); c.set(qn("w:val"), "1F3A5F")
    u = OxmlElement("w:u"); u.set(qn("w:val"), "dotted")
    rp.append(c); rp.append(u); r.append(rp)
    tt = OxmlElement("w:t"); tt.text = text; tt.set(qn("xml:space"), "preserve"); r.append(tt)
    h.append(r)
    e = OxmlElement("w:bookmarkEnd"); e.set(qn("w:id"), str(bm_id))
    return [s, h, e]


def add_text_with_links(par, text):
    """Add text to a paragraph, turning the first story use of each science phrase into a link."""
    todo = [(n, ph) for n, ph in LINKS if n not in linked_done and ph in text]
    if not todo:
        par.add_run(text)
        return
    # handle one link at a time, left to right
    todo.sort(key=lambda x: text.index(x[1]))
    pos = 0
    for n, ph in todo:
        i = text.index(ph, pos) if ph in text[pos:] else -1
        if i < 0:
            continue
        if i > pos:
            par.add_run(text[pos:i])
        bid[0] += 1
        for el in _link_el(ph, "tt%d" % n, bid[0], "src%d" % n):
            par._p.append(el)
        linked_done.add(n)
        pos = i + len(ph)
    if pos < len(text):
        par.add_run(text[pos:])


chapters = [b.split("\n")[0] for b in blocks if b.startswith("Chapter ")]
TOC = [("Meet the Characters", "cast")] + [(c, "ch%d" % (i + 1)) for i, c in enumerate(chapters)] + [
    ("Coming Next: Book 2", "next"), ("Hopsi\u2019s True Things", "true"), ("For Parents and Teachers", "guide"), ("About the Author", "about"), ("Also by Anneliese Tauber", "alsoby")]

# cover (inline, ~1000x1600)
sec.top_margin = sec.bottom_margin = Inches(0.2)
sec.left_margin = sec.right_margin = Inches(0.2)
p = d.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(0)
p.add_run().add_picture(COVER, height=Inches(8.5))
sec2 = d.add_section()
for a in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
    setattr(sec2, a, Inches(0.8))

# title page
for _ in range(5):
    d.add_paragraph()
centered(TITLE1, 34, True, color=NAVY)
centered(TITLE2, 34, True, color=RGBColor(0xA8, 0x4B, 0x2A))
centered("Book 1 of " + SERIES, 15, italic=True, after=30)
centered(AUTHOR, 18)
pagebreak()
heading("Contents", 1, "toc")
for t, a in TOC:
    p = d.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    link(p, t, a)
pagebreak()

heading("Meet the Characters", 1, "cast")
cast = [
    ("Hal (Halbert) Millersen", "8. The miller\u2019s son. Brown, tousled hair, chalk on his fingers, a slate under his arm. Dry and steady. Says \u201cAcceptable\u201d when things are bad."),
    ("Titch Mitchel", "8. Dark brown hair under a red knitted cap. Quick, practical, always has a crust. Tears bread into even pieces."),
    ("Scally, the wee scallywag", "A little devil: short sharp pointed ears, messy hair with a leaf in it, a grin like a lit match. Makes bargains and runs up walls. His true name is the puzzle of the book."),
    ("Hopsi (Elfie Didlhop)", "9. An elf: long swept-back ears, long wavy light-brown hair, a needle and a bowl of hops. Only says what is true."),
    ("Tabitha Smoke", "8. Curly black hair in a bun. Warm and watchful. Asks the one question nobody else thought of."),
    ("Gucci", "A baby wolf with a grumpy face and a very big appetite for crusts."),
    ("Old Miller Millersen", "A ghost, with a flour-sack cloak and a dusty hat. Polite, lonely, and Hal\u2019s great-great-grandfather."),
    ("The Soft Voice", "A tall thin stranger with sweets, a smile too wide and a quicker way. Never take it."),
]
for n, t in cast:
    p = d.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(9)
    r = p.add_run(n)
    r.bold = True
    p.add_run("\n" + t)
centered("Ages 5\u20139. Little ones can follow the whole story from the pictures alone.", 12, italic=True)
pagebreak()

ci = 0
in_true = False
first_chapter = True
for b in blocks:
    if b == "---":
        continue
    if b.startswith("Chapter "):
        ci += 1
        if not first_chapter:
            toc_link()
            pagebreak()
        first_chapter = False
        heading(b, 2, "ch%d" % ci)
        continue
    if b == "The End":
        centered("The End", 18, True).paragraph_format.space_before = Pt(18)
        toc_link()
        pagebreak()
        heading("Coming Next", 1, "next")
        centered("The Step Stories, Book 2", 14, italic=True, after=6)
        centered("A New Path", 30, True, color=RGBColor(0xA8, 0x4B, 0x2A), after=18)
        for line in [
            "The mill wheel is turning. The three white stones still keep their line. The way back to the step is open.",
            "But a boy in the pond winked. A small brown bird will not leave. And the rectangular eggs in the long grass are not all hatched.",
            "Hal, Scally, Hopsi, Titch, Tabitha and Gucci set out along the stream, on a path none of them has walked before, to find where the water goes and what else is true.",
            "Coming soon.",
        ]:
            centered(line, 14, italic=(line == "Coming soon."), after=12)
        toc_link()
        pagebreak()
        continue
    if b.startswith("Hopsi\u2019s True Things"):
        heading("Hopsi\u2019s True Things", 1, "true")
        in_true = True
        continue
    if len(b) < 40 and "\n" not in b and not re.search(r"[a-z]", b):
        centered(b, 16, True, after=12)
        continue
    if b.startswith("Little ones take after"):
        p = d.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.4)
        lines = b.split("\n")
        for i, l in enumerate(lines):
            r = p.add_run(l)
            r.italic = True
            if i < len(lines) - 1:
                r.add_break()
        continue
    p = d.add_paragraph()
    if in_true:
        lead_bold(p, b.replace("\n", " "))
        mm = re.match(r"^(\d+)\. ", b)
        if mm:
            n = int(mm.group(1))
            bid[0] += 1
            st = OxmlElement("w:bookmarkStart"); st.set(qn("w:id"), str(bid[0])); st.set(qn("w:name"), "tt%d" % n)
            en = OxmlElement("w:bookmarkEnd"); en.set(qn("w:id"), str(bid[0]))
            p._p.insert(0, st); p._p.insert(1, en)
            if n in linked_done:
                p.add_run(" ")
                hl = OxmlElement("w:hyperlink"); hl.set(qn("w:anchor"), "src%d" % n)
                rr = OxmlElement("w:r"); rp2 = OxmlElement("w:rPr")
                cc = OxmlElement("w:color"); cc.set(qn("w:val"), "A84B2A"); uu = OxmlElement("w:u"); uu.set(qn("w:val"), "single")
                rp2.append(cc); rp2.append(uu); rr.append(rp2)
                tx = OxmlElement("w:t"); tx.text = "\u21a9 Back to the story"; rr.append(tx); hl.append(rr)
                p._p.append(hl)
    else:
        add_text_with_links(p, b.replace("\n", " "))
    for k, v in place.items():
        if b.startswith(k):
            pic(v)
            break

toc_link()
pagebreak()
heading("For Parents and Teachers", 1, "guide")


def para(t, bold_lead=None):
    p = d.add_paragraph()
    if bold_lead:
        r = p.add_run(bold_lead + " ")
        r.bold = True
    p.add_run(t)


para("This is a fantasy adventure to read aloud at ages 5\u20139. Children from about 5 can follow the whole story from the pictures alone (a \u201cpicture walk\u201d: turn the pages and ask what is happening), and older children can read it by themselves. Every fact the elf Hopsi states is true, and every piece of magic is marked as magic. The children test things before they believe them. The book promises curiosity, new words and good scientific habits, not exam results.")
heading("The habits of a scientist in the story", 2, center=False)
for lead, t in [
    ("\u201cHow do we know?\u201d", "Hal asks it. The answer is always: look, count, test, or make a model."),
    ("Write it down.", "Hal\u2019s slate is a lab notebook: the lever test, the grasshopper counts, the bent stick, the moon shapes."),
    ("Make a model.", "An apple and a lantern show why the moon changes shape."),
    ("Looks can fool you.", "A fish is deeper than it looks, a stick looks bent in water, and a mirror swaps left and right."),
    ("Say only what you know.", "Hopsi says only what is true, and \u201cI don\u2019t know\u201d counts."),
    ("Notice what others miss.", "Tabitha spots the firm path from the reeds."),
    ("Guess first, then test.", "Sparefunk\u2019s fake fly fails, and Hal\u2019s first lever try takes ten fingers."),
]:
    para(t, lead)
heading("What each chapter teaches", 2, center=False)
for lead, t in [
    ("The Stuck Wheel:", "a problem to solve against the clock; the idea of small print."),
    ("The Lord of the Riverbank:", "levers (guess, test, be wrong), a fair test, hail and rainbows."),
    ("The Ghost in the Wall:", "tunnelling, as a marble that sometimes turns up beyond a hill, and why the sun shines."),
    ("The Toad of the Ford:", "genes as a recipe, variation, counting data, why the colour that fits becomes common, poisonous plants, telling someone where you are going, and what to do when a friendly stranger asks you to go with them: No, Go, Yell, Tell."),
    ("The Pike and the Torn Tag:", "light bends when it slows in water, quickest is not shortest, a stick test and a puddle race."),
    ("The Howl and the Moon:", "why the moon changes shape, with an apple and a lantern, why wolves really howl, and a scientist’s “I do not know”."),
    ("The Hot Night:", "what a fire needs, how to smother a small one, why a bonfire must be put out properly, and why a grown-up must always be fetched."),
    ("The Girl with the Apples:", "a welcome gift when words are not allowed, careful observation, and tides: why the sea rises and falls twice a day."),
    ("The Quick Way and the True Way:", "reading evidence, reversed writing in a mirror, an honest answer about a crack of light, and why a cloud of flour dust burns all at once."),
    ("The True Name:", "reading mirror writing, and telling the truth when it is hard."),
    ("Sunday:", "sunlight is about eight minutes old."),
]:
    para(t, lead)
heading("Staying safe: what this book models", 2, center=False)
for lead, tx in [
    ("Strangers.", "The woman with the puppy and the Soft Voice are two strangers. Neither is spotted by looks. The warning signs are a request to keep a secret, offers of sweets or pets, a claim that a parent sent them, or a child being asked for help. The children say No, go to a safe place (the baker), yell, and tell. Agree a family code word, like marmalade."),
    ("Fire.", "Shout, get out, get a grown-up. A fire needs fuel, heat and air. Never leave a bonfire burning; drown it, stir it, drown it again. Smoke is the real danger, so stay low and stay out. Please make a family escape plan together."),
    ("Water.", "Reach or throw, never go in after a friend. Never go to a pond or stream without a grown-up."),
    ("Roads.", "Stop, look and listen before you cross."),
    ("Plants, medicines and cleaning things.", "Never taste berries, leaves, mushrooms, medicines or cleaning products unless a grown-up says so. Foxglove is poisonous."),
    ("Animals.", "Let a strange animal sniff a flat hand first, and never grab one. Ask the owner."),
    ("Hot things.", "Turn pot handles inward. Keep small hands away from stoves and kettles."),
    ("Telling someone.", "Hal writes where he is going and when he will be back. Go in pairs, and never go anywhere alone in the dark."),
    ("Screens and online.", "There are no screens in this story, but the Soft Voice is exactly the person who messages a child online, offering something nice and asking for a secret. The same rules apply: no, tell a grown-up, and never keep that kind of secret."),
]:
    para(tx, lead)
heading("Talk about it", 2, center=False)
for q in [
    "Why did Hal write his results down? What would have happened if he hadn\u2019t?",
    "Why is a fish deeper than it looks? What else looks different from how it is?",
    "If the shortest way is the slowest, how would you plan a journey?",
    "What shape is the moon tonight? Where will it be in a week?",
    "Who made the fire go wrong, and how did everyone put it right? What would you do if you saw a fire?",
    "Which part of the story is magic and which is true? How can you tell?",
    "Who noticed something the others missed? What did they notice?",
]:
    para(q)
heading("Words to grow", 2, center=False)
para("lever, force, hail, rainbow, inherit, variation, camouflage, bend (refraction), reflection, tunnelling, phase, tide.")
heading("A note for adults", 2, center=False)
para("Light travels about three-quarters as fast in water as in air, so \u201ca quarter slower\u201d is a child-sized number. A fish is deeper than it looks when you look into the water from above. Sunlight takes about eight minutes and twenty seconds to reach Earth. The recipe picture of genes is an analogy: real genes are made of DNA. Tunnelling is real quantum behaviour of very small particles, and it matters in how the Sun releases energy; the marble picture is only a way in. The moon\u2019s phases come from how much of its sunlit half we can see, not from the Earth\u2019s shadow. Tides come from the Moon\u2019s gravity pulling more strongly on the near side of the Earth than the far side, with the Sun adding to it at full and new moon; the dough picture is a simple way in. A mirror swaps front and back, which is why writing reads the wrong way round in a mirror; the Mirror Mill itself is magic. Stranger safety is not about looks: a stranger can be a man or a woman and can seem kind. The signs are a request to keep a secret, offers of sweets or pets, a claim that a parent sent them, or a child being asked for help. Please agree a family code word, like marmalade in the story. A bonfire must be drowned, stirred and drowned again until the ashes are cold to the hand. A cloud of flour dust, unlike a heap of flour, can burn all at once; this is real, and it has burned down mills, so the flash in the story is for reading about and never for trying. These are seed ideas, meant to make children curious, not to be tested at school yet.")
toc_link()
pagebreak()
heading("About the Author", 1, "about")
para("Anneliese Tauber grew up in Bavaria, where the old stories have mills in them and the mills have ghosts. Today she is a mother of four, and her children have taught her that the best bedtime stories answer the questions they ask, and keep the real answers true. The Step Stories are for curious children and the grown-ups who read to them.")
centered("If your child enjoyed this book, a short, kind review on Amazon helps other curious families find it. Thank you.", 13, italic=True)
toc_link()
pagebreak()
heading("Also by Anneliese Tauber", 1, "alsoby")
centered("Bavarian Cooking Made Simple", 15, True, after=6)
centered("Classic dishes from a Bavarian kitchen, written so that any cook can enjoy a taste of Bavaria.", 12, italic=True)
toc_link()
centered("Typeface in the e-book edition: Comic Neue, © The Comic Neue Project Authors, used under the SIL Open Font License 1.1.", 10, italic=True)
d.core_properties.title = TITLE1 + " " + TITLE2
d.core_properties.author = AUTHOR
import datetime
now = datetime.datetime.now(datetime.timezone.utc)
d.core_properties.created = now
d.core_properties.modified = now
d.core_properties.last_modified_by = AUTHOR
print("UNLINKED:", [n for n, _ in LINKS if n not in linked_done])
dst = os.environ.get("OUTDOCX") or os.path.join(B, "The_Wee_Scallywag_and_the_Gentle_Elf_Book1.docx")
d.save(dst)
print(dst, os.path.getsize(dst) // 1024, "KB")
