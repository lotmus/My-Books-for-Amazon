# -*- coding: utf-8 -*-
"""Append Book 3 Part II (bipolar) and Part III (below 20 nm). Does not rebuild Part I."""

import math
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

OUT = Path(r"C:\Users\lomus\OneDrive\My Books for Amazon\EE - Series\EE3\Semiconductor_Physics_and_Devices_Book3.docx")

VT = 0.026
BETA = 100.0
IB = 10e-6
IC = BETA * IB
IE = IC + IB
IS = 1e-15
VBE = VT * math.log(IC / IS + 1.0)


def set_run(run, name, size, bold=False, color=None):
    run.bold = bold
    run.font.name = name
    run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = run._element.makeelement(qn("w:rFonts"), {})
        rpr.append(rfonts)
    for a in ("w:ascii", "w:hAnsi", "w:cs"):
        rfonts.set(qn(a), name)


def add(doc, text, style="Normal"):
    p = doc.add_paragraph(text, style=style)
    if p.runs:
        if style == "Heading 1":
            set_run(p.runs[0], "Amazon Ember", 18, True, RGBColor(0, 0, 255))
        elif style == "Heading 2":
            set_run(p.runs[0], "Amazon Ember", 18, True, RGBColor(0, 0, 255))
        else:
            set_run(p.runs[0], "Calibri", 11, False, RGBColor(0, 0, 0))
    return p


def main():
    if not OUT.exists():
        raise SystemExit(f"missing {OUT} — do not invent Part I")
    doc = Document(OUT)
    add(doc, "Part II. The Bipolar Transistor", "Heading 1")
    add(
        doc,
        "Part I stopped at the diode. This part is the two-junction device that sits on that physics: "
        "the bipolar junction transistor, in this book's own words, checked against Tietze/Schenk 16th, Sedra/Smith, and Gray/Meyer. "
        "It is not a paste of those books. The constants stay the ones Part I declared: VT = 26 mV, IS = 1 fA unless a new value is named.",
    )
    add(doc, "Chapter 8. Two Junctions, One Current", "Heading 1")
    add(
        doc,
        "An NPN transistor is an n-type emitter, a p-type base, and an n-type collector. Forward-bias the emitter-base junction and reverse-bias the collector-base junction. "
        "Electrons injected into the thin base mostly survive to the collector. The small base current is the ones that recombine, plus the holes that went the other way. "
        "Kirchhoff still holds: IE = IC + IB. Current gain β = IC/IB is large because the base is thin and lightly doped relative to the emitter.",
    )
    add(
        doc,
        f"Declared picture: β = {BETA:.0f}, IB = 10 µA. Then IC = {IC*1e3:.2f} mA and IE = {IE*1e3:.2f} mA. "
        f"The emitter-base voltage that produces that collector current, from the same diode law as Chapter 5, is "
        f"VBE = VT ln(IC/IS + 1) = {VBE:.3f} V with IS = 1 fA. That is about 0.72 V, the usual hand value of 0.7 V with the book's own constants underneath.",
    )
    add(
        doc,
        "A PNP is the same story with the signs flipped. SiGe HBTs raise β and fT by grading germanium in the base; they are still bipolar devices, not a new law. "
        "GaAs and InP HBTs were named in Chapter 7. They belong in a radio, not in a 5 V logic family.",
    )
    add(doc, "Key idea. The collector current is an amplified copy of the base current only while the junctions stay biased that way. Saturate the collector-base junction and β collapses.", "Normal")

    add(doc, "Chapter 9. Bias, Early Effect, and the Hybrid-π", "Heading 1")
    add(
        doc,
        "A useful bias network sets IE with a resistor in the emitter, or with a current source, so β variation does not steal the operating point. "
        "That is Book 2's analog chapter seen from the device side. Tietze/Schenk 16th remains the circuit-technique check. "
        "The Early effect is the collector voltage modulating the base width, so IC rises slightly with VCE. A finite Early voltage VA makes the output resistance ro ≈ VA/IC. "
        "Declared: IC = 1 mA, VA = 100 V, then ro = 100 kΩ. That resistor sits in parallel with the load in a small-signal model.",
    )
    add(
        doc,
        "The hybrid-π model at a declared IC = 1 mA: gm = IC/VT = 1 mA / 26 mV = 38.5 mS. rπ = β/gm = 100 / 0.0385 ≈ 2.60 kΩ. "
        "Those two numbers, plus ro, are the bipolar amplifier for Book 2's 'ideal op-amp insides' later in this volume. "
        "This chapter stops at the model. The op-amp as a system is Book 2.",
    )
    add(doc, "Key idea. gm = IC/VT. Raise the current and the transconductance rises. The thermal voltage is still 26 mV.", "Normal")

    add(doc, "Part III. Below 20 nm", "Heading 1")
    add(doc, "Chapter 10. The Fin, Then the Sheet", "Heading 1")
    add(
        doc,
        "Below about 20 nm, a planar MOSFET cannot control the channel from one gate. The fin stands the channel on edge so the gate wraps three sides. "
        "The next nodes replace the fin with a gate-all-around nanosheet, then a stacked CFET, an nFET over a pFET, when the fin has no more height to give. "
        "Weste and Harris still teach how you draw the circuit. You count fins instead of a continuous width. "
        "Taur and Ning, 3rd edition, is the device book. Veendrick is the survey of what else changes: leakage, RC in the back end, variability, yield. "
        "The foundry PDK outranks any textbook at these nodes. Gate-all-around and CFET past the fin are papers and IRDS, not a second Taur.",
    )
    add(
        doc,
        "AlScN as a piezo film that a CMOS line will run is named here as a material. The 2026 MEMS products that use it are Book 9. "
        "Do not turn this chapter into a phone-microphone catalog.",
    )
    add(doc, "Key idea. The transistor shrank until the package and the memory stack became the product. That sentence is finished in Book 9.", "Normal")

    add(doc, "Appendix E. Sources for Parts II and III", "Heading 1")
    add(doc, "Tietze, Schenk, and Gamm, 16th edition (2019), bipolar and amplifier chapters.")
    add(doc, "Sedra and Smith; Gray and Meyer (or Gray, Hurst, Lewis, and Meyer) for analog bipolar and IC.")
    add(doc, "Taur and Ning, Fundamentals of Modern VLSI Devices, 3rd. Veendrick, Nanometer CMOS ICs. Saha on FinFET. The PDK for anything you will actually tape out.")

    # Fix stale 21-book pointers in existing appendix C if present
    for p in doc.paragraphs:
        t = p.text
        new = t
        new = new.replace("Radio matching and lines belong to Book 4. Noise figures and receiver cascades belong to Book 5. Modulation belongs to Book 6. Phase-locked loops belong to Book 7. Microwave lines and the historical Radiation Laboratory volumes belong to Book 8. Antennas belong to Book 10. Emissions and grounding belong to Book 11. Sampling, multirate filters, and a receiver you can code belong to Book 14. Linear regulators and switch-mode supplies belong to Book 15. Bode plots as a general language belong to Book 16. High-speed copper, the hot loop of a buck converter, and the feedback divider of an LDO belong to Books 15 and 18 together.",
                          "Matching and antennas belong to Book 4. Modulation and SDR belong to Book 5. Noise, PLLs, and high-power RF belong to Book 6. Emissions belong to Book 7. Regulators and magnetics belong to Book 8. Packaging and 2026 systems belong to Book 9. Bode plots as a general language belong to Book 2.")
        new = new.replace("The power-amplifier design, the 28 V layout, and the 650 V switch are Books 17 and 15.",
                          "The power-amplifier design and the 28 V layout are Book 6. The 650 V switch is Book 8.")
        if new != t and p.runs:
            p.runs[0].text = new
            for r in p.runs[1:]:
                r.text = ""

    doc.save(OUT)
    print(f"appended Part II/III to {OUT}")
    print(f"VBE {VBE:.4f} IC_mA {IC*1e3:.2f} gm_mS {1e-3/VT*1e3:.1f}")


if __name__ == "__main__":
    main()
