# -*- coding: utf-8 -*-
"""Book 3 completion (2026-10-01): Parts IV-VI (Chapters 12-19).

Input: Book 3 already rebuilt by build_book3_parts2_3.py (Parts I-III, Ch 1-11).
Inserts Part IV (Ch 12-14, inside the amplifier), Part V (Ch 15-17, CMOS as a
chip), Part VI (Ch 18-19, light and power) before Appendix A in the house style,
adds answers (Appendix B) and constants (Appendix A), rewrites Appendix C as the
handoff to the series, extends Appendix E's sources, and updates front matter.
Usage: python build_book3_parts4_6.py <in.docx> <out.docx>   (DOCX only)
"""
import copy, sys
from pathlib import Path
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC = Path(__file__).resolve().parent / "src"
CHS = [f"ch{n}" for n in range(12, 20)]
CONSTS = [
    "Noise temperature for Part IV: 290 K, so 4kT = 1.60×10⁻²⁰ J (the same reference as Book 2). Device equations keep VT = 26 mV.",
    "Channel-noise factor of a long-channel MOSFET: γ = 2/3.",
    "Pelgrom matching constant, working value for a recent process: AVt = 3 mV·µm.",
    "Critical field for avalanche, hand values: silicon 3×10⁵ V/cm, 4H-SiC 2.5×10⁶ V/cm, GaN 3.3×10⁶ V/cm.",
    "Photon energy and wavelength: λ (nm) ≈ 1240 / E (eV).",
    "Human-body ESD model: 100 pF through 1.5 kΩ.",
]
APPC = [
    "Appendix C. Where This Book Hands Off",
    "This book now runs from the crystal to the chip: carriers and junctions (Part I), the bipolar transistor (Part II), the MOSFET to the nanosheet (Part III), the amplifier built from them (Part IV), CMOS as a chip (Part V), and the junction as a source and detector of light and as a high-voltage switch (Part VI). Each later book of the series starts from a device this one has explained and stops explaining it.",
    "Book 4 takes the transistor's small-signal model and the GaAs, GaN, and InP devices of Chapter 7 into matching networks, transmission lines, and antennas. Book 5 uses the converter and sample-and-hold limits of Book 2 and the noise of Chapter 14 to build modems and software-defined radios. Book 6 uses the load line of Chapter 7, the noise of Chapter 14, and the power-device limits of Chapter 19 for the transmitter chain, the PLL, and the receiver noise cascade. Book 7 treats the ESD and latch-up of Chapter 17 as system-level immunity and test, together with emissions. Book 8 takes the power MOSFET, superjunction, SiC, GaN, and IGBT of Chapter 19 into converters, magnetics, and the grid. Book 9 picks up the chip of Part V at the package: interposers, high-bandwidth memory built from the DRAM cells of Chapter 16, the 875 A of a modern accelerator, and MEMS.",
    "Two habits travel with the reader. First, every quantity in this book was computed from a declared hand value; when a vendor model or a process design kit gives a different number, the vendor number wins for that part. Second, below about 20 nm the PDK outranks every textbook, including this one.",
]
SOURCES = [
    "Part IV: Gray, Hurst, Lewis, and Meyer, Analysis and Design of Analog Integrated Circuits; Razavi, Design of Analog CMOS Integrated Circuits; Johns and Martin, Analog Integrated Circuit Design. Pelgrom, Duinmaijer, and Welbers, \"Matching Properties of MOS Transistors\" (IEEE JSSC, 1989). Motchenbacher and Connelly, Low-Noise Electronic System Design, for Chapter 14.",
    "Part V: Rabaey, Chandrakasan, and Nikolić, Digital Integrated Circuits; Weste and Harris, CMOS VLSI Design. Jacob, Ng, and Wang, Memory Systems, for Chapter 16. Amerasekera and Duvvury, ESD in Silicon Integrated Circuits; the JEDEC/ESDA JS-001 (HBM), JS-002 (CDM), and JESD78 (latch-up) standards; Black's 1969 electromigration papers.",
    "Part VI: Sze and Ng, Physics of Semiconductor Devices, for recombination, photodetectors, and solar cells; Schubert, Light-Emitting Diodes; Baliga, Fundamentals of Power Semiconductor Devices, for Chapter 19.",
]

def strip_runs(p):
    for c in list(p):
        if c.tag in (qn("w:r"), qn("w:hyperlink"), qn("w:bookmarkStart"), qn("w:bookmarkEnd")):
            p.remove(c)
    for a in list(p.attrib):
        if a.endswith("paraId") or a.endswith("textId"):
            del p.attrib[a]
    return p

def make_p(tpl, text):
    p = strip_runs(copy.deepcopy(tpl))
    r0 = next((r for r in tpl.findall(qn("w:r")) if r.find(qn("w:t")) is not None), None)
    r = OxmlElement("w:r")
    if r0 is not None and r0.find(qn("w:rPr")) is not None:
        r.append(copy.deepcopy(r0.find(qn("w:rPr"))))
    t = OxmlElement("w:t"); t.set(qn("xml:space"), "preserve"); t.text = text
    r.append(t); p.append(r)
    return p

def set_text(par, new):
    for r in par.runs[1:]:
        r._r.getparent().remove(r._r)
    par.runs[0].text = new

def main(inp, outp):
    doc = Document(inp)
    P = doc.paragraphs
    H1 = lambda p: p.style is not None and p.style.name == "Heading 1"
    idx = lambda pred: next(i for i, p in enumerate(P) if pred(p))
    if any(H1(p) and p.text.startswith("Chapter 12.") for p in P):
        sys.exit("already rebuilt")
    idx(lambda p: H1(p) and p.text.startswith("Chapter 11."))
    i5 = idx(lambda p: H1(p) and p.text.startswith("Chapter 5."))
    T = {"H1": P[i5]._p, "P": P[i5 + 1]._p, "H2": P[i5 + 2]._p, "EQ": P[i5 + 4]._p}
    iwe = idx(lambda p: p.text == "Worked Example 5.1")
    T["WEP"], T["WES"] = P[iwe + 1]._p, P[iwe + 2]._p
    T["KEY"] = P[idx(lambda p: p.text.startswith("Key idea. Forward current"))]._p
    T["PR"] = P[idx(lambda p: p.text.startswith("1. Same diode."))]._p
    ipart = idx(lambda p: H1(p) and p.text.startswith("Part II."))
    T["PART"], T["PARTP"] = P[ipart]._p, P[ipart + 1]._p
    iab = idx(lambda p: p.style.name == "Heading 2" and p.text == "Chapter 11" and P.index(p) > iwe)
    T["AH"], T["AP"] = P[iab]._p, P[iab + 1]._p
    iA = idx(lambda p: H1(p) and p.text.startswith("Appendix A."))
    T["CONST"] = P[iA + 1]._p
    T = {k: copy.deepcopy(v) for k, v in T.items()}
    anchor = P[iA]._p
    answers, words = {}, {}
    for ch in CHS:
        prn = 0; answers[ch] = []
        for line in (SRC / f"{ch}.txt").read_text(encoding="utf-8").splitlines():
            if not line.strip():
                continue
            tag, _, b = line.partition(":"); b = b.strip()
            if tag == "ANS":
                answers[ch].append(b); continue
            if tag == "PR":
                prn += 1; b = f"{prn}. {b}"
            if tag not in T:
                raise ValueError(f"{ch}: unknown tag {tag}")
            anchor.addprevious(make_p(T[tag], b))
        assert prn == len(answers[ch]), ch
    # Appendix A constants
    P = doc.paragraphs
    iB = idx(lambda p: H1(p) and p.text.startswith("Appendix B."))
    last = P[iB - 1]._p
    for c in CONSTS:
        p = make_p(T["CONST"], c); last.addnext(p); last = p
    # Appendix B answers
    P = doc.paragraphs
    iC = idx(lambda p: H1(p) and p.text.startswith("Appendix C."))
    a = P[iC]._p
    for ch in CHS:
        a.addprevious(make_p(T["AH"], f"Chapter {int(ch[2:])}"))
        for k, s in enumerate(answers[ch], 1):
            a.addprevious(make_p(T["AP"], f"{k}. {s}"))
    # Appendix C rewrite
    P = doc.paragraphs
    iC = idx(lambda p: H1(p) and p.text.startswith("Appendix C."))
    iD = idx(lambda p: H1(p) and p.text.startswith("Appendix D."))
    set_text(P[iC], APPC[0])
    body_tpl = copy.deepcopy(P[iC + 1]._p)
    for p in P[iC + 1:iD]:
        p._p.getparent().remove(p._p)
    anchorD = doc.paragraphs[iD - (iD - iC - 1)]._p
    for t in APPC[1:]:
        anchorD.addprevious(make_p(body_tpl, t))
    # Appendix E sources
    P = doc.paragraphs
    iE = idx(lambda p: H1(p) and p.text.startswith("Appendix E."))
    set_text(P[iE], "Appendix E. Sources for Parts II to VI")
    last = P[-1]._p; stpl = copy.deepcopy(P[iE + 1]._p)
    for s in SOURCES:
        p = make_p(stpl, s); last.addnext(p); last = p
    # front matter
    P = doc.paragraphs
    for p in P[:12]:
        t = p.text
        if t.startswith("Parts I–III"):
            set_text(p, "Parts I–VI  —  From the Crystal to the Chip")
        elif t.startswith("Copyright © 2026"):
            set_text(p, t.replace("This file holds Parts I through III of a longer book. Parts IV through VI are not in it yet.",
                                  "This file holds the complete book, Parts I through VI."))
        elif H1(p) and t == "How to Use This Part":
            set_text(p, "How to Use This Book")
        elif t.startswith("Book 1 taught the circuit laws"):
            set_text(p, t.replace("This part opens those blocks. By the end of it you can say",
                                  "This book opens those blocks. By the end of Part I you can say"))
        elif t.startswith("Two numbers are fixed for the whole part"):
            set_text(p, t.replace("fixed for the whole part", "fixed for the whole book"))
        elif t.startswith("Electron and hole mobilities"):
            set_text(p, t.replace("Whenever a problem in this part", "Whenever a problem in this book"))
        elif t.startswith("Later parts of this book take up"):
            set_text(p, "Part II takes up the bipolar transistor and Part III the MOSFET, down to the fin and the nanosheet. "
                        "Part IV builds the op-amp of Book 2 from those transistors and derives the noise they make. Part V treats "
                        "CMOS as a chip: gate delay and power, memory cells, and the failures that end a chip's life. Part VI turns "
                        "the junction into a source and detector of light and into a high-voltage switch. The circuit check for "
                        "Parts II to IV is Tietze, Schenk, and Gamm, Halbleiter-Schaltungstechnik, 16th edition (2019), whose op-amp "
                        "chapter was rewritten for that edition. High-speed boards, radio layout, switching regulators, and "
                        "software-defined radio are later books in this series, named in Appendix C so this book does not pretend "
                        "to cover them.")
    doc.save(outp); print("saved", outp)

if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
