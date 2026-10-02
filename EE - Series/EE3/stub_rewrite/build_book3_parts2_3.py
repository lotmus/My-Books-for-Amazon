# -*- coding: utf-8 -*-
"""Book 3 stub rewrite (2026-10-01).

Removes the templated filler chapters fill_to_200.py appended ("Chapter <slogan>"
with "Row 1..16" families) and the short Part II/III sketches append_part2.py left
after the appendices (Chapters 8-10, ~200 words each). Inserts full Part II
(Chapters 8-9, bipolar) and Part III (Chapters 10-11, MOSFET to nanosheet) before
Appendix A in Part I's house style, adds their answers to Appendix B and their
constants to Appendix A, keeps Appendix E (sources for Parts II-III) last, and
updates the front-matter and Appendix C/D sentences that said those parts were missing.
Usage: python build_book3_parts2_3.py <in.docx> <out.docx>   (DOCX only)
"""
import copy, re, sys
from pathlib import Path
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC = Path(__file__).resolve().parent / "src"
CHS = ["ch08", "ch09", "ch10", "ch11"]

def strip_runs(p):
    for c in list(p):
        if c.tag in (qn("w:r"), qn("w:hyperlink"), qn("w:bookmarkStart"), qn("w:bookmarkEnd")):
            p.remove(c)
    for a in list(p.attrib):
        if a.endswith("paraId") or a.endswith("textId"):
            del p.attrib[a]
    return p

def make_p(tpl, text):
    p = strip_runs(copy.deepcopy(tpl))
    r0 = next((r for r in tpl.findall(qn("w:r")) if r.find(qn("w:t")) is not None), None)
    r = OxmlElement("w:r")
    if r0 is not None and r0.find(qn("w:rPr")) is not None:
        r.append(copy.deepcopy(r0.find(qn("w:rPr"))))
    t = OxmlElement("w:t"); t.set(qn("xml:space"), "preserve"); t.text = text
    r.append(t); p.append(r)
    return p

def set_text(par, new):
    for r in par.runs[1:]:
        r._r.getparent().remove(r._r)
    par.runs[0].text = new

def main(inp, outp):
    doc = Document(inp)
    P = doc.paragraphs
    H1 = lambda p: p.style is not None and p.style.name == "Heading 1"
    idx = lambda pred: next(i for i, p in enumerate(P) if pred(p))
    if any(H1(p) and p.text.startswith("Chapter 11.") for p in P):
        sys.exit("already rebuilt")
    T = {"H1": P[idx(lambda p: H1(p) and p.text.startswith("Chapter 5."))]._p}
    i5 = idx(lambda p: H1(p) and p.text.startswith("Chapter 5."))
    T["P"], T["H2"], T["EQ"] = P[i5 + 1]._p, P[i5 + 2]._p, P[i5 + 4]._p
    iwe = idx(lambda p: p.text == "Worked Example 5.1")
    T["WEP"], T["WES"] = P[iwe + 1]._p, P[iwe + 3 - 1]._p
    T["KEY"] = P[idx(lambda p: p.text.startswith("Key idea. Forward current"))]._p
    T["PR"] = P[idx(lambda p: p.text.startswith("1. Same diode."))]._p
    ipart = idx(lambda p: H1(p) and p.text.startswith("Part II."))
    T["PART"], T["PARTP"] = P[ipart]._p, P[ipart + 1]._p
    iab5 = idx(lambda p: p.style.name == "Heading 2" and p.text == "Chapter 5" and P.index(p) > iwe)
    T["AH"], T["AP"] = P[iab5]._p, P[iab5 + 1]._p
    iA = idx(lambda p: H1(p) and p.text.startswith("Appendix A."))
    T["CONST"] = P[iA + 1]._p
    # keep Appendix E (sources for Parts II and III) as deep copies
    iE = idx(lambda p: H1(p) and p.text.startswith("Appendix E."))
    iStub = idx(lambda p: H1(p) and re.match(r"Chapter [^0-9]", p.text))
    keepE = [copy.deepcopy(p._p) for p in P[iE:iStub]]
    T = {k: copy.deepcopy(v) for k, v in T.items()}
    body = doc.element.body
    el, removed = P[ipart]._p, 0
    while el is not None:
        nxt = el.getnext()
        if el.tag != qn("w:sectPr"):
            body.remove(el); removed += 1
        el = nxt
    print("removed elements:", removed)
    P = doc.paragraphs
    anchor = P[idx(lambda p: H1(p) and p.text.startswith("Appendix A."))]._p
    answers = {}
    for ch in CHS:
        prn = 0; answers[ch] = []; title = None
        for line in (SRC / f"{ch}.txt").read_text(encoding="utf-8").splitlines():
            if not line.strip():
                continue
            tag, _, b = line.partition(":"); b = b.strip()
            if tag == "ANS":
                answers[ch].append(b); continue
            if tag == "PR":
                prn += 1; b = f"{prn}. {b}"
            if tag == "H1":
                title = b
            tpl = T["H2"] if tag == "H2" else T[tag]
            anchor.addprevious(make_p(tpl, b))
        answers[ch].insert(0, title)
    # Appendix A constants
    P = doc.paragraphs
    iA = idx(lambda p: H1(p) and p.text.startswith("Appendix A."))
    iB = idx(lambda p: H1(p) and p.text.startswith("Appendix B."))
    set_text(P[iA], "Appendix A. Constants Used in This File")
    last = P[iB - 1]._p
    for c in ["Declared bipolar transistor (Part II): β = 100, IS = 1 fA, VA = 100 V, operating point IC = 1 mA, which gives VBE = 0.718 V, gm = 38.5 mS, rπ = 2.60 kΩ, ro = 100 kΩ.",
              "VBE temperature coefficient at constant IC, hand value: −2 mV/°C.",
              "Oxide permittivity εox = 3.9 ε0 = 3.45×10⁻¹³ F/cm (Part III).",
              "Electron saturation velocity in silicon, hand value: vsat = 1.0×10⁷ cm/s.",
              "Symbol note: VT is the thermal voltage; the MOSFET threshold is Vth."]:
        p = make_p(T["CONST"], c); last.addnext(p); last = p
    # Appendix B answers
    P = doc.paragraphs
    iC = idx(lambda p: H1(p) and p.text.startswith("Appendix C."))
    a = P[iC]._p
    for ch in CHS:
        num = int(ch[2:])
        a.addprevious(make_p(T["AH"], f"Chapter {num}"))
        for k, s in enumerate(answers[ch][1:], 1):
            a.addprevious(make_p(T["AP"], f"{k}. {s}"))
    # Appendix C / D / front matter sentences
    P = doc.paragraphs
    iC = idx(lambda p: H1(p) and p.text.startswith("Appendix C."))
    set_text(P[iC], "Appendix C. What This File Does Not Yet Cover")
    set_text(P[iC + 1], P[iC + 1].text.replace(
        "Part II is the bipolar transistor. Part III is the MOSFET and the CMOS gate. Part IV is",
        "Part II, the bipolar transistor, and Part III, the MOSFET through the fin and the nanosheet, are now in this file. Still to come: Part IV is"))
    for p in P:
        if "that are not in this file yet" in p.text:
            set_text(p, p.text.replace("for the bipolar and MOS chapters that are not in this file yet",
                                       "for the bipolar and MOS chapters of Parts II and III"))
        if p.text.startswith("Copyright © 2026") and "Parts II through VI" in p.text:
            set_text(p, p.text.replace("This file is Part I of a longer book. Parts II through VI are not in it yet.",
                                       "This file holds Parts I through III of a longer book. Parts IV through VI are not in it yet."))
        if p.text.strip() == "Part I  —  The Crystal and the Diode":
            set_text(p, "Parts I–III  —  The Crystal, the Diode, and the Transistor")
    # re-append Appendix E at the end of the body
    sect = body.find(qn("w:sectPr"))
    for e in keepE:
        strip_ids = e
        if sect is not None:
            sect.addprevious(e)
        else:
            body.append(e)
    doc.save(outp); print("saved", outp)

if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
