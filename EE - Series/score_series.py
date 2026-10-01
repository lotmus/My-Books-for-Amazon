# -*- coding: utf-8 -*-
"""Score the nine-book series: length, outline coverage, one-home numbers."""

import math
import re
from pathlib import Path
from docx import Document

ROOT = Path(r"C:\Users\lomus\OneDrive\My Books for Amazon\EE - Series")
WPP = 280  # Kindle words per page, this score's declared constant
TARGET_PAGES = 200

BOOKS = [
    (1, ROOT / "EE1" / "Foundations_Book1.docx", ROOT / "Foundations_Book1.docx"),
    (2, ROOT / "EE2" / "Circuits_Components_and_Control_Book2.docx", None),
    (3, ROOT / "EE3" / "Semiconductor_Physics_and_Devices_Book3.docx", None),
    (4, ROOT / "EE4" / "RF_Microwave_and_Antennas_Book4.docx", None),
    (5, ROOT / "EE5" / "Communications_Wireless_and_SDR_Book5.docx", None),
    (6, ROOT / "EE6" / "Transceivers_and_High_Power_RF_Book6.docx", None),
    (7, ROOT / "EE7" / "EMC_Simulation_and_Test_Book7.docx", None),
    (8, ROOT / "EE8" / "Power_and_Energy_Book8.docx", None),
    (9, ROOT / "EE9" / "Packaging_Layout_and_Emerging_Book9.docx", None),
]

CHECKS = {
    1: [("10 V", "resonance companion or body"), ("nine", "series count")],
    2: [("5 ms", "tracker"), ("40 nV", "Johnson"), ("follower", "buffer split")],
    3: [("26 mV", "VT"), ("FinFET", "below 20 nm"), ("bipolar", "Part II")],
    4: [("Friis", "link"), ("mmWave", "not a separate book")],
    5: [("320 MHz", "Wi-Fi 7"), ("modem", "bits to waveform")],
    6: [("10.13", "Ropt GaN"), ("3.13", "Ropt LDMOS"), ("0.794", "1 dB")],
    7: [("CISPR", "method"), ("7layers", "lab"), ("Hermon", "lab")],
    8: [("1.70", "rack gap"), ("skin", "magnetics")],
    9: [("875", "current"), ("55", "HBM ratio"), ("AlScN", "MEMS")],
}


def words_of(path: Path) -> tuple[int, int, str]:
    if not path.exists():
        return 0, 0, "MISSING"
    doc = Document(path)
    w = 0
    h1 = 0
    blob = []
    for p in doc.paragraphs:
        t = p.text.strip()
        if t:
            w += len(t.split())
            blob.append(t)
        st = p.style.name if p.style else ""
        if st == "Heading 1":
            h1 += 1
    for tbl in doc.tables:
        for row in tbl.rows:
            for c in row.cells:
                w += len(c.text.split())
                blob.append(c.text)
    return w, h1, "\n".join(blob)


def score_length(pages: float) -> float:
    # 200 pages = 10. Partial credit, cap 10.
    return max(0.0, min(10.0, 10.0 * pages / TARGET_PAGES))


def main():
    print(f"Declared: {WPP} words/page, target {TARGET_PAGES} pages.")
    print(f"{'Bk':<4}{'words':>8}{'pp':>7}{'H1':>5}{'len':>6}{'cov':>6}{'home':>6}{'tot':>6}  file")
    totals = []
    for n, primary, fallback in BOOKS:
        candidates = [p for p in (primary, fallback) if p is not None and p.exists()]
        if not candidates:
            print(f"{n:<4}{'':>8}{'':>7}{'':>5}{'0':>6}{'0':>6}{'0':>6}{'0':>6}  MISSING")
            totals.append(0)
            continue
        scored = [(words_of(p)[0], p) for p in candidates]
        path = max(scored, key=lambda x: x[0])[1]
        w, h1, blob = words_of(path)
        pages = w / WPP
        slen = score_length(pages)
        keys = CHECKS.get(n, [])
        hits = sum(1 for k, _ in keys if k.lower() in blob.lower() or k in blob)
        shome = 10.0 * hits / max(1, len(keys))
        # coverage: heading count vs a full book (~25 H1 for 200 pages). Cap 10.
        scov = max(0.0, min(10.0, 10.0 * h1 / 25.0))
        # total: length 40%, coverage 30%, one-home 30%
        tot = 0.4 * slen + 0.3 * scov + 0.3 * shome
        print(f"{n:<4}{w:8d}{pages:7.1f}{h1:5d}{slen:6.1f}{scov:6.1f}{shome:6.1f}{tot:6.1f}  {path.name}")
        missing = [lab for k, lab in keys if k.lower() not in blob.lower() and k not in blob]
        if missing:
            print(f"     missing home checks: {', '.join(missing)}")
        totals.append(tot)
    series = sum(totals) / max(1, len(totals))
    print()
    print(f"Series mean (equal books): {series:.1f} / 10")
    print("Length score is honest: spines cannot fake 200 pages.")
    print("Coverage is Heading-1 density. One-home is required numbers/phrases.")


if __name__ == "__main__":
    main()
