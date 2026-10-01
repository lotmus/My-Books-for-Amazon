# -*- coding: utf-8 -*-
"""Book 1 companion: phasors, one resonance example, Fourier / Laplace / wavelets."""

import math
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUT = Path(__file__).resolve().parent / "Complex_Circuit_Analysis.docx"

SQRT2 = math.sqrt(2)
W0 = 1000.0
R = 30.0
L = 0.040
C = 25e-6
V_RMS = 10.0  # 14.14 V peak is 10*sqrt(2), rounded
I_RMS = V_RMS / R
V_L = I_RMS * (W0 * L)
P = I_RMS ** 2 * R
Q_L = I_RMS ** 2 * (W0 * L)
Q_SERIES = (W0 * L) / R
BW = W0 / Q_SERIES

W2 = 2000.0
X2 = W2 * L - 1.0 / (W2 * C)
Z2 = math.hypot(R, X2)
ANG2 = math.degrees(math.atan2(X2, R))
I2 = V_RMS / Z2
P2 = I2 ** 2 * R
Q2 = I2 ** 2 * X2

B1 = 4 / math.pi
B3 = 4 / (3 * math.pi)
B5 = 4 / (5 * math.pi)
T1 = 8 / (math.pi ** 2)
T3 = T1 / 9


def font(run, name, size, bold=False, color=None):
    run.bold = bold
    run.font.name = name
    run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = run._element.makeelement(qn("w:rFonts"), {})
        rpr.append(rfonts)
    for a in ("w:ascii", "w:hAnsi", "w:cs"):
        rfonts.set(qn(a), name)


def shade(style, name, size, bold, color, align, before=0, after=8, page_break=False):
    style.font.name = name
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = color
    pf = style.paragraph_format
    pf.alignment = align
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.page_break_before = page_break
    pf.line_spacing = 1.15
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = style.element.makeelement(qn("w:rFonts"), {})
        rpr.append(rfonts)
    for a in ("w:ascii", "w:hAnsi", "w:cs"):
        rfonts.set(qn(a), name)


def add_p(doc, text, center=False, italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.15
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    font(run, "Calibri", 11, color=RGBColor(0, 0, 0))
    run.italic = italic


def build():
    doc = Document()
    for s in doc.sections:
        s.top_margin = Inches(1)
        s.bottom_margin = Inches(1)
        s.left_margin = Inches(1)
        s.right_margin = Inches(1)
    blue = RGBColor(0x00, 0x00, 0xFF)
    shade(doc.styles["Normal"], "Calibri", 11, False, RGBColor(0, 0, 0), WD_ALIGN_PARAGRAPH.LEFT)
    shade(doc.styles["Heading 1"], "Amazon Ember", 18, True, blue, WD_ALIGN_PARAGRAPH.CENTER, after=12, page_break=True)
    shade(doc.styles["Heading 2"], "Amazon Ember", 18, True, blue, WD_ALIGN_PARAGRAPH.CENTER, before=14, after=8)

    add_p(doc, "Complex Circuit Analysis", center=True)
    title = doc.paragraphs[-1]
    font(title.runs[0], "Amazon Ember", 28, bold=True, color=blue)
    add_p(doc, "A companion to Book 1, Part III", center=True)
    add_p(doc, "Phasors, one resonant loop, and when Fourier is the wrong tool", center=True)
    add_p(doc, "Lothar J. Musiol", center=True)
    add_p(
        doc,
        "Copyright © 2026 Lothar J. Musiol. The numbers in the worked example are exact for the stated parts. "
        "Do not retune them.",
        center=True,
        italic=True,
    )

    doc.add_heading("1. A sinusoid becomes a number", 1)
    add_p(
        doc,
        "Steady-state AC analysis is circuit analysis with jω in place of a derivative. "
        "A source vm cos(ωt + φ) is written as an RMS phasor, V = Vrms at angle φ. "
        "Power work uses RMS. Time is restored only at the end: "
        "v(t) is the real part of √2 V exp(jωt). Differentiation along a sinusoid is multiplication by jω, "
        "so a differential equation becomes algebra. Kirchhoff's laws still hold. Mesh and node analysis "
        "are the DC methods with an impedance where a resistance used to sit.",
    )
    add_p(doc, "Ohm's law in this language is V = Z I.")
    add_p(doc, "A resistor is Z = R. An inductor is Z = jωL. A capacitor is Z = 1/(jωC) = −j/(ωC).")
    add_p(doc, "Series elements add as impedances. Parallel elements add as admittances, Y = 1/Z.")
    add_p(
        doc,
        "Write Z = R + jX. Positive X is inductive. Negative X is capacitive. "
        "The magnitude is the square root of R squared plus X squared. The angle is the inverse tangent of X/R. "
        "The current lags the voltage by that angle when X is positive, and leads it when X is negative.",
    )
    add_p(doc, "Key idea. One frequency, steady state, and a complex number for each voltage and current. That is the whole method.")

    doc.add_heading("2. The loop that cancels", 1)
    add_p(
        doc,
        "A source 14.14 cos(1000 t) volts drives R = 30 Ω, L = 40 mH, and C = 25 μF in series. "
        "The peak 14.14 V is 10√2, so the RMS phasor is 10 V at angle 0. ω = 1000 rad/s.",
    )
    add_p(
        doc,
        f"ZL = jωL = j{W0 * L:.0f} Ω. "
        f"ZC = 1/(jωC) = −j{1/(W0 * C):.0f} Ω. "
        "They are equal and opposite. Z = 30 Ω. "
        f"Check the resonance formula: 1/√(LC) = {1/math.sqrt(L * C):.0f} rad/s. The drive sits on it.",
    )
    add_p(
        doc,
        f"I = 10/30 = 1/3 A at angle 0, which is {I_RMS:.4f} A RMS. "
        f"In time, i(t) = {I_RMS * SQRT2:.3f} cos(1000 t) A.",
    )
    add_p(
        doc,
        f"VR = I R = 10 V at angle 0. "
        f"VL = I · j40 = {V_L:.2f} V at +90°. "
        f"VC = {V_L:.2f} V at −90°. "
        "Add them: 10 + j(40/3) − j(40/3) = 10 V, which is the source. "
        "The inductor and the capacitor each stand taller than the source. They cancel each other. "
        "That is resonance. It is not a broken Kirchhoff law.",
    )
    add_p(
        doc,
        f"Complex power S = V I* = {P:.2f} + j0 VA. All of it is real, and it lands in R: "
        f"I squared times R = {P:.2f} W. "
        f"The inductor takes +{Q_L:.2f} var. The capacitor takes −{Q_L:.2f} var. "
        "Net reactive power at the source is zero. Ideal L and C shuttle energy. They do not keep it.",
    )
    add_p(
        doc,
        f"Quality factor of this series loop: Q = ω0 L / R = {Q_SERIES:.3f}, which is 4/3. "
        f"The same number from 1/(ω0 R C). "
        f"The half-power bandwidth is ω0/Q = {BW:.0f} rad/s.",
    )
    add_p(
        doc,
        f"Move the same parts to ω = 2000 rad/s and the cancellation is gone. "
        f"ZL = j80 Ω, ZC = −j20 Ω, Z = 30 + j60 Ω. "
        f"Magnitude {Z2:.2f} Ω at +{ANG2:.2f}°. "
        f"Current {I2:.4f} A at −{ANG2:.2f}°. "
        f"Real power falls to {P2:.2f} W. Reactive power at the source is +{Q2:.2f} var. "
        "That is the general case: write each Z(jω), add, divide.",
    )
    add_p(doc, "Key idea. At series resonance the source sees only R. Off resonance it sees R + jX, and it must supply vars.")

    doc.add_heading("3. Sums of sines", 1)
    add_p(
        doc,
        "Fourier's series says a periodic waveform is a sum of sines and cosines at integer multiples of one fundamental. "
        "A square wave keeps only the odd harmonics, the 1st, 3rd, 5th, and so on, and each falls as 1/n. "
        "For a wave that swings between −A and +A, the coefficient of harmonic n (n odd) is 4A/(nπ). "
        f"With A = 1, the fundamental is {B1:.3f}, the third is {B3:.3f}, and the fifth is {B5:.3f}. "
        "Those are 1, 1/3, and 1/5 of the fundamental, which is the 1/n rule with the even lines missing.",
    )
    add_p(
        doc,
        "A triangle keeps that same odd set and falls as 1/n squared. "
        f"With the usual peak of 1, the fundamental is 8/π squared, which is {T1:.3f}, "
        f"and the third is one ninth of that, {T3:.3f}. "
        "The high harmonics are much weaker than in the square wave. "
        "That is why a triangle looks gentler: fewer high frequencies, softer edges, less of the harsh distortion a square wave carries. "
        "A sawtooth keeps every integer harmonic, odd and even, and falls only as 1/n. It is as bright in its highs as the square wave, and it is not odd-only.",
    )
    add_p(
        doc,
        "The Fourier transform extends the same idea to signals that do not repeat. "
        "Frequency becomes a continuous curve instead of a comb of lines. "
        "A single pulse maps to a sinc spectrum. White noise is flat. "
        "For a periodic current you still read discrete lines. For a one-shot pulse you read a continuous shape.",
    )
    add_p(
        doc,
        "The FFT is the fast algorithm for that transform on samples. "
        "A direct sum is on the order of N squared operations. The FFT is on the order of N log N, "
        "because it splits the record into smaller pieces and combines them. "
        "For 1024 samples that is 1024 squared, about a million operations, against 1024 times 10, about ten thousand. "
        "That gap is why a spectrum analyzer, an OFDM modem, and a software radio can keep up with the samples.",
    )

    doc.add_heading("4. Laplace, and what a wavelet is for", 1)
    add_p(
        doc,
        "Laplace is Fourier with a complex frequency s = σ + jω. The extra real part carries growth, decay, and initial conditions. "
        "Differentiation becomes multiplication by s, so a circuit equation or a feedback loop becomes algebra. "
        "Stability is where the poles sit in the s-plane. A control loop is designed by placing those poles. "
        "There is no standard wavelet Bode plot. Phasor analysis is this same algebra restricted to the imaginary axis, after the transient has died.",
    )
    add_p(
        doc,
        "Fourier is the natural language of a linear time-invariant circuit. A sinusoid in is a sinusoid out, at the same frequency, with a new amplitude and phase. "
        "That is why an impedance, a filter, an antenna, and an emission mask are written as H(jω). "
        "The harmonics of a square wave, the bandwidth of an op-amp, and an FCC limit are Fourier statements.",
    )
    add_p(
        doc,
        "Wavelets keep time and frequency together. Fourier tells you which frequencies are present and throws away when they occurred. "
        "A wavelet uses a short window at high frequency and a long window at low frequency, so a spike stays a spike. "
        "They are the tool for a spark, a speech burst, an edge in a photograph, and any signal that is not stationary. "
        "They do not hand you a single frequency axis or a transfer function.",
    )
    add_p(
        doc,
        "A rectangular current in a switch-mode supply is the Fourier case. "
        "It is close to a square wave, so it carries a comb of harmonics, and those harmonics leave the converter and disturb the mains. "
        "An LC filter is a low-pass impedance. It attenuates the high harmonics and leaves the DC, or the line-frequency current, that the supply meant to draw. "
        "You design that filter with Z(jω), not with a mother wavelet.",
    )
    add_p(
        doc,
        "A single step or spike on that same mains is the wavelet case. "
        "An FFT of a long record smears the step across the whole spectrum and loses the moment it happened. "
        "A wavelet can say both that a high-frequency feature occurred and when it occurred. "
        "Use it to find the glitch. Use Fourier to decide whether the switcher's steady hash is inside the mask.",
    )
    add_p(
        doc,
        "Key idea. Fourier for steady periodic content. Laplace for transients and loops. "
        "The FFT when you are measuring. Wavelets when the event's when matters. "
        "Electronics is linear and time-invariant most of the time, which is why the first three carry the design and the fourth is an extra tool.",
    )
    add_p(
        doc,
        "Phasors assume one frequency and steady state. Switch-on transients, clipping, and a circuit whose parts change with time need Laplace or a time-step simulation. "
        "A distorted but periodic wave is handled by giving each harmonic its own Z(jnω) and adding the results.",
    )

    doc.save(OUT)
    print(f"wrote {OUT}")
    print(f"Q {Q_SERIES:.6f} BW {BW:.3f}")
    print(f"I2 {I2:.6f} P2 {P2:.4f} Q2 {Q2:.4f} Z2 {Z2:.4f}")


if __name__ == "__main__":
    build()
