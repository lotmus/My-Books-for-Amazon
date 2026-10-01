# -*- coding: utf-8 -*-
"""Grow Books 4-9. Does not touch Book 2 (merged Core Circuits) or Book 1."""

import math
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
BLUE = RGBColor(0x00, 0x00, 0xFF)


def set_run_font(run, name, size_pt, bold=False, color=None):
    run.bold = bold
    run.font.name = name
    run.font.size = Pt(size_pt)
    if color is not None:
        run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = run._element.makeelement(qn("w:rFonts"), {})
        rpr.append(rfonts)
    for a in ("w:ascii", "w:hAnsi", "w:cs"):
        rfonts.set(qn(a), name)


def shade(style, font, size, bold, color, align, before=0, after=8, page_break=False):
    style.font.name = font
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = color
    pf = style.paragraph_format
    pf.alignment = align
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.page_break_before = page_break
    pf.line_spacing = 1.15
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = style.element.makeelement(qn("w:rFonts"), {})
        rpr.append(rfonts)
    for a in ("w:ascii", "w:hAnsi", "w:cs"):
        rfonts.set(qn(a), font)


def add_p(doc, text, center=False, italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.15
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    set_run_font(run, "Calibri", 11, color=RGBColor(0, 0, 0))
    run.italic = italic


def h(doc, text, level=1):
    doc.add_heading(text, level=level)


def new_doc(folder, filename, title, subtitle):
    folder.mkdir(parents=True, exist_ok=True)
    doc = Document()
    for s in doc.sections:
        s.top_margin = Inches(1)
        s.bottom_margin = Inches(1)
        s.left_margin = Inches(1)
        s.right_margin = Inches(1)
    shade(doc.styles["Normal"], "Calibri", 11, False, RGBColor(0, 0, 0), WD_ALIGN_PARAGRAPH.LEFT)
    shade(doc.styles["Heading 1"], "Amazon Ember", 18, True, BLUE, WD_ALIGN_PARAGRAPH.CENTER, after=12, page_break=True)
    shade(doc.styles["Heading 2"], "Amazon Ember", 18, True, BLUE, WD_ALIGN_PARAGRAPH.CENTER, before=14, after=8)
    add_p(doc, title, center=True)
    set_run_font(doc.paragraphs[-1].runs[0], "Amazon Ember", 28, bold=True, color=BLUE)
    add_p(doc, subtitle, center=True)
    add_p(doc, "Lothar J. Musiol", center=True)
    add_p(
        doc,
        "Copyright © 2026 Lothar J. Musiol. Nine-book series. Worked numbers are computed here. "
        "Book 2 is a separate merged manuscript. Do not paste Cripps, Pozar, Ott, or Sklar.",
        center=True,
        italic=True,
    )
    return doc, folder / filename


def book4(doc):
    c = 2.99792458e8
    f = 2.45e9
    lam = c / f
    r = 100.0
    friis = (lam / (4 * math.pi * r)) ** 2
    z0 = 50.0
    er = 3.55
    h_m = 0.001524  # 60 mil
    # microstrip ballpark W/h for ~50 ohm on er=3.55
    add_p(doc, "Book 4 of 9. Old 4, 8, 9, and 10. Millimeter wave is this book at a shorter wavelength.", center=True, italic=True)
    h(doc, "Chapter 1. One Microwave Book")
    add_p(doc, "Bowick is the on-ramp: matching, the Smith chart, a first filter, a first amplifier. Pozar, or Steer volumes 1 to 3, is the microwave text. Book 1 already taught the Smith chart. This book uses it.")
    add_p(doc, "Millimeter wave uses Maxwell's equations. What changes is the loss per millimeter, the waveguide size, the beamwidth, and how much the etch and the bond-wire count. There is no second mmWave volume.")
    add_p(doc, "Key idea. Shorter wavelength is not a new subject. It is the same subject with less forgiveness.")
    h(doc, "Chapter 2. Lines and Matching")
    add_p(doc, "A TEM line is characterized by Z0 and delay. A 50 Ω system is a convention, not a law. A quarter-wave transformer matches R to Z0 when its own impedance is sqrt(R Z0).")
    zq = math.sqrt(25 * 50)
    add_p(doc, f"Worked: match 25 Ω to 50 Ω. Zq = √(25×50) = {zq:.2f} Ω. At 2.45 GHz, λ = {lam*100:.2f} cm in free space. On a board with √er = √{er:.2f} = {math.sqrt(er):.3f}, the guided wavelength is {lam/math.sqrt(er)*100:.2f} cm, so a quarter-wave is {lam/math.sqrt(er)/4*100:.2f} cm. Laminate numbers at 10 GHz are Book 9. This chapter only needs that RO4000-class material exists and FR-4 is not it at this frequency.")
    add_p(doc, "L-section matching and stubs are the rest of Bowick. Cavities and waveguide filters appear when microstrip dies from voltage and heat. Those high-power filters sit after a PA; the PA is Book 6.")
    h(doc, "Chapter 3. The Antenna")
    add_p(doc, "Gain, beam, polarization, sidelobes, scan. An array's element spacing is about λ/2 if you do not want grating lobes in visible space. An active electronically scanned array puts a small amplifier at each element. That amplifier is still Book 6.")
    add_p(doc, f"Friis: Pr/Pt = Gt Gr (λ/4πR)². At 2.45 GHz, λ = {lam:.4f} m. At R = 100 m, isotropic, the space factor is {friis:.3e}, which is {10*math.log10(friis):.1f} dB. Antenna gains multiply that factor. Book 6 owns transmitter power and the statement that 1 dB of feed loss is 1 dB of EIRP.")
    add_p(doc, "Key idea. This book supplies Gt and Gr. Book 6 supplies Pt. Book 7 supplies the legal mask on what left the antenna.")
    h(doc, "Chapter 4. Arrays and mmWave")
    add_p(doc, "A 28 GHz wavelength is 10.7 mm. A 77 GHz automotive radar wavelength is 3.9 mm. Patch antennas become millimeters. Bond wires become inductors you cannot ignore. PTFE or RO3003-class boards, thin copper, tight etch: specified in Book 9, used here.")
    add_p(doc, "Practice. 1. Repeat the quarter-wave example at 10 GHz on the same er. 2. A 10 dBi antenna at each end of the 100 m link: by how many decibels does received power rise? 3. Why is mmWave not its own book in this series?")


def book5(doc):
    full = 7125 - 5925
    n_full = full // 320
    leftover = full - n_full * 320
    add_p(doc, "Book 5 of 9. Old 6, 12, and 14.", center=True, italic=True)
    h(doc, "Chapter 1. Bits Become a Waveform")
    add_p(doc, "A modem turns bits into a waveform and the waveform back into bits. The constellation, the pulse shape, the code, and the equalizer live here. The filter bank in the phone, the PA, and the antenna often do not. Book 6 owns watts. Book 7 owns the chamber. Book 4 owns the antenna gain.")
    add_p(doc, "A square wave's odd harmonics falling as 1/n are derived in Book 1. This book uses that fact: OFDM is many slow sinusoids instead of one fast square-ish symbol. The FFT that builds the symbol is Lyons's territory, written here, not a second DSP volume.")
    h(doc, "Chapter 2. The Band Plan, 2026")
    add_p(doc, "Wi-Fi is unlicensed indoor and campus radio. Cellular is licensed, plus some unlicensed NR-U. 2.4 GHz is crowded. 5 GHz is the workhorse. 6 GHz, 5925–7125 MHz where open, is 160 and 320 MHz channels.")
    add_p(doc, f"Full 6 GHz span = {full} MHz. 320 MHz fits {n_full} times, leftover {leftover} MHz. A 500–600 MHz national slice holds one 320 MHz channel. Wi-Fi 7 paper peak ~23 Gbit/s, 4096-QAM, multi-link. Wi-Fi 8 (~2028) is reliability, not another QAM doubling.")
    add_p(doc, "5G FR1: 410 MHz–7.125 GHz. FR2: 24.25–71 GHz. 5G-Advanced squeezes those bands. 6G toward 2030 wants 7.125–24.25 GHz. Sub-THz 90–300 GHz is a study item, not a 2026 phone. Wi-Fi 6 GHz and cellular U6G sit next to each other; indoor AFC and power limits are how they share.")
    h(doc, "Chapter 3. SDR as a Lab, Not a Religion")
    add_p(doc, "Sampling, decimation, and a receiver you can code belong here. GNU Radio and PySDR are laboratories. They are not the textbook. Digital control as a subject is Book 2. This book only owes sampling.")
    add_p(doc, "Key idea. The waveform is this book. The license is the band table in Chapter 2. The mask you fail in a chamber is Book 7.")
    h(doc, "Practice")
    add_p(doc, "1. How many 320 MHz channels fit in 1200 MHz? 2. Who owns 6 GHz indoor, Wi-Fi or 5G, in this book's split? 3. Why does a chamber failure on a Wi-Fi gadget usually not implicate the LDPC decoder?")


def book6(doc):
    r_gan = (50.0 - 5.0) ** 2 / 200.0
    r_ldmos = (28.0 - 3.0) ** 2 / 200.0
    loss = 10 ** (-0.1)
    nf_db = 3.0
    t0 = 290.0
    t_e = t0 * (10 ** (nf_db / 10) - 1)
    add_p(doc, "Book 6 of 9. Old 5, 7, 13, and 17.", center=True, italic=True)
    h(doc, "Chapter 1. The Chain")
    add_p(doc, "Device → load line → class → match → bias → heat → combine → isolate → harmonic filter → line → duplexer → antenna → EIRP, exposure, mask. Book 3 owns why GaN stands the voltage. Book 4 owns Gt. Book 7 owns the test. This book owns the watts.")
    add_p(doc, "LDMOS to ~3 GHz. GaN-on-SiC ~0.5–40 GHz. GaAs HBT/pHEMT for drivers and a few watts. InP for fT, not kilowatts. Tubes when combining cost explodes. Doherty is the cellular default. Pulse radar is duty-cycle heat.")
    h(doc, "Chapter 2. Load Line and Feed Loss")
    add_p(doc, f"Ropt ≈ (VDD − Vknee)²/(2P). 100 W, 50 V, 5 V knee: {r_gan:.2f} Ω. 28 V, 3 V knee: {r_ldmos:.2f} Ω. Ratios to 50 Ω: {50/r_gan:.1f} and {50/r_ldmos:.0f}. One decibel of feed = factor {loss:.3f}. 100 W becomes 79.4 W at the antenna.")
    h(doc, "Chapter 3. Noise and the PLL")
    add_p(doc, f"A 3 dB noise figure at T0 = 290 K is an excess temperature Te = T0(F−1) = {t_e:.0f} K. Cascades are Friis's formula for noise factor, McClaning's job, in this volume because the LNA sits in the same radio as the PA. A PLL is one chapter: Gardner or Best, then Razavi on a chip. It is not a separate book.")
    h(doc, "Chapter 4. Radar")
    add_p(doc, "Received power falls as 1/R^4. Antenna gains from Book 4 go in the numerator. Duty cycle and pulse compression buy detectability you cannot buy with CW power. The modulator pulse transformer, if the last stage is a tube, is Book 8 shelf C named from here.")
    add_p(doc, "Key idea. Raise the rail and matching gets easier. One decibel after the last amplifier is one decibel of EIRP.")
    h(doc, "Practice")
    add_p(doc, "1. Repeat Ropt at 200 W, 50 V, 5 V knee. 2. A 2 dB feed: what fraction of PA power reaches the antenna? 3. Why is a PLL not Book 2's control chapter even though both use Bode plots?")


def book7(doc):
    """Retired stub. The manuscript is EE7/build_book7.py."""
    raise RuntimeError(
        "Book 7 is built by EE7/build_book7.py. Do not overwrite that manuscript from this stub."
    )


def book8(doc):
    def waste(p, eta):
        return p * (1.0 / eta - 1.0)
    gap80 = waste(80, 0.96) - waste(80, 0.98)
    gap45 = waste(45, 0.96) - waste(45, 0.98)
    skin60 = 66 / math.sqrt(60)
    skin100k = 66 / math.sqrt(1e5)
    add_p(doc, "Book 8 of 9. Old 15. Three professions live here except high-power RF, which is Book 6.", center=True, italic=True)
    h(doc, "Chapter 1. The Rack")
    add_p(doc, f"Waste = P(1/η−1). 80 kW at 96% vs 98%: {gap80:.2f} kW extra, {gap80:.2f} MW per 1000 racks. 45 kW average: {gap45:.2f} kW. Book 9 describes the hall. It does not recompute this. 700 W / 0.8 V = 875 A is Book 9.")
    add_p(doc, "JLL 2026: ~100 GW new capacity 2026–2030, ~$30 million per megawatt AI fit-out. Planning numbers.")
    h(doc, "Chapter 2. LDO versus Switcher")
    add_p(doc, "An LDO is a controlled resistor. Dropout, PSRR, and ceramic-cap stability are the craft (Rincón-Mora, Basso on loops, TI notes). A buck is a switched LC. The hot loop (input cap, switch, diode or sync FET) must be small. Isolated flyback and offline are Pressman plus McLyman. Erickson is the plant equations. The 650 V GaN switch is this book, never InP.")
    h(doc, "Chapter 3. Four Magnetics")
    add_p(doc, f"Skin depth ≈ 66/√f mm. 60 Hz: {skin60:.1f} mm. 100 kHz: {skin100k:.2f} mm. Line-frequency steel, switch-mode ferrite, pulse transformers, RF baluns: four shelves. McLyman sizes. Kazimierczuk explains frequency. Sevick starts when the part is RF. The SMD inductor as a catalog part is Book 2.")
    h(doc, "Chapter 4. Grid Codes and HV")
    add_p(doc, "US 60 Hz, EU/China 50 Hz, Japan split, Korea 60 Hz. Codes do not travel with Cripps. Shelf C is insulation, impulse, pulsed power. Radar modulators borrow a pulse transformer from here and return the RF to Book 6.")
    add_p(doc, "Key idea. Name the rack power you assumed. Name the magnetics shelf before you name the core.")
    h(doc, "Practice")
    add_p(doc, "1. Repeat the 96/98 gap at 100 kW. 2. Why is a 50 Hz transformer with a SiC inverter a harmonics problem? 3. Which book owns a 100 W GaN PA at 3.5 GHz?")


def book9(doc):
    i700 = 700 / 0.8
    ddr = 6400e6 * 8 / 1e9
    ratio = 2.8e12 / (ddr * 1e9)
    f_hz = 10e9
    c = 2.99792458e8
    a_fr4 = 8.685889 * (math.pi * f_hz / c) * math.sqrt(4.5) * 0.020
    a_ro = 8.685889 * (math.pi * f_hz / c) * math.sqrt(3.55) * 0.0037
    titan = 0.46 * 0.46
    can = 3.2 * 2.5
    add_p(doc, "Book 9 of 9. Old 18 and 21.", center=True, italic=True)
    h(doc, "Chapter 1. Bandwidth per Watt")
    add_p(doc, f"HBM4 stack 2.8 TB/s vs one DDR5-6400 channel {ddr:.1f} GB/s is {ratio:.0f}×. 700 W on a declared 0.8 V rail is {i700:.0f} A. CXL does not replace HBM. CoWoS-L through about 2028. zHBM and on-die DRAM are roadmap.")
    h(doc, "Chapter 2. Four Words")
    add_p(doc, "Flip-chip is a join. Chip-scale is a size (~1.2× die). Lead frame is a skeleton (SOIC, QFN, TO). Interposer is extra wiring. Fan-out is a package, not an interposer. You can flip without an interposer. You almost never put a silicon interposer down without a flip or hybrid bond on it.")
    h(doc, "Chapter 3. FR-4 and Better")
    add_p(doc, f"FR-4 is default. Buy better for loss, Dk stability, heat, or density. At 10 GHz, dielectric ballpark: FR-4 εr=4.5 Df=0.020 → {a_fr4:.1f} dB/m ({a_fr4/39.37:.2f} dB/in). RO4000-class εr=3.55 Df=0.0037 → {a_ro:.1f} dB/m. Spec Dk/Df at the frequency you care about. Same 'FR-4' from two mills is not the same Df. Module substrates (BT, ABF, LTCC, glass, silicon, fan-out) are not this motherboard.")
    h(doc, "Chapter 4. Layout Habits")
    add_p(doc, "Johnson and Bogatin: a continuous reference under an RF or fast digital trace, via fence when needed, short ground into a plane, digital return off the RF island, PA matching treated as microwave. The hot loop of a buck is Book 8 named from this copper.")
    h(doc, "Chapter 5. The Hall and the Robot")
    add_p(doc, "Rack 40–100 kW, cited average ~45 kW, new builds toward 80% liquid. Conversion arithmetic is Book 8. Scarce inputs: power, HBM, packaging, optics. VLA model plus Book 2's 200 Hz tracker. Helix 2.5, Digit 5, Gemini Robotics 2, GR00T are a 2026 snapshot.")
    h(doc, "Chapter 6. MEMS, 2026")
    add_p(doc, f"AlScN on 8-inch silicon is the process shift. Titan {titan:.3f} mm² vs 1210 can {can:.1f} mm² ≈ {can/titan:.0f}× smaller. Q~1e5 at 70 MHz is 700 Hz wide. Consumer IMUs grew a learning hook. Navigation IMUs: 175 °C, ~1 °/h on the datasheet. Only 2026 survey. Book 3 names the film. Book 5 names the radio that buys the BAW.")
    add_p(doc, "Key idea. A modern accelerator is a memory system with compute attached. A modern robot is a model ahead of a mechanism.")
    h(doc, "Practice")
    add_p(doc, "1. Four HBM stacks at 2.8 TB/s: aggregate, versus four DDR5-6400 channels. 2. 1000 W at 0.7 V: current. 3. Name one rung CXL may replace and one it may not.")


def main():
    jobs = [
        ("EE4", "RF_Microwave_and_Antennas_Book4.docx", "RF, Microwave, and Antennas", "Book 4 of 9", book4),
        ("EE5", "Communications_Wireless_and_SDR_Book5.docx", "Communications, Wireless, and SDR", "Book 5 of 9", book5),
        ("EE6", "Transceivers_and_High_Power_RF_Book6.docx", "Transceivers and High-Power RF", "Book 6 of 9", book6),
        # Book 7 is built by EE7/build_book7.py. Do not regenerate it here.
        ("EE8", "Power_and_Energy_Book8.docx", "Power and Energy", "Book 8 of 9", book8),
        ("EE9", "Packaging_Layout_and_Emerging_Book9.docx", "Packaging, Layout, and Emerging Systems", "Book 9 of 9", book9),
    ]
    for folder, filename, title, subtitle, fill in jobs:
        doc, path = new_doc(ROOT / folder, filename, title, subtitle)
        fill(doc)
        doc.save(path)
        print("wrote", path)


if __name__ == "__main__":
    main()
