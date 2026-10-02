# -*- coding: utf-8 -*-
"""Shared docx helpers for the nine-book EE series."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

BLUE = RGBColor(0x00, 0x00, 0xFF)
BLACK = RGBColor(0, 0, 0)


def set_run(run, name, size, bold=False, color=None):
    run.bold = bold
    run.font.name = name
    run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = run._element.makeelement(qn("w:rFonts"), {})
        rpr.append(rfonts)
    for a in ("w:ascii", "w:hAnsi", "w:cs"):
        rfonts.set(qn(a), name)


def shade(style, font, size, bold, color, align, before=0, after=8, page_break=False):
    style.font.name = font
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = color
    pf = style.paragraph_format
    pf.alignment = align
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.page_break_before = page_break
    pf.line_spacing = 1.15
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = style.element.makeelement(qn("w:rFonts"), {})
        rpr.append(rfonts)
    for a in ("w:ascii", "w:hAnsi", "w:cs"):
        rfonts.set(qn(a), font)


def add_p(doc, text, center=False, italic=False, style="Normal"):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.15
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    if style == "Heading 1":
        set_run(run, "Amazon Ember", 18, True, BLUE)
    elif style == "Heading 2":
        set_run(run, "Amazon Ember", 18, True, BLUE)
    else:
        set_run(run, "Calibri", 11, False, BLACK)
        run.italic = italic
    return p


def new_book(path: Path, title: str, subtitle: str):
    path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    for s in doc.sections:
        s.top_margin = Inches(1)
        s.bottom_margin = Inches(1)
        s.left_margin = Inches(1)
        s.right_margin = Inches(1)
    shade(doc.styles["Normal"], "Calibri", 11, False, BLACK, WD_ALIGN_PARAGRAPH.LEFT)
    shade(doc.styles["Heading 1"], "Amazon Ember", 18, True, BLUE, WD_ALIGN_PARAGRAPH.CENTER, after=12, page_break=True)
    shade(doc.styles["Heading 2"], "Amazon Ember", 18, True, BLUE, WD_ALIGN_PARAGRAPH.CENTER, before=14, after=8)
    add_p(doc, title, center=True)
    set_run(doc.paragraphs[-1].runs[0], "Amazon Ember", 28, True, BLUE)
    add_p(doc, subtitle, center=True)
    add_p(doc, "Lothar J. Musiol", center=True)
    add_p(
        doc,
        "Copyright © 2026 Lothar J. Musiol. Kindle edition. "
        "Original prose. American spelling. Worked numbers are computed in the builder. "
        "Nine-book Electrical Engineering Series.",
        center=True,
        italic=True,
    )
    return doc


def open_existing(path: Path) -> Document:
    if not path.exists():
        raise FileNotFoundError(path)
    return Document(str(path))


def word_count(doc: Document) -> int:
    n = 0
    for p in doc.paragraphs:
        n += len(p.text.split())
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                n += len(c.text.split())
    return n
