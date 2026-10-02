# -*- coding: utf-8 -*-
"""Book 2 stub rewrite (2026-10-01).

Removes the templated filler chapters that fill_to_200.py appended after
Appendix E ("Chapter <slogan>" headings with "Row 1..16" worked families) and
inserts real Chapters 11-23 before Appendix A, in the house style of Chapters 1-9
(Worked Example / Problem / Solution, Key Idea, Common Misconception, figures,
numbered Practice Problems). Extends TOC, Appendix A table, Appendix B solutions,
Appendix C glossary (re-sorted), Appendix D summaries, Appendix E reading.

Usage: python build_book2_ch11_23.py <in.docx> <out.docx>
Chapter sources: src/chNN.txt ; figures: fig/*.png (made by make_figures.py).
DOCX only. Do not run on a file that has already been rebuilt (it checks).
"""
import copy, re, sys
from pathlib import Path
from docx import Document
from docx.shared import Inches
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = Path(__file__).resolve().parent
SRC, FIG = HERE / "src", HERE / "fig"
CHAPTERS = list(range(11, 24))
CH10_SUMMARY = ("A 200 Hz joint tracker has a 5 ms period. The classical loop runs underneath any "
                "robot model, uses the same Bode habits as Chapter 5, and answers to an independent "
                "safety controller that has the final say over the joint.")
READING = [
    ("Karl Johan Åström and Richard M. Murray, Feedback Systems: An Introduction for Scientists and Engineers — ",
     "the open-text companion to Chapters 17–19: loops, PID, frequency-domain margins, and their limits, at the next level of depth."),
    ("Norman S. Nise, Control Systems Engineering — ",
     "a standard undergraduate control text for Chapters 17 and 18, with many worked compensator designs."),
    ("Gene F. Franklin, J. David Powell, and Michael Workman, Digital Control of Dynamic Systems — ",
     "the reference for Chapter 19's sampling, hold, and discretization questions."),
    ("Ron Mancini (editor), Op Amps for Everyone — ",
     "a free Texas Instruments design reference covering feedback, noise, DC errors, and active filters (Chapters 11–16 and 21)."),
    ("Eric Bogatin, Signal and Power Integrity — Simplified — ",
     "the natural next step after Chapter 23's decoupling and target-impedance discussion."),
]

def parse(path):
    items = []
    for line in path.read_text(encoding="utf-8").splitlines():
        if not line.strip():
            continue
        tag, _, body = line.partition(":")
        items.append((tag.strip(), body.strip()))
    return items

def strip_runs(p):
    for child in list(p):
        if child.tag in (qn("w:r"), qn("w:hyperlink"), qn("w:bookmarkStart"), qn("w:bookmarkEnd")):
            p.remove(child)
    for a in list(p.attrib):
        if a.endswith("paraId") or a.endswith("textId"):
            del p.attrib[a]
    return p

def base_rpr(tpl):
    for r in tpl.findall(qn("w:r")):
        if r.find(qn("w:t")) is not None:
            rpr = r.find(qn("w:rPr"))
            if rpr is None:
                return None
            rpr = copy.deepcopy(rpr)
            for tagname in ("w:b", "w:bCs"):
                e = rpr.find(qn(tagname))
                if e is not None:
                    rpr.remove(e)
            return rpr
    return None

def make_p(tpl, segments):
    """segments: list of (text, bold)."""
    p = strip_runs(copy.deepcopy(tpl))
    rpr0 = base_rpr(tpl)
    for text, bold in segments:
        r = OxmlElement("w:r")
        rpr = copy.deepcopy(rpr0) if rpr0 is not None else OxmlElement("w:rPr")
        if bold:
            b = OxmlElement("w:b"); bcs = OxmlElement("w:bCs")
            pos = 0
            for i, ch in enumerate(rpr):
                if ch.tag in (qn("w:rStyle"), qn("w:rFonts")):
                    pos = i + 1
            rpr.insert(pos, bcs); rpr.insert(pos, b)
        if len(rpr):
            r.append(rpr)
        t = OxmlElement("w:t"); t.set(qn("xml:space"), "preserve"); t.text = text
        r.append(t); p.append(r)
    return p

def main(inp, outp):
    doc = Document(inp)
    P = doc.paragraphs
    def find(pred, start=0):
        for i in range(start, len(P)):
            if pred(P[i]):
                return i
        raise LookupError
    isH1 = lambda p: p.style is not None and p.style.name == "Heading 1"
    if any(isH1(p) and p.text.startswith("Chapter 11 ") for p in P):
        sys.exit("already rebuilt; refusing to run twice")
    # templates (Chapters 1-9 house style)
    T = {}
    T["H1"] = P[find(lambda p: isH1(p) and p.text.startswith("Chapter 1 —"))]._p
    T["H2"] = P[find(lambda p: p.style.name == "Heading 2" and p.text.startswith("1.1"))]._p
    i_we = find(lambda p: p.text.startswith("Worked Example 1.1"))
    T["WE"], T["PROB"], T["SOL"] = P[i_we]._p, P[i_we+1]._p, P[i_we+2]._p
    T["FIGP"], T["CAP"] = P[i_we+3]._p, P[i_we+4]._p
    T["P"] = P[i_we-2]._p  # body paragraph before the example
    i_key = find(lambda p: p.text == "Key Idea", i_we)
    T["KEYL"], T["KEY"] = P[i_key]._p, P[i_key+1]._p
    i_mis = find(lambda p: p.text == "Common Misconception", i_we)
    T["MISL"], T["MIS"] = P[i_mis]._p, P[i_mis+1]._p
    T["EQ"] = P[find(lambda p: p.text.startswith("NMʜ = "))]._p
    i_pp = find(lambda p: p.style.name == "List Paragraph")
    T["PP"] = P[i_pp]._p
    numid_tpl = T["PP"].find(qn("w:pPr")).find(qn("w:numPr")).find(qn("w:numId")).get(qn("w:val"))
    # --- 1. delete appended filler chapters (H1 "Chapter " + non-digit) to end of body
    i_stub = find(lambda p: isH1(p) and re.match(r"Chapter [^0-9]", p.text))
    body = doc.element.body
    el = P[i_stub]._p
    removed = 0
    while el is not None:
        nxt = el.getnext()
        if el.tag != qn("w:sectPr"):
            body.remove(el); removed += 1
        el = nxt
    print("removed filler elements:", removed)
    P = doc.paragraphs
    # numbering
    numbering = doc.part.numbering_part.element
    absid = None
    for num in numbering.findall(qn("w:num")):
        if num.get(qn("w:numId")) == numid_tpl:
            absid = num.find(qn("w:abstractNumId")).get(qn("w:val"))
    next_num = max(int(n.get(qn("w:numId"))) for n in numbering.findall(qn("w:num"))) + 1
    def new_num():
        nonlocal next_num
        num = OxmlElement("w:num"); num.set(qn("w:numId"), str(next_num))
        a = OxmlElement("w:abstractNumId"); a.set(qn("w:val"), absid); num.append(a)
        lo = OxmlElement("w:lvlOverride"); lo.set(qn("w:ilvl"), "0")
        so = OxmlElement("w:startOverride"); so.set(qn("w:val"), "1"); lo.append(so); num.append(lo)
        numbering.append(num); next_num += 1
        return str(next_num - 1)
    # --- 2. build chapters
    anchor = P[find(lambda p: isH1(p) and p.text.startswith("Appendix A"))]._p
    sols, sums, gloss, forms, titles = {}, {}, [], [], []
    for n in CHAPTERS:
        items = parse(SRC / f"ch{n}.txt")
        numid = None
        sols[n] = []
        for tag, b in items:
            new = []
            if tag == "H1":
                p = make_p(T["H1"], [(b, False)]); titles.append(b); new.append(p)
            elif tag == "H2":
                new.append(make_p(T["H2"], [(b, False)]))
            elif tag == "P":
                new.append(make_p(T["P"], [(b, False)]))
            elif tag == "EQ":
                new.append(make_p(T["EQ"], [(b, False)]))
            elif tag == "WE":
                new.append(make_p(T["WE"], [(b, True)]))
            elif tag == "PROB":
                new.append(make_p(T["PROB"], [("Problem: ", True), (b, False)]))
            elif tag == "SOL":
                new.append(make_p(T["SOL"], [("Solution: ", True), (b, False)]))
            elif tag == "KEY":
                bold, _, rest = b.partition("||")
                new.append(make_p(T["KEYL"], [("Key Idea", True)]))
                new.append(make_p(T["KEY"], [(bold.strip() + " ", True), (rest.strip(), False)]))
            elif tag == "MIS":
                new.append(make_p(T["MISL"], [("Common Misconception", True)]))
                new.append(make_p(T["MIS"], [(b, False)]))
            elif tag == "FIG":
                fname, _, cap = b.partition("|")
                fp = strip_runs(copy.deepcopy(T["FIGP"]))
                tmp = doc.add_paragraph()
                run = tmp.add_run(); run.add_picture(str(FIG / fname.strip()), width=Inches(4.79))
                fp.append(run._r); tmp._p.getparent().remove(tmp._p)
                new.append(fp)
                new.append(make_p(T["CAP"], [(cap.strip(), False)]))
            elif tag == "PP":
                if numid is None:
                    numid = new_num()
                p = make_p(T["PP"], [(b, False)])
                p.find(qn("w:pPr")).find(qn("w:numPr")).find(qn("w:numId")).set(qn("w:val"), numid)
                new.append(p)
            elif tag == "SOLN":
                sols[n].append(b)
            elif tag == "SUM":
                sums[n] = b
            elif tag == "GLOS":
                a, _, c = b.partition("||"); gloss.append((a.strip(), c.strip()))
            elif tag == "FORM":
                a, _, c = b.partition("||"); forms.append((a.strip(), c.strip()))
            else:
                raise ValueError(f"ch{n}: unknown tag {tag}")
            for p in new:
                anchor.addprevious(p)
    P = doc.paragraphs
    # --- 3. TOC
    i_toc9 = find(lambda p: p.text == "Chapter 9 — Inductors" and not isH1(p))
    tpl = P[i_toc9]._p; prev = tpl
    for t in ["Chapter 10 — The Loop Under the Model"] + titles:
        p = make_p(tpl, [(t, False)]); prev.addnext(p); prev = p
    # --- 4. How This Book Is Organized
    P = doc.paragraphs
    i_org = find(lambda p: p.text.startswith("Chapters 1–2 build the vocabulary"))
    org = P[i_org]
    old = org.runs[0].text if len(org.runs) == 1 else None
    full = org.text.replace("Chapters 7–9 close the book by returning to", "Chapters 7–9 return to")
    newp = make_p(org._p, [(full, False)]); org._p.addprevious(newp); org._p.getparent().remove(org._p)
    extra = make_p(newp, [(
        "Chapter 10 introduces the control loop that runs underneath a robot model. Chapters 11–16 return to the op-amp "
        "with real numbers: feedback as a quantitative tool, followers and capacitive loads, bandwidth and slew rate, "
        "noise, DC errors, and comparators. Chapters 17–19 build classical control on that foundation, from plant and "
        "controller through PID and lead-lag compensation to the digital loop and its sampling delay. Chapters 20–23 "
        "close the book with the practical side of real parts: reading an inductor datasheet, designing active filters "
        "to standard values, and the heat, drift, and impedance behavior of real resistors and capacitors.", False)])
    newp.addnext(extra)
    # --- 5. Appendix A table
    tbl = doc.tables[-1]
    last = tbl.rows[-1]._tr
    for a, c in forms:
        tr = copy.deepcopy(last)
        for tc, text in zip(tr.findall(qn("w:tc")), (a, c)):
            ps = tc.findall(qn("w:p"))
            for extra_p in ps[1:]:
                tc.remove(extra_p)
            rs = ps[0].findall(qn("w:r"))
            for r in rs[1:]:
                ps[0].remove(r)
            rs[0].find(qn("w:t")).text = text
        tbl._tbl.append(tr)
    # --- 6. Appendix B solutions
    P = doc.paragraphs
    i_c = find(lambda p: isH1(p) and p.text.startswith("Appendix C"))
    i_b10 = find(lambda p: p.text == "Chapter 10 Solutions")
    h2tpl, stpl = P[i_b10]._p, P[i_b10 + 1]._p
    anchorC = P[i_c]._p
    for n in CHAPTERS:
        anchorC.addprevious(make_p(h2tpl, [(f"Chapter {n} Solutions", False)]))
        for k, s in enumerate(sols[n], 1):
            anchorC.addprevious(make_p(stpl, [(f"{k}. {s}", False)]))
    # --- 7. Appendix C glossary, merged and sorted
    P = doc.paragraphs
    i_c = find(lambda p: isH1(p) and p.text.startswith("Appendix C"))
    i_d = find(lambda p: isH1(p) and p.text.startswith("Appendix D"))
    entries, gtpl = [], P[i_c + 1]._p
    for p in P[i_c + 1:i_d]:
        term, _, rest = p.text.partition(" — ")
        entries.append((term, rest))
    have = {e[0].lower() for e in entries}
    for a, c in gloss:
        if a.lower() not in have:
            entries.append((a, c))
    old_ps = [p._p for p in P[i_c + 1:i_d]]
    anchorD = P[i_d]._p
    for e in old_ps:
        e.getparent().remove(e)
    for term, rest in sorted(entries, key=lambda e: e[0].lower()):
        anchorD.addprevious(make_p(gtpl, [(term + " — ", True), (rest, False)]))
    # --- 8. Appendix D summaries
    P = doc.paragraphs
    i_e = find(lambda p: isH1(p) and p.text.startswith("Appendix E"))
    i_d9 = find(lambda p: p.text == "Chapter 9" and p.style.name == "Heading 2")
    dh, dp = P[i_d9]._p, P[i_d9 + 1]._p
    anchorE = P[i_e]._p
    anchorE.addprevious(make_p(dh, [("Chapter 10", False)])); anchorE.addprevious(make_p(dp, [(CH10_SUMMARY, False)]))
    for n in CHAPTERS:
        anchorE.addprevious(make_p(dh, [(f"Chapter {n}", False)]))
        anchorE.addprevious(make_p(dp, [(sums[n], False)]))
    # --- 9. Appendix E reading
    P = doc.paragraphs
    i_e = find(lambda p: isH1(p) and p.text.startswith("Appendix E"))
    lastp = P[-1]._p
    rtpl = P[i_e + 2]._p
    for a, c in READING:
        p = make_p(rtpl, [(a, True), (c, False)]); lastp.addnext(p); lastp = p
    doc.save(outp)
    print("saved", outp, "chapters", titles[0], "...", titles[-1])

if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
