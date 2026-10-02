# -*- coding: utf-8 -*-
"""Book 2 extension (2026-10-01): Chapters 24-28 and title-page fix.

Input must be the Book 2 file already rebuilt by build_book2_ch11_23.py
(Chapters 1-23). Inserts Chapters 24-28 before Appendix A in the same house
style, extends TOC, "How This Book Is Organized", Appendix A table, Appendix B
solutions, Appendix C glossary (re-sorted), Appendix D summaries, Appendix E
reading, and replaces the old title "CORE CIRCUITS AND COMPONENTS" on the title
page with "CIRCUITS, COMPONENTS, AND CONTROL". Refuses to run twice.
Usage: python build_book2_ch24_28.py <in.docx> <out.docx>
"""
import copy, sys
from pathlib import Path
from docx import Document
from docx.shared import Inches
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from b2lib import parse, strip_runs, make_p

HERE = Path(__file__).resolve().parent
ROOT = Path(__file__).resolve().parents[1]  # series root (scripts/ is one level down)
SRC, FIG = ROOT / "chapters" / "Book2" / "ch24_28" / "src", ROOT / "chapters" / "Book2" / "ch24_28" / "fig"
CHAPTERS = list(range(24, 29))
TITLE = "CIRCUITS, COMPONENTS, AND CONTROL"
READING = [
    ("Charles Kitchin and Lew Counts, A Designer's Guide to Instrumentation Amplifiers — ",
     "a free Analog Devices guide that goes beyond Chapter 24 into input protection, RFI rectification, and single-supply operation."),
    ("Walt Kester (editor), The Data Conversion Handbook — ",
     "the Analog Devices reference behind Chapter 26: sample-and-hold architectures, jitter, and converter drive circuits."),
    ("Ned Mohan, Electric Drives: An Integrative Approach — ",
     "the next step after Chapter 28, from the DC motor drive to AC machines and their control."),
]
ORG_OLD = "Chapters 20–23 close the book with the practical side of real parts"
ORG_NEW = "Chapters 20–23 turn to the practical side of real parts"
ORG_ADD = (" Chapters 24–28 close the book with five building blocks that join its two halves: the instrumentation "
           "amplifier, state feedback with observers, the sample-and-hold circuit, the voltage reference, and the "
           "H-bridge motor drive written as a plant for the loops of Chapters 17–19.")

def main(inp, outp):
    doc = Document(inp)
    P = doc.paragraphs
    def find(pred, start=0):
        for i in range(start, len(P)):
            if pred(P[i]):
                return i
        raise LookupError
    isH1 = lambda p: p.style is not None and p.style.name == "Heading 1"
    if any(isH1(p) and p.text.startswith("Chapter 24 ") for p in P):
        sys.exit("already extended; refusing to run twice")
    find(lambda p: isH1(p) and p.text.startswith("Chapter 23 "))
    T = {}
    T["H1"] = P[find(lambda p: isH1(p) and p.text.startswith("Chapter 1 —"))]._p
    T["H2"] = P[find(lambda p: p.style.name == "Heading 2" and p.text.startswith("1.1"))]._p
    i_we = find(lambda p: p.text.startswith("Worked Example 1.1"))
    T["WE"], T["PROB"], T["SOL"] = P[i_we]._p, P[i_we+1]._p, P[i_we+2]._p
    T["FIGP"], T["CAP"] = P[i_we+3]._p, P[i_we+4]._p
    T["P"] = P[i_we-2]._p
    i_key = find(lambda p: p.text == "Key Idea", i_we)
    T["KEYL"], T["KEY"] = P[i_key]._p, P[i_key+1]._p
    i_mis = find(lambda p: p.text == "Common Misconception", i_we)
    T["MISL"], T["MIS"] = P[i_mis]._p, P[i_mis+1]._p
    T["EQ"] = P[find(lambda p: p.text.startswith("NMʜ = "))]._p
    T["PP"] = P[find(lambda p: p.style.name == "List Paragraph")]._p
    numid_tpl = T["PP"].find(qn("w:pPr")).find(qn("w:numPr")).find(qn("w:numId")).get(qn("w:val"))
    # title page
    tp = P[0]
    assert tp.text == "CORE CIRCUITS AND COMPONENTS", tp.text
    tp.runs[0].text = TITLE
    for r in tp.runs[1:]:
        r.text = ""
    # numbering
    numbering = doc.part.numbering_part.element
    absid = None
    for num in numbering.findall(qn("w:num")):
        if num.get(qn("w:numId")) == numid_tpl:
            absid = num.find(qn("w:abstractNumId")).get(qn("w:val"))
    next_num = max(int(n.get(qn("w:numId"))) for n in numbering.findall(qn("w:num"))) + 1
    def new_num():
        nonlocal next_num
        num = OxmlElement("w:num"); num.set(qn("w:numId"), str(next_num))
        a = OxmlElement("w:abstractNumId"); a.set(qn("w:val"), absid); num.append(a)
        lo = OxmlElement("w:lvlOverride"); lo.set(qn("w:ilvl"), "0")
        so = OxmlElement("w:startOverride"); so.set(qn("w:val"), "1"); lo.append(so); num.append(lo)
        numbering.append(num); next_num += 1
        return str(next_num - 1)
    anchor = P[find(lambda p: isH1(p) and p.text.startswith("Appendix A"))]._p
    sols, sums, gloss, forms, titles = {}, {}, [], [], []
    for n in CHAPTERS:
        numid = None; sols[n] = []
        for tag, b in parse(SRC / f"ch{n}.txt"):
            new = []
            if tag == "H1":
                new.append(make_p(T["H1"], [(b, False)])); titles.append(b)
            elif tag in ("H2", "P", "EQ"):
                new.append(make_p(T[tag], [(b, False)]))
            elif tag == "WE":
                new.append(make_p(T["WE"], [(b, True)]))
            elif tag == "PROB":
                new.append(make_p(T["PROB"], [("Problem: ", True), (b, False)]))
            elif tag == "SOL":
                new.append(make_p(T["SOL"], [("Solution: ", True), (b, False)]))
            elif tag == "KEY":
                bold, _, rest = b.partition("||")
                new.append(make_p(T["KEYL"], [("Key Idea", True)]))
                new.append(make_p(T["KEY"], [(bold.strip() + " ", True), (rest.strip(), False)]))
            elif tag == "MIS":
                new.append(make_p(T["MISL"], [("Common Misconception", True)]))
                new.append(make_p(T["MIS"], [(b, False)]))
            elif tag == "FIG":
                fname, _, cap = b.partition("|")
                fp = strip_runs(copy.deepcopy(T["FIGP"]))
                tmp = doc.add_paragraph()
                run = tmp.add_run(); run.add_picture(str(FIG / fname.strip()), width=Inches(4.79))
                fp.append(run._r); tmp._p.getparent().remove(tmp._p)
                new.append(fp); new.append(make_p(T["CAP"], [(cap.strip(), False)]))
            elif tag == "PP":
                if numid is None:
                    numid = new_num()
                p = make_p(T["PP"], [(b, False)])
                p.find(qn("w:pPr")).find(qn("w:numPr")).find(qn("w:numId")).set(qn("w:val"), numid)
                new.append(p)
            elif tag == "SOLN":
                sols[n].append(b)
            elif tag == "SUM":
                sums[n] = b
            elif tag == "GLOS":
                a, _, c = b.partition("||"); gloss.append((a.strip(), c.strip()))
            elif tag == "FORM":
                a, _, c = b.partition("||"); forms.append((a.strip(), c.strip()))
            else:
                raise ValueError(f"ch{n}: unknown tag {tag}")
            for p in new:
                anchor.addprevious(p)
        assert len(sols[n]) == sum(1 for t, _ in parse(SRC / f"ch{n}.txt") if t == "PP"), n
    P = doc.paragraphs
    # TOC
    i_toc = find(lambda p: p.text.startswith("Chapter 23 —") and not isH1(p))
    prev = P[i_toc]._p
    for t in titles:
        p = make_p(prev, [(t, False)]); prev.addnext(p); prev = p
    # organization paragraph
    P = doc.paragraphs
    i_org = find(lambda p: p.text.startswith("Chapter 10 introduces the control loop"))
    org = P[i_org]
    assert ORG_OLD in org.text
    full = org.text.replace(ORG_OLD, ORG_NEW) + ORG_ADD
    newp = make_p(org._p, [(full, False)]); org._p.addprevious(newp); org._p.getparent().remove(org._p)
    # Appendix A table
    tbl = doc.tables[-1]; last = tbl.rows[-1]._tr
    for a, c in forms:
        tr = copy.deepcopy(last)
        for tc, text in zip(tr.findall(qn("w:tc")), (a, c)):
            ps = tc.findall(qn("w:p"))
            for extra_p in ps[1:]:
                tc.remove(extra_p)
            rs = ps[0].findall(qn("w:r"))
            for r in rs[1:]:
                ps[0].remove(r)
            rs[0].find(qn("w:t")).text = text
        tbl._tbl.append(tr)
    # Appendix B
    P = doc.paragraphs
    i_c = find(lambda p: isH1(p) and p.text.startswith("Appendix C"))
    i_b = find(lambda p: p.text == "Chapter 23 Solutions")
    h2tpl, stpl = P[i_b]._p, P[i_b + 1]._p
    anchorC = P[i_c]._p
    for n in CHAPTERS:
        anchorC.addprevious(make_p(h2tpl, [(f"Chapter {n} Solutions", False)]))
        for k, s in enumerate(sols[n], 1):
            anchorC.addprevious(make_p(stpl, [(f"{k}. {s}", False)]))
    # Appendix C glossary merge
    P = doc.paragraphs
    i_c = find(lambda p: isH1(p) and p.text.startswith("Appendix C"))
    i_d = find(lambda p: isH1(p) and p.text.startswith("Appendix D"))
    entries, gtpl = [], copy.deepcopy(P[i_c + 1]._p)
    for p in P[i_c + 1:i_d]:
        term, _, rest = p.text.partition(" — ")
        entries.append((term, rest))
    have = {e[0].lower() for e in entries}
    for a, c in gloss:
        if a.lower() not in have:
            entries.append((a, c)); have.add(a.lower())
    old_ps = [p._p for p in P[i_c + 1:i_d]]
    anchorD = P[i_d]._p
    for e in old_ps:
        e.getparent().remove(e)
    for term, rest in sorted(entries, key=lambda e: e[0].lower()):
        anchorD.addprevious(make_p(gtpl, [(term + " — ", True), (rest, False)]))
    # Appendix D summaries
    P = doc.paragraphs
    i_e = find(lambda p: isH1(p) and p.text.startswith("Appendix E"))
    i_d23 = find(lambda p: p.text == "Chapter 23" and p.style.name == "Heading 2")
    dh, dp = P[i_d23]._p, P[i_d23 + 1]._p
    anchorE = P[i_e]._p
    for n in CHAPTERS:
        anchorE.addprevious(make_p(dh, [(f"Chapter {n}", False)]))
        anchorE.addprevious(make_p(dp, [(sums[n], False)]))
    # Appendix E reading
    P = doc.paragraphs
    i_e = find(lambda p: isH1(p) and p.text.startswith("Appendix E"))
    lastp = P[-1]._p; rtpl = P[i_e + 2]._p
    for a, c in READING:
        p = make_p(rtpl, [(a, True), (c, False)]); lastp.addnext(p); lastp = p
    doc.save(outp)
    print("saved", outp, titles)

if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
