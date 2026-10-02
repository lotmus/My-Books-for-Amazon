# -*- coding: utf-8 -*-
"""Rebuild one of EE Books 4-7 (seven-book layout) from written chapter sources (2026-10-01).

Keeps the file's styles, page setup, and title lines; removes every paragraph
after the title block (the four sketch chapters and all fill_to_200.py filler);
writes: copyright line, subtitle, "How to Use This Book", Contents, chapters
(src/chNN.txt), Appendix A (constants), Appendix B (answers), Appendix C (sources).
Tags: PART (optional part heading before H1) H1 H2 P EQ WEP WES KEY PR ANS. front.txt: TITLE/SUB/COPY/HOW/CONST/SRC lines.
Usage: python build_book.py <bookdir> <in.docx> <out.docx>
    e.g. python scripts/build_book.py chapters/Book4 chapters/Book4/template_stub.docx out.docx    (DOCX only)
"""
import sys, re
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

def lines(path):
    out = []
    for ln in Path(path).read_text(encoding="utf-8").splitlines():
        if ln.strip():
            t, _, b = ln.partition(":"); out.append((t.strip(), b.strip()))
    return out

def main(bookdir, inp, outp):
    bd = Path(bookdir)
    front = lines(bd / "front.txt")
    F = {}
    for t, b in front:
        F.setdefault(t, []).append(b)
    doc = Document(inp)
    body = doc.element.body
    P = doc.paragraphs
    first_h1 = next(i for i, p in enumerate(P) if p.style.name == "Heading 1")
    if any(p.style.name == "Heading 1" and p.text == "How to Use This Book" for p in P):
        sys.exit("already rebuilt")
    keep = P[:3]
    keep[0].runs[0].text = F["TITLE"][0]
    for r in keep[0].runs[1:]: r.text = ""
    keep[1].runs[0].text = F["SUB"][0]
    for r in keep[1].runs[1:]: r.text = ""
    removed = 0
    for el in list(body):
        if el.tag == qn("w:sectPr"):
            continue
        if any(el is k._p for k in keep):
            continue
        body.remove(el); removed += 1
    print("removed elements:", removed)
    def para(text="", style="Normal", align=None, bold_prefix=None, italic=False):
        p = doc.add_paragraph(style=style)
        if bold_prefix:
            r = p.add_run(bold_prefix); r.bold = True
        if text:
            r = p.add_run(text); r.italic = italic
        if align is not None:
            p.alignment = align
        return p
    para(F["COPY"][0], align=WD_ALIGN_PARAGRAPH.CENTER)
    if "SUBTITLE" in F:
        para(F["SUBTITLE"][0], align=WD_ALIGN_PARAGRAPH.CENTER, italic=True)
    para("How to Use This Book", "Heading 1")
    for h in F["HOW"]:
        para(h)
    chapters = sorted(bd.glob("src/ch*.txt"))
    titles = [next(b for t, b in lines(c) if t == "H1") for c in chapters]
    toc = []
    for c in chapters:
        toc += [b for t, b in lines(c) if t in ("PART", "H1")]
    para("Contents", "Heading 1")
    for t in toc + ["Appendix A. Constants and Formulas Used in This Book",
                       "Appendix B. Answers to the Practice Problems", "Appendix C. Sources"]:
        para(t)
    answers = []
    for c in chapters:
        items = lines(c); prn = 0; ans = []
        for t, b in items:
            if t == "PART":
                para(b, "Heading 1")
            elif t == "H1":
                para(b, "Heading 1")
            elif t == "H2":
                para(b, "Heading 2")
            elif t == "P":
                para(b)
            elif t == "EQ":
                para(b, align=WD_ALIGN_PARAGRAPH.CENTER)
            elif t == "WEP":
                para(b, bold_prefix="Problem. ")
            elif t == "WES":
                para(b, bold_prefix="Solution. ")
            elif t == "KEY":
                assert b.startswith("Key idea."), c
                para(b[len("Key idea."):].strip(), bold_prefix="Key idea. ")
            elif t == "PR":
                prn += 1; para(b, bold_prefix=f"{prn}. ")
            elif t == "ANS":
                ans.append(b)
            else:
                raise ValueError(f"{c.name}: unknown tag {t}")
        assert prn == len(ans), (c.name, prn, len(ans))
        num = re.match(r"Chapter (\d+)", next(b for t, b in items if t == "H1")).group(1)
        answers.append((num, ans))
    para("Appendix A. Constants and Formulas Used in This Book", "Heading 1")
    for c in F["CONST"]:
        para(c)
    para("Appendix B. Answers to the Practice Problems", "Heading 1")
    for num, ans in answers:
        para(f"Chapter {num}", "Heading 2")
        for k, a in enumerate(ans, 1):
            para(a, bold_prefix=f"{k}. ")
    para("Appendix C. Sources", "Heading 1")
    for s in F["SRC"]:
        para(s)
    doc.core_properties.title = F["TITLE"][0]
    doc.core_properties.author = "Lothar J. Musiol"
    doc.save(outp)
    print("saved", outp, len(titles), "chapters")

if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2], sys.argv[3])
