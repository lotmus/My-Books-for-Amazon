"""Build the Kindle/KDP Word master for *Life, Actually* (Physics, Actually series)
from the Markdown sources in chapters/.  Docx only.

Usage:  python scripts/build_book.py [output.docx]

House style follows Physics, Actually Vol. 2 (the 6x9 KDP volume) and Vol. 3 (Part pages):
6x9 in, Georgia body justified, Amazon Ember bold blue headings, Title-style chapter
headings "Chapter N: Title" with outline level 0, Heading 2 sections, hyperlinked
Contents with PAGEREF page numbers, "Back to Contents" link after every chapter,
running header with the book title and a page number in the footer.

Markdown conventions
  # PART I — Name          Part page (lines up to the first ## are the Part's opening note)
  ## 12. Title            numbered chapter
  ## Prologue: Title      unnumbered Title-style heading (Prologue, Epilogue, Glossary ...)
  ### Section            Heading 2
  ![Figure 5. Caption](figures/figs/fig05.jpg)
  | table |, - list item, **bold**, *italic*
  %%TITLEPAGE%% %%COPYRIGHT%% %%TOC%% %%ALSO%% %%INDEX%% %%ALSOBY%%   generated blocks
"""
import io, os, re, sys, glob, json
from docx import Document
from docx.shared import Pt, Inches, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
SRC = os.path.join(ROOT, "chapters")
OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(ROOT, "Life, Actually.docx")

SERIES = "Physics, Actually"
TITLE = "Life, Actually"
SUBTITLE = "From the First Cell to the Edited Genome and the Search for Life Elsewhere"
AUTHOR = "Lothar J. Musiol"
BODY_FONT, HEAD_FONT = "Georgia", "Amazon Ember"
BLUE = RGBColor(0x00, 0x00, 0xFF)
ALSO = [
    "Volume 1 — Motion, Forces, Time, and Relativity",
    "Volume 2 — Gravity, Cosmology, and the Limits of Spacetime",
    "Volume 3 — The Standard Model, Chaos, and the Edge of Knowledge",
]
# Back matter: the canonical "Also by Lothar J. Musiol" list (same in every book; see
# notes/ALSO_BY - canonical list.md at the repo root). Titles as on each master's title page.
ALSO_BY = [
    ('Physics, Actually', [
        'Physics, Actually, Volume 1: Motion, Forces, Time, and Relativity',
        'Physics, Actually, Volume 2: Gravity, Cosmology, and the Limits of Spacetime',
        'Physics, Actually, Volume 3: The Standard Model, Chaos, and the Edge of Knowledge',
        'Life, Actually: From the First Cell to the Edited Genome and the Search for Life Elsewhere',
    ]),
    ('Math, Actually', [
        'Math, Actually, Volume 1: From Arithmetic to Calculus',
        'Math, Actually, Volume 2: From Multivariable Calculus to Set Theory & Logic',
        'Math, Actually, Volume 3: From Differential Equations to Abstract Algebra',
        'Math, Actually, Volume 4: From Category Theory to the Frontier',
    ]),
    ('Quanta, Actually', [
        'Quanta, Actually, Volume 1: The Quantum World',
        'Quanta, Actually, Volume 2: The Quantum Conversation',
        'Quanta, Actually, Volume 3: Complete Quantum Electrodynamics Course',
    ]),
    ('Science Sparks', [
        'Science Sparks: Physics, Life, and Mathematics — The Same Few Rules, Told in Highlights',
    ]),
    ('Look First', [
        'Look First, Volume 1: The Universe Has No Now',
        'Look First, Volume 2: A Trip Is Not a New Life',
    ]),
    ('Electrical Engineering Series', [
        'Foundations of Electronics (Book 1)',
        'Circuits, Components, and Control (Book 2)',
        'Semiconductor Physics and Devices (Book 3)',
        'RF, Microwave, and Transceivers (Book 4)',
        'Communications, Wireless, and SDR (Book 5)',
        'Power and Energy (Book 6)',
        'Packaging, Layout, EMC, and Test (Book 7)',
    ]),
    ('History', [
        "The Dolphins' View of History",
    ]),
    ('Fiction', [
        "The Murder That Hadn't Happened Yet (The Relativistic Investigation Bureau, Book 1)",
        'The Warning That Was Sent Too Late (The Relativistic Investigation Bureau, Book 2)',
        'Schrödinger’s Paperwork (Lolly Wren’s Curious Science Adventures, Book 1)',
        'The Permitted Options (Lolly Wren’s Curious Science Adventures, Book 2)',
        'Protocol Flamingo (The Invasion Storybooks, Book 1), as George Herbert Fontaine',
    ]),
    ('How-To', [
        'Your First Book That Sells',
        'Your First YouTube Channel That Rocks',
    ]),
]
COPYRIGHT = [
    SERIES,
    f"{TITLE} — {SUBTITLE}",
    "",
    "Copyright © 2026 Lothar J. Musiol",
    "All rights reserved.",
    "",
    "No part of this publication may be reproduced, distributed, or transmitted in any form or by any means, including photocopying, recording, or other electronic or mechanical methods, without the prior written permission of the author, except in the case of brief quotations embodied in critical reviews and certain other noncommercial uses permitted by copyright law.",
    "",
    "This book is intended for general educational and informational purposes. Every effort has been made to ensure the accuracy of the science described; some topics discussed represent active or speculative areas of research, and this is noted in the text where relevant. It reflects published research as of October 2026.",
    "",
    "Photographs are credited in the Photo Credits section at the back of the book. Diagrams are original.",
    "",
    "First edition.",
    "Physics, Actually series",
]

# ---------------------------------------------------------------- parsing
def load_blocks():
    files = sorted(glob.glob(os.path.join(SRC, "[0-9][0-9]_*.md")))
    blocks = []
    for f in files:
        lines = open(f, encoding="utf-8").read().splitlines()
        i = 0
        while i < len(lines):
            line = lines[i].rstrip()
            if not line.strip(): i += 1; continue
            m = re.match(r"^%%(\w+)%%$", line.strip())
            if m: blocks.append(("gen", m.group(1))); i += 1; continue
            if line.startswith("# "):
                m = re.match(r"# PART ([IVXLC]+) — (.*)", line)
                blocks.append(("part", m.group(1), m.group(2).strip())); i += 1; continue
            if line.startswith("## "):
                t = line[3:].strip(); m = re.match(r"(\d+)\.\s+(.*)", t)
                if m: blocks.append(("chapter", int(m.group(1)), m.group(2)))
                else: blocks.append(("unnum", t))
                i += 1; continue
            if line.startswith("### "):
                blocks.append(("h2", line[4:].strip())); i += 1; continue
            if line.startswith("!["):
                m = re.match(r"!\[Figure (\d+)\. (.*)\]\((.*)\)\s*$", line)
                blocks.append(("fig", int(m.group(1)), m.group(2), m.group(3))); i += 1; continue
            if line.startswith("|"):
                rows = []
                while i < len(lines) and lines[i].startswith("|"):
                    r = lines[i].strip()
                    if not re.match(r"^\|\s*-", r): rows.append([c.strip() for c in r.strip("|").split("|")])
                    i += 1
                blocks.append(("table", rows)); continue
            if re.match(r"^\s*-\s+", line):
                items = []
                while i < len(lines) and re.match(r"^\s*-\s+", lines[i]):
                    items.append(re.sub(r"^\s*-\s+", "", lines[i]).strip()); i += 1
                blocks.append(("list", items)); continue
            para = [line.strip()]; i += 1
            while i < len(lines) and lines[i].strip() and not re.match(r"^(#|!\[|\||%%|\s*-\s)", lines[i]):
                para.append(lines[i].strip()); i += 1
            blocks.append(("p", " ".join(para)))
    return blocks

# ---------------------------------------------------------------- inline formatting
TOKEN = re.compile(r"(\*\*.+?\*\*|\*[^*\n]+?\*)")
def add_runs(par, text, bold=False, italic=False, size=None):
    pos = 0
    for m in TOKEN.finditer(text):
        if m.start() > pos: _run(par, text[pos:m.start()], bold, italic, size)
        tok = m.group(0)
        if tok.startswith("**"): add_runs(par, tok[2:-2], True, italic, size)
        else: add_runs(par, tok[1:-1], bold, not italic, size)
        pos = m.end()
    if pos < len(text): _run(par, text[pos:], bold, italic, size)
def _run(par, s, bold, italic, size):
    if not s: return
    r = par.add_run(s); r.bold = bold or None; r.italic = italic or None
    if size: r.font.size = Pt(size)
def plain(t): return re.sub(r"\*+", "", t)

# ---------------------------------------------------------------- docx helpers
_bm = [10]
def bookmark(par, name):
    _bm[0] += 1
    s = OxmlElement("w:bookmarkStart"); s.set(qn("w:id"), str(_bm[0])); s.set(qn("w:name"), name)
    e = OxmlElement("w:bookmarkEnd"); e.set(qn("w:id"), str(_bm[0]))
    pPr = par._p.find(qn("w:pPr"))
    if pPr is not None: pPr.addnext(s)
    else: par._p.insert(0, s)
    par._p.append(e)
def link(par, text, anchor, color=True):
    h = OxmlElement("w:hyperlink"); h.set(qn("w:anchor"), anchor); h.set(qn("w:history"), "1")
    r = OxmlElement("w:r"); rpr = OxmlElement("w:rPr")
    if color:
        c = OxmlElement("w:color"); c.set(qn("w:val"), "0563C1"); rpr.append(c)
        u = OxmlElement("w:u"); u.set(qn("w:val"), "single"); rpr.append(u)
    r.append(rpr); t = OxmlElement("w:t"); t.text = text; t.set(qn("xml:space"), "preserve"); r.append(t)
    h.append(r); par._p.append(h)
def field(par, instr, shown="1"):
    def fc(kind):
        r = OxmlElement("w:r"); f = OxmlElement("w:fldChar"); f.set(qn("w:fldCharType"), kind); r.append(f); return r
    par._p.append(fc("begin"))
    r = OxmlElement("w:r"); it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = f" {instr} "; r.append(it); par._p.append(r)
    par._p.append(fc("separate"))
    r = OxmlElement("w:r"); t = OxmlElement("w:t"); t.text = shown; r.append(t); par._p.append(r)
    par._p.append(fc("end"))
def outline0(par):
    pPr = par._p.get_or_add_pPr(); o = OxmlElement("w:outlineLvl"); o.set(qn("w:val"), "0"); pPr.append(o)
def page_break(doc):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(0); p.add_run().add_break(WD_BREAK.PAGE)
def set_font(style, name, size=None, bold=None, italic=None, color=None):
    style.font.name = name
    if size: style.font.size = Pt(size)
    if bold is not None: style.font.bold = bold
    if italic is not None: style.font.italic = italic
    if color is not None: style.font.color.rgb = color
    rpr = style.element.get_or_add_rPr(); rf = rpr.find(qn("w:rFonts"))
    if rf is None: rf = OxmlElement("w:rFonts"); rpr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"): rf.set(qn(a), name)
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if rf.get(qn(a)) is not None: del rf.attrib[qn(a)]

def setup_styles(doc):
    from docx.enum.style import WD_STYLE_TYPE
    n = doc.styles["Normal"]; set_font(n, BODY_FONT, 11)
    pf = n.paragraph_format; pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY; pf.space_before = Pt(0); pf.space_after = Pt(8); pf.line_spacing = 1.06; pf.first_line_indent = Inches(0)
    t = doc.styles["Title"]; set_font(t, HEAD_FONT, 20, True, False, BLUE)
    t.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT; t.paragraph_format.space_before = Pt(36); t.paragraph_format.space_after = Pt(15); t.paragraph_format.keep_with_next = True
    tp = t.element.find(qn("w:pPr"))
    if tp is not None:
        for b in tp.findall(qn("w:pBdr")): tp.remove(b)
    h1 = doc.styles["Heading 1"]; set_font(h1, HEAD_FONT, 18, True, False, BLUE)
    h1.paragraph_format.space_before = Pt(24); h1.paragraph_format.space_after = Pt(6); h1.paragraph_format.keep_with_next = True; h1.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    h2 = doc.styles["Heading 2"]; set_font(h2, HEAD_FONT, 18, True, False, BLUE)
    h2.paragraph_format.space_before = Pt(14); h2.paragraph_format.space_after = Pt(4); h2.paragraph_format.keep_with_next = True; h2.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    def mk(name, base="Normal"):
        try: st = doc.styles[name]
        except KeyError: st = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        st.base_style = doc.styles[base]; return st
    cap = mk("Figure Caption"); set_font(cap, BODY_FONT, 9.5, italic=True); cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER; cap.paragraph_format.space_after = Pt(14)
    li = mk("List Item"); li.paragraph_format.left_indent = Inches(0.35); li.paragraph_format.first_line_indent = Inches(-0.2); li.paragraph_format.space_after = Pt(3); li.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    toc = mk("TOC Entry"); toc.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT; toc.paragraph_format.space_after = Pt(2)
    tocp = mk("TOC Part"); set_font(tocp, BODY_FONT, 11, bold=True); tocp.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT; tocp.paragraph_format.space_before = Pt(8); tocp.paragraph_format.space_after = Pt(2)
    small = mk("Copyright Text"); set_font(small, BODY_FONT, 9); small.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT; small.paragraph_format.space_after = Pt(6)
    idx = mk("Index Entry"); set_font(idx, BODY_FONT, 9.5); idx.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT; idx.paragraph_format.space_after = Pt(1); idx.paragraph_format.left_indent = Inches(0.2); idx.paragraph_format.first_line_indent = Inches(-0.2)

def word_picture(path):
    from PIL import Image
    im = Image.open(path); im.load()
    if im.mode == "RGBA":
        bg = Image.new("RGB", im.size, (255, 255, 255)); bg.paste(im, mask=im.split()[-1]); im = bg
    elif im.mode != "RGB": im = im.convert("RGB")
    if max(im.size) > 1800: im.thumbnail((1800, 1800), Image.Resampling.LANCZOS)
    buf = io.BytesIO(); jpg = path.lower().endswith((".jpg", ".jpeg"))
    im.save(buf, format="JPEG" if jpg else "PNG", **({"quality": 85, "optimize": True} if jpg else {})); buf.seek(0)
    return buf

def add_footer_header(section):
    section.footer.is_linked_to_previous = False; section.header.is_linked_to_previous = False
    fp = section.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER; field(fp, "PAGE", "1")
    hp = section.header.paragraphs[0]; hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = hp.add_run(TITLE); r.italic = True; r.font.size = Pt(9); r.font.name = BODY_FONT

# ---------------------------------------------------------------- index
def build_index(blocks):
    """Index by chapter number. Terms live in notes/index_terms.txt, one per line:
       Display term | regex (optional; defaults to the term itself, case-insensitive)."""
    path = os.path.join(ROOT, "notes", "index_terms.txt")
    if not os.path.exists(path): return []
    terms = []
    for line in open(path, encoding="utf-8"):
        line = line.strip()
        if not line or line.startswith("#"): continue
        disp, _, rx = line.partition("|"); disp = disp.strip(); rx = rx.strip() or re.escape(disp)
        terms.append((disp, re.compile(r"(?<![\w-])(?:" + rx + r")(?![\w-])", re.I)))
    texts = {}; cur = None
    for b in blocks:
        if b[0] == "chapter": cur = str(b[1])
        elif b[0] == "unnum":
            t = b[1].split(":")[0]; cur = t if t in ("Prologue", "Epilogue") else None
        elif b[0] == "part": cur = None
        if cur and b[0] in ("p", "list", "table", "h2"):
            s = b[1] if isinstance(b[1], str) else " ".join(x if isinstance(x, str) else " ".join(x) for x in b[1])
            texts.setdefault(cur, []).append(plain(s))
    out = []
    for disp, rx in terms:
        hits = [k for k, v in texts.items() if rx.search(" ".join(v))]
        def key(k): return (0, 0) if k == "Prologue" else ((2, 0) if k == "Epilogue" else (1, int(k)))
        hits.sort(key=key)
        if hits: out.append((disp, hits))
    out.sort(key=lambda x: re.sub(r"[^a-z0-9 ]", "", x[0].lower()))
    return out

# ---------------------------------------------------------------- build
TOC_PAGES = json.load(open(os.path.join(ROOT, "notes", "toc_pages.json"))) if os.path.exists(os.path.join(ROOT, "notes", "toc_pages.json")) else {}

def build():
    blocks = load_blocks()
    doc = Document(); setup_styles(doc)
    sec = doc.sections[0]; sec.page_width, sec.page_height = Inches(6), Inches(9)
    for m in ("left_margin", "right_margin"): setattr(sec, m, Inches(0.75))
    sec.top_margin = sec.bottom_margin = Inches(0.8)
    # anchors
    toc = []
    for b in blocks:
        if b[0] == "part": toc.append(("part", f"Part {b[1]} — {b[2]}", f"part_{b[1]}"))
        elif b[0] == "chapter": toc.append(("ch", f"Chapter {b[1]}: {plain(b[2])}", f"chapter_{b[1]}"))
        elif b[0] == "unnum" and b[1] not in ("Contents",): toc.append(("ch", plain(b[1]), "ch_" + re.sub(r"\W+", "_", plain(b[1]).split(":")[0].lower()).strip("_")))
    anchors = {t: a for k, t, a in toc}
    figs = []; in_body = False; prev = None; open_chapter = False
    def close_chapter():
        p = doc.add_paragraph(); link(p, "↑ Back to Contents", "chcontents")
    for b in blocks:
        k = b[0]
        if k == "gen":
            g = b[1]
            if g == "TITLEPAGE":
                doc.add_paragraph(SERIES, style="Title").alignment = WD_ALIGN_PARAGRAPH.CENTER
                for _ in range(3): doc.add_paragraph()
                h = doc.add_paragraph(TITLE, style="Heading 1"); h.alignment = WD_ALIGN_PARAGRAPH.CENTER
                s = doc.add_paragraph(); s.alignment = WD_ALIGN_PARAGRAPH.CENTER; r = s.add_run(SUBTITLE); r.italic = True; r.font.size = Pt(13)
                for _ in range(2): doc.add_paragraph()
                a = doc.add_paragraph("A Volume in the Physics, Actually Series", style="Heading 2"); a.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for _ in range(3): doc.add_paragraph()
                a = doc.add_paragraph(AUTHOR, style="Heading 2"); a.alignment = WD_ALIGN_PARAGRAPH.CENTER
                page_break(doc)
            elif g == "COPYRIGHT":
                for line in COPYRIGHT: doc.add_paragraph(line, style="Copyright Text")
                page_break(doc)
            elif g == "TOC":
                h = doc.add_paragraph("Contents", style="Heading 1"); bookmark(h, "chcontents")
                for kind, text, anc in toc:
                    p = doc.add_paragraph(style="TOC Part" if kind == "part" else "TOC Entry")
                    if kind == "ch":
                        tabs = p._p.get_or_add_pPr(); tb = OxmlElement("w:tabs"); t = OxmlElement("w:tab")
                        t.set(qn("w:val"), "right"); t.set(qn("w:leader"), "dot"); t.set(qn("w:pos"), "6480"); tb.append(t); tabs.append(tb)
                        p.paragraph_format.left_indent = Inches(0.2)
                    link(p, text, anc)
                    if kind == "ch":
                        p.add_run().add_tab(); field(p, f"PAGEREF {anc} \\h", str(TOC_PAGES.get(anc, 1)))
                page_break(doc)
            elif g == "ALSO":
                doc.add_paragraph("Also in This Series", style="Heading 1")
                for line in ALSO: doc.add_paragraph(line).alignment = WD_ALIGN_PARAGRAPH.LEFT
                # body section starts here: running header + page numbers
                new = doc.add_section(WD_SECTION.NEW_PAGE); add_footer_header(new)
                sec.header.is_linked_to_previous = False
            elif g == "ALSOBY":
                for group, titles in ALSO_BY:
                    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    p.paragraph_format.space_before = Pt(8); p.paragraph_format.space_after = Pt(2)
                    p.paragraph_format.keep_with_next = True; p.paragraph_format.first_line_indent = Inches(0)
                    p.add_run(group).bold = True
                    for title in titles:
                        q = doc.add_paragraph(); q.alignment = WD_ALIGN_PARAGRAPH.LEFT
                        q.paragraph_format.left_indent = Inches(0.25); q.paragraph_format.first_line_indent = Inches(0)
                        q.paragraph_format.space_before = Pt(0); q.paragraph_format.space_after = Pt(1)
                        q.add_run(title)
            elif g == "INDEX":
                for disp, hits in build_index(blocks):
                    p = doc.add_paragraph(style="Index Entry")
                    add_runs(p, disp); p.add_run(", ")
                    for j, hnum in enumerate(hits):
                        if j: p.add_run(", ")
                        link(p, hnum, f"chapter_{hnum}" if hnum.isdigit() else "ch_" + hnum.lower(), color=False)
            prev = "gen"; continue
        if k == "part":
            if open_chapter: close_chapter(); open_chapter = False
            page_break(doc)
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before = Pt(100); p.paragraph_format.space_after = Pt(10)
            r = p.add_run(f"PART {b[1]}"); r.bold = True; r.font.size = Pt(20); r.font.name = HEAD_FONT; r.font.color.rgb = BLUE
            bookmark(p, f"part_{b[1]}"); outline0(p)
            q = doc.add_paragraph(); q.alignment = WD_ALIGN_PARAGRAPH.CENTER; r = q.add_run(b[2]); r.italic = True; r.font.size = Pt(14)
            prev = "part"; continue
        if k in ("chapter", "unnum"):
            if open_chapter: close_chapter()
            page_break(doc)
            text = f"Chapter {b[1]}: {plain(b[2])}" if k == "chapter" else plain(b[1])
            h = doc.add_paragraph(style="Title"); add_runs(h, text if k == "unnum" else f"Chapter {b[1]}: {b[2]}")
            outline0(h); bookmark(h, anchors.get(text) or f"chapter_{b[1]}")
            open_chapter = True; prev = "head"; continue
        if k == "h2":
            doc.add_paragraph(style="Heading 2").add_run(plain(b[1])) if "*" not in b[1] else add_runs(doc.add_paragraph(style="Heading 2"), b[1])
            prev = "head"; continue
        if k == "fig":
            n, caption, rel = b[1], b[2], b[3]
            path = os.path.join(ROOT, rel.replace("/", os.sep))
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.keep_with_next = True; p.paragraph_format.space_before = Pt(8)
            if os.path.exists(path):
                shape = p.add_run().add_picture(word_picture(path), width=Inches(4.5))
                shape._inline.docPr.set("descr", f"Figure {n}. {plain(caption)}"[:250]); shape._inline.docPr.set("title", f"Figure {n}")
            else:
                p.add_run(f"[Figure {n} missing: {rel}]").italic = True; figs.append(("MISSING", n))
            c = doc.add_paragraph(style="Figure Caption"); c.add_run(f"Figure {n}. ").bold = True; add_runs(c, caption)
            figs.append(n); prev = "fig"; continue
        if k == "table":
            rows = b[1]; ncol = max(len(r) for r in rows)
            t = doc.add_table(rows=len(rows), cols=ncol); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
            for ri, row in enumerate(rows):
                for ci in range(ncol):
                    cell = t.cell(ri, ci); par = cell.paragraphs[0]; par.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT; par.paragraph_format.space_after = Pt(2)
                    add_runs(par, row[ci] if ci < len(row) else "", bold=(ri == 0), size=8.5)
            doc.add_paragraph(); prev = "table"; continue
        if k == "list":
            for it in b[1]:
                p = doc.add_paragraph(style="List Item"); p.add_run("•  "); add_runs(p, it)
            prev = "list"; continue
        if k == "p":
            p = doc.add_paragraph(); add_runs(p, b[1])
            if prev == "part":
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs: r.italic = True
                p.paragraph_format.left_indent = p.paragraph_format.right_indent = Inches(0.4)
                p.paragraph_format.space_before = Pt(30)
                continue
            prev = "p"; continue
    if open_chapter: close_chapter()
    # Word refreshes PAGEREF page numbers in the Contents when the file is opened
    settings = doc.settings.element; uf = OxmlElement("w:updateFields"); uf.set(qn("w:val"), "true"); settings.append(uf)
    from datetime import datetime, timezone
    cp = doc.core_properties
    cp.title = TITLE; cp.subject = SUBTITLE; cp.author = AUTHOR; cp.last_modified_by = AUTHOR
    cp.category = SERIES + " series"; cp.keywords = "life, evolution, origin of life, DNA, genetics, CRISPR, genome, astrobiology, exobiology"
    cp.comments = f"{TITLE}. A volume in the {SERIES} series by {AUTHOR}."
    stamp = datetime(2026, 10, 2, tzinfo=timezone.utc); cp.created = stamp; cp.modified = stamp
    doc.save(OUT)
    nfig = [f for f in figs if isinstance(f, int)]
    print("saved", OUT); print("chapters:", sum(1 for b in blocks if b[0] == "chapter"), "figures:", len(nfig), "missing:", [f for f in figs if not isinstance(f, int)])
    if nfig != list(range(1, len(nfig) + 1)): print("WARNING figure numbers not continuous:", nfig)

if __name__ == "__main__":
    build()
