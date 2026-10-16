import sys,zipfile,re
from docx import Document
p=sys.argv[1]
d=Document(p)
s=d.sections[0]; print('page',s.page_width.inches,s.page_height.inches,s.left_margin.inches)
print('tables',len(d.tables),'inline pics',len(d.inline_shapes))
h1=[x for x in d.paragraphs if x.style.name=='Heading 1']; print('H1',len(h1),[x.text for x in h1[:2]],[x.text for x in h1[21:25]],[x.text for x in h1[-9:]])
st=d.styles['Heading 1']; print(st.font.size.pt,st.font.color.rgb,st.paragraph_format.alignment,st.paragraph_format.page_break_before)
st=d.styles['Heading 2']; print(st.font.size.pt,st.font.color.rgb,st.paragraph_format.alignment)
z=zipfile.ZipFile(p); print('testzip',z.testzip()); x=z.read('word/document.xml').decode('utf8')
print('hyperlinks',x.count('<w:hyperlink'),'bookmarks',x.count('w:bookmarkStart'))
anchors=set(re.findall(r'w:anchor="([^"]+)"',x)); names=set(re.findall(r'w:name="(_Toc\d+)"',x)); print('unresolved',len(anchors-names))
print([pp.text for pp in d.paragraphs if pp.text.startswith('Copyright')])
empt=sum(1 for pp in d.paragraphs if not pp.text.strip() and not pp._p.xpath('.//w:drawing')); print('empty paras',empt)
