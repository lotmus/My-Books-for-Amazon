import re, sys, os, glob
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

MD=sys.argv[1]; OUT=sys.argv[2]; COVER=sys.argv[3]
BLUE=RGBColor(0x1F,0x4E,0x79); FONT='Georgia'
doc=Document()
sec=doc.sections[0]
sec.page_width=Inches(3.6); sec.page_height=Inches(4.8)
for a in ('left_margin','right_margin','top_margin','bottom_margin'): setattr(sec,a,Inches(0.3))
sec.header_distance=Inches(0.1); sec.footer_distance=Inches(0.1)
st=doc.styles
def setfont(style,size,bold=None,italic=None,color=None):
    f=style.font; f.name=FONT; f.size=Pt(size)
    rpr=style.element.get_or_add_rPr(); rf=rpr.find(qn('w:rFonts'))
    if rf is None: rf=OxmlElement('w:rFonts'); rpr.append(rf)
    for k in ('w:ascii','w:hAnsi','w:cs','w:eastAsia'): rf.set(qn(k),FONT)
    if bold is not None: f.bold=bold
    if italic is not None: f.italic=italic
    if color is not None: f.color.rgb=color
n=st['Normal']; setfont(n,10)
n.paragraph_format.space_after=Pt(6); n.paragraph_format.space_before=Pt(0); n.paragraph_format.line_spacing=1.12
lang=OxmlElement('w:lang'); lang.set(qn('w:val'),'en-GB'); n.element.get_or_add_rPr().append(lang)
h1=st['Heading 1']; setfont(h1,16,True,False,BLUE)
h1.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER; h1.paragraph_format.page_break_before=True
h1.paragraph_format.space_before=Pt(14); h1.paragraph_format.space_after=Pt(12); h1.paragraph_format.keep_with_next=True
h2=st['Heading 2']; setfont(h2,11,True,False,BLUE)
h2.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT; h2.paragraph_format.space_before=Pt(10); h2.paragraph_format.space_after=Pt(4); h2.paragraph_format.keep_with_next=True
def mk(name,base='Normal'):
    try: s=st[name]
    except KeyError: s=st.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
    s.base_style=st[base]; return s
part=mk('Part Title','Heading 1'); setfont(part,18,True,False,BLUE)
part.paragraph_format.space_before=Pt(70); part.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
part.element.get_or_add_pPr().append(OxmlElement('w:outlineLvl')); part.element.pPr.find(qn('w:outlineLvl')).set(qn('w:val'),'0')
q=mk('Quote Block'); setfont(q,9); q.paragraph_format.left_indent=Inches(0.15); q.paragraph_format.right_indent=Inches(0.1); q.paragraph_format.space_after=Pt(2)
sb=mk('Scene Break'); sb.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER; sb.paragraph_format.space_before=Pt(6); sb.paragraph_format.space_after=Pt(6)
ctr=mk('Centered'); ctr.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
toc1=mk('TOC 1'); setfont(toc1,9); toc1.paragraph_format.space_after=Pt(2); toc1.paragraph_format.left_indent=Inches(0.12)
tocp=mk('TOC Part'); setfont(tocp,9,True); tocp.paragraph_format.space_before=Pt(5); tocp.paragraph_format.space_after=Pt(2)
cop=mk('Small'); setfont(cop,8); cop.paragraph_format.space_after=Pt(4)
tt=mk('Novel Title'); setfont(tt,24,True,False,BLUE); tt.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER; tt.paragraph_format.space_before=Pt(60); tt.paragraph_format.space_after=Pt(8)
tsub=mk('Subtitle'); setfont(tsub,11,False,True); tsub.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
grp=mk('Group'); setfont(grp,9,True); grp.paragraph_format.space_before=Pt(5); grp.paragraph_format.space_after=Pt(2); grp.paragraph_format.keep_with_next=True
item=mk('Also By Item'); setfont(item,8.5); item.paragraph_format.space_after=Pt(2); item.paragraph_format.left_indent=Inches(0.1)

TOK=re.compile(r'(\*\*|\*|\^\{[^}]*\}|_\{[^}]*\})')
def add_runs(p,text):
    b=i=False
    for part_ in TOK.split(text):
        if part_=='**': b=not b
        elif part_=='*': i=not i
        elif part_ and part_[:2] in ('^{','_{') and part_.endswith('}'):
            r=p.add_run(part_[2:-1]); r.bold=b or None; r.italic=i or None
            if part_[0]=='^': r.font.superscript=True
            else: r.font.subscript=True
        elif part_:
            r=p.add_run(part_); r.bold=b or None; r.italic=i or None
bm_id=[100]
def heading(text,style='Heading 1'):
    p=doc.add_paragraph(style=style); add_runs(p,text)
    bm_id[0]+=1; name=f'_Toc{bm_id[0]}'
    s=OxmlElement('w:bookmarkStart'); s.set(qn('w:id'),str(bm_id[0])); s.set(qn('w:name'),name)
    e=OxmlElement('w:bookmarkEnd'); e.set(qn('w:id'),str(bm_id[0]))
    p._p.insert(1 if p._p.pPr is not None else 0,s); p._p.append(e)
    return name
def para(text,style=None,align=None):
    p=doc.add_paragraph(style=style); add_runs(p,text)
    if align: p.alignment=align
    return p
def pagebreak_para(style='Normal'):
    p=doc.add_paragraph(style=style); p.paragraph_format.page_break_before=True; return p
def link_entry(text,anchor,style):
    p=doc.add_paragraph(style=style)
    h=OxmlElement('w:hyperlink'); h.set(qn('w:anchor'),anchor); h.set(qn('w:history'),'1')
    r=OxmlElement('w:r'); rpr=OxmlElement('w:rPr'); rs=OxmlElement('w:rStyle'); 
    c=OxmlElement('w:color'); c.set(qn('w:val'),'1F4E79'); rpr.append(c)
    u=OxmlElement('w:u'); u.set(qn('w:val'),'single'); rpr.append(u); r.append(rpr)
    t=OxmlElement('w:t'); t.text=text; t.set(qn('xml:space'),'preserve'); r.append(t); h.append(r); p._p.append(h)


cth=mk('Contents Heading'); setfont(cth,16,True,False,BLUE); cth.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
cth.paragraph_format.page_break_before=True; cth.paragraph_format.space_after=Pt(10)
txt=open(MD,encoding='utf-8').read()
blocks=re.split(r'\n(?=## )',txt)
title_block=blocks[0]; blocks=blocks[1:]
chapters=[];back={}
for b in blocks:
    lines=b.split('\n'); h=lines[0][3:].strip(); body=[l for l in lines[1:]]
    if h in ('A Note to Readers','About the Author') or h.startswith('Also by'): back[h]=body
    else: chapters.append((h,body))
# cover
p=doc.add_paragraph(style='Centered'); p.paragraph_format.space_after=Pt(0)
p.add_run().add_picture(COVER,width=Inches(2.55))
pp=pagebreak_para('Novel Title'); pp.paragraph_format.space_before=Pt(50); pp.add_run('The Department of Historical Corrections')
para('A Novel','Subtitle'); para('','Normal'); para('**Lothar J. Musiol**','Centered')
pc=pagebreak_para('Small'); pc.add_run('Copyright © 2026 Lothar J. Musiol. All rights reserved.')
para('This is a work of fiction. Historical settings are used imaginatively; the Department, its technology and its people are invented, and any resemblance to actual persons, living or dead, or to any department that has ever misplaced a country, is coincidental.','Small')
para('No part of this book may be reproduced without written permission, except for brief quotations in reviews.','Small')
para('First edition, 2026.','Small')
entries=[]
ch=pagebreak_para('Contents Heading'); ch.add_run('Contents')
toc_marker=doc.add_paragraph(style='Normal')
for h,body in chapters:
    nm=heading(h); entries.append(('ch',h,nm))
    for l in body:
        s_=l.rstrip()
        if not s_.strip(): continue
        if s_.strip() in ('* * *','---'): para('* * *','Scene Break'); continue
        if s_.startswith('### '): para(s_[4:].strip(),'Heading 2'); continue
        para(s_.strip())
def section(h,body,style_para=None):
    nm=heading(h); entries.append(('ch',h,nm))
    return nm
nm=heading('A Note to Readers'); entries.append(('ch','A Note to Readers',nm))
for l in back['A Note to Readers']:
    if l.strip(): para(l.strip())
nm=heading('About the Author'); entries.append(('ch','About the Author',nm))
for l in back['About the Author']:
    if l.strip(): para(l.strip())
nm=heading('Also by Lothar J. Musiol'); entries.append(('ch','Also by Lothar J. Musiol',nm))
for l in back[[k for k in back if k.startswith('Also by')][0]]:
    l=l.strip()
    if not l: continue
    if l.startswith('**'): para(l.strip('*'),'Group')
    elif l.startswith('- '): para(l[2:],'Also By Item')
for kind,text,nm in entries:
    link_entry(text,nm,'TOC 1')
body_el=doc.element.body
ents=[e for e in body_el if e.tag==qn('w:p')]
last=ents[-len(entries):]
for e in last: body_el.remove(e)
idx=list(body_el).index(toc_marker._p)
for k,e in enumerate(last): body_el.insert(idx+1+k,e)
body_el.remove(toc_marker._p)
cp=doc.core_properties; cp.title='The Department of Historical Corrections'; cp.author='Lothar J. Musiol'; cp.language='en-GB'
doc.save(OUT); print('saved',OUT,len(chapters),'chapters')
