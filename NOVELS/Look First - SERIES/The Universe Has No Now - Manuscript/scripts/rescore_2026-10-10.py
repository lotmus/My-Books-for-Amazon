"""The Universe Has No Now: 10 Oct 2026 fixes, applied in place to export/The_Universe_Has_No_Now.docx
(the docx carries back matter the markdown build does not, so it is edited directly).
Usage: python rescore_2026-10-10.py SRC DST FIG09_JPG"""
import docx, copy, sys
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
SRC, DST, FIG = sys.argv[1:4]
d = docx.Document(SRC); P = d.paragraphs
def one(pred):
    h = [i for i, p in enumerate(P) if pred(p)]; assert len(h) == 1, h; return h[0]
# 1. straight apostrophe
i = one(lambda p: "Chapter 24's proper subject" in p.text)
for t in P[i]._p.iter(qn('w:t')):
    if t.text and "24's" in t.text: t.text = t.text.replace("24's", "24\u2019s")
# 2. Figure 9 in colour
i = one(lambda p: p.text.startswith('Figure 9. The microwave sky'))
img = P[i - 1]; rid = [b.get(qn('r:embed')) for b in img._p.iter(qn('a:blip'))][0]
part = d.part.related_parts[rid]; assert part.partname.endswith('.jpg'); part._blob = open(FIG, 'rb').read()
cap = 'Figure 9. The microwave sky as an oval map: blue patches are slightly cooler, red patches slightly warmer. This is a measurement, not a snapshot from a camera.'
ts = list(P[i]._p.iter(qn('w:t'))); ts[0].text = cap
for t in ts[1:]: t.text = ''
for dp in img._p.iter(qn('wp:docPr')):
    dp.set('descr', 'Planck map of the microwave sky: an oval speckled with blue (cooler) and red (warmer) patches, with a labelled cooler-to-warmer bar.')
# 3. Also by page (canonical list, without this book), after About the Author
ab = one(lambda p: p.text == 'ABOUT THE AUTHOR' and p.style.name == 'Heading 1')
last = P[ab + 2]; assert last.text.startswith('He writes to make')
GROUPS = [
 ('Physics, Actually', ['Physics, Actually, Volume 1: Motion, Forces, Time, and Relativity', 'Physics, Actually, Volume 2: Gravity, Cosmology, and the Limits of Spacetime', 'Physics, Actually, Volume 3: The Standard Model, Chaos, and the Edge of Knowledge', 'Life, Actually: From the First Cell to the Edited Genome and the Search for Life Elsewhere']),
 ('More Actually Books', ['Gravity, Actually: Einstein’s Relativity in Plain Words: Falling, Time, Black Holes, and the Waves We Finally Heard', 'The Universe, Actually: What Humanity Currently Understands About Reality']),
 ('Math, Actually', ['Math, Actually, Volume 1: From Arithmetic to Calculus', 'Math, Actually, Volume 2: From Multivariable Calculus to Set Theory & Logic', 'Math, Actually, Volume 3: From Differential Equations to Abstract Algebra', 'Math, Actually, Volume 4: From Category Theory to the Frontier']),
 ('Quantum, Actually', ['Quantum, Actually, Volume 1: Questions and Answers from the Double Slit to the Superconducting Wire', 'Quantum, Actually, Volume 2: A QED Course']),
 ('Science Sparks', ['Science Sparks: Physics, Life, and Mathematics \u2014 The Same Few Rules, Told in Highlights']),
 ('Look First', ['Look First, Volume 2: A Trip Is Not a New Life']),
 ('Electrical Engineering Series', ['Foundations of Electronics (Book 1)', 'Circuits, Components, and Control (Book 2)', 'Semiconductor Physics and Devices (Book 3)', 'RF, Microwave, and Transceivers (Book 4)', 'Communications, Wireless, and SDR (Book 5)', 'Power and Energy (Book 6)', 'Packaging, Layout, EMC, and Test (Book 7)']),
 ('History', ['The Dolphins\u2019 View of History']),
 ('Fiction', ['The Murder That Hadn\u2019t Happened Yet (The Relativistic Investigation Bureau, Book 1), Lothar J. Musiol writing as Spezala Genara Relavi (S.G.R.)',
              'The Warning That Was Sent Too Late (The Relativistic Investigation Bureau, Book 2), Lothar J. Musiol writing as Spezala Genara Relavi (S.G.R.)',
              'Schr\u00f6dinger\u2019s Paperwork (Lolly Wren\u2019s Curious Science Adventures, Book 1), Lothar J. Musiol writing as Quentin Edmund Darling (Q.E.D.)',
              'The Permitted Options (Lolly Wren\u2019s Curious Science Adventures, Book 2), Lothar J. Musiol writing as Quentin Edmund Darling (Q.E.D.)',
              'Protocol Flamingo (The Invasion Storybooks, Book 1), Lothar J. Musiol writing as George Herbert Fontaine',
              'In Love with Murder, Actually: A Derailed Mystery in Provence, Lothar J. Musiol writing as Romuald Livingston']),
 ('In German', ['Auswandern \u2013 Ab in die USA, Kanada, Australien oder Neuseeland?']),
]
def clone(src, text, bold=False, center=False, space_before=None):
    p = copy.deepcopy(src._p)
    for el in list(p):
        if el.tag not in (qn('w:pPr'),): p.remove(el)
    for el in list(p.iter(qn('w:bookmarkStart'))) + list(p.iter(qn('w:bookmarkEnd'))): el.getparent().remove(el)
    r = OxmlElement('w:r')
    if bold:
        rp = OxmlElement('w:rPr'); b = OxmlElement('w:b'); rp.append(b); r.append(rp)
    t = OxmlElement('w:t'); t.set(qn('xml:space'), 'preserve'); t.text = text; r.append(t); p.append(r)
    ppr = p.find(qn('w:pPr'))
    if ppr is None: ppr = OxmlElement('w:pPr'); p.insert(0, ppr)
    if center:
        jc = OxmlElement('w:jc'); jc.set(qn('w:val'), 'center'); ppr.append(jc)
    if space_before is not None:
        sp = OxmlElement('w:spacing'); sp.set(qn('w:before'), str(space_before)); sp.set(qn('w:after'), '60'); ppr.append(sp)
    return p
head = clone(P[ab], 'ALSO BY LOTHAR J. MUSIOL')
bid = max(int(b.get(qn('w:id'))) for b in d.element.body.iter(qn('w:bookmarkStart'))) + 1
bs = OxmlElement('w:bookmarkStart'); bs.set(qn('w:id'), str(bid)); bs.set(qn('w:name'), 'alsoBy'); be = OxmlElement('w:bookmarkEnd'); be.set(qn('w:id'), str(bid))
head.find(qn('w:pPr')).addnext(bs); head.append(be)
seq = [head]
for g, titles in GROUPS:
    seq.append(clone(last, g, bold=True, center=True, space_before=240))
    for t in titles: seq.append(clone(last, t, center=True, space_before=0))
anchor = last._p
for el in seq: anchor.addnext(el); anchor = el
# TOC line
toc = one(lambda p: p.style.name.startswith('TOC') and p.text == 'ABOUT THE AUTHOR')
new = copy.deepcopy(P[toc]._p); P[toc]._p.addnext(new)
h = new.find(qn('w:hyperlink')); h.set(qn('w:anchor'), 'alsoBy'); ts = list(h.iter(qn('w:t'))); ts[0].text = 'ALSO BY LOTHAR J. MUSIOL'
for t in ts[1:]: t.text = ''
# 4. copyright-page credit note read like a file note
i = one(lambda p: p.text.startswith('Photograph credits are not printed under the figures.'))
ts = list(P[i]._p.iter(qn('w:t'))); ts[0].text = 'Photograph credits are not printed under the figures. They are listed here, so a copy that never sees the store page still names its sources.'
for t in ts[1:]: t.text = ''
d.save(DST); print('saved', len(d.paragraphs))
