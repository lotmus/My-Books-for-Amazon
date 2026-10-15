"""The Permitted Options: 10 Oct 2026 strict rescore fixes.
Bibliography completed and linked (keyword links, no visible URLs), in-text citations linked to entries,
About the Author / Also by made Heading 2 chapters with TOC entries. Usage: python rescore_2026-10-10.py SRC DST"""
import docx, copy, re, sys
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
SRC, DST = sys.argv[1], sys.argv[2]
d = docx.Document(SRC); P = d.paragraphs
RT = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink"
BMID = [max(int(b.get(qn('w:id'))) for b in d.element.body.iter(qn('w:bookmarkStart')))]
def linkrpr(base=None):
    rp = copy.deepcopy(base) if base is not None else OxmlElement('w:rPr')
    for c in list(rp):
        if c.tag in (qn('w:rStyle'), qn('w:color'), qn('w:u')): rp.remove(c)
    c = OxmlElement('w:color'); c.set(qn('w:val'), '0563C1'); rp.append(c)
    u = OxmlElement('w:u'); u.set(qn('w:val'), 'single'); rp.append(u); return rp
def mkrun(text, rpr=None):
    r = OxmlElement('w:r')
    if rpr is not None and len(rpr): r.append(copy.deepcopy(rpr))
    t = OxmlElement('w:t'); t.set(qn('xml:space'), 'preserve'); t.text = text; r.append(t); return r
def mklink(text, target, base=None):
    h = OxmlElement('w:hyperlink')
    if target.startswith('#'): h.set(qn('w:anchor'), target[1:])
    else: h.set(qn('r:id'), d.part.relate_to(target, RT, is_external=True))
    h.set(qn('w:history'), '1'); h.append(mkrun(text, linkrpr(base))); return h
def bookmark(p, name):
    BMID[0] += 1
    a = OxmlElement('w:bookmarkStart'); a.set(qn('w:id'), str(BMID[0])); a.set(qn('w:name'), name)
    b = OxmlElement('w:bookmarkEnd'); b.set(qn('w:id'), str(BMID[0]))
    ppr = p._p.find(qn('w:pPr'))
    (ppr.addnext(a) if ppr is not None else p._p.insert(0, a)); p._p.append(b)
TOK = re.compile(r'\[([^\]]+)\]\(([^)\s]+)\)')
def setp(p, markup):
    for el in list(p._p):
        if el.tag not in (qn('w:pPr'), qn('w:bookmarkStart'), qn('w:bookmarkEnd')): p._p.remove(el)
    ends = [el for el in p._p if el.tag == qn('w:bookmarkEnd')]
    for e in ends: p._p.remove(e)
    pos = 0
    for m in TOK.finditer(markup):
        if m.start() > pos: p._p.append(mkrun(markup[pos:m.start()]))
        p._p.append(mklink(m.group(1), m.group(2))); pos = m.end()
    if pos < len(markup): p._p.append(mkrun(markup[pos:]))
    for e in ends: p._p.append(e)
def wrap(i, phrase, anchor):
    p = P[i]
    for r in p._p.iter(qn('w:r')):
        if r.getparent().tag == qn('w:hyperlink'): continue
        t = r.find(qn('w:t'))
        if t is None or not t.text or phrase not in t.text: continue
        a, b = t.text.split(phrase, 1); rpr = r.find(qn('w:rPr'))
        r1 = mkrun(a, rpr); h = mklink(phrase, '#' + anchor, rpr); r3 = mkrun(b, rpr)
        r.addprevious(r1); r.addprevious(h); r.addprevious(r3); r.getparent().remove(r)
        if not a: r1.getparent().remove(r1)
        if not b: r3.getparent().remove(r3)
        return
    raise AssertionError((i, phrase, p.text[:120]))
def bibidx(prefix):
    h = [i for i in range(2262, 2290) if P[i].text.startswith(prefix)]; assert len(h) == 1, (prefix, h); return h[0]

DOI = 'https://doi.org/'
# ---- existing entries: rebuild those without a title link, add bookmarks
OLD = {
 'Almheiri': ('amps2012', 'Almheiri, A., Marolf, D., Polchinski, J., and Sully, J. (2012). [Black holes: complementarity or firewalls?](https://arxiv.org/abs/1207.3123) Journal of High Energy Physics, 2013, 62; preprint 2012.'),
 'Casimir, H.': ('casimir1948', 'Casimir, H. B. G. (1948). [On the attraction between two perfectly conducting plates](https://dwc.knaw.nl/DL/publications/PU00018547.pdf). Proceedings of the Koninklijke Nederlandse Akademie van Wetenschappen B51, 793–795.'),
 'Coleman': ('coleman1980', None),
 'Einstein, A. Letter': ('einstein1926', 'Einstein, A. Letter to Max Born, 4 December 1926. Published in [The Born–Einstein Letters](https://openlibrary.org/works/OL1214249W).'),
 'Einstein, A., Podolsky': ('epr1935', None), 'Esaki': ('esaki1958', None), 'Fuller': ('fuller1962', None),
 'Gödel': ('godel1931', 'Gödel, K. (1931). [On formally undecidable propositions of Principia Mathematica and related systems](' + DOI + '10.1007/BF01700692). Monatshefte für Mathematik und Physik 38, 173–198.'),
 'Hamming': ('hamming1980', 'Hamming, R. W. (1980). [The unreasonable effectiveness of mathematics](' + DOI + '10.1080/00029890.1980.11994966). The American Mathematical Monthly 87, 81–90.'),
 'Hofstadter': ('hofstadter1979', 'Hofstadter, D. (1979). [Gödel, Escher, Bach: an Eternal Golden Braid](https://openlibrary.org/works/OL716850W). Basic Books.'),
 'Kruskal': ('kruskal1960', 'Kruskal, M. D. (1960). [Maximal extension of Schwarzschild metric](' + DOI + '10.1103/PhysRev.119.1743). Physical Review 119, 1743–1745.'),
 'Lamoreaux': ('lamoreaux1997', None), 'Landauer': ('landauer1961', None),
 'Nash, J. (1951)': ('nash1951', 'Nash, J. (1951). [Non-cooperative games](' + DOI + '10.2307/1969529). Annals of Mathematics 54, 286–295.'),
 'Page, D. N. (1993)': ('page1993', None),
 'Schwarzschild': ('schwarzschild1916', 'Schwarzschild, K. (1916). On the gravitational field of a mass point according to Einstein’s theory. Sitzungsberichte der Königlich Preussischen Akademie der Wissenschaften, 189–196. [English translation](https://arxiv.org/abs/physics/9905030), 1999.'),
 'Szekeres': ('szekeres1960', None), 'Wigner': ('wigner1960', None),
}
for pre, (key, markup) in OLD.items():
    i = bibidx(pre)
    if markup: setp(P[i], markup)
    names = [b.get(qn('w:name')) for b in P[i]._p.iter(qn('w:bookmarkStart'))]
    if 'bib_' + key not in names: bookmark(P[i], 'bib_' + key)
# ---- new entries (alphabetical insertion)
NEW = [
 ('barbour1999', 'Barbour, J. (1999). [The End of Time](https://openlibrary.org/works/OL8003307W). Weidenfeld & Nicolson.'),
 ('bekenstein1972', 'Bekenstein, J. D. (1972). [Black holes and the second law](' + DOI + '10.1007/BF02757029). Lettere al Nuovo Cimento 4, 737–740.'),
 ('bennett1982', 'Bennett, C. H. (1982). [The thermodynamics of computation: a review](' + DOI + '10.1007/BF02084158). International Journal of Theoretical Physics 21, 905–940.'),
 ('berut2012', 'Bérut, A., Arakelyan, A., Petrosyan, A., Ciliberto, S., Dillenschneider, R., and Lutz, E. (2012). [Experimental verification of Landauer’s principle linking information and thermodynamics](' + DOI + '10.1038/nature10872). Nature 483, 187–189.'),
 ('bressi2002', 'Bressi, G., Carugno, G., Onofrio, R., and Ruoso, G. (2002). [Measurement of the Casimir force between parallel metallic surfaces](' + DOI + '10.1103/PhysRevLett.88.041804). Physical Review Letters 88, 041804.'),
 ('cubitt2015', 'Cubitt, T. S., Pérez-García, D., and Wolf, M. M. (2015). [Undecidability of the spectral gap](' + DOI + '10.1038/nature16059). Nature 528, 207–211.'),
 ('er1935', 'Einstein, A., and Rosen, N. (1935). [The particle problem in the general theory of relativity](' + DOI + '10.1103/PhysRev.48.73). Physical Review 48, 73–77.'),
 ('hawking1974', 'Hawking, S. W. (1974). [Black hole explosions?](' + DOI + '10.1038/248030a0) Nature 248, 30–31.'),
 ('ji2020', 'Ji, Z., Natarajan, A., Vidick, T., Wright, J., and Yuen, H. (2020). [MIP* = RE](https://arxiv.org/abs/2001.04383). Preprint.'),
 ('ks1967', 'Kochen, S., and Specker, E. P. (1967). [The problem of hidden variables in quantum mechanics](' + DOI + '10.1512/iumj.1968.17.17004). Journal of Mathematics and Mechanics 17, 59–87.'),
 ('rfc6962', 'Laurie, B., Langley, A., and Kasper, E. (2013). [Certificate Transparency](https://www.rfc-editor.org/rfc/rfc6962). RFC 6962, Internet Engineering Task Force.'),
 ('loftus1974', 'Loftus, E. F., and Palmer, J. C. (1974). [Reconstruction of automobile destruction: an example of the interaction between language and memory](' + DOI + '10.1016/S0022-5371(74)80011-3). Journal of Verbal Learning and Verbal Behavior 13, 585–589.'),
 ('maldacena1997', 'Maldacena, J. (1997). [The large N limit of superconformal field theories and supergravity](https://arxiv.org/abs/hep-th/9711200). Advances in Theoretical and Mathematical Physics 2 (1998), 231–252.'),
 ('mohideen1998', 'Mohideen, U., and Roy, A. (1998). [Precision measurement of the Casimir force from 0.1 to 0.9 μm](' + DOI + '10.1103/PhysRevLett.81.4549). Physical Review Letters 81, 4549–4552.'),
 ('moreva2014', 'Moreva, E., Brida, G., Gramegna, M., Giovannetti, V., Maccone, L., and Genovese, M. (2014). [Time from quantum entanglement: an experimental illustration](' + DOI + '10.1103/PhysRevA.89.052122). Physical Review A 89, 052122.'),
 ('nash1950', 'Nash, J. (1950). [Equilibrium points in n-person games](' + DOI + '10.1073/pnas.36.1.48). Proceedings of the National Academy of Sciences 36, 48–49.'),
 ('pagewootters1983', 'Page, D. N., and Wootters, W. K. (1983). [Evolution without evolution: dynamics described by stationary observables](' + DOI + '10.1103/PhysRevD.27.2885). Physical Review D 27, 2885–2892.'),
 ('pawlowski2009', 'Pawłowski, M., Paterek, T., Kaszlikowski, D., Scarani, V., Winter, A., and Żukowski, M. (2009). [Information causality as a physical principle](' + DOI + '10.1038/nature08400). Nature 461, 1101–1104.'),
 ('pr1994', 'Popescu, S., and Rohrlich, D. (1994). [Quantum nonlocality as an axiom](' + DOI + '10.1007/BF02058098). Foundations of Physics 24, 379–385.'),
 ('robertson1929', 'Robertson, H. P. (1929). [The uncertainty principle](' + DOI + '10.1103/PhysRev.34.163). Physical Review 34, 163–164.'),
 ('stu1993', 'Susskind, L., Thorlacius, L., and Uglum, J. (1993). [The stretched horizon and black hole complementarity](' + DOI + '10.1103/PhysRevD.48.3743). Physical Review D 48, 3743–3761.'),
 ('szilard1929', 'Szilard, L. (1929). [On the decrease of entropy in a thermodynamic system by the intervention of intelligent beings](' + DOI + '10.1007/BF01341281). Zeitschrift für Physik 53, 840–856.'),
 ('weinberg1987', 'Weinberg, S. (1987). [Anthropic bound on the cosmological constant](' + DOI + '10.1103/PhysRevLett.59.2607). Physical Review Letters 59, 2607–2610.'),
]
def sortkey(s): return s.replace('Ö','O').replace('ö','o').replace('é','e').replace('ł','l').lower()
tmpl = P[bibidx('Lamoreaux')]
for key, markup in NEW:
    plain = TOK.sub(lambda m: m.group(1), markup)
    PP = d.paragraphs; s = [k for k, p in enumerate(PP) if p.text == 'Bibliography' and p.style.name == 'Heading 2'][0]
    e = [k for k in range(s, len(PP)) if PP[k].text.startswith('If you enjoyed the story')][0]
    after = None
    for k in range(s + 2, e):
        if sortkey(PP[k].text) < sortkey(plain): after = PP[k]
    new = copy.deepcopy(tmpl._p)
    for el in list(new):
        if el.tag != qn('w:pPr'): new.remove(el)
    after._p.addnext(new)
    np_ = docx.text.paragraph.Paragraph(new, after._parent); setp(np_, markup); bookmark(np_, 'bib_' + key)
P = d.paragraphs
# ---- in-text first-mention links (lectures)
CITES = [
 ('The Casimir effect. Hendrik Casimir', 'Hendrik Casimir', 'bib_casimir1948'),
 ('falling off as the fourth power', 'Lamoreaux in 1997', 'bib_lamoreaux1997'),
 ('What Einstein actually objected to.', 'a 1926 letter to Max Born', 'bib_einstein1926'),
 ('Hidden variables, and where Einstein', 'The 1935 EPR paper', 'bib_epr1935'),
 ('Why it fails: the firewall argument.', 'Almheiri, Marolf, Polchinski and Sully', 'bib_amps2012'),
 ('Complementarity. Susskind', 'Susskind, Thorlacius and Uglum', 'bib_stu1993'),
 (None, 'Bekenstein’s 1972', 'bib_bekenstein1972'),
 (None, 'Maldacena, 1997', 'bib_maldacena1997'),
 ('The first incompleteness theorem.', 'Gödel proved in 1931', 'bib_godel1931'),
 (None, 'Cubitt, Pérez-García and Wolf', 'bib_cubitt2015'),
 ('The Hofstadter strand.', 'Hofstadter, 1979', 'bib_hofstadter1979'),
 ('Chapter Eight brings in Nash', '2020 result', 'bib_ji2020'),
 (None, 'Nash’s 1950 result', 'bib_nash1950'),
 (None, 'Page’s 1993 result', 'bib_page1993'),
 (None, 'Robertson widened this in 1929', 'bib_robertson1929'),
 (None, 'Elizabeth Loftus and John Palmer', 'bib_loftus1974'),
 (None, 'Hamming’s 1980 reply', 'bib_hamming1980'),
 (None, 'Popescu and Rohrlich', 'bib_pr1994'),
 (None, 'The End of Time', 'bib_barbour1999'),
 (None, 'Hawking’s 1974 calculation', 'bib_hawking1974'),
 (None, 'Pawłowski et al., 2009', 'bib_pawlowski2009'),
 (None, 'Page–Wootters', 'bib_pagewootters1983'),
 (None, 'Moreva et al., 2014', 'bib_moreva2014'),
 (None, 'Weinberg’s 1987', 'bib_weinberg1987'),
 (None, 'Szilard’s 1929', 'bib_szilard1929'),
 (None, 'Landauer (1961)', 'bib_landauer1961'),
 (None, 'Bennett (1982)', 'bib_bennett1982'),
 (None, 'Bérut et al., 2012', 'bib_berut2012'),
 (None, 'Kochen and Specker', 'bib_ks1967'),
 (None, 'Leo Esaki', 'bib_esaki1958'),
 (None, 'first exact solution to Einstein’s field equations', 'bib_schwarzschild1916'),
 (None, 'Nathan Rosen’s 1935 paper', 'bib_er1935'),
 (None, 'Fuller and Wheeler’s 1962 analysis', 'bib_fuller1962'),
 (None, 'Mohideen and Roy', 'bib_mohideen1998'),
 (None, 'Bressi and colleagues (2002)', 'bib_bressi2002'),
 (None, 'RFC 6962', 'bib_rfc6962'),
]
lec0 = [k for k, p in enumerate(P) if p.text.startswith('For Curious Readers: The Lectures') and p.style.name == 'Heading 2'][0]
gl = [k for k, p in enumerate(P) if p.text == 'Glossary' and p.style.name == 'Heading 2'][0]
for start, phrase, anc in CITES:
    hits = [k for k in range(lec0, gl) if phrase in P[k].text and (start is None or P[k].text.startswith(start))]
    assert hits, phrase
    wrap(hits[0], phrase, anc)
# Wigner: link the paper's quoted title
k = [k for k in range(lec0, gl) if 'Wigner argued' in P[k].text][0]
wrap(k, 'The Unreasonable Effectiveness of Mathematics in the Natural Sciences', 'bib_wigner1960')
# Kruskal and Szekeres (split)
k = [k for k in range(lec0, gl) if 'Kruskal and Szekeres' in P[k].text][0]
wrap(k, 'Kruskal', 'bib_kruskal1960'); wrap(k, 'Szekeres', 'bib_szekeres1960')
# ---- About the Author / Also by: Heading 2 chapters + TOC entries
toc_bib = [k for k, p in enumerate(P) if k < 60 and p.text == 'Bibliography'][0]
for title, bm in (('Also by Lothar J. Musiol', 'alsoBy'), ('About the Author', 'aboutAuthor')):
    k = [k for k, p in enumerate(P) if k > gl and p.text == title][0]
    P[k].style = d.styles['Heading 2']
    ppr = P[k]._p.find(qn('w:pPr'))
    for c in list(ppr):
        if c.tag in (qn('w:spacing'), qn('w:jc')): ppr.remove(c)
    for r in P[k].runs:
        rp = r._r.find(qn('w:rPr'))
        if rp is not None: r._r.remove(rp)
    bookmark(P[k], bm)
    new = copy.deepcopy(P[toc_bib]._p); P[toc_bib]._p.addnext(new)
    np_ = docx.text.paragraph.Paragraph(new, P[toc_bib]._parent)
    h = new.find(qn('w:hyperlink')); h.set(qn('w:anchor'), bm); h.find('.//' + qn('w:t')).text = title
# order: Bibliography, About the Author, Also by  (inserted in reverse)
d.save(DST); print('saved', len(d.paragraphs))
