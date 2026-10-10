# -*- coding: utf-8 -*-
"""Assemble Science Sparks from the markdown sources.

Usage:
    python build_almanac.py [output.docx]                  Kindle edition (hyperlinked)
    python build_almanac.py --print [output.docx]          6x9 print edition: page numbers and
                                                           index (Word fills them in if installed)
    python build_almanac.py --volume N [output.docx]       one volume of the three-volume set
                                                           (combine with --print if wanted)
Sources: 10_ to 50_*.md are the five sections (chapters numbered continuously through the
whole book), 90_glossary.md is the glossary.
Figures come from make_figures.py (run it first if figures/ is empty).
"""
import copy
import os
import re
import sys
from urllib.parse import quote_plus

import docx
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))          # scripts/
BASE = os.path.dirname(HERE)                                  # book root
CH_DIR = os.path.join(BASE, 'chapters')                       # the .md text
FIG_DIR = os.path.join(BASE, 'figures')
DEFAULT_OUT = os.path.join(HERE, "build", "almanac_build.docx")

TITLE = "Science Sparks"
SUBTITLE = "Physics, Life, and Mathematics — The Same Few Rules, Told in Highlights"
AUTHOR = "Lothar J. Musiol"

BODY_FONT = "Georgia"
HEADLINE_FONT = "Georgia"
BLACK = RGBColor(0, 0, 0)
GREY = RGBColor(0x55, 0x55, 0x55)
HEADING_BLUE = RGBColor(0x1F, 0x4E, 0x9F)   # one medium blue for every heading level
TOC_LEVELS = 2          # main contents: sections and chapters
BOOKMARK_LEVELS = 4     # every heading down to level 4 can be a link target

PRINT = False           # set from the command line

# Five sections. Each file holds numbered chapters ("## N. Title"); the numbers run
# continuously through the whole book, so "Chapter N" in the text is unambiguous.
PARTS = [
    ("Mathematics",
     "The language the rest of the book is written in: from counting to matrices, one idea at a "
     "time, with the engineer's working tools for logarithms, sinusoids and complex numbers set "
     "where they belong. Most chapters end with a worked example.",
     ["10_mathematics.md"]),
    ("Physics and the Cosmos",
     "Motion, energy, time, light, gravity, and what all of it does when you leave it running "
     "for thirteen billion years; then the ageing body and the road to the Moon and Mars, "
     "held to the same standard of evidence.",
     ["20_physics_and_cosmos.md"]),
    ("Quantum Physics",
     "The rules that run underneath everything else: first the working parts, then the same "
     "ground taken deeper, then phase and potentials, and finally the full machinery of "
     "quantum electrodynamics.",
     ["30_quantum_physics.md"]),
    ("Biology",
     'Physics and chemistry that copy themselves: the same rules of energy, information, feedback and probability, now running on a wet planet for four billion years, and the three-billion-letter text in which the results were written.',
     ["40_biology.md"]),
    ("Science in the Novels",
     "Where the plot stops so someone can explain relativity, and the explanation turns out "
     "to be the best part.",
     ["50_science_in_the_novels.md"]),
]

VOLUMES = {
    1: ("Volume 1: Mathematics", [0]),
    2: ("Volume 2: Physics, the Cosmos, and the Quantum", [1, 2, 4]),
    3: ("Volume 3: Biology", [3]),
}

GLOSSARY_FILE = "90_glossary.md"
GLOSS_REF_RE = re.compile(r'\bChapter (\d+)\b')

# Also by Lothar J. Musiol: (group, [(title, store search or None)]). The canonical list, rebuilt on
# 2 Oct 2026 from the title pages of the masters in the repository (one entry per real book; same
# list in every book's Also-by page, see notes/ALSO_BY - canonical list.md at the repo root).
# A search string of None means "title + author name".
# Rebuilt 10 Oct 2026 from notes/ALSO_BY - canonical list.md. The book itself is left out; Kevin Drew
# Peters books never appear on a Lothar J. Musiol page.
ALSO_BY = [
    ('Physics, Actually', [
        ('Physics, Actually, Volume 1: Motion, Forces, Time, and Relativity', None),
        ('Physics, Actually, Volume 2: Gravity, Cosmology, and the Limits of Spacetime', None),
        ('Physics, Actually, Volume 3: The Standard Model, Chaos, and the Edge of Knowledge', None),
        ('Life, Actually: From the First Cell to the Edited Genome and the Search for Life Elsewhere', None),
    ]),
    ('More Actually Books', [
        ('Gravitation, Actually: A Short Book on Gravity', None),
        ('The Universe, Actually: What Humanity Currently Understands About Reality', None),
    ]),
    ('Math, Actually', [
        ('Math, Actually, Volume 1: From Arithmetic to Calculus', None),
        ('Math, Actually, Volume 2: From Multivariable Calculus to Set Theory & Logic', None),
        ('Math, Actually, Volume 3: From Differential Equations to Abstract Algebra', None),
        ('Math, Actually, Volume 4: From Category Theory to the Frontier', None),
    ]),
    ('Quantum, Actually', [
        ('Quantum, Actually, Volume 1: Questions and Answers from the Double Slit to the Superconducting Wire', None),
        ('Quantum, Actually, Volume 2: A QED Course', None),
    ]),
    ('Look First', [
        ('Look First, Volume 1: The Universe Has No Now', None),
        ('Look First, Volume 2: A Trip Is Not a New Life', None),
    ]),
    ('Electrical Engineering Series', [
        ('Foundations of Electronics (Book 1)', 'Foundations of Electronics'),
        ('Circuits, Components, and Control (Book 2)', 'Circuits, Components, and Control'),
        ('Semiconductor Physics and Devices (Book 3)', 'Semiconductor Physics and Devices'),
        ('RF, Microwave, and Transceivers (Book 4)', 'RF, Microwave, and Transceivers'),
        ('Communications, Wireless, and SDR (Book 5)', 'Communications, Wireless, and SDR'),
        ('Power and Energy (Book 6)', 'Power and Energy'),
        ('Packaging, Layout, EMC, and Test (Book 7)', 'Packaging, Layout, EMC, and Test'),
    ]),
    ('History', [
        ("The Dolphins' View of History", None),
    ]),
    ('Fiction', [
        ("The Murder That Hadn't Happened Yet (The Relativistic Investigation Bureau, Book 1), Lothar J. Musiol writing as Spezala Genara Relavi (S.G.R.)", "The Murder That Hadn't Happened Yet"),
        ('The Warning That Was Sent Too Late (The Relativistic Investigation Bureau, Book 2), Lothar J. Musiol writing as Spezala Genara Relavi (S.G.R.)', 'The Warning That Was Sent Too Late'),
        ("Schrödinger's Paperwork (Lolly Wren's Curious Science Adventures, Book 1), Lothar J. Musiol writing as Quentin Edmund Darling (Q.E.D.)", "Schrödinger's Paperwork"),
        ("The Permitted Options (Lolly Wren's Curious Science Adventures, Book 2), Lothar J. Musiol writing as Quentin Edmund Darling (Q.E.D.)", 'The Permitted Options Lolly Wren'),
        ('Protocol Flamingo (The Invasion Storybooks, Book 1), Lothar J. Musiol writing as George Herbert Fontaine', 'Protocol Flamingo George Herbert Fontaine'),
        ('In Love with Murder, Actually: A Comic Mystery, Lothar J. Musiol writing as Romuald Livingston', 'In Love with Murder, Actually Romuald Livingston'),
    ]),
    ('In German', [
        ('Auswandern – Ab in die USA, Kanada, Australien oder Neuseeland?', None),
    ]),
]

# chapter number -> (png name, caption). Placed after the chapter's first paragraph.
FIGURES = {
    30: ("entropy_coins", "Entropy is counting. Of all the ways 100 coins can land, almost every one "
                         "is near half heads; the tidy extremes are a rounding error."),
    35: ("light_cone", "A light cone. Only events inside your future cone can be affected by what "
                      "you do now; 'elsewhere' cannot be reached by any signal."),
    68: ("hubble", "The farther the galaxy, the faster it recedes: the signature of space itself "
                   "stretching. Schematic data."),
    164: ("feynman", "The simplest Feynman diagram: two electrons repel by exchanging a photon. "
                    "A bookkeeping device for a calculation, not a photograph."),
    56: ("bands", "Energy bands. The size of the gap decides between insulator, semiconductor "
                  "and conductor; silicon's gap is about 1.1 electron-volts."),
    67: ("cosmic_timeline", "The universe's calendar, on a logarithmic scale. Most of the "
                            "interesting physics happened before anyone was around to be bored by it."),
    145: ("double_slit", "Interference. With no record of which opening each particle used, hits "
                         "pile up in stripes; keep that record and the stripes go (Chapter 149)."),
    147: ("stern_gerlach", "Stern and Gerlach, 1922: silver atoms through a magnet land in two "
                           "spots, not a smear. Angular momentum comes in bins."),
    148: ("fourier_budget", "The shared budget. A wave that is narrow in time is wide in pitch, and "
                            "the reverse; position and momentum keep the same books."),
    151: ("bell_ceiling", "Bell's ceiling. Local prewritten answers cannot score above 2; quantum "
                          "mechanics reaches 2.83, and experiments agree with quantum mechanics."),
    152: ("decoherence", "Decoherence: the more the environment learns, the faster the stripes "
                         "fade. Schematic."),
    155: ("zeno", "The quantum Zeno effect. Check a slowly changing system often enough and it "
                  "almost never gets around to changing."),
    156: ("hawking", "Hawking temperature falls as mass rises. A black hole of one solar mass is "
                     "far colder than the microwave sky, so today it gains more than it loses."),
    259: ("drift_selection", "Luck and selection. In a small population a neutral variant wanders "
                             "until it is lost or fixed; a modest advantage in a large one wins "
                             "almost every time. Simulated."),
    293: ("codons", "The genetic code is redundant: 64 three-letter words, 20 amino acids and a "
                    "stop signal. Some meanings have six spellings, two have only one."),
    7: ("unit_circle", "Sine and cosine are the two shadows of a point going round a circle "
                         "of radius 1."),
    16: ("tangent", "The derivative is the slope of the tangent line: for y = x squared at x = 1, "
                     "the slope is 2."),
    17: ("area", "The integral is accumulated area: under y = x squared from 0 to 3, the area "
                  "is exactly 9."),
    22: ("square_wave", "Fourier's claim, tested: enough sine waves add up to a square wave. The "
                         "overshoot at the corners never fully goes away."),
    334: ("simultaneity", "Two events that are simultaneous for you are not simultaneous for a "
                          "moving observer: for them, B happens first. Neither of you is wrong."),
}

TOKEN_RE = re.compile(r'(\*\*.+?\*\*|\*[^*\n]+?\*)')
STATUS_RE = re.compile(r'^\*Status: (.+?)\*$')
LEVEL_RE = re.compile(r'^\*Level: (.+?)\*$')
BOX_RE = re.compile(r'^> (.+)$')
LABEL_DASH_RE = re.compile(r'^(Prologue|Epilogue|Appendix(?: [A-Z0-9]+)?) — (.+)$')
PART_ROMAN_RE = re.compile(r'^Part (I|II|III|IV|V|VI|VII)\b')
NUM_HEADING_RE = re.compile(r'^(\d+)\.\s')
LABEL_NUM_RE = re.compile(r'^(Chapter|Storey|Lesson|Lecture|Appendix)\s+(\d+)\b')
LABEL_RANGE_RE = re.compile(r'^(?:Chapters|Lessons|Lectures)\s+(\d+)\s*[–-]\s*(\d+)\b')
CHAPTER_NUM_RE = re.compile(r'\bChapter \d+\b')
TIMES_RE = re.compile(r'(?<=[\d)]) x (?=[\d(])')
TITLE_LINK_STOPLIST = {'Prologue', 'Epilogue', 'Foreword', 'Afterword', 'Contents',
                       'About the Author', 'Appendix', 'Glossary', 'Index', 'Also by the Author', 'Also by Lothar J. Musiol',
                       'Physics and the Cosmos', 'Quantum Physics', 'Science in the Novels'}

HEADINGS = []           # (level, text, bookmark)
GLOBAL_TITLE_LINKS = {}
BODY_PARAGRAPHS = []    # (paragraph, segment or 'GLOSSARY')
GLOBAL_SEGMENT = {}     # ('num', chapter number) -> bookmark, for the whole book
INDEX_TERMS = []
_bookmark_counter = [0]
_figure_counter = [0]


def normalize_heading(text):
    m = LABEL_DASH_RE.match(text)
    if m and ':' not in m.group(2):
        return '%s: %s' % (m.group(1), m.group(2))
    return text


def plain(text):
    return text.replace('**', '').replace('*', '')


# ---------- styles ----------
def set_style_color(style, color):
    """Set an explicit RGB colour and drop any theme colour, which Word would prefer."""
    style.font.color.rgb = color
    c = style.element.get_or_add_rPr().find(qn('w:color'))
    for a in ('w:themeColor', 'w:themeShade', 'w:themeTint'):
        if c is not None and c.get(qn(a)) is not None:
            del c.attrib[qn(a)]


def set_style_font(style, name, size=None, bold=None, italic=None, color=None):
    style.font.name = name
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.append(rfonts)
    for a in ('w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs'):
        rfonts.set(qn(a), name)
    for a in ('w:asciiTheme', 'w:hAnsiTheme', 'w:eastAsiaTheme', 'w:cstheme'):
        if rfonts.get(qn(a)) is not None:
            del rfonts.attrib[qn(a)]
    if size is not None:
        style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    if italic is not None:
        style.font.italic = italic
    if color is not None:
        style.font.color.rgb = color


def page_geometry(sec):
    if PRINT:
        sec.page_width, sec.page_height = Inches(6), Inches(9)
        sec.left_margin, sec.right_margin = Inches(0.875), Inches(0.5)
        sec.top_margin, sec.bottom_margin = Inches(0.7), Inches(0.7)
    else:
        sec.left_margin = sec.right_margin = Inches(1.1)
        sec.top_margin = sec.bottom_margin = Inches(1.0)
    sec.header_distance = sec.footer_distance = Inches(0.35)


def setup_styles(doc):
    normal = doc.styles['Normal']
    set_style_font(normal, BODY_FONT, 10.5 if PRINT else 11)
    pf = normal.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.first_line_indent = Inches(0.25)
    pf.space_before = pf.space_after = Pt(0)
    pf.line_spacing = 1.25
    pf.widow_control = True

    # Every heading level (section, chapter, subheading) is centered and the same medium blue.
    spec = {  # level: (size, bold, italic, align, space_before, space_after)
        1: (26, True, False, WD_ALIGN_PARAGRAPH.CENTER, 150, 18),
        2: (18, True, False, WD_ALIGN_PARAGRAPH.CENTER, 48, 14),
        3: (14, True, False, WD_ALIGN_PARAGRAPH.CENTER, 22, 6),
        4: (11.5, True, True, WD_ALIGN_PARAGRAPH.CENTER, 14, 4),
    }
    for lvl, (size, bold, italic, align, before, after) in spec.items():
        h = doc.styles['Heading %d' % lvl]
        set_style_font(h, HEADLINE_FONT, size, bold=bold, italic=italic)
        set_style_color(h, HEADING_BLUE)
        h.paragraph_format.alignment = align
        h.paragraph_format.first_line_indent = Inches(0)
        h.paragraph_format.page_break_before = (lvl == 2)   # every chapter opens a page
        h.paragraph_format.space_before = Pt(before)
        h.paragraph_format.space_after = Pt(after)
        h.paragraph_format.keep_with_next = True

    for lvl, (indent, size, bold, before) in enumerate(
            [(0.0, 11, True, 9), (0.25, 10.5, False, 1), (0.5, 10, False, 0)], start=1):
        st = doc.styles.add_style('toc %d' % lvl, WD_STYLE_TYPE.PARAGRAPH)
        st.base_style = normal
        set_style_font(st, BODY_FONT, size, bold=bold, italic=False)
        st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        st.paragraph_format.first_line_indent = Inches(0)
        st.paragraph_format.left_indent = Inches(indent)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(0)
        st.paragraph_format.line_spacing = 1.1

    page_geometry(doc.sections[0])
    if PRINT:
        settings = doc.settings.element
        mm = OxmlElement('w:mirrorMargins')
        settings.insert(0, mm)


# ---------- fields, links, sections ----------
def add_field(paragraph, instr, placeholder=''):
    def fld(t):
        r = OxmlElement('w:r')
        c = OxmlElement('w:fldChar')
        c.set(qn('w:fldCharType'), t)
        r.append(c)
        return r
    p = paragraph._p
    p.append(fld('begin'))
    r = OxmlElement('w:r')
    it = OxmlElement('w:instrText')
    it.set(qn('xml:space'), 'preserve')
    it.text = ' %s ' % instr
    r.append(it)
    p.append(r)
    p.append(fld('separate'))
    r = OxmlElement('w:r')
    t = OxmlElement('w:t')
    t.text = placeholder
    r.append(t)
    p.append(r)
    p.append(fld('end'))


def field_runs(instr, base_rpr=None, placeholder='0'):
    out = []
    for kind in ('begin', 'instr', 'separate', 'text', 'end'):
        r = OxmlElement('w:r')
        if base_rpr is not None:
            r.append(copy.deepcopy(base_rpr))
        if kind == 'instr':
            it = OxmlElement('w:instrText')
            it.set(qn('xml:space'), 'preserve')
            it.text = ' %s ' % instr
            r.append(it)
        elif kind == 'text':
            t = OxmlElement('w:t')
            t.text = placeholder
            r.append(t)
        else:
            c = OxmlElement('w:fldChar')
            c.set(qn('w:fldCharType'), kind)
            r.append(c)
        out.append(r)
    return out


def add_external_link(paragraph, text, url):
    rid = paragraph.part.relate_to(
        url, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',
        is_external=True)
    h = OxmlElement('w:hyperlink')
    h.set(qn('r:id'), rid)
    r = OxmlElement('w:r')
    rpr = OxmlElement('w:rPr')
    if not PRINT:
        c = OxmlElement('w:color')
        c.set(qn('w:val'), '0000FF')
        u = OxmlElement('w:u')
        u.set(qn('w:val'), 'single')
        rpr.append(c)
        rpr.append(u)
    i = OxmlElement('w:i')
    rpr.append(i)
    r.append(rpr)
    t = OxmlElement('w:t')
    t.set(qn('xml:space'), 'preserve')
    t.text = text
    r.append(t)
    h.append(r)
    paragraph._p.append(h)


def store_url(title, search=None):
    q = search if search else ('%s %s' % (title, AUTHOR)).strip()
    return 'https://www.amazon.com/s?k=%s&i=stripbooks' % quote_plus(q)


def new_section(doc, header_text=''):
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    page_geometry(sec)
    for hf in (sec.header, sec.footer):
        hf.is_linked_to_previous = False
        for p in hf.paragraphs:
            for r in list(p.runs):
                r._r.getparent().remove(r._r)
    hp = sec.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if header_text:
        r = hp.add_run(header_text)
        r.italic = True
        r.font.size = Pt(8.5)
        r.font.color.rgb = GREY
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if header_text:
        add_field(fp, 'PAGE', '1')
    return sec


def add_runs_from_markdown(paragraph, text):
    for part in TOKEN_RE.split(text):
        if not part:
            continue
        if part.startswith('**') and part.endswith('**') and len(part) > 4:
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith('*') and part.endswith('*') and len(part) > 2:
            paragraph.add_run(part[1:-1]).italic = True
        else:
            paragraph.add_run(part)


def flush_left(p, space_before=0, space_after=0, italic=False, align=None, text=None, size=None):
    pf = p.paragraph_format
    pf.first_line_indent = Inches(0)
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    if align is not None:
        p.alignment = align
    if text is not None:
        r = p.add_run(text)
        r.italic = italic
        if size:
            r.font.size = Pt(size)
    return p


def register_heading(p, level, text):
    if level > BOOKMARK_LEVELS:
        return None
    _bookmark_counter[0] += 1
    bid = _bookmark_counter[0]
    name = "_Toc%06d" % (900000 + bid)
    start = OxmlElement('w:bookmarkStart')
    start.set(qn('w:id'), str(bid))
    start.set(qn('w:name'), name)
    end = OxmlElement('w:bookmarkEnd')
    end.set(qn('w:id'), str(bid))
    ppr = p._p.find(qn('w:pPr'))
    if ppr is not None:
        ppr.addnext(start)
    else:
        p._p.insert(0, start)
    p._p.append(end)
    clean = plain(text)
    HEADINGS.append((level, clean, name))
    if (clean not in TITLE_LINK_STOPLIST and len(clean.split()) >= 3
            and not NUM_HEADING_RE.match(clean) and not LABEL_NUM_RE.match(clean)
            and not LABEL_RANGE_RE.match(clean) and level <= 2):
        GLOBAL_TITLE_LINKS.setdefault(clean, name)
    m = PART_ROMAN_RE.match(clean)
    if m:
        GLOBAL_TITLE_LINKS.setdefault('Part %s' % m.group(1), name)
    return name


def heading(doc, level, text):
    p = doc.add_paragraph(style='Heading %d' % level)
    add_runs_from_markdown(p, text)
    return register_heading(p, level, text)


def body(doc, text, segment=None, first=False):
    p = doc.add_paragraph()
    add_runs_from_markdown(p, text)
    if first:
        p.paragraph_format.first_line_indent = Inches(0)
    BODY_PARAGRAPHS.append((p, segment if segment is not None else {}))
    return p


def link_paragraph(doc_or_before, level, text, anchor, before=True):
    if before:
        p = doc_or_before.insert_paragraph_before(style='toc %d' % level)
    else:
        p = doc_or_before.add_paragraph(style='toc %d' % level)
    h = OxmlElement('w:hyperlink')
    h.set(qn('w:anchor'), anchor)
    h.set(qn('w:history'), '1')
    r = OxmlElement('w:r')
    t = OxmlElement('w:t')
    t.set(qn('xml:space'), 'preserve')
    t.text = text
    r.append(t)
    h.append(r)
    p._p.append(h)
    return p


def add_figure(doc, png, caption):
    path = os.path.join(FIG_DIR, png + '.png')
    if not os.path.exists(path):
        return
    _figure_counter[0] += 1
    p = doc.add_paragraph()
    flush_left(p, space_before=10, space_after=2, align=WD_ALIGN_PARAGRAPH.CENTER)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(path, width=Inches(4.2 if PRINT else 4.8))
    c = doc.add_paragraph()
    flush_left(c, space_after=10, align=WD_ALIGN_PARAGRAPH.CENTER)
    r = c.add_run('Figure %d. ' % _figure_counter[0])
    r.bold = True
    r.font.size = Pt(9)
    r = c.add_run(caption)
    r.italic = True
    r.font.size = Pt(9)


def status_line(doc, label, tag='STATUS  '):
    p = doc.add_paragraph()
    flush_left(p, space_before=0, space_after=6)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(tag)
    r.bold = True
    r.font.size = Pt(8)
    r.font.color.rgb = GREY
    r = p.add_run(label)
    r.italic = True
    r.font.size = Pt(9)
    r.font.color.rgb = GREY
    return p


def box_paragraph(doc, text, segment):
    """A boxed aside ('> **Label.** text'): indented, ruled on the left, lightly shaded."""
    p = doc.add_paragraph()
    add_runs_from_markdown(p, text)
    pf = p.paragraph_format
    pf.first_line_indent = Inches(0)
    pf.left_indent = Inches(0.3)
    pf.right_indent = Inches(0.2)
    pf.space_before = Pt(8)
    pf.space_after = Pt(8)
    ppr = p._p.get_or_add_pPr()
    bdr = OxmlElement('w:pBdr')
    left = OxmlElement('w:left')
    for k, v in (('w:val', 'single'), ('w:sz', '18'), ('w:space', '8'), ('w:color', '1F4E9F')):
        left.set(qn(k), v)
    bdr.append(left)
    shd = OxmlElement('w:shd')
    for k, v in (('w:val', 'clear'), ('w:color', 'auto'), ('w:fill', 'EEF2F8')):
        shd.set(qn(k), v)
    # keep schema order: pBdr and shd come before tabs, spacing, ind, jc ...
    after = ('w:tabs', 'w:suppressAutoHyphens', 'w:kinsoku', 'w:wordWrap', 'w:overflowPunct',
             'w:topLinePunct', 'w:autoSpaceDE', 'w:autoSpaceDN', 'w:bidi', 'w:adjustRightInd',
             'w:snapToGrid', 'w:spacing', 'w:ind', 'w:contextualSpacing', 'w:mirrorIndents',
             'w:suppressOverlap', 'w:jc', 'w:textDirection', 'w:textAlignment',
             'w:textboxTightWrap', 'w:outlineLvl', 'w:divId', 'w:cnfStyle', 'w:rPr',
             'w:sectPr', 'w:pPrChange')
    ppr.insert_element_before(bdr, *after)
    ppr.insert_element_before(shd, *after)
    BODY_PARAGRAPHS.append((p, segment))
    return p


# ---------- markdown rendering ----------
def render_markdown(doc, path):
    """Render one section file: '## N. Title' chapters, '###' subheadings."""
    segment = GLOBAL_SEGMENT
    pending_fig = None
    first_para = False
    with open(path, 'r', encoding='utf-8') as f:
        lines = f.read().splitlines()
    for raw in lines:
        line = raw.rstrip()
        if not line.strip():
            continue
        hashes = len(line) - len(line.lstrip('#'))
        if hashes and line[hashes:hashes + 1] == ' ':
            text = normalize_heading(line[hashes:].strip())
            level = min(max(hashes, 2), 4)
            bookmark = heading(doc, level, text)
            m = NUM_HEADING_RE.match(text)
            pending_fig = None
            if bookmark and m and level == 2:
                segment[('num', int(m.group(1)))] = bookmark
                pending_fig = FIGURES.get(int(m.group(1)))
            first_para = True
            continue
        sm = STATUS_RE.match(line.strip())
        if sm:
            status_line(doc, sm.group(1))
            continue
        lm = LEVEL_RE.match(line.strip())
        if lm:
            BODY_PARAGRAPHS.append((status_line(doc, lm.group(1), 'LEVEL  '), segment))
            continue
        bm = BOX_RE.match(line.strip())
        if bm:
            box_paragraph(doc, bm.group(1), segment)
            continue
        if line.strip() in ('---', '***', '* * *'):
            p = doc.add_paragraph()
            flush_left(p, space_before=8, space_after=8, align=WD_ALIGN_PARAGRAPH.CENTER, text='*  *  *')
            continue
        if line.lstrip().startswith(('- ', '* ')):
            p = doc.add_paragraph(style='List Bullet')
            add_runs_from_markdown(p, line.lstrip()[2:].strip())
            p.paragraph_format.first_line_indent = Inches(0)
            BODY_PARAGRAPHS.append((p, segment))
            continue
        if line.startswith('**Worked example.**'):
            line = TIMES_RE.sub(' \u00d7 ', line)
            p = body(doc, line, segment, first=True)
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(6)
            continue
        body(doc, line, segment, first=first_para)
        first_para = False
        if pending_fig:
            add_figure(doc, *pending_fig)
            pending_fig = None


def render_glossary(doc):
    path = os.path.join(CH_DIR, GLOSSARY_FILE)
    if not os.path.exists(path):
        return
    new_section(doc, 'Glossary')
    heading(doc, 1, 'Glossary')
    with open(path, 'r', encoding='utf-8') as f:
        lines = f.read().splitlines()
    for raw in lines:
        line = raw.strip()
        if not line or line.startswith('# '):
            continue
        if line.startswith('## '):
            p = doc.add_paragraph(style='Heading 4')
            p.add_run(line[3:].strip())
            continue
        if line.startswith('**'):
            m = re.match(r'\*\*(.+?)\.\*\*', line)
            if m:
                INDEX_TERMS.append(m.group(1))
            p = doc.add_paragraph()
            add_runs_from_markdown(p, line)
            p.paragraph_format.first_line_indent = Inches(-0.2)
            p.paragraph_format.left_indent = Inches(0.2)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
            BODY_PARAGRAPHS.append((p, 'GLOSSARY'))
        else:
            if PRINT and line.startswith('*Each entry'):
                line = ('*Each entry names the place where the idea is properly taught, '
                        'with the page where that chapter begins.*')
            p = doc.add_paragraph()
            add_runs_from_markdown(p, line)
            flush_left(p, space_after=8)


# ---------- linking ----------
def _text_run(text, base_rpr):
    r = OxmlElement('w:r')
    if base_rpr is not None:
        r.append(copy.deepcopy(base_rpr))
    t = OxmlElement('w:t')
    t.set(qn('xml:space'), 'preserve')
    t.text = text
    r.append(t)
    return r


def _link_run(text, base_rpr, bookmark):
    h = OxmlElement('w:hyperlink')
    h.set(qn('w:anchor'), bookmark)
    h.set(qn('w:history'), '1')
    r = OxmlElement('w:r')
    rpr = OxmlElement('w:rPr')
    if base_rpr is not None:
        for child in base_rpr:
            rpr.append(copy.deepcopy(child))
    if not PRINT:
        c = OxmlElement('w:color')
        c.set(qn('w:val'), '0000FF')
        u = OxmlElement('w:u')
        u.set(qn('w:val'), 'single')
        rpr.append(c)
        rpr.append(u)
    r.append(rpr)
    t = OxmlElement('w:t')
    t.set(qn('xml:space'), 'preserve')
    t.text = text
    r.append(t)
    h.append(r)
    return h


def linkify_document():
    title_keys = sorted(GLOBAL_TITLE_LINKS.keys(), key=len, reverse=True)
    title_pattern = '|'.join(re.escape(k) for k in title_keys)
    body_re = re.compile('(?:' + '|'.join(p for p in (title_pattern, CHAPTER_NUM_RE.pattern) if p) + ')')
    count = 0
    for paragraph, segment in BODY_PARAGRAPHS:
        scan_re = GLOSS_REF_RE if segment == 'GLOSSARY' else body_re
        for run in list(paragraph.runs):
            text = run.text
            if not text or not scan_re.search(text):
                continue
            pieces, pos, found = [], 0, False
            for m in scan_re.finditer(text):
                token = m.group(0)
                numbered = bool(CHAPTER_NUM_RE.fullmatch(token))
                bm = GLOBAL_TITLE_LINKS.get(token)
                if bm is None and numbered:
                    bm = GLOBAL_SEGMENT.get(('num', int(token.rsplit(' ', 1)[1])))
                if bm is None:
                    continue
                s_, e = m.span()
                if s_ > pos:
                    pieces.append((text[pos:s_], None, False))
                pieces.append((token, bm, numbered))
                pos, found = e, True
            if not found:
                continue
            if pos < len(text):
                pieces.append((text[pos:], None, False))
            r = run._r
            parent = r.getparent()
            idx = list(parent).index(r)
            base_rpr = r.find(qn('w:rPr'))
            parent.remove(r)
            for piece, bm, numbered in pieces:
                if not piece:
                    continue
                nodes = [_link_run(piece, base_rpr, bm)] if bm else [_text_run(piece, base_rpr)]
                if bm and PRINT and numbered:
                    nodes.append(_text_run(' (p. ', base_rpr))
                    nodes.extend(field_runs('PAGEREF %s \\h' % bm, base_rpr))
                    nodes.append(_text_run(')', base_rpr))
                for node in nodes:
                    parent.insert(idx, node)
                    idx += 1
                if bm:
                    count += 1
    return count


def mark_index_terms():
    """Print edition: one XE entry per glossary term per book, at its first appearance."""
    if not INDEX_TERMS:
        return 0
    patterns = [(t, re.compile(r'\b%s\b' % re.escape(t), re.I)) for t in INDEX_TERMS]
    seen = set()
    n = 0
    for paragraph, segment in BODY_PARAGRAPHS:
        if segment == 'GLOSSARY' or not segment:
            continue
        text = paragraph.text
        for term, pat in patterns:
            key = (term, id(segment))
            if key in seen or not pat.search(text):
                continue
            seen.add(key)
            for node in field_runs('XE "%s"' % term.replace('"', ''), None, ''):
                paragraph._p.append(node)
            n += 1
    return n


# ---------- front and back matter ----------
FOREWORD = [
    "Energy is never created and never lost. It only changes form, like money moving between "
    "accounts, except that no central bank can print more of it and no decree can change the "
    "rule. The universe keeps the books, and this book is about what those books record.",
    "It makes one promise: a small number of rules keep turning up everywhere. Energy is moved "
    "around and written down, never made. Statistics governs crowds without saying anything about "
    "individuals. What you measure depends on how you are moving, while a few quantities stay the "
    "same for everyone. Copies are never exact, and the errors add up to everything that ever "
    "lived. Waves add before they are counted, and knowing has a price.",
    "The book runs in five sections, and the order is deliberate. Mathematics comes first, because "
    "it is the language everything else is written in. Physics and the Cosmos uses that language "
    "on motion, energy, time, gravity and the history of the universe, and then holds claims "
    "about living longer and settling the Moon and Mars to the same standard. Quantum Physics "
    "goes underneath, to the rules that run atoms, light and chips. Biology shows what those "
    "rules produced on one wet planet in four billion years. Science in the Novels closes the "
    "book with relativity told as a detective story. The Afterword collects the rules and points "
    "back to where each one turned up.",
    "The chapters are numbered straight through the book, so a reference such as Chapter "
    "148 means the same thing wherever you meet it, and on a Kindle you can tap it. You can "
    "read the sections in any order, but three runs reward reading front to back: Mathematics, "
    "the opening course of Quantum Physics (Chapters 144 to 156), and the closing course "
    "on quantum electrodynamics (Chapters 223 to 252).",
    "This edition is updated through October 2026; claims that are likely to change are marked "
    "\"as of 2026\".",
    "Every chapter carries a status line under its title. It tells you how sure science is about "
    "the chapter's main claims. Mathematics is proved rather than measured, so its chapters are "
    "all Settled. There are four labels:",
]
STATUS_LEGEND = [
    ("Settled", "tested many times; textbook material."),
    ("Strange but solid", "just as well established, and still counterintuitive."),
    ("Serious but unconfirmed", "mainstream work that has not yet been decided by evidence."),
    ("Speculative", "interesting, labeled, and far ahead of the evidence."),
]
FOREWORD_CLOSE = (
    "When a chapter mixes levels, the status line says which part is which. The glossary at the "
    "back defines the terms that trip people up and sends you to the chapter that teaches each one."
)

WHO_FOR = [
    "**This book is for curious adults and technically minded readers who want a rigorous but "
    "readable map of mathematics, physics, life and the questions science has not yet answered.**",
    "It is not a textbook, not a proof course and not an exhaustive encyclopedia. There are no "
    "problem sets, the mathematics is explained rather than proved, and whole fields are left out "
    "on purpose. Think of it as an atlas: 354 short, numbered chapters, each of which can be read "
    "on its own, with cross-references that tell you where an idea came from and where it goes next.",
    "Each chapter carries two labels under its title. The status line says how sure science is "
    "about the chapter's claims; the level line says how much of the book you need to have read "
    "first. Both are explained in the Introduction. Every chapter ends with a one-sentence "
    "summary, and the Reader Routes after the Introduction suggest paths through the book for "
    "five different interests.",
    "Where a topic deserves more room than an atlas can give it, a short note points to the "
    "deeper books in the same family: Physics, Actually (three volumes), Life, Actually, "
    "Math, Actually (four volumes) and Quantum, Actually (two volumes).",
    "*Updated through October 2026. Statements that are likely to change, such as mission "
    "schedules, drug approvals and the status of open measurements, are marked \"as of 2026\".*",
]
LEVEL_INTRO = (
    "Under the status line sits a second label, the level. It tells you how much of the book a "
    "chapter assumes. There are five levels:"
)
LEVEL_LEGEND = [
    ("Start here", "no background needed; a good place to enter the book or a section."),
    ("Foundation", "introduces ideas that later chapters rely on, and assumes nothing earlier."),
    ("Builds on earlier chapters", "readable once you know the chapters it names; on a Kindle "
     "you can tap them."),
    ("Technical", "uses equations or formal machinery, and names the chapters to read first."),
    ("Optional deep dive", "rewarding, but nothing later depends on it; skip it freely."),
]
LEVEL_CLOSE = (
    "The two labels answer different questions. A chapter can be Settled and Technical (the "
    "Dirac equation, Chapter 241) or Start here and partly Speculative (Chapter 57). A few chapters also carry short "
    "boxed notes: *Why this matters*, *Engineer's view*, *Common wrong picture*, and in the "
    "hardest chapters *Skip if you only want the idea*, which gives the main point in a few "
    "lines so you can move on. Each part, and each large block within a part, ends with a short "
    "Rules Check that shows where the book's recurring rules came up: conservation, symmetry, "
    "probability, feedback and emergence, evidence beating intuition, and models that are useful "
    "but incomplete. Start wherever your curiosity is loudest."
)
ROUTES_INTRO = (
    "You can read straight through, and the order is deliberate. But the book is built as an "
    "atlas, and most readers will get more from following a route. Five suggestions follow; the "
    "chapter numbers are the ones printed in the headings, and every route can be left at any point."
)
ROUTES = [
    ("The big picture (about thirty chapters).",
     "Chapter 26 (How Sure Are We, Really?), Chapter 29 (Energy), Chapter 30 (the direction of "
     "time), Chapter 35 and Chapter 36 (relativity and gravity), Chapter 46 (Symmetry), Chapter 55 "
     "(when a crazy idea is science), Chapters 57–58 (the frontier), Chapter 59, Chapters 67–69 "
     "(the Big Bang, expansion and the microwave background), Chapter 105 (The Honest Ending), "
     "Chapters 144–156 (the opening course in quantum physics), Chapter 253 (What Is Life?), "
     "Chapter 259 (Natural Selection), Chapter 261 (DNA) and Chapter 285 (the start of genetics)."),
    ("Quantum physics.",
     "Chapters 144–156 first: thirteen short chapters that build the working rules from a state "
     "as a weighted list to Hawking radiation. Then Chapters 157–168 take the same ground deeper, "
     "including fields (Chapter 162), QCD (Chapter 163) and renormalization (Chapter 164); "
     "Chapters 169–172 are optional philosophy and alternatives. Chapters 173–222 retell "
     "electromagnetism through phase and potentials, from a single electron to a superconducting "
     "circuit. Chapters 223–252 are the full course in quantum electrodynamics; it assumes the "
     "Mathematics section, especially Chapter 11, Chapter 12 and Chapter 25. Finish with Chapter "
     "354 (Interpretations on Trial)."),
    ("Science for science fiction.",
     "Relativity and time: Chapter 35, Chapter 36 and Chapters 60–66 (no universal now, the block "
     "universe). Black holes: Chapters 80–84, then Chapter 156. Wormholes and time travel: "
     "Chapter 43 and Chapters 93–96. Multiverses: Chapters 97–101. Other worlds and interstellar "
     "travel: Chapters 85–92. A reality check on the Moon and Mars: Chapters 126–143. Quantum "
     "interpretation: Chapter 153 and Chapter 354. Then Science in the Novels, Chapters 333–353, "
     "where the physics is told as a detective story."),
    ("The engineer refreshing fundamentals.",
     "Mathematics: Chapters 7–12 (trigonometry, the sinusoid, logarithms, the decibel, complex "
     "numbers, Euler's formula), Chapter 14, Chapters 16–20, Chapter 22 (Fourier and the FFT), "
     "Chapter 23 (Laplace) and Chapter 25 (matrices). Waves and fields: Chapter 49 (Light), "
     "Chapter 52 (Optics), Chapter 229 and Chapter 230 (Maxwell's equations and potentials), and "
     "Chapters 173–186. Quantum and QED: Chapters 144–148, Chapters 231–237, then Chapters "
     "162–164 for the overview and Chapters 245–252 for the machinery. Devices: Chapter 56 "
     "(Semiconductors) and Chapters 213–217 (from QED to a superconducting circuit)."),
    ("Life and evolution.",
     "Chapters 253–257 (what life is and how it may have started), Chapters 258–273 (evolution, "
     "from Darwin to evolution you can watch), Chapters 274–284 (brains, minds, AI and "
     "exobiology), then genetics, told as one continuous story: Chapters 285–299 (the molecule), "
     "Chapters 300–304 (Mendel and his shadow), Chapters 305–320 (the tools, cloning and food), "
     "and Chapters 321–332 (the genome, testing and the ethics of editing). For the human body, "
     "add Chapters 107–125 on ageing and medicine."),
]

PREFACE = [
    "I spent more than forty years in the semiconductor industry, first in Germany and then "
    "mostly in the United States, working with companies that were trying to turn physics into "
    "products. That work teaches one habit more firmly than any other: before you believe a "
    "result, ask how it was measured, and ask what you would have seen if it were false.",
    "A chip either works or it does not. A wafer either yields or it does not. Nobody in a "
    "fabrication plant is impressed by an elegant explanation of why a process should work; "
    "they want the data, the error bars, and the run that failed. Over the years I found that "
    "the same three questions sort almost every claim in science as well as they sort claims "
    "on a production line. What was actually observed? How many independent times? And what "
    "would the theory have forbidden?",
    "Those three questions are where the status lines in this book come from. Settled means "
    "observed, repeated by people who were trying to break it, and still standing. Strange but "
    "solid means the same, with the added insult that it contradicts common sense. Serious but "
    "unconfirmed means competent people are working on it and the evidence has not yet voted. "
    "Speculative means somebody had an idea, which is how everything starts, and nothing more "
    "has happened yet.",
    "I studied physics alongside the engineering, deeply enough to know where I am simplifying "
    "and where I am not. Where this book simplifies, it tries to say so once, clearly, and "
    "move on. Where it does not know, it says that too. A field guide that pretends every bird "
    "has been identified is no use in a real forest.",
    "These days I divide my time between San Clemente in California and Passau in Bavaria, "
    "where Austria starts at the end of the garden and eleven chickens supervise the writing. "
    "The chickens apply the three questions to everything I bring out of the house. They accept "
    "no claim without evidence of food. I have tried to hold this book to a similar standard.",
]

AFTERWORD_RULES = [
    ("The books balance.", "Energy, charge and momentum are conserved because the laws do not "
     "change from place to place or moment to moment (Chapter 29 and Chapter 46). "
     "The same bookkeeping turns up as charge conservation in every circuit, and as the "
     "accounting of a black hole's entropy (Chapter 156)."),
    ("Crowds are predictable; individuals are not.", "Entropy's arrow is a statement about "
     "overwhelming numbers (Chapter 30). So is radioactive decay, so is the Born rule "
     "(Chapter 144), and so is the spread of a variant through a population (Chapter 259)."),
    ("Descriptions change; invariants do not.", "Observers disagree about time and length and "
     "agree about the interval (Chapter 334). Coordinates, gauges and phases can be chosen "
     "freely; the physics that survives every choice is the physics that is real."),
    ("Copies are never exact.", "DNA copies itself with a small error rate, and the errors are "
     "the raw material of evolution (Chapter 295). An unknown quantum state cannot "
     "be copied at all (Chapter 154). Culture copies too, with its own error rate, and "
     "language and ideas evolve in the same way (Chapter 276)."),
    ("Waves add before they are counted.", "Amplitudes add and then get squared "
     "(Chapter 145). Rotating arrows add in an alternating-current circuit (Chapter 12), Fourier "
     "components add into a square wave (Chapter 22), and Feynman's paths add into the one "
     "classical route."),
    ("Knowing has a price.", "A which-path record erases the stripes (Chapter 149); position and "
     "momentum share one budget (Chapter 148); an eavesdropper leaves fingerprints. Information "
     "is physical, and the universe charges for it."),
]
AFTERWORD_CLOSE = (
    "The book began with counting and ended with a detective arguing about when a murder "
    "happened. In between, the same six rules kept turning up. They are not the whole of science, and some of "
    "the chapters that lean on them are labeled Speculative for good reasons. But if you can "
    "spot them the next time you read a headline about physics, biology or a new chip, this "
    "book has done what it set out to do. The universe keeps the books. You can read them now."
)

ABOUT = (
    "Lothar J. Musiol is a graduate of Munich University of Applied Sciences and has spent many "
    "years in the semiconductor industry. Meanwhile, he studied science with enough depth to know "
    "exactly where he simplifies in his explanations, and where he does not. He gained most of his "
    "experience in the United States, where he supported several start-ups. Today, he makes good "
    "use of his dual citizenship and divides his time between San Clemente (California) and Passau "
    "(Bavaria). His goal is to present complex relationships more clearly and engagingly than many "
    "treatments manage, without obscuring their true complexity."
)


def front_matter(doc, volume_label):
    flush_left(doc.add_paragraph(), space_before=160, align=WD_ALIGN_PARAGRAPH.CENTER)
    t = doc.paragraphs[-1].add_run(TITLE)
    t.bold = True
    t.font.size = Pt(32)
    t.font.name = HEADLINE_FONT
    if volume_label:
        flush_left(doc.add_paragraph(), space_before=10, align=WD_ALIGN_PARAGRAPH.CENTER,
                   text=volume_label, size=16)
    flush_left(doc.add_paragraph(), space_before=24, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER,
               text=SUBTITLE, size=14)
    flush_left(doc.add_paragraph(), space_before=90, align=WD_ALIGN_PARAGRAPH.CENTER,
               text=AUTHOR, size=14)

    new_section(doc)
    flush_left(doc.add_paragraph(), space_before=330 if PRINT else 300,
               align=WD_ALIGN_PARAGRAPH.CENTER, text="Copyright © 2026 %s." % AUTHOR, size=9)
    flush_left(doc.add_paragraph(), space_before=4, align=WD_ALIGN_PARAGRAPH.CENTER,
               text="All rights reserved.", size=9)
    flush_left(doc.add_paragraph(), space_before=10, align=WD_ALIGN_PARAGRAPH.CENTER,
               text="Figures drawn by the author for this book.", size=9)
    flush_left(doc.add_paragraph(), space_before=10, align=WD_ALIGN_PARAGRAPH.CENTER,
               text="Updated through October 2026.", size=9)

    new_section(doc)
    p = doc.add_paragraph()
    flush_left(p, space_before=60, space_after=18, align=WD_ALIGN_PARAGRAPH.CENTER)
    r = p.add_run("Contents")
    r.font.color.rgb = HEADING_BLUE
    r.bold = True
    r.font.size = Pt(24)
    r.font.name = HEADLINE_FONT
    toc_placeholder = doc.add_paragraph()

    new_section(doc, 'Who This Book Is For')
    heading(doc, 1, "Who This Book Is For")
    for i, para in enumerate(WHO_FOR):
        body(doc, para, first=(i == 0))

    new_section(doc, 'How an Engineer Decides What to Trust')
    heading(doc, 1, "Preface: How an Engineer Decides What to Trust")
    for i, para in enumerate(PREFACE):
        body(doc, para, first=(i == 0))

    new_section(doc, 'Introduction')
    heading(doc, 1, "Introduction: The Universe Keeps the Books")
    for i, para in enumerate(FOREWORD):
        body(doc, para, first=(i == 0))
    for name, desc in STATUS_LEGEND:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.first_line_indent = Inches(0)
        p.add_run(name).bold = True
        p.add_run(': ' + desc)
    body(doc, FOREWORD_CLOSE, first=True)
    body(doc, LEVEL_INTRO, first=True)
    for name, desc in LEVEL_LEGEND:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.first_line_indent = Inches(0)
        p.add_run(name).bold = True
        p.add_run(': ' + desc)
    body(doc, LEVEL_CLOSE, first=True)

    new_section(doc, 'How to Use This Book')
    heading(doc, 1, "How to Use This Book: Reader Routes")
    body(doc, ROUTES_INTRO, first=True)
    for name, text in ROUTES:
        p = doc.add_paragraph()
        flush_left(p, space_before=8)
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.add_run(name + ' ').bold = True
        p.add_run(text)
        BODY_PARAGRAPHS.append((p, {}))
    return toc_placeholder


def back_matter(doc):
    new_section(doc, 'Afterword')
    heading(doc, 1, "Afterword: The Same Few Rules")
    body(doc, "Here are the rules this book promised, and where each one turned up.", first=True)
    for name, text in AFTERWORD_RULES:
        p = doc.add_paragraph()
        flush_left(p, space_before=8)
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.add_run(name + ' ').bold = True
        p.add_run(text)
        BODY_PARAGRAPHS.append((p, {}))
    p = body(doc, AFTERWORD_CLOSE, first=True)
    p.paragraph_format.space_before = Pt(12)

    render_glossary(doc)

    if PRINT:
        new_section(doc, 'Index')
        heading(doc, 1, "Index")
        p = doc.add_paragraph()
        flush_left(p)
        add_field(p, 'INDEX \\c "2" \\z "1033"', 'Update fields to build the index.')

    new_section(doc, 'About the Author')
    heading(doc, 1, "About the Author")
    body(doc, ABOUT, first=True)

    new_section(doc, 'Also by Lothar J. Musiol')
    heading(doc, 1, "Also by Lothar J. Musiol")
    for group, titles in ALSO_BY:
        p = doc.add_paragraph()
        flush_left(p, space_before=10, space_after=2)
        p.add_run(group).bold = True
        for title, search in titles:
            q = doc.add_paragraph()
            flush_left(q, space_after=1, align=WD_ALIGN_PARAGRAPH.LEFT)
            q.paragraph_format.left_indent = Inches(0.25)
            add_external_link(q, title, store_url(title, search))
    p = doc.add_paragraph()
    flush_left(p, space_before=14, align=WD_ALIGN_PARAGRAPH.LEFT)
    p.add_run('Every title: ').italic = True
    add_external_link(p, 'search for %s on Amazon' % AUTHOR, store_url(''))


# ---------- build ----------
def update_with_word(docx_path):
    try:
        import win32com.client
    except ImportError:
        print("pywin32 missing; open the file in Word and press Ctrl+A, F9 to fill page numbers.")
        return
    word = win32com.client.DispatchEx('Word.Application')
    word.Visible = False
    word.DisplayAlerts = 0
    try:
        d = word.Documents.Open(os.path.abspath(docx_path))
        for _ in range(2):
            d.Fields.Update()
            for toc in d.TablesOfContents:
                toc.Update()
            for ix in d.Indexes:
                ix.Update()
        d.Save()
        pages = d.ComputeStatistics(2)
        d.Close(False)
        print("Word updated fields; pages:", pages)
    finally:
        word.Quit()


def build(out_path, volume=None):
    doc = docx.Document()
    setup_styles(doc)
    volume_label = VOLUMES[volume][0] if volume else ''
    doc.core_properties.title = TITLE + (': ' + volume_label if volume_label else '')
    doc.core_properties.subject = SUBTITLE
    doc.core_properties.author = AUTHOR

    toc_placeholder = front_matter(doc, volume_label)

    part_indexes = VOLUMES[volume][1] if volume else range(len(PARTS))
    for pi in part_indexes:
        part_title, blurb, files = PARTS[pi]
        new_section(doc, part_title)
        heading(doc, 1, part_title)
        BODY_PARAGRAPHS.append((flush_left(doc.add_paragraph(), space_before=6, space_after=24,
                                           italic=True, align=WD_ALIGN_PARAGRAPH.CENTER,
                                           text=blurb), {}))
        for item in files:
            render_markdown(doc, os.path.join(CH_DIR, item))

    back_matter(doc)

    links = linkify_document()
    xe = mark_index_terms() if PRINT else 0

    if PRINT:
        p = toc_placeholder
        add_field(p, 'TOC \\o "1-2" \\h \\z \\u', 'Update fields to build the contents.')
    else:
        for level, text, anchor in HEADINGS:
            if level <= TOC_LEVELS:
                link_paragraph(toc_placeholder, level, text, anchor)
        toc_placeholder._p.getparent().remove(toc_placeholder._p)

    os.makedirs(os.path.dirname(os.path.abspath(out_path)), exist_ok=True)
    doc.save(out_path)
    print("Saved:", out_path)
    print("Words (rough):", sum(len(p.text.split()) for p in doc.paragraphs))
    print("Contents entries:", sum(1 for h in HEADINGS if h[0] <= TOC_LEVELS))
    print("Cross-reference links:", links, " Figures:", _figure_counter[0],
          " Glossary terms:", len(INDEX_TERMS), " Index marks:", xe)
    if PRINT:
        update_with_word(out_path)


if __name__ == '__main__':
    args = sys.argv[1:]
    vol = None
    if '--print' in args:
        PRINT = True
        args.remove('--print')
    if '--verbose' in args:
        args.remove('--verbose')
    if '--volume' in args:
        i = args.index('--volume')
        vol = int(args[i + 1])
        del args[i:i + 2]
    build(args[0] if args else DEFAULT_OUT, vol)
