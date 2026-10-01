"""Build Your First YouTube Channel That Sells.docx in the same shape as the Kindle book."""
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor, Twips
from docx.text.paragraph import Paragraph

ROOT = Path(__file__).parent
FIG = ROOT / "figures"
ORDER = [
    (1, ROOT / "manuscript" / "Sell One Thing.md"),
    (2, ROOT / "manuscript" / "The Click Is the Whole Business.md"),
    (3, ROOT / "manuscript" / "Nobody Can Tell What It Cost.md"),
    (4, ROOT / "manuscript" / "The Videos That Do the Selling.md"),
    (5, ROOT / "manuscript" / "Where the Money Actually Comes From.md"),
    (6, ROOT / "manuscript" / "Six Weeks to a Live Offer.md"),
    (7, ROOT / "manuscript" / "Say This.md"),
    (8, ROOT / "manuscript" / "The Channel Workbook.md"),
    (9, ROOT / "manuscript" / "When Nothing Sells.md"),
]
SERIES = "A working guide for creators who want a channel, not a hobby."
TOKEN = re.compile(r"(\*\*[^*]+?\*\*|\*[^*]+?\*|`[^`]+`|\[[^\]]+?\]\([^)]+?\))")


def shade(cell, fill):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def set_cell_width(cell, inches):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcW = OxmlElement("w:tcW")
    dxa = str(int(inches * 1440))
    tcW.set(qn("w:w"), dxa)
    tcW.set(qn("w:type"), "dxa")
    tcPr.append(tcW)


def add_hyperlink(paragraph, text, url):
    r_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    rStyle = OxmlElement("w:rStyle")
    rStyle.set(qn("w:val"), "Hyperlink")
    rPr.append(rStyle)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "22")
    rPr.append(sz)
    rFonts = OxmlElement("w:rFonts")
    rFonts.set(qn("w:ascii"), "Calibri")
    rFonts.set(qn("w:hAnsi"), "Calibri")
    rPr.append(rFonts)
    run.append(rPr)
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    run.append(t)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_inlines(paragraph, text):
    for part in TOKEN.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
            run.font.name = "Calibri"
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Calibri"
        elif part.startswith("*") and part.endswith("*"):
            run = paragraph.add_run(part[1:-1])
            run.italic = True
            run.font.name = "Calibri"
        elif part.startswith("["):
            match = re.match(r"\[([^\]]+)\]\(([^)]+)\)", part)
            if match:
                add_hyperlink(paragraph, match.group(1), match.group(2))
            else:
                run = paragraph.add_run(part)
                run.font.name = "Calibri"
        else:
            run = paragraph.add_run(part)
            run.font.name = "Calibri"


def is_italic_block(text):
    return text.startswith("*") and text.endswith("*") and not text.startswith("**")


def add_picture(doc, rel, alt, base):
    src = (base / rel).resolve()
    png = src.with_suffix(".png")
    path = png if png.exists() else src
    if path.suffix.lower() not in {".png", ".jpg", ".jpeg"}:
        paragraph = doc.add_paragraph()
        run = paragraph.add_run(alt)
        run.italic = True
        return
    paragraph = doc.add_paragraph()
    paragraph.alignment = 1
    run = paragraph.add_run()
    run.add_picture(str(path), width=Inches(5.6))
    doc_pr = run._r.xpath(".//wp:docPr")
    if doc_pr:
        doc_pr[0].set("descr", alt)


def add_table(doc, rows):
    cols = max(len(row) for row in rows)
    widths = column_widths(cols)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.autofit = False
    table.allow_autofit = False
    table.style = "Table Grid"
    tbl = table._tbl
    tblPr = tbl.tblPr if tbl.tblPr is not None else OxmlElement("w:tblPr")
    tblW = OxmlElement("w:tblW")
    total = int(sum(widths) * 1440)
    tblW.set(qn("w:w"), str(total))
    tblW.set(qn("w:type"), "dxa")
    tblPr.append(tblW)
    for i, row in enumerate(rows):
        for j in range(cols):
            cell = table.rows[i].cells[j]
            set_cell_width(cell, widths[j])
            cell.text = ""
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(2)
            paragraph.paragraph_format.space_before = Pt(2)
            text = row[j] if j < len(row) else ""
            add_inlines(paragraph, text)
            for run in paragraph.runs:
                run.font.size = Pt(10.5)
                run.font.name = "Calibri"
                if i == 0:
                    run.bold = True
            if i == 0:
                shade(cell, "F3F1EA")
    doc.add_paragraph()


def column_widths(cols):
    if cols == 4:
        return [1.05, 1.2, 2.2, 1.15]
    if cols == 3:
        return [1.5, 1.5, 2.6]
    share = 5.6 / cols
    return [share] * cols


def parse_blocks(text):
    lines = text.replace("\r\n", "\n").split("\n")
    blocks = []
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line.strip():
            i += 1
            continue
        if line.strip() == "---":
            i += 1
            continue
        if line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                raw = lines[i].strip().strip("|")
                cells = [cell.strip() for cell in raw.split("|")]
                if not all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells):
                    rows.append(cells)
                i += 1
            blocks.append(("table", rows))
            continue
        if line.startswith("- "):
            items = []
            while i < len(lines) and lines[i].startswith("- "):
                items.append(lines[i][2:].strip())
                i += 1
            blocks.append(("bullets", items))
            continue
        if re.match(r"\d+\. ", line):
            items = []
            while i < len(lines) and re.match(r"\d+\. ", lines[i]):
                items.append(re.sub(r"^\d+\. ", "", lines[i]).strip())
                i += 1
            blocks.append(("numbers", items))
            continue
        if line.startswith("!["):
            match = re.match(r"!\[([^\]]*)\]\(([^)]+)\)", line.strip())
            if match:
                blocks.append(("image", match.group(2), match.group(1)))
                i += 1
                continue
        if line.startswith(">"):
            blocks.append(("callout", line.lstrip(">").strip()))
            i += 1
            continue
        paragraph = [line]
        i += 1
        while i < len(lines) and lines[i].strip() and not lines[i].startswith(("#", "|", "- ", "!", ">")) and not re.match(r"\d+\. ", lines[i]) and lines[i].strip() != "---":
            paragraph.append(lines[i].rstrip())
            i += 1
        blocks.append(("p", " ".join(paragraph).strip()))
    return blocks


def apply_page(section):
    section.page_width = Inches(7)
    section.page_height = Inches(10)
    section.top_margin = section.bottom_margin = Inches(0.7)
    section.left_margin = section.right_margin = Inches(0.7)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)


def set_running_head(section, text):
    apply_page(section)
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    header = section.header.paragraphs[0]
    header.clear()
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header.add_run(text)
    run.italic = True
    run.font.name = "Calibri"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
    footer = section.footer.paragraphs[0]
    footer.clear()
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    footer.add_run()._r.append(begin)
    footer.add_run()._r.append(instr)
    footer.add_run()._r.append(end)
    for item in footer.runs:
        item.font.name = "Calibri"
        item.font.size = Pt(9)


def add_body(doc, blocks, number, base):
    skipped_series = False
    for block in blocks:
        kind = block[0]
        if kind == "p":
            text = block[1]
            if text.startswith("# "):
                title = text[2:].strip()
                section = doc.add_section(WD_SECTION.NEW_PAGE)
                set_running_head(section, f"{number}  {title}")
                doc.add_paragraph(f"{number} {title}", "Heading 1")
                continue
            if text.startswith("## "):
                doc.add_paragraph(text[3:].strip(), "Heading 2")
                continue
            if not skipped_series and text.strip("*") == SERIES.strip("."):
                skipped_series = True
                continue
            if not skipped_series and "working guide for creators" in text:
                skipped_series = True
                continue
            paragraph = doc.add_paragraph()
            if is_italic_block(text):
                run = paragraph.add_run(text[1:-1])
                run.italic = True
                run.font.name = "Calibri"
            else:
                add_inlines(paragraph, text)
        elif kind == "bullets":
            for item in block[1]:
                paragraph = doc.add_paragraph(style="List Bullet")
                add_inlines(paragraph, item)
        elif kind == "numbers":
            num_id = new_restarted_num(doc)
            for item in block[1]:
                paragraph = doc.add_paragraph(style="List Number")
                if num_id:
                    force_num(paragraph, num_id)
                add_inlines(paragraph, item)
        elif kind == "callout":
            add_callout(doc, block[1])
        elif kind == "table":
            add_table(doc, block[1])
        elif kind == "image":
            add_picture(doc, block[1], block[2], base)


def add_callout(doc, text):
    table = doc.add_table(rows=1, cols=1)
    table.autofit = False
    table.allow_autofit = False
    cell = table.cell(0, 0)
    set_cell_width(cell, 5.6)
    shade(cell, "F3F1EA")
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.space_before = Pt(2)
    add_inlines(paragraph, text)
    for run in paragraph.runs:
        run.font.size = Pt(11)
        run.font.name = "Calibri"
    doc.add_paragraph()


def new_restarted_num(doc):
    numbering = doc.part.numbering_part._element
    nums = numbering.findall(qn("w:num"))
    if not nums:
        return None
    source = next((n for n in nums if n.get(qn("w:numId")) == "5"), nums[-1])
    abstract = source.find(qn("w:abstractNumId"))
    if abstract is None:
        return None
    new_id = max(int(n.get(qn("w:numId"))) for n in nums) + 1
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(new_id))
    abs_el = OxmlElement("w:abstractNumId")
    abs_el.set(qn("w:val"), abstract.get(qn("w:val")))
    num.append(abs_el)
    override = OxmlElement("w:lvlOverride")
    override.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:startOverride")
    start.set(qn("w:val"), "1")
    override.append(start)
    num.append(override)
    numbering.append(num)
    return str(new_id)


def force_num(paragraph, num_id):
    pPr = paragraph._p.get_or_add_pPr()
    existing = pPr.find(qn("w:numPr"))
    if existing is not None:
        pPr.remove(existing)
    numPr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    nid = OxmlElement("w:numId")
    nid.set(qn("w:val"), num_id)
    numPr.append(ilvl)
    numPr.append(nid)
    pPr.append(numPr)


def insert_contents(doc):
    headings = [p for p in doc.paragraphs if p.style.name in ("Heading 1", "Heading 2")]
    first = next(p for p in headings if p.style.name == "Heading 1")
    contents_el = OxmlElement("w:p")
    first._p.addprevious(contents_el)
    contents = Paragraph(contents_el, first._parent)
    contents.style = doc.styles["Heading 1"]
    contents.add_run("Contents")
    last = contents
    for i, heading in enumerate(headings):
        start = OxmlElement("w:bookmarkStart")
        start.set(qn("w:id"), str(i + 1))
        start.set(qn("w:name"), f"toc{i}")
        heading._p.insert(0, start)
        end = OxmlElement("w:bookmarkEnd")
        end.set(qn("w:id"), str(i + 1))
        heading._p.append(end)
        p_el = OxmlElement("w:p")
        last._p.addnext(p_el)
        last = Paragraph(p_el, first._parent)
        last.style = doc.styles["Normal"]
        if heading.style.name == "Heading 2":
            last.paragraph_format.left_indent = Inches(0.3)
        link = OxmlElement("w:hyperlink")
        link.set(qn("w:anchor"), f"toc{i}")
        run = OxmlElement("w:r")
        rPr = OxmlElement("w:rPr")
        rFonts = OxmlElement("w:rFonts")
        rFonts.set(qn("w:ascii"), "Calibri")
        rFonts.set(qn("w:hAnsi"), "Calibri")
        rPr.append(rFonts)
        sz = OxmlElement("w:sz")
        sz.set(qn("w:val"), "20" if heading.style.name == "Heading 2" else "22")
        rPr.append(sz)
        color = OxmlElement("w:color")
        color.set(qn("w:val"), "1F4E79")
        rPr.append(color)
        u = OxmlElement("w:u")
        u.set(qn("w:val"), "single")
        rPr.append(u)
        run.append(rPr)
        t = OxmlElement("w:t")
        t.text = heading.text
        run.append(t)
        link.append(run)
        p_el.append(link)


def main():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(7)
    section.page_height = Inches(10)
    section.top_margin = section.bottom_margin = Inches(0.7)
    section.left_margin = section.right_margin = Inches(0.7)
    for name in ["Normal", "Title", "Subtitle", "Heading 1", "Heading 2"]:
        doc.styles[name].font.name = "Calibri"
        doc.styles[name].font.color.rgb = RGBColor(0, 0, 0)
    doc.styles["Normal"].font.size = Pt(11)
    doc.styles["Normal"].paragraph_format.space_after = Pt(8)
    doc.styles["Normal"].paragraph_format.line_spacing = 1.15
    doc.styles["Title"].font.size = Pt(32)
    doc.styles["Heading 1"].font.size = Pt(21)
    doc.styles["Heading 1"].paragraph_format.page_break_before = False
    doc.styles["Heading 2"].font.size = Pt(16)
    doc.styles["Heading 2"].paragraph_format.space_before = Pt(16)
    doc.styles["Heading 2"].paragraph_format.space_after = Pt(8)
    section.different_first_page_header_footer = True
    set_running_head(section, "Start here")

    doc.add_paragraph("Your First YouTube\nChannel That Sells", "Title")
    doc.add_paragraph(
        "How to Choose One Offer, Make the Videos That Sell It, and Get Paid",
        "Subtitle",
    )
    doc.add_paragraph(SERIES)
    doc.add_paragraph("Updated September 2026")

    doc.add_paragraph("Start here", "Heading 1")
    for text in START:
        paragraph = doc.add_paragraph()
        add_inlines(paragraph, text)

    for number, path in ORDER:
        add_body(doc, parse_blocks(path.read_text(encoding="utf-8")), number, path.parent)

    gloss = doc.add_section(WD_SECTION.NEW_PAGE)
    set_running_head(gloss, "Glossary")
    doc.add_paragraph("Glossary", "Heading 1")
    glossary_intro = doc.add_paragraph()
    add_inlines(
        glossary_intro,
        "Short definitions for terms this book uses. A definition here is not a substitute for the platform page cited in the chapter.",
    )
    for term, definition in GLOSSARY:
        paragraph = doc.add_paragraph()
        run = paragraph.add_run(term + ". ")
        run.bold = True
        run.font.name = "Calibri"
        add_inlines(paragraph, definition)

    sources = doc.add_section(WD_SECTION.NEW_PAGE)
    set_running_head(sources, "Sources")
    doc.add_paragraph("Official sources and updates", "Heading 1")
    for text in SOURCES:
        paragraph = doc.add_paragraph()
        if text.startswith("*") and text.endswith("*"):
            run = paragraph.add_run(text[1:-1])
            run.italic = True
            run.font.name = "Calibri"
        else:
            add_inlines(paragraph, text)

    insert_contents(doc)
    doc.core_properties.title = "Your First YouTube Channel That Sells"
    doc.core_properties.author = "Kevin Drew Peters"
    doc.core_properties.subject = "A working guide for creators who want a channel, not a hobby"
    out = ROOT / "Your First YouTube Channel That Sells.docx"
    doc.save(out)
    words = sum(len(p.text.split()) for p in doc.paragraphs)
    print(out)
    print("Words", words)


START = [
    "A channel that sells has one thing a stranger can buy. Views, subscribers, and a monetization badge are not that thing. This book is the path from no offer to a small set of videos that point at one price.",
    "Read it in order if you are starting. If you already post, use the last chapter to find the earliest break, then go to the chapter that fixes that break.",
    "Chapter 1 names the buyer and the single offer, and puts a price on a page before you film.",
    "Chapter 2 wins the click: titles, descriptions, tags, captions, and a posting rhythm you can keep. Where a YouTube help page is cited, that page is the source. Anything else about ranking is labeled as creator consensus.",
    "Chapter 3 makes the videos cheaply enough that you can continue, and tells you which licenses you still have to read.",
    "Chapter 4 is the first eight videos, each with one job, including the video that states the offer.",
    "Chapter 5 is how money actually arrives: platform programs, their thresholds, and affiliate links. Ad revenue is extra. It is not the offer.",
    "Chapter 6 is six weeks to get that offer in front of strangers. Chapter 7 is wording you can paste and then make true. Chapter 8 is the notebook. Chapter 9 is what to fix when nothing sells.",
    "Character limits, export settings, and eligibility rules are facts you can check in the product. They move. The planning numbers in this book are arithmetic, not promises of income. This is not legal, tax, or official platform advice.",
]

SOURCES = [
    "These notes support the claims in the chapters. Planning methods, worksheets, and creator-consensus tactics are editorial guidance, not platform requirements. The sample conversion figures are arithmetic placeholders, not measured rates. Accessed September 2026.",
    "*Chapters 1, 4, and 6 through 9 are a selling path: one offer, eight videos, a six-week calendar, wording, a workbook, and a diagnosis when nothing sells. They state no ranking formula.*",
    "*Chapter 2. Title, description, and tag limits, captions, thumbnails, the title-and-thumbnail test, audience retention, recommendation signals, the recommendation system, Shorts discovery, Shorts analytics, and series playlists are cited to YouTube’s own help pages in the chapter. The eight jobs are the only production load in this book. YouTube does not publish that set as a quota. Four of them have to be long videos, because a Short cannot carry the link. Advertiser-friendly guidelines are cited for the existence of a sensitive-content category. The exact words that trigger it are not published.*",
    "*Chapter 3. Tool names and URLs come from working bookmark lists. Export bitrates, codec, and sample rate are YouTube’s recommended upload encoding settings. The Audio Library location is YouTube’s own help page. A higher-resolution upload getting a better playback encode is creator talk, and the chapter says so. Pricing, licenses, and features change. Specific personal projects were not used.*",
    "*Chapter 5. YouTube Partner Program requirements were checked against YouTube’s own monetization guidance. Other platforms on the list were checked against the linked help page in September 2026. Instagram, Dailymotion, and Vimeo have no number in the chapter. One URL, repeated, is the rule. The FTC pages are US guidance.*",
    "Independent guide. Not affiliated with or endorsed by YouTube or any other platform named here. Not legal or tax advice. Check current official terms before acting.",
    "Pages cited in the chapters, so you can open them:",
    "[YouTube recommended upload encoding settings](https://support.google.com/youtube/answer/1722171) — bitrate, H.264, AAC-LC, 48 kHz.",
    "[YouTube Audio Library](https://support.google.com/youtube/answer/3376882) — where the library lives, and that its tracks are the ones YouTube calls copyright-safe.",
    "[YouTube Partner Program eligibility](https://support.google.com/youtube/answer/72851), [monetization products](https://support.google.com/youtube/answer/94522), [which links are clickable](https://support.google.com/youtube/answer/13748639), [cards](https://support.google.com/youtube/answer/6140493), [end screens](https://support.google.com/youtube/answer/6388789).",
    "[Instagram, editing your profile](https://help.instagram.com/936495066470190/) — lists adding a website to the profile. [Instagram link sticker](https://help.instagram.com/192168966243613) — a sticker on an organic Story can send a tap to a website.",
    "[TikTok, adding a website to your profile](https://support.tiktok.com/en/getting-started/setting-up-your-profile/adding-a-website-to-your-profile) — whether the control appears is on that page.",
    "[FTC Endorsement Guides, 16 CFR Part 255](https://www.ecfr.gov/current/title-16/chapter-I/subchapter-B/part-255) and [Disclosures 101 for Social Media Influencers](https://www.ftc.gov/business-guidance/resources/disclosures-101-social-media-influencers).",
    "[Decoder interview with Marques Brownlee, The Verge, January 2021](https://www.theverge.com/22231657/mkbhd-marques-brownlee-interview-youtube-creator-influencer-decoder). [Recode Media transcript, April 2018](https://www.vox.com/2018/4/16/17241282/transcript-youtube-creator-marques-brownlee-mkbhd). [Andrew Rea, Mashed](https://www.mashed.com/612523/andrew-rea-tells-us-how-binging-with-babish-got-started-exclusive-interview/). [Ali Abdaal, Mixergy](https://mixergy.com/interviews/youtubes-most-popular-productivity-creator/). [Hannah Hart, The Verge, 19 October 2016](https://www.theverge.com/2016/10/19/13315924/hannah-hart-interview-youtube-buffering-my-drunk-kitchen).",
]

GLOSSARY = [
    ("Eight jobs", "The only set of videos this book asks you to finish: who it is for, the problem in their words, the method, proof, the comparison, one objection, the offer, and which video to watch next. The click chapter and the six-week calendar use this same set."),
    ("Phone test", "On your phone, open the offer link the way a stranger will, from the description, and finish a payment or a booking. The step where you stall is the step a buyer will abandon. Refund the test if the tool allows it."),
    ("Ticket", "One workshop on one date, with the date, the length, and the price on the page. After that date, sell the recording as the file, or stop. Do not leave both for sale."),
    ("Monthly pass", "The next file or the next call, billed by you, on your page. The page says what arrives each month and when the next charge happens. A video site’s own membership button can wait until you have cleared that site’s gate."),
    ("Sponsor", "One company pays you for one video. Say so in the video and next to the link. Do not also pitch your own file in that video."),
    ("Advanced features", "The Studio switch that makes an address in a long-form description, and in a long-form comment, clickable. Phone verification comes first. A Short’s description and comments stay unclickable after it is on."),
    ("Qualified watch hours", "Public long-form viewing that YouTube counts toward the hour bars. Hours watched in the Shorts feed do not count. Private, unlisted, deleted, and ad-campaign views do not count."),
    ("Expanded program", "In countries where YouTube has opened it, the earlier gate: 500 subscribers, three public uploads in 90 days, and either 3,000 long-form hours in a year or 3 million Shorts views in 90 days. Fan funding and Shopping. Not a share of watch-page ads."),
    ("YouTube Partner Program", "The higher gate: 1,000 subscribers and either 4,000 long-form hours in 12 months or 10 million qualified Shorts views in 90 days. This is the gate that adds watch-page ads, Shorts Feed ads, and YouTube Premium revenue. The follower-floor chart’s YouTube bar is this gate."),
    ("End screen", "An element in the last 5 to 20 seconds of a video at least 25 seconds long. One that opens a site outside YouTube requires the Partner Program."),
    ("Info card", "A small panel attached to the video. One that opens a site outside YouTube also requires the Partner Program."),
    ("AdSense", "The Google account YouTube uses to pay a channel it has accepted. Meeting a subscriber number does not open it. You apply, and YouTube reviews the channel."),
    ("H.264", "The video codec named in YouTube’s recommended upload settings for an MP4. The audio codec named beside it is AAC-LC."),
    ("9:16", "The vertical frame for a Short: 1080 pixels wide by 1920 tall. A horizontal video is 16:9, 1920 by 1080."),
    ("Affiliate link", "A link that pays you if the viewer buys. Say so next to the link. The FTC pages in the sources list are the U.S. disclosure guidance this book points at."),
    ("Closed captions", "A text track the viewer can turn on or off. The words can be searched. They are not burned into the picture."),
    ("Burned-in captions", "Words that are part of the picture. They stay on. They are a different choice from closed captions."),
    ("Series playlist", "A playlist YouTube can feature as the next video while someone is watching one of yours. The account has to be verified, the videos have to be yours, and a video can sit in only one series playlist."),
    ("Fader", "The volume slider in an editor. A music fader at about a tenth to a fifth of the way up is a position on that slider, not a measurement of loudness."),
    ("Text-to-speech", "A tool that reads a script aloud. Usable when the voice is not a clone of someone else. Cloning someone else’s voice is a consent question the tool does not answer for you."),
    ("Stock license", "The terms on one clip or track. “Free” and “free to use commercially” are different sentences. Read the line on that file before you publish it, including on a second site."),
    ("Offer page", "The page where a stranger sees the price and pays or books. The video is not that page unless the platform gives you a product shelf you are allowed to use."),
    ("Short", "A vertical video, 1080 by 1920. On YouTube, an address in a Short’s description or comments is not clickable. A Short can point at a long video. It cannot be the checkout."),
    ("Qualified Shorts views", "Public views of Shorts in the Shorts feed that YouTube counts toward the Shorts bars: 3 million in 90 days for the expanded program, or 10 million in 90 days for the Partner Program. They do not fill the long-form hour bars."),
]


if __name__ == "__main__":
    main()
