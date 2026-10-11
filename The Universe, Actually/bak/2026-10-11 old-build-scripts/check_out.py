import docx, re, sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
d = docx.Document('../The_Universe_Actually__REV13.docx')
t = [p.text for p in d.paragraphs]
print(len(t), sum(len(x.split()) for x in t))
BS = chr(92)
bad = [x[:100] for x in t if BS in x or '**' in x or '$' in x]
print('leftover markup', len(bad), bad[:6])
print('artifacts', [x[:80] for x in t if re.search(r'earlier guides|Isaac the', x)][:3])
print('empty paras', sum(1 for x in t if not x.strip()))
# headings
import collections
hs = [(p.style.name if p.style is not None else '', p.text) for p in d.paragraphs if p.style is not None and p.style.name.startswith('Heading')]
print(collections.Counter(s for s, _ in hs))
dup = collections.Counter(x for s, x in hs if s == 'Heading 3')
print('duplicate H3 titles', [(k, v) for k, v in dup.items() if v > 1])
