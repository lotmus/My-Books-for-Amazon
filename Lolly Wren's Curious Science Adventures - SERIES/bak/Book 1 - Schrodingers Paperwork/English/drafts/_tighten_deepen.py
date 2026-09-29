# -*- coding: utf-8 -*-
"""Deepen cuts on current BOOK_1 to reach ~2500-4000w removed vs BEFORE."""
from __future__ import annotations

import re
import shutil
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

ROOT = Path(r"D:\My Books for Amazon\Schrodingers_Paperwork")
SRC = ROOT / "Schrodingers_Paperwork_BOOK_1.docx"
BAK = ROOT / "Schrodingers_Paperwork_BOOK_1_BEFORE_TIGHTEN.docx"
LOG = ROOT / "_tighten_log.txt"

PROTECT = [
    "impersonation of my address",
    "lie required electricity",
    "did you get the idea that you were the assistant",
    "There is no lid",
    "accounts are the scarce",
    "Then do not tell me nothing was destroyed",
    "That is not mine",
    "no view from nowhere",
    "cast no shadow",
    "DOMESTIC CONSOLIDATION",
    "OPTIONAL LINEN",
    "EQUIVALENT RESIDENCE",
]


def set_paragraph_text(paragraph, text: str) -> None:
    for t in paragraph._p.findall(".//" + qn("w:t")):
        t.text = ""
    if paragraph.runs:
        paragraph.runs[0].text = text
        r0 = paragraph.runs[0]._r
        t0 = r0.find(qn("w:t"))
        if t0 is None:
            t0 = r0.makeelement(qn("w:t"), {})
            r0.append(t0)
        t0.text = text
        if text.startswith(" ") or text.endswith(" "):
            t0.set(qn("xml:space"), "preserve")
    else:
        paragraph.add_run(text)


def words(s: str) -> int:
    return len(s.split())


def chapter_of(idx: int) -> int | None:
    for ch, a, b in [
        (1, 109, 260), (3, 459, 658), (5, 816, 954), (11, 1502, 1599),
        (12, 1599, 1695), (13, 1695, 1799), (14, 1799, 1897), (18, 2012, 2089),
    ]:
        if a <= idx < b:
            return ch
    return None


def novel_wc(doc: Document) -> int:
    return sum(len(p.text.split()) for p in doc.paragraphs[109:2089])


def protected(text: str) -> bool:
    if text.startswith("Rule:") or text.startswith("→") or text.startswith("Chapter "):
        return True
    low = text.lower()
    for p in PROTECT:
        if p.lower() in low:
            return True
    return False


# Hard deepen: index -> much shorter rewrite (from CURRENT text)
DEEP: dict[int, str] = {
    # Ch18 heat-death / outlook — cut harder while keeping scarce accounts + lid elsewhere
    2026: (
        "“The stars are the first thing to understand. There are stars now. This is not the normal condition. "
        "Star formation peaked long ago; the gas will run down; in a hundred trillion years or so the last star will form and the sky will go out. "
        "Then a very long time of embers. Almost all of history is afterwards. "
        "Then, if the protons are not eternal — and we do not know — even the embers dissolve."
    ),
    2027: (
        "“And then it is the age of the black holes, and they do exactly what your 4C did in a car park in Surrey: they evaporate. "
        "The largest will take on the order of 10 to the power of 100 years. But they go. "
        "After that: a thin, cold, dark space, expanding, and nothing in it that could be called an event."
    ),
    2032: (
        "“Your Ministry, Ms Wren, was doing the same thing, badly, and without knowing it. It wished a district to be simple. "
        "And it discovered what every living system discovers: you cannot simply have that. You must pay for it, and the disorder goes somewhere else. "
        "So they built a hole and put it in there. That was thermodynamics by people who had not been told they were doing thermodynamics, "
        "and it ended in a fire because they never asked where the bill was going. "
        "Every institution that promises to make things simple is proposing to move complexity somewhere you cannot see. Always, and by law. "
        "That is the sentence I should like you to take away.”"
    ),
    2034: (
        "“And so — the outlook,” said Schrottfinger. “Yes. In the long run: the embers, the holes, the dark. I shall not soften it. "
        "But consider the enormous middle: a very long period with order still left to be spent, in which anything that can arrange to be paid for elsewhere may persist, and think, and ask questions."
    ),
    2030: (
        "“Now,” said Erich Schrottfinger. “The interesting part. Every physicist my age was taught that the universe runs down. "
        "Order becomes disorder; the coffee goes cold. Perfectly true. And when I was younger I looked at a living thing and thought: how does that happen. "
        "Because a living thing does the opposite — holds a pattern against the current for eighty years."
    ),
    2031: (
        "“And the answer is not that life defies the law. The answer is that a living thing survives by feeding on order — "
        "taking in structure, excreting disorder, paying the bill locally and passing the cost outward. "
        "A living thing is a place where order is temporarily concentrated because it has arranged to be paid for elsewhere. You have spent a year on it."
    ),
    2017: (
        "“Don’t,” said Schrottfinger. “It is empty. It was always a joke, Ms Wren — I invented it to be ridiculous, "
        "and within twenty years it was on coffee mugs. You make one sarcastic remark about a cat and it outlives your equation.”"
    ),
    2020: (
        "Schrottfinger pushed it aside. “Now. You have finished a thing, and you want somebody to tell you what it was all for, "
        "and Fainrose cannot tell you because she is a working physicist and von Wittenberg cannot tell you because he will start on the eleven dimensions. "
        "So it falls to me, who am not respectable.”"
    ),
    2025: (
        "“Here is the outlook. Your Ministry believed it was managing a country. It was interfering, in a small and provincial way, with something that is not finished. "
        "I mean the universe is not finished, and everything you have lived through this year took place in what future arrangements of matter will regard as the opening seconds."
    ),
    2043: (
        "“One last thing. You will be tempted, in your new post, to find out which of them was right — "
        "de Brévanne and his wave, Keinschein and his agreement, Tegmark and his branches, von Wittenberg and his strings. "
        "You will want the answer, because you spent ten years in electronics and you want to know what the circuit is.”"
    ),

    # Ch11 Keinschein deepen
    1548: (
        "“I do not drink blood, Mrs Chain. Blood is a courier.” He laced his old hands together. "
        "“When a question is put to the world, and the world is obliged to produce one answer, the other answers are not filed. They cease, and in ceasing they release something.” "
        "He shrugged. “I take that. I have taken it since the first measurement.”"
    ),
    1554: (
        "“Nine thousand dwellings,” he said, “each forced from a rich state into a thin one, tens of thousands of alternatives extinguished per address, per week, with certificates. "
        "It is not a meal. A meal is small. This is industry. A national programme with a budget line.”"
    ),
    1561: (
        "“Houses I can dine on for a century. But memory is the record of which branch was taken — how the universe keeps its accounts.” He tapped the cover. "
        "“If you collapse memory by questionnaire, you falsify the ledger. Then nobody can tell what happened — including me, including the system trying to repair itself.”"
    ),
    1563: (
        "“Yes,” said Keinschein, with real approval. “Your Dr Fainrose is right that the country can be recovered from its correlations. "
        "But if you ask leading questions about what people remember, and each answer becomes true as it is given, "
        "the agreements will be manufactured, the syndrome will be clean, and the code will report itself intact — around nothing.”"
    ),
    1565: (
        "“It would be a successful recovery of the wrong state. Certified. Permanently. With no residue for anyone to appeal to, and no residue for anyone to eat.” "
        "He sat back. “You have made me pedagogical. I resent it.”"
    ),
    1569: (
        "“Entirely for my own reasons,” said Keinschein. “Mrs Chain, I have never once acted from virtue. "
        "But my reasons and yours are pointing the same way this month.” He glanced at the boxes. "
        "“Also, I am the only one of you who has read Phase Two — invited into it this morning by a Deputy Under-Secretary who filled in a form about his mother.”"
    ),
    1574: (
        "“You write,” said Keinschein, “as though a measurement were an event: a thing that happens at a time, after which there is one answer.” "
        "He shook his head. “That is the working assumption of every laboratory on this planet. It is not true. "
        "When your Mr Venn scanned Mrs Chain’s mantelpiece, what became definite for Mr Venn did not become definite for the mantelpiece, or the county, or me. "
        "There is no single moment at which the universe decides — only a widening region of agreement, spreading at the speed of gossip.”"
    ),
    1576: (
        "“He froze it for the grid,” said Keinschein. “Not for the woman remembering her sister. She was inside. Different question. Different answer.” "
        "He rose. “There is no view from nowhere, Ms Wren. Your Ministry believes it is standing outside the room. It is not. There is no outside. "
        "When it finally asks what everybody remembers, it will not be reading the answer.”"
    ),
    1518: (
        "He had a great disorganised nimbus of white hair, a moustache of magnificent indifference, a cardigan that had outlived several arguments, and no shoes. "
        "His face was one everyone in the room had seen on posters and school laboratory walls. The resemblance produced not recognition but a sort of vertigo."
    ),
    1507: (
        "“Phase Two doesn’t take memories.” She put the folder on the table. "
        "“It can’t. Not directly. A memory cannot be read out by external survey. It only becomes definite when the person is asked.”"
    ),
    1509: (
        "“So they ask,” said Lolly. “And the asking is the whole mechanism. If a citizen fills in a form, describes the memory, signs the declaration — "
        "then the memory becomes whatever the form says, because the form was the measurement.” She turned the folder round. “Look at the title of the volume.”"
    ),

    # Ch12 deepen speeches
    1640: (
        "“No,” said de Brévanne. “That is my point. You are treating the pre-existence of the name as evidence that the observation was worthless. "
        "It is evidence that the observation was correct.” He let that sit. "
        "“One may receive a piece of physics and possess a fact. Or one may arrive at it and possess the reason, and go on to the next thing, which has no name yet. "
        "The second is the only kind that is any use in a crisis, and it is the only kind you have.”"
    ),
    1656: (
        "“Suppose the house always had one definite configuration. One actual wallpaper. An actual state, at every moment.” De Brévanne raised one gloved finger. "
        "“And suppose there exists a wave: a real physical wave, spread over every configuration the house might have had. "
        "Suppose the actual configuration is guided by that wave — steered, never anywhere but somewhere, yet always told where to go next by a thing that fills the whole space of alternatives.”"
    ),
    1661: (
        "“He did not create a fact. There was already a fact.” De Brévanne’s voice sharpened. "
        "“He deformed the guidance. He altered the wave steering that woman’s household. The actual configuration went somewhere thin. "
        "When Mr Flint tore the cable out, the deformation ceased. The flowers came back because the wave that had always contained them was permitted, briefly, to matter again.”"
    ),
    1664: (
        "“It is an entirely different crime,” said de Brévanne, “and this is the only thing I came to say.” He turned to face her. "
        "“Ms Sloan is building a case that the Ministry falsified records — true, and small. "
        "Dr Fainrose is building a case that they interfered with lawful evolution — truer, and larger. "
        "But if what I have described is so, then everything they have done — the queue, the scanner, the certificates, and now the memories — was not an intrusion into knowledge.”"
    ),
    1630: (
        "“In 1911 my brother served as secretary to a congress of physicists,” said de Brévanne, “and brought home the papers. I read them because I read everything.” "
        "He turned the hat over. “I had no laboratory, no standing. What I had was the habit of a historian: when two authorities disagree, one does not choose. "
        "One looks for the document they have both misread.”"
    ),
    1632: (
        "“Light,” said de Brévanne, “which everybody knew was a wave, and which Einstein had shown behaved in some circumstances as though it were a great number of small hard things.” "
        "He shrugged. “The physicists were embarrassed. I was not, because I was not a physicist, and so I took it seriously in both directions at once.”"
    ),
    1634: (
        "“Everyone was asking how a wave could be a particle.” He looked at Lolly. "
        "“Nobody was asking whether a particle might be a wave. It required no apparatus and no genius — only somebody outside the room, with the wrong training.”"
    ),
    1683: (
        "“You said you were a physicist for eleven minutes, once.” He was already turning away. "
        "“I was a historian until twenty-two, and a physicist for sixty years afterward, and never once both. "
        "That is why I could not see what your bureaucracy was doing to those houses and you could. Do not stand on a wall about it.”"
    ),
    1679: (
        "“There are people who would recognise the shape of it,” said de Brévanne. “There is no one who has said it about a country.” "
        "He rose. “You will now go upstairs and say it to Dr Fainrose. She will tell you it is unfalsifiable metaphysics and a hundred years out of fashion. "
        "Then she will be unable to sleep, because it also specifies a procedure, and hers does not.”"
    ),
    1624: (
        "“I was a physicist,” she said. “For about eleven minutes. Then electronics for ten years, then Transit Reconciliation. "
        "Three days ago I started noticing things in a notebook, and every one turns out to be a standard result with somebody’s name on it.” "
        "She did not much care how it sounded. “I write that the house is larger than the form and feel clever until somebody names it.”"
    ),
    1605: (
        "The woman in Sub-District 6 had a classification: P-9, pending status, for a citizen the survey had caught in the middle of something — "
        "a name half-said, a journey not yet agreed to — and then required to sign regardless. "
        "Beside the box, on file after file, a second field read no action required. The log stood at ninety-one that morning, and was climbing."
    ),
    1601: (
        "Fainrose had a view, delivered to the syndrome table. “They have taken the one signature off the directive,” she said. "
        "“Now it has no author, and whoever stands up in that inquiry and says what it did will be thanked for the disruption and not the truth. "
        "The cheapest move there is, and they will make it every time.”"
    ),
    1644: (
        "“Dr Fainrose is formidable and will save perhaps nine thousand houses this week,” said de Brévanne. "
        "“But she cannot hear the question she was trained out of hearing. That is what training costs.” He picked up the hat again. "
        "“Keinschein told you there is no view from nowhere, and that nothing is definite until agreement has spread. It frightened you.”"
    ),

    # Ch14 deepen toward holography
    1842: (
        "Lolly closed the notebook. “Professor von Wittenberg,” she said. "
        "“I will write down anything that tells me what they dumped into 4C. I will not take dictation on eleven dimensions until then. "
        "In three days I have been told four incompatible accounts of what is real. And you have now told me everything is made of eleven-dimensional membranes nobody can see.” "
        "She looked at the little loop. “I have a woman in this district who was stopped in the middle of remembering her sister’s name. What am I supposed to do with strings?”"
    ),
    1853: (
        "“In 1996 two colleagues took a particular class of black hole in string theory, counted the configurations of membranes and strings, "
        "and found that the number of states gave exactly the Bekenstein-Hawking entropy. Exactly. With the correct factor.” He said it very quietly. "
        "“That is the first time anyone counted the microscopic states of a black hole and got the entropy formula back exactly.”"
    ),
    1824: (
        "“Ah,” said von Wittenberg. “That is the part that persuaded me. For fifty years the scandal was that our theory of the very small and our theory of gravity would not sit in the same room. "
        "Put them together and the answers came out infinite. Every attempt failed.”"
    ),
    1857: (
        "“Then the question becomes whether the radiation carries them out,” said von Wittenberg. "
        "“The area law suggests something outrageous: that everything inside a region can be encoded on its boundary. The account is on the surface. "
        "There is a precise version in which gravity in a certain space is equivalent to a theory without gravity on its edge. Two descriptions. Same physics.”"
    ),
    1861: (
        "“Yes,” said Leonhard von Wittenberg. “And so the current view — not proven, but the view — is that evaporation does not destroy the information. "
        "It is scrambled almost beyond recovery, but it is not annihilated. Unitarity survives. The books balance.”"
    ),
    1848: (
        "“Alfred telephoned me on Tuesday and shouted at me for eleven minutes.” Von Wittenberg’s mouth twitched. "
        "“His concern is that your Ministry has used a small engineered black hole as a disposal channel for suppressed alternatives, "
        "that the object is evaporating, and that when it has gone there will be no record of what was destroyed. He calls this the falsification of the ledger.”"
    ),
    1814: (
        "Von Wittenberg smiled very slightly. “That is the standard answer and it is the best-tested answer in the history of thought,” he said. "
        "“It is also, I believe, wrong, and I have believed it for about fifty years and been unable to prove it.”"
    ),
    1869: (
        "He picked up his pencil. “I mention it only because everybody who has come to talk to you this week arrived from somewhere else, "
        "and I do not think that is a coincidence. The people who ask what a thing is are usually people who were not trained to take it for granted.”"
    ),
    1879: (
        "“Because if the information is on the horizon, then it comes out in the radiation, over the entire remaining lifetime of the object.” "
        "She was already up. “Which means there is a window. And the window closes when 4C does.” She looked at her. "
        "“Lolly. How long have we been treating the evaporation as the catastrophe?”"
    ),

    # Ch13 deepen
    1745: (
        "Everyone looked at him. Bellew had not moved. “Werner says the question has no meaning. Max says it has too many answers, all of them real. "
        "Between the pair of you there’s not one thing a Board of Inquiry could go and check.” He tapped the table once. "
        "“You’ve both got theories in which nobody can ever be wrong. I’ve spent my career objecting to that.”"
    ),
    1724: (
        "Tegmark had his hand half up. “See, this is where I get off the bus,” he said. "
        "“You say the question generates no numbers. But the equation generates numbers, and my position — a minority one — "
        "is that the mathematical structure isn’t a description of reality, it is the reality.”"
    ),
    1710: (
        "The second arrived talking. He was younger by decades, sandy-haired, tablet in one hand and coffee in the other. "
        "“—if the equations are the ontology then this whole crisis is a bookkeeping dispute — sorry, hello, Max Tegmark, cosmology, "
        "I’ve read the Fainrose paper, Dr Fainrose it’s an honour, your syndrome table is gorgeous.”"
    ),
    1740: (
        "“That would be the defence,” said Beatrix. “That is what they convened him for — a physics under which the harm is relocated rather than caused. "
        "You will see it in the Board’s response: no alternative is ever lost, and therefore no citizen has suffered a loss, and therefore Phase Two may proceed.”"
    ),
    1767: (
        "“And if the correlations exceed Bellew’s ceiling,” said Lolly, “then they weren’t separate. It’s a number. "
        "You can measure it in Sub-District 6 this afternoon and get a figure that no committee can call metaphysics. "
        "It just says: if these were independent systems being independently surveyed, this correlation is impossible.”"
    ),
    1757: (
        "“That’s the point of it, Mr Chairman.” Bellew leaned forward. "
        "“If both things were true — properties in advance, no influence at a distance — then certain correlations have a ceiling. "
        "And if nature ever exceeds that ceiling, then one of your two comfortable beliefs is finished.”"
    ),
    1723: (
        "Heitmann looked at her with no unkindness at all, which was worse. “Then she was harmed,” he said. "
        "“I do not dispute the harm. I dispute the ontology of the harm. She was harmed as a citizen, by a procedure, and you should prosecute the procedure. "
        "Do not ask me to certify that a wave function was injured.”"
    ),
    1747: (
        "“The formalism is grand,” said Bellew. “I’m talking about the word. Measurement. "
        "It sounds like a special class of event the rest of nature is exempt from. There isn’t. "
        "A surveyor is a physical system. A form is a physical system. There is no magic in a clipboard.”"
    ),
    1785: (
        "“Your Duc’s account and Max’s account and Werner’s silence — you’ll be tempted to pick one.” The pale eyes were unblinking. "
        "“Don’t pick yet. Until then hold all three, and use whichever one tells you where to put the probe. That’s not cowardice, that’s method.”"
    ),

    # Ch1 deepen
    126: (
        "There came a brisk knock at the open office door, immediately followed by a woman entering without waiting for permission. "
        "She was in her late sixties, silver hair pinned back, in a dark coat and sensible shoes, "
        "with the expression of a woman who had accepted the collapse of standards in public life as regrettable "
        "but had no intention of accepting the collapse of space-time."
    ),
    113: (
        "Gideon Wigglesworth was one of the few people who visited Lolly there voluntarily, and one of the few she trusted with an impossible problem. "
        "They had shared a workbench for the better part of a decade. Gideon was very clever, painfully clumsy, and so agreeable that even the Ministry’s furniture seemed reluctant to injure him. "
        "He appeared carrying two coffees, caught his shoe on the threshold, recovered, and handed Lolly one of them."
    ),
    111: (
        "Lolly looked up from her tea. This was unfortunate for two reasons. She had not yet had enough of it to be tolerant, "
        "and the tea itself had the colour and emotional range of varnished disappointment. The terminal clicked again. "
        "Lolly sighed in the manner of a woman who had once trained in theoretical physics and now worked in Transit Reconciliation — "
        "celestial harmony, then local government."
    ),
    110: (
        "There are, in the management of large and failing systems, certain sounds one learns to distrust. "
        "The metallic cough of a heating vent. The clerical thump of an internal memorandum. "
        "And the soft clicking of the complaint terminal on Lolly Wren’s desk when it had decided that the morning was not going to proceed properly."
    ),
    171: (
        "Lolly reached for a pencil. Most officers used regulation pens, permanent and chained to the desk. Lolly preferred pencil. "
        "In the margin of the printed complaint slip she wrote: |journey⟩ = a|not yet taken⟩ + b|already arrived⟩ "
        "She stopped, considered, and added: with a and b being annoyingly nonzero."
    ),
    217: (
        "She opened her bottom drawer and took out a thin green notebook labelled PRIVATE / UNHELPFUL. "
        "It contained several pages of equations, four pages of complaints about the new software, "
        "and one pressed bus ticket from a journey that had taken place twice but only on alternate Tuesdays."
    ),
    257: (
        "Mrs Chain picked up her handbag. “Well,” she said, “that sounds dreadful.” Lolly reached for her coat. “Yes,” she said. “I rather think it does.” "
        "The printer produced a second sheet all by itself: FOR THE AVOIDANCE OF CONFUSION. "
        "Lolly shut her eyes. This was how serious things generally began. With stationery."
    ),
    253: (
        "Lolly looked from the form to the flickering complaint file, to the notebook, to the pale messenger, and at last back to Mrs Chain, "
        "who had arrived somewhere before deciding to go there. She wished very much that the tea had been better."
    ),

    # Ch5 deepen
    817: (
        "Forty minutes after the interruption of Mrs Elspeth Chain’s domestic convergence, Beatrix Sloan was in an unlicensed grey car with Lolly Wren, Jago Flint, "
        "a confiscated module, Mrs Chain’s umbrella, and a woman whose house was still fluctuating. "
        "This was witness selection — the third response to institutional catastrophe she considered proper, after documentation and containment."
    ),
    846: (
        "Before anyone could knock, the front door opened. Stephanie Fainrose was not tall, and not young, and had a stillness that made other people feel over-articulated. "
        "She wore a dark cardigan, narrow trousers, and spectacles that looked like a final administrative judgment upon the human face."
    ),
    843: (
        "Beyond the gate stood a long, low brick house with a slate roof and the defensive air of a property that had spent years discouraging governments. "
        "An old dish antenna had been bolted to one chimney. On the lawn stood a brass armillary sphere and what looked like a weather vane but was almost certainly measuring something less forgivable."
    ),
    852: (
        "They entered. The house smelled of paper, tea, solder, and intellectual impatience. Books lined every wall in layers of active use. "
        "A long table beneath the window held microscopes, mugs, and what Lolly realised was half of a dismantled detector array."
    ),
    893: (
        "Fainrose continued. “Now imagine some administrative imbecile decides this complexity is inefficient. "
        "They force it into a reduced pattern selected for stability and cost. They don’t understand which details support the rest. "
        "So they flatten the system before it has finished cohering.”"
    ),
    841: (
        "At length they turned off the main road into a narrow lane. A wrought-iron gate appeared where Lolly would have sworn there had been no gate. "
        "It bore a sign: THE INSTITUTE FOR UNHELPFUL CLARITY Visitors by prior argument only No refunds for conclusions Jago braked."
    ),
    837: (
        "They drove west. The city thinned gradually. Office blocks gave way to terraces, then detached houses, then stretches of overmanaged green. "
        "The morning had become one of those bright English afternoons determined to deny the possibility of disaster."
    ),
    880: (
        "She turned from the screen. “When a system evolves naturally,” she said, “it doesn’t jump from richness to paperwork. "
        "They’re imposing an end state without understanding the dynamics that produce one.”"
    ),
    921: (
        "Fainrose called up a satellite map, then a network diagram overlaid with municipal infrastructure. "
        "At the centre of one regional cluster blinked a small black icon. Next to it, in yellow: Singularity Asset 4C Status: Diminishing"
    ),

    # Ch3 deepen house-tour gags
    496: (
        "Lolly sat by the window and watched the city fail to remain itself in small indecent increments. "
        "A barber became an orthodontist. A florist became a closed florist. A Victorian pub blurred into a branch office of the Civic Harmony Board."
    ),
    521: (
        "Semi-detached houses stood in rows behind low hedges. At second glance the violations emerged: numbers fractionally out of sequence, "
        "curtains matching too well, identical stone birdbaths. A cat sat on a wall with the air of an animal generated from survey data."
    ),
    522: (
        "And the clocks were wrong. One house displayed eleven twenty-three, another eleven nineteen, a third almost noon — "
        "drawn from different but administratively compatible afternoons."
    ),
    520: (
        "When they got off at Chain Terrace, the weather had the strained brightness of something recently approved by committee. "
        "The street was neat, calm, and appalling."
    ),
    524: (
        "On the far side of the street stood a modest cream-coloured house with blue trim. "
        "Lolly might have called it cheerful had it not looked as if someone had described cheerfulness to a machine and accepted the first draft."
    ),
    564: "Lolly followed. It was a good room: armchair, bookshelves, two lamps, a sideboard.",
    566: (
        "The armchair was too symmetrical. The books were arranged by spine colour. "
        "Three identical coasters sat by the window, because no real household has ever believed in identical coasters."
    ),
    576: (
        "There, too, things were nearly right and therefore unbearable. The kettle was stainless steel instead of enamel. "
        "The tea caddy had been labelled HOT BEVERAGE INFUSION MATERIALS. Three motivational mugs hung from hooks."
    ),
    550: "People really living in a place introduced minor frictions: shoes with opinions, keys of varying utility, a newspaper not quite where it should be.",
    642: (
        "The clock vanished, reappeared, and turned into a photograph of a man with spectacles. "
        "The daffodils browned into geraniums, then both at once. The carpet shifted from stripes to swirls to bare boards and back."
    ),
    653: (
        "Lolly looked down at the humming module, at the cracked mirror, at the photograph still trying to become a clock, "
        "and then back at the man in the doorway. “Oh,” she said. “Was that yours.”"
    ),
    477: (
        "They took the bus because Beatrix did not trust Ministry pool cars, taxis were on strike, "
        "and Mrs Chain would not pay to revisit her own compromised life."
    ),
    519: (
        "Gideon got off at the next stop. “Find out who signed for it,” Lolly told him. "
        "He nodded, walked into the pole by the doors, and was gone before she could thank him."
    ),
    460: (
        "By half past eleven, Lolly Wren decided that watching reality subdivide before lunch had become personal. "
        "She followed Beatrix out of Queue Management with Mrs Chain beside her and Jago materialising wherever least helpful."
    ),
    633: (
        "Lolly took the device in both hands. It was warm. Too warm. “It’s local,” she said. "
        "“They’re installing reality correction hardware in people’s houses.”"
    ),
    625: (
        "Too late. She struck the mirror smartly across the middle. The text vanished. The mirror cracked, "
        "and for the first time all day something in the house seemed honestly itself."
    ),
    600: (
        "Beatrix was already moving again, opening cupboards and checking drawers. “In here,” she said."
    ),
    493: (
        "They boarded. The driver had the hollow expression of a man ordered all morning to carry passengers to places that had become merely adjacent."
    ),
    538: (
        "Lolly turned back. The gate was there — hinges, latch — but its history was gone: the dent, the flake of blue paint, the droop Mrs Chain had resented for years."
    ),
}


def soft_trim(text: str) -> str | None:
    """Light algorithmic trim for medium descriptive paras: drop one em-dash aside."""
    if protected(text):
        return None
    if words(text) < 48:
        return None
    # remove one secondary em-dash clause
    m = re.search(r"\s+[—–-]\s+[^—–.]{12,80}(?=[.,;]|\s+[A-Z“\"])", text)
    if not m:
        # try comma which-clause
        m = re.search(r", which [^,]{15,70},", text)
    if not m:
        return None
    new = (text[: m.start()] + text[m.end() :]).replace("  ", " ").strip()
    # fix doubled punctuation
    new = re.sub(r"\s+,", ",", new)
    new = re.sub(r"\.\.", ".", new)
    if words(new) >= words(text) or words(new) < words(text) * 0.5:
        return None
    if words(text) - words(new) < 4:
        return None
    return new


def main() -> None:
    before_doc = Document(str(BAK))
    before = novel_wc(before_doc)
    doc = Document(str(SRC))

    applied = []
    # hard deep edits
    lock_phrases = [
        "lie required electricity",
        "impersonation of my address",
        "did you get the idea that you were the assistant",
        "accounts are the scarce",
        "there is no lid",
        "then do not tell me nothing was destroyed",
        "that is not mine",
        "no view from nowhere",
        "cast no shadow",
        "he cast no shadow",
    ]
    for idx, new in DEEP.items():
        old = doc.paragraphs[idx].text
        old_l, new_l = old.lower(), new.lower()
        skip = False
        for phrase in lock_phrases:
            if phrase in old_l and phrase not in new_l:
                skip = True
                break
        if skip:
            continue
        if words(new) >= words(old):
            continue
        set_paragraph_text(doc.paragraphs[idx], new)
        applied.append({"idx": idx, "ch": chapter_of(idx), "cut": words(old) - words(new),
                        "old100": old[:100], "new100": new[:100], "how": "deep"})

    # soft trim remaining medium paras in focus chapters
    focus = []
    for a, b in [(109, 260), (459, 658), (816, 954), (1502, 1599), (1599, 1695),
                 (1695, 1799), (1799, 1897), (2012, 2089)]:
        focus.extend(range(a, b))
    for idx in focus:
        if idx in DEEP:
            continue
        old = doc.paragraphs[idx].text
        if not old or protected(old):
            continue
        new = soft_trim(old)
        if not new:
            continue
        set_paragraph_text(doc.paragraphs[idx], new)
        applied.append({"idx": idx, "ch": chapter_of(idx), "cut": words(old) - words(new),
                        "old100": old[:100], "new100": new[:100], "how": "soft"})

    after = novel_wc(doc)
    removed = before - after
    doc.save(str(SRC))
    shutil.copy2(SRC, ROOT / "Schrodingers_Paperwork_BOOK_1_KINDLE.docx")
    shutil.copy2(SRC, ROOT / "Schrodingers_Paperwork_BOOK_1_Photo_Edition.docx")

    after_doc = Document(str(SRC))
    body = "\n".join(p.text for p in after_doc.paragraphs[109:2089])
    locks = {
        "impersonation of my address": "impersonation of my address" in body,
        "DOMESTIC CONSOLIDATION VISIT": "DOMESTIC CONSOLIDATION VISIT" in body,
        "lie required electricity": "lie required electricity" in body,
        "assistant closer": "did you get the idea that you were the assistant" in body,
        "no view from nowhere": "no view from nowhere" in body.lower(),
        "That is not mine": "That is not mine" in body,
        "cast no shadow / He cast no shadow": ("cast no shadow" in body or "He cast no shadow" in body),
        "Mrs Chain sister correction": "Then do not tell me nothing was destroyed" in body,
        "accounts are the scarce": "accounts are the scarce" in body,
        "There is no lid": "There is no lid" in body,
    }

    # best cuts vs BEFORE
    samples = []
    for a, b in [(109, 260), (459, 658), (816, 954), (1502, 1599), (1599, 1695),
                 (1695, 1799), (1799, 1897), (2012, 2089)]:
        for i in range(a, b):
            ob = before_doc.paragraphs[i].text
            na = after_doc.paragraphs[i].text
            if ob != na:
                c = words(ob) - words(na)
                if c > 0:
                    samples.append({"ch": chapter_of(i), "cut": c, "old100": ob[:100], "new100": na[:100]})
    best = sorted(samples, key=lambda x: -x["cut"])
    seen = set()
    uniq = []
    for s in best:
        if s["old100"] in seen:
            continue
        seen.add(s["old100"])
        uniq.append(s)
    best5 = uniq[:5]

    ch_lines = []
    chapters = set()
    for ch, a, b in [
        (1, 109, 260), (3, 459, 658), (5, 816, 954), (11, 1502, 1599),
        (12, 1599, 1695), (13, 1695, 1799), (14, 1799, 1897), (18, 2012, 2089),
    ]:
        bw = sum(len(before_doc.paragraphs[i].text.split()) for i in range(a, b))
        aw = sum(len(after_doc.paragraphs[i].text.split()) for i in range(a, b))
        if bw != aw:
            chapters.add(ch)
        ch_lines.append(f"  Ch{ch}: {bw} -> {aw} (-{bw - aw}w)")

    lines = [
        "Schrodinger's Paperwork -- tighten pass",
        f"Source: {SRC.name}",
        "Backup: Schrodingers_Paperwork_BOOK_1_BEFORE_TIGHTEN.docx",
        f"Novel words before: {before}",
        f"Novel words after:  {after}",
        f"Words removed:      {removed}",
        f"Paragraphs updated this deepen: {len(applied)}",
        f"Chapters touched:   {sorted(chapters)}",
        "Synced: KINDLE + Photo_Edition (copied from BOOK_1)",
        "",
        "Per-chapter (before -> after, cut):",
        *ch_lines,
        "",
        "Continuity locks present:",
    ]
    for k, ok in locks.items():
        lines.append(f"  [{'OK' if ok else 'MISSING'}] {k}")
    lines.append("")
    lines.append("Sample of 5 best cuts (old->new, first 100 chars):")
    for i, c in enumerate(best5, 1):
        lines.append(f"{i}. Ch{c['ch']} -{c['cut']}w")
        lines.append(f"   OLD: {c['old100']}")
        lines.append(f"   NEW: {c['new100']}")

    LOG.write_text("\n".join(lines) + "\n", encoding="utf-8")
    print(LOG.read_text(encoding="utf-8").encode("ascii", "replace").decode("ascii"))
    if removed < 2500:
        print(f"WARNING: only removed {removed}w; target was 2500-4000")


if __name__ == "__main__":
    main()
