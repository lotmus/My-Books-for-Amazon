# -*- coding: utf-8 -*-
from pathlib import Path
from lxml import etree
import zipfile

LIVE = Path(
    r"D:\My Books for Amazon"
    r"\Lolly Wren's Curious Science Adventures - SERIES"
    r"\Book 1 - Schrodingers Paperwork"
    r"\Schrodingers_Paperwork_BOOK_1_KINDLE_FINAL.docx"
)
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
out = Path(__file__).with_name("dump_hl_xml.txt")
lines = []

with zipfile.ZipFile(LIVE) as z:
    xml = z.read("word/document.xml")
    styles = z.read("word/styles.xml")
root = etree.fromstring(xml)
paras = root.findall(f".//{W}p")

def dump_p(i, label):
    p = paras[i]
    lines.append(f"\n===== {label} p{i} =====")
    texts = "".join(t.text or "" for t in p.findall(f".//{W}t"))
    lines.append(texts)
    # pretty children
    xmls = etree.tostring(p, encoding="unicode")
    # shorten
    if len(xmls) > 8000:
        xmls = xmls[:8000] + "\n...TRUNC..."
    lines.append(xmls)

dump_p(571, "CH3 pointer")
dump_p(1302, "CH10 pointer")
dump_p(2032, "epilogue pointer")

# find a simple glossary wrap in novel if any, else a lecture term
# look for hyperlink with anchor=glPointer
n = 0
for i, p in enumerate(paras):
    for h in p.findall(f".//{W}hyperlink"):
        if h.get(f"{W}anchor") == "glPointer":
            dump_p(i, "existing glPointer wrap")
            n += 1
            break
    if n >= 3:
        break
if n == 0:
    # any glossary wrap in lectures
    for i, p in enumerate(paras):
        for h in p.findall(f".//{W}hyperlink"):
            a = h.get(f"{W}anchor") or ""
            if a.startswith("gl") and i > 2400:
                dump_p(i, f"sample gl wrap {a}")
                n += 1
                break
        if n >= 2:
            break

# hyperlink style
sroot = etree.fromstring(styles)
for st in sroot.findall(f".//{W}style"):
    name = st.find(f"{W}name")
    if name is not None and (name.get(f"{W}val") or "").lower() in ("hyperlink", "followedhyperlink"):
        lines.append("\n===== style =====")
        lines.append(etree.tostring(st, encoding="unicode")[:2000])

# How to Read youtube mention
from docx import Document
d = Document(str(LIVE))
lines.append("\n===== How to Read / youtube mentions in front matter =====")
for i in range(80, 130):
    t = d.paragraphs[i].text
    if t.strip():
        lines.append(f"{i}|{t[:300]}")

# backlog proton wavelength
lines.append("\n===== backlog 2855-2875 =====")
for i in range(2855, 2876):
    t = d.paragraphs[i].text
    if t.strip():
        lines.append(f"{i}|{t}")

# Lesson 10 going deeper pointer sentence 2531
lines.append("\n===== 2520-2535 =====")
for i in range(2520, 2536):
    t = d.paragraphs[i].text
    if t.strip():
        lines.append(f"{i}|{t[:400]}")

out.write_text("\n".join(lines), encoding="utf-8")
print("wrote", out)
