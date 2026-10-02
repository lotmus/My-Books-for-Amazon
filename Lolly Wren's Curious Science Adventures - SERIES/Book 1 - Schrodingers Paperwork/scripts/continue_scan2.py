# -*- coding: utf-8 -*-
from pathlib import Path
from lxml import etree
import zipfile

LIVE = Path(
    r"c:\Users\lomus\OneDrive\My Books for Amazon"
    r"\Lolly Wren's Curious Science Adventures - SERIES"
    r"\Book 1 - Schrodingers Paperwork"
    r"\Schrodingers_Paperwork_BOOK_1_KINDLE_FINAL.docx"
)
NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
      "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships"}
out = Path(__file__).with_name("continue_scan2.txt")
lines = []

with zipfile.ZipFile(LIVE) as z:
    xml = z.read("word/document.xml")
    rels = z.read("word/_rels/document.xml.rels")
root = etree.fromstring(xml)
relroot = etree.fromstring(rels)
relmap = {}
for rel in relroot:
    rid = rel.get("Id")
    tgt = rel.get("Target")
    relmap[rid] = tgt

paras = root.findall(".//w:p", NS)
lines.append(f"xml paras {len(paras)}")

# youtube hyperlinks
lines.append("=== hyperlinks containing youtube or http ===")
for i, p in enumerate(paras):
    texts = "".join(t.text or "" for t in p.findall(".//w:t", NS))
    for h in p.findall(".//w:hyperlink", NS):
        rid = h.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
        tgt = relmap.get(rid, "")
        if tgt and ("youtube" in tgt.lower() or tgt.startswith("http")):
            htext = "".join(t.text or "" for t in h.findall(".//w:t", NS))
            if i < 2100 or "watch?v=" in tgt:
                lines.append(f"p{i} rid={rid} tgt={tgt[:80]} visible={htext[:80]!r} para={texts[:120]!r}")

lines.append("\n=== para 238, 904, 1104, 1407, 1810, 1302, 571 ===")
from docx import Document
d = Document(str(LIVE))
for i in (238, 904, 905, 1103, 1104, 1105, 1406, 1407, 1408, 1808, 1809, 1810, 1811, 1301, 1302, 570, 571, 2028, 2029):
    lines.append(f"\n--- {i} ---\n{d.paragraphs[i].text}")

# Lesson 12 heading
lines.append("\n=== Lesson 12 block headings ===")
for i, p in enumerate(d.paragraphs):
    t = p.text.strip()
    if t.startswith("Lesson 12") or (i > 2550 and i < 2650 and t.startswith("Going deeper")):
        if "12" in t or (2550 < i < 2650 and "Going deeper" in t):
            lines.append(f"{i} {t[:200]}")

out.write_text("\n".join(lines), encoding="utf-8")
print("wrote", out)
