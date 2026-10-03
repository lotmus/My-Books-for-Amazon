# -*- coding: utf-8 -*-
from pathlib import Path
from docx import Document

LIVE = Path(
    r"D:\My Books for Amazon"
    r"\Lolly Wren's Curious Science Adventures - SERIES"
    r"\Book 1 - Schrodingers Paperwork"
    r"\Schrodingers_Paperwork_BOOK_1_KINDLE_FINAL.docx"
)
d = Document(str(LIVE))
out = Path(__file__).with_name("voice_context2.txt")
lines = []

def dump(a, b, label):
    lines.append(f"\n===== {label} {a}-{b} =====")
    for i in range(a, min(b+1, len(d.paragraphs))):
        t = d.paragraphs[i].text
        if t.strip():
            lines.append(f"{i}|{t[:500]}")

dump(330, 360, "Desk Three / Beatrix intro")
dump(620, 655, "Venn intro")
dump(880, 940, "Ch6 Hawking / thermal")
dump(1685, 1710, "von Wittenberg voice")
dump(1918, 1930, "Schrottfinger")

# remaining unhurried/pitched/leaning
lines.append("\n===== leftover unhurried/leaning/pitched low =====")
for i, p in enumerate(d.paragraphs[:2100]):
    t = p.text
    low = t.lower()
    if any(k in low for k in (
        "unhurried", "pitched low", "leaning in", "leaning forward", "before they",
        "before she'd decided", "before they’d decided", "quietly and so flatly",
        "so softly",
    )):
        lines.append(f"{i}|{t[:320]}")

out.write_text("\n".join(lines), encoding="utf-8")
print("wrote", out)
