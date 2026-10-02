# -*- coding: utf-8 -*-
"""Cast-wide voice pass + grey recasts + Unruh employed in novel/lectures."""
from __future__ import annotations

import shutil
import zipfile
from copy import deepcopy
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from lxml import etree

SERIES = Path(
    r"c:\Users\lomus\OneDrive\My Books for Amazon"
    r"\Lolly Wren's Curious Science Adventures - SERIES"
    r"\Book 1 - Schrodingers Paperwork"
)
LIVE = SERIES / "Schrodingers_Paperwork_BOOK_1_KINDLE_FINAL.docx"
OUT = SERIES / "bak" / "Schrodingers_Paperwork_BOOK_1_KINDLE_FINAL.docx.next"
WS = Path(
    r"c:\Users\lomus\OneDrive\My Books for Amazon"
    r"\Schrodingers_Paperwork"
    r"\Schrodingers_Paperwork_BOOK_1_KINDLE_FINAL.docx"
)
GUIDE = SERIES / "notes" / "MANUSCRIPT_GUIDE.md"
GUIDE_WS = Path(
    r"c:\Users\lomus\OneDrive\My Books for Amazon"
    r"\Schrodingers_Paperwork"
    r"\MANUSCRIPT_GUIDE.md"
)
CHAR = Path(
    r"c:\Users\lomus\OneDrive\My Books for Amazon"
    r"\Schrodingers_Paperwork"
    r"\CHARACTER_AND_PLACE_GUIDE.md"
)
CHAR_SERIES = SERIES / "notes" / "CHARACTER_AND_PLACE_GUIDE.md"

LDQ = "\u201c"
RDQ = "\u201d"
LSQ = "\u2018"
RSQ = "\u2019"
EM = "\u2014"

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"


def set_text(p, new: str) -> None:
    runs = p.runs
    if not runs:
        p.add_run(new)
        return
    runs[0].text = new
    for r in runs[1:]:
        r.text = ""
    texts = p._p.findall(".//" + qn("w:t"))
    joined = "".join(t.text or "" for t in texts)
    if joined != new:
        for i, t in enumerate(texts):
            t.text = new if i == 0 else ""


def replace_once(p, old: str, new: str) -> bool:
    full = p.text
    if old not in full:
        return False
    for r in p.runs:
        if old in (r.text or ""):
            r.text = (r.text or "").replace(old, new, 1)
            return True
    set_text(p, full.replace(old, new, 1))
    return True


def insert_after(paragraph, text: str) -> Paragraph | None:
    if paragraph is None:
        print("SKIP insert, anchor missing")
        return None
    new_el = deepcopy(paragraph._p)
    for child in list(new_el):
        if child.tag != qn("w:pPr"):
            new_el.remove(child)
    paragraph._p.addnext(new_el)
    new_p = Paragraph(new_el, paragraph._parent)
    new_p.add_run(text)
    return new_p


def must(cond, msg):
    if not cond:
        print("MISSING", msg)
        return False
    return True


def find_para(doc, pred, msg):
    for p in doc.paragraphs:
        if pred(p.text):
            return p
    print("MISSING PARA", msg)
    return None


def req(p, old, new, label):
    if p is None:
        return "SKIP para " + label
    if not replace_once(p, old, new):
        print("MISSING TEXT", label)
        return "SKIP text " + label
    return label


def make_hyperlink(anchor: str, text: str):
    hl = etree.Element(f"{W}hyperlink")
    hl.set(f"{W}anchor", anchor)
    hl.set(f"{W}history", "1")
    r = etree.SubElement(hl, f"{W}r")
    rpr = etree.SubElement(r, f"{W}rPr")
    rstyle = etree.SubElement(rpr, f"{W}rStyle")
    rstyle.set(f"{W}val", "Hyperlink")
    t = etree.SubElement(r, f"{W}t")
    t.text = text
    return hl


def make_run(text: str):
    r = etree.Element(f"{W}r")
    t = etree.SubElement(r, f"{W}t")
    if text[:1].isspace() or text[-1:].isspace():
        t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    t.text = text
    return r


def wrap_first_unruh(p) -> bool:
    full = "".join(t.text or "" for t in p.findall(f".//{W}t"))
    # don't wrap the glossary heading itself
    if full.strip().startswith("Unruh effect"):
        return False
    idx = full.find("Unruh")
    if idx < 0:
        return False
    # already wrapped?
    for h in p.findall(f"{W}hyperlink"):
        vis = "".join(x.text or "" for x in h.findall(f".//{W}t"))
        if vis == "Unruh" and h.get(f"{W}anchor") == "glUnruh":
            return True
    phrase = "Unruh"
    prefix, suffix = full[:idx], full[idx + len(phrase) :]
    for child in list(p):
        if child.tag != f"{W}pPr":
            p.remove(child)
    if prefix:
        p.append(make_run(prefix))
    p.append(make_hyperlink("glUnruh", phrase))
    if suffix:
        p.append(make_run(suffix))
    return True


def prose_pass(doc) -> list[str]:
    log = []

    # --- voice ---
    log.append(
        req(
            find_para(doc, lambda t: "flat, unhurried register" in t, "Lolly voice"),
            "Lolly answered in the flat, unhurried register she"
            + RSQ
            + "d perfected for delivering bad news to people who outranked her "
            + EM
            + " pitched low enough that panicking rooms tended to quiet down just to catch the rest of the sentence.",
            "Lolly answered in the register she"
            + RSQ
            + "d perfected for delivering bad news to people who outranked her: already edited, as if the sentence had spent a night in the green notebook first.",
            "voice/Lolly: notebook-edited, not pitched-low rooms",
        )
    )

    log.append(
        req(
            find_para(doc, lambda t: "pitched low enough that people found themselves leaning in" in t, "Jago voice"),
            "His voice matched the rest of him: unhurried, faintly amused, and pitched low enough that people found themselves leaning in before they"
            + RSQ
            + "d decided the sentence was worth it.",
            "His voice matched the rest of him: cheerful, faintly amused, and carrying slightly too much information, as if the sentence were a parcel he was pleased to have got through in one piece.",
            "voice/Jago: courier-parcel, not leaning-in",
        )
    )

    log.append(
        req(
            find_para(doc, lambda t: "unhurried satisfaction of someone who had waited his whole career" in t, "Priddy warrant"),
            "the unhurried satisfaction of someone who had waited his whole career",
            "the professional satisfaction of someone who had waited his whole career",
            "voice/Priddy: professional satisfaction, not unhurried",
        )
    )

    log.append(
        req(
            find_para(doc, lambda t: "built for dictation, not conversation" in t, "Priddy dictation"),
            "His voice was flat and precise, built for dictation, not conversation, as though every sentence to come was already halfway to being a statement.",
            "His voice was built for dictation, not conversation, as though every sentence to come was already halfway to being a statement.",
            "voice/Priddy: dictation (Heisenburger keeps flat monotone)",
        )
    )

    log.append(
        req(
            find_para(doc, lambda t: "same low, even voice she used for threats" in t, "Beatrix car voice"),
            "When Beatrix finally spoke it was in the same low, even voice she used for threats and directions alike "
            + EM
            + " precise enough that people found themselves complying before they"
            + RSQ
            + "d decided to.",
            "When Beatrix finally spoke it was in the voice she used for minutes: each clause already a finding, and the room rearranged itself around the full stop.",
            "voice/Beatrix: minutes/findings, not complying-before-deciding",
        )
    )

    log.append(
        req(
            find_para(doc, lambda t: "quieter than the doorway suggested" in t, "Fainrose voice"),
            "Her voice was quieter than the doorway suggested it should be, precise rather than raised, each word set down like an instrument she"
            + RSQ
            + "d already calibrated.",
            "Her voice was quieter than the doorway suggested it should be, each word set down like an instrument she"
            + RSQ
            + "d already calibrated.",
            "voice/Fainrose: calibrated instrument (precise freed for others)",
        )
    )

    log.append(
        req(
            find_para(doc, lambda t: "raising his voice were one more form of clumsiness" in t, "Gideon voice"),
            "He spoke, even now, in the same soft, unhurried register he used for everything, as though raising his voice were one more form of clumsiness he had long since trained himself out of.",
            "He spoke, even now, as if volume were a thing you could trip over: carefully, and a little behind the thought, which always arrived intact even when the cup didn"
            + RSQ
            + "t.",
            "voice/Gideon: volume as something you can trip over",
        )
    )

    log.append(
        req(
            find_para(doc, lambda t: "dry as old paper" in t, "Eilstein voice"),
            "The voice was low, dry as old paper, and carried the ghost of an accent too old to belong to any country still drawing its own borders.",
            "The voice carried the ghost of an accent too old to belong to any country still drawing its own borders, and a courtesy so complete it was faintly alarming.",
            "voice/Eilstein: old accent + alarming courtesy (dry stays Mrs Chain)",
        )
    )

    log.append(
        req(
            find_para(doc, lambda t: "slow, exact, and unhurried in a way that made Lolly realise" in t, "de Broccoli voice"),
            "His English was slow, exact, and unhurried in a way that made Lolly realise how quickly everybody else had been speaking, these last three days.",
            "His English was slow and exact, and made Lolly realise how quickly everybody else had been speaking, these last three days.",
            "voice/de Broccoli: slow exact, not unhurried-template",
        )
    )

    log.append(
        req(
            find_para(doc, lambda t: "Belfast accent had gone flat with decades" in t, "Bellboy accent"),
            "His Belfast accent had gone flat with decades of being right in rooms that resented it",
            "His Belfast accent had been worn smooth by decades of being right in rooms that resented it",
            "voice/Bellboy: worn-smooth Belfast (flat kept for Heisenburger)",
        )
    )

    log.append(
        req(
            find_para(doc, lambda t: "spoke so softly that Lolly found herself leaning forward" in t, "Wittenberg intro"),
            "and he spoke so softly that Lolly found herself leaning forward.",
            "and he spoke without competing with the room, so that she had to come to the sentence rather than the other way around.",
            "voice/von Wittenberg: room does not compete; listener comes to him",
        )
    )

    log.append(
        req(
            find_para(doc, lambda t: "so quietly and so flatly that Lolly nearly missed the size" in t, "Wittenberg 1995"),
            "He said it so quietly and so flatly that Lolly nearly missed the size of it.",
            "He said it without raising anything, so that the size of the claim had to do the volume.",
            "voice/von Wittenberg: claim does the volume",
        )
    )

    # --- grey ---
    log.append(
        req(
            find_para(doc, lambda t: "respectable grey building in Coldharrow" in t, "Ministry grey"),
            "a respectable grey building in Coldharrow",
            "a respectable soot-softened stone building in Coldharrow",
            "grey/Ministry: soot-softened stone",
        )
    )
    # Slam pass already recast Beatrix to charcoal wool. Do not revert that sentence.
    if any("Dark grey; badge clipped with punitive exactness" in p.text for p in doc.paragraphs):
        log.append(
            req(
                find_para(doc, lambda t: "Dark grey; badge clipped with punitive exactness" in t, "Beatrix suit"),
                "Dark grey; badge clipped with punitive exactness.",
                "Charcoal; badge clipped with punitive exactness.",
                "grey/Beatrix: charcoal suit",
            )
        )
    else:
        log.append("grey/Beatrix: already charcoal; left the slam wording")
    log.append(
        req(
            find_para(doc, lambda t: "slim case of brushed grey metal" in t, "Venn case"),
            "a slim case of brushed grey metal",
            "a slim case of brushed metal",
            "grey/Venn case: brushed metal (teal named later)",
        )
    )
    log.append(
        req(
            find_para(doc, lambda t: "warm grey module in her hands" in t, "warm module"),
            "the warm grey module in her hands",
            "the warm module in her hands",
            "grey/module 2: the object, not the paint",
        )
    )
    log.append(
        req(
            find_para(doc, lambda t: "little grey thing wireless" in t, "Jago module"),
            "this little grey thing wireless",
            "this little humming thing wireless",
            "grey/Jago: humming thing",
        )
    )
    log.append(
        req(
            find_para(doc, lambda t: "turned the grey module over in her hands" in t, "module over"),
            "turned the grey module over in her hands",
            "turned the module over in her hands",
            "grey/module 3: the module",
        )
    )
    log.append(
        req(
            find_para(doc, lambda t: "left a grey smudge across the pad of her thumb" in t, "graphite"),
            "left a grey smudge across the pad of her thumb",
            "left a graphite smudge across the pad of her thumb",
            "grey/Lolly thumb: graphite",
        )
    )
    log.append(
        req(
            find_para(doc, lambda t: "brushed-grey case" in t, "Priddy Venn case"),
            "brushed-grey case",
            "teal case",
            "grey/Priddy: teal case (matches hall table + epilogue)",
        )
    )
    log.append(
        req(
            find_para(doc, lambda t: "unlicensed grey car" in t, "pool car"),
            "an unlicensed grey car",
            "an unlicensed pool car",
            "grey/car: pool car",
        )
    )
    log.append(
        req(
            find_para(doc, lambda t: "holding the grey module in both hands" in t, "back seat module"),
            "holding the grey module in both hands",
            "holding the module in both hands",
            "grey/module 4: the module",
        )
    )
    log.append(
        req(
            find_para(doc, lambda t: "a narrow grey tie, and gloves" in t, "de Broccoli tie"),
            "a narrow grey tie, and gloves",
            "a narrow slate tie, and gloves",
            "grey/de Broccoli: slate tie",
        )
    )
    log.append(
        req(
            find_para(doc, lambda t: "immaculately grey, his posture" in t, "Heisenburger"),
            "immaculately grey, his posture",
            "charcoal from collar to cuff, his posture",
            "grey/Heisenburger: charcoal, not default grey man",
        )
    )
    log.append(
        req(
            find_para(doc, lambda t: "grey, no-nonsense perm" in t, "Prosser perm"),
            "the sort of grey, no-nonsense perm that had outlasted several fashions",
            "the sort of steel-wool, no-nonsense perm that had outlasted several fashions",
            "grey/Prosser: steel-wool perm (moon smudge keeps grey)",
        )
    )
    log.append(
        req(
            find_para(doc, lambda t: "the flat, permanent grey that London kept" in t, "London rain"),
            "the flat, permanent grey that London kept in reserve",
            "the flat, permanent pewter that London kept in reserve",
            "grey/London: pewter (form-sky keeps grey)",
        )
    )

    # --- Unruh employed ---
    p906 = find_para(doc, lambda t: "Thermal," in t and "Radiation with a temperature" in t, "thermal note")
    insert_after(
        p906,
        "Fainrose added, without looking up, "
        + LDQ
        + "There is a cousin, if you collect cousins. Accelerate a thermometer through empty space and it runs warm, though a colleague sitting still will swear the room is cold. Unruh. Hawking is the same disagreement, written on a horizon. I do not recommend verifying it in a Ministry car."
        + RDQ,
    )
    log.append("Unruh/novel: Fainrose names the cousin over the temperature curve")

    p2377 = find_para(
        doc,
        lambda t: "Hawking temperature of roughly 60 billionths" in t,
        "L6 going deeper body",
    )
    must(
        replace_once(
            p2377,
            "yet the complete mechanism requires quantum gravity and is not settled in a simple textbook picture.",
            "yet the complete mechanism requires quantum gravity and is not settled in a simple textbook picture. "
            "A closely related result, Unruh"
            + RSQ
            + "s 1976 calculation, says an accelerating observer measures a thermal bath in a vacuum an inertial observer still calls empty: the same observer-dependence of "
            + LDQ
            + "nothing"
            + RDQ
            + " that this lecture"
            + RSQ
            + "s vacuum is built on, without needing a horizon. The Backlog keeps the civic translation (legacy emptiness).",
        ),
        "L6 Unruh clause",
    )
    log.append("Unruh/L6 Going deeper: 1976 cousin, backlog keeps civic phrase")

    sticks = find_para(
        doc,
        lambda t: t.startswith("\u2022") and "Black-hole evaporation raises the question" in t,
        "L6 last stick",
    )
    insert_after(
        sticks,
        "\u2022 Unruh: an accelerating observer finds warmth in a vacuum a stationary one still calls empty: Hawking"
        + RSQ
        + "s cousin, without a horizon.",
    )
    log.append("Unruh/L6 What sticks: one bullet")

    p2701 = find_para(
        doc,
        lambda t: "Hawking radiation is better understood as radiation predicted by quantum fields" in t,
        "L14b popular version",
    )
    must(
        replace_once(
            p2701,
            "Hawking radiation is better understood as radiation predicted by quantum fields in curved spacetime.",
            "Hawking radiation is better understood as radiation predicted by quantum fields in curved spacetime. "
            "The Unruh effect is the flat-spacetime cousin: accelerate, and empty space runs a temperature. Lesson 6 named it; the Backlog keeps the Ministry"
            + RSQ
            + "s phrase.",
        ),
        "L14b Unruh",
    )
    log.append("Unruh/L14b: named as Hawking's flat-spacetime cousin")

    return log


def xml_pass(docx_path: Path) -> list[str]:
    buf = BytesIO(docx_path.read_bytes())
    with zipfile.ZipFile(buf, "r") as zin:
        infos = {i.filename: zin.read(i.filename) for i in zin.infolist()}
    root = etree.fromstring(infos["word/document.xml"])
    n = 0
    markers = (
        "Ministry car",
        "1976 calculation",
        "without a horizon",
        "flat-spacetime cousin",
    )
    for p in root.findall(f".//{W}p"):
        t = "".join(x.text or "" for x in p.findall(f".//{W}t"))
        if "Unruh" in t and any(m in t for m in markers):
            if wrap_first_unruh(p):
                n += 1
    infos["word/document.xml"] = etree.tostring(
        root, xml_declaration=True, encoding="UTF-8", standalone=True
    )
    out_buf = BytesIO()
    with zipfile.ZipFile(out_buf, "w", zipfile.ZIP_DEFLATED) as zout:
        with zipfile.ZipFile(docx_path, "r") as zin:
            for item in zin.infolist():
                data = infos.get(item.filename, zin.read(item.filename))
                zout.writestr(item, data)
    docx_path.write_bytes(out_buf.getvalue())
    return [f"glUnruh wraps applied in {n} paragraphs"]


def verify(path: Path) -> None:
    d = Document(str(path))
    blob = "\n".join(p.text for p in d.paragraphs)
    checks = {
        "lolly notebook": "spent a night in the green notebook first" in blob,
        "jago parcel": "parcel he was pleased to have got through" in blob,
        "beatrix minutes": "voice she used for minutes" in blob,
        "priddy dictation": "built for dictation, not conversation" in blob,
        "no pitched low rooms": "pitched low enough that panicking rooms" not in blob,
        "no leaning in": "leaning in before they" not in blob,
        "no complying before": "complying before they" not in blob,
        "gideon trip": "volume were a thing you could trip over" in blob,
        "eilstein courtesy": "courtesy so complete it was faintly alarming" in blob,
        "chain dry kept": "low and dry, built" in blob,
        "heisenburger flat": "same flat monotone" in blob,
        "wittenberg volume": "size of the claim had to do the volume" in blob,
        "ministry stone": "soot-softened stone building" in blob,
        "no respectable grey": "respectable grey building" not in blob,
        "earl grey kept": "Earl Grey" in blob,
        "module first grey": "small grey module" in blob,
        "moon smudge kept": "sort of grey smudge where it ought to be" in blob,
        "form sky kept": "exact grey of a form still awaiting a signature" in blob,
        "graphite": "graphite smudge across the pad of her thumb" in blob,
        "teal case": "teal case" in blob,
        "unruh novel": "I do not recommend verifying it in a Ministry car" in blob,
        "unruh L6 1976": "Unruh" in blob and "1976 calculation" in blob,
        "unruh L14b": "The Unruh effect is the flat-spacetime cousin" in blob,
        "backlog kept": "legacy emptiness" in blob,
        "no dry old paper": "dry as old paper" not in blob,
        "no unhurried register": "unhurried register" not in blob,
    }
    failed = [k for k, v in checks.items() if not v]
    if failed:
        raise SystemExit("VERIFY FAILED: " + ", ".join(failed))

    # leftover default greys in novel (first 2100), excluding keepers
    leftover = []
    keep_sub = (
        "Earl Grey",
        "small grey module",
        "grey smudge where it ought to be",
        "exact grey of a form",
    )
    for i, p in enumerate(d.paragraphs[:2100]):
        t = p.text
        if "grey" not in t.lower() and "gray" not in t.lower():
            continue
        if any(k in t for k in keep_sub) or "Earl Grey" in t:
            continue
        leftover.append(f"{i}:{t[:80]}")
    if leftover:
        print("LEFTOVER GREY", " | ".join(leftover))
    print("VERIFY OK paras", len(d.paragraphs))


GUIDE_ADDENDUM = """
## Voice, grey, Unruh pass (28 Sep 2026)

Recast-not-cut. Applied to the live Kindle file.

**Cast-wide voice (intros only; dialogue already character-led)**
Each overlapping "flat/quiet/precise/unhurried/lean-in" intro now has one acoustic job:
- Lolly: already-edited, as if the sentence spent a night in the green notebook.
- Mrs Chain: low and dry, built to carry a queue (kept; dry is hers).
- Jago: cheerful courier-parcel, slightly too much information.
- Priddy: dictation, not conversation (warrant-card satisfaction is professional, not unhurried).
- Beatrix: the voice she uses for minutes; each clause already a finding.
- Fainrose: quieter than the doorway; words set down like a calibrated instrument.
- Gideon: volume as something you can trip over.
- Eilstein: old accent plus a courtesy so complete it is faintly alarming (dry paper given back to Mrs Chain).
- de Broccoli: slow and exact.
- Bellboy: Belfast worn smooth by being right in rooms that resented it.
- Heisenburger: keeps the flat monotone (emphasis set to zero).
- von Wittenberg: does not compete with the room; the size of the claim does the volume.
- Schrottfinger: Vienna gravel/delight (untouched).

**Grey**
Atmospheric greys recast to specific jobs. Kept on purpose: Earl Grey; the module's first sighting; Mrs Prosser's moon-smudge; the strip of sky the colour of an unsigned form. Venn's case now matches the teal already named in the hall and the Epilogue.

**Unruh**
Employed, not parked. Fainrose names it over the Ch6 temperature curve (cousin to Hawking; do not check in a Ministry car). Lesson 6 Going deeper carries the 1976 clause; a What-sticks bullet; Lesson 14b names the flat-spacetime cousin. Glossary and Backlog ("legacy emptiness") kept.

"""


def update_guide(path: Path) -> None:
    if not path.exists():
        return
    text = path.read_text(encoding="utf-8")
    if "Voice, grey, Unruh pass (28 Sep 2026)" in text:
        return
    path.write_text(text.rstrip() + "\n" + GUIDE_ADDENDUM, encoding="utf-8")


def patch_character_guide(path: Path) -> None:
    if not path.exists():
        return
    t = path.read_text(encoding="utf-8")
    reps = [
        (
            '- **Sound:** Not directly described (she\'s the narrating consciousness for most of the book).',
            '- **Sound:** When she speaks up the chain: already edited, "as if the sentence had spent a night in the green notebook first." (Ch1)',
        ),
        (
            '- **Sound:** Speaks in a "soft, unhurried register... as though raising his voice were one more form\n  of clumsiness he had long since trained himself out of." (Ch9)',
            '- **Sound:** Speaks as if volume were a thing you could trip over: "carefully, and a little behind the thought, which always arrived intact even when the cup didn\u2019t." (Ch9)',
        ),
        (
            '- **Sound:** "Quieter than the doorway suggested it should be, precise rather than raised, each word\n  set down like an instrument she\'d already calibrated. It carried the particular weariness that\n  comes of spending a career making catastrophic news sound like routine paperwork." (Ch5)',
            '- **Sound:** "Quieter than the doorway suggested it should be, each word set down like an instrument she\'d already calibrated. It carried the particular weariness that comes of spending a career making catastrophic news sound like routine paperwork." (Ch5)',
        ),
        (
            '  once been patient and had since overcome it. Dark grey; badge clipped with punitive exactness."',
            '  once been patient and had since overcome it. Charcoal; badge clipped with punitive exactness."',
        ),
        (
            '- **Sound:** Not separately described from her manner \u2014 dialogue carries her voice (clipped,\n  declarative, unadorned).',
            '- **Sound:** "The voice she used for minutes: each clause already a finding." Dialogue stays clipped, declarative, unadorned. (Ch2, Ch5)',
        ),
        (
            '- **Sound:** "Unhurried, faintly amused, and pitched low enough that people found themselves leaning\n  in before they\'d decided the sentence was worth it." (Ch2)',
            '- **Sound:** "Cheerful, faintly amused, and carrying slightly too much information, as if the sentence were a parcel he was pleased to have got through in one piece." (Ch2)',
        ),
        (
            '- **Sound:** "Flat and precise, pitched for the page rather than the room, as though every sentence to\n  come was already halfway to being a statement." (Ch4)',
            '- **Sound:** Built for dictation, not conversation, "as though every sentence to come was already halfway to being a statement." (Ch4)',
        ),
        (
            '  case of brushed grey metal and a card badge with a silver spiral. (Ch4)',
            '  case of brushed metal (later teal) and a card badge with a silver spiral. (Ch4)',
        ),
        (
            '- **Sound:** "Low, dry as old paper, and carried the ghost of an accent too old to belong to any\n  country still drawing its own borders." (Ch11)',
            '- **Sound:** Ghost of an accent too old for any country still drawing its own borders, "and a courtesy so complete it was faintly alarming." (Ch11)',
        ),
        (
            '- **Sound:** "Slow, exact, and unhurried in a way that made Lolly realise how quickly everybody else\n  had been speaking." (Ch12)',
            '- **Sound:** "Slow and exact," which makes Lolly realise how quickly everybody else has been speaking. (Ch12)',
        ),
        (
            '  a narrow grey\n  tie, and gloves. He held a hat." His handshake is "quite cold" and never warms. (Ch12)',
            '  a narrow slate\n  tie, and gloves. He held a hat." His handshake is "quite cold" and never warms. (Ch12)',
        ),
        (
            '- **Looks:** "Tall, hollow-cheeked, immaculately grey, his posture that of someone who had spent\n  decades being the cleverest person in dangerous rooms. He carried nothing." (Ch13)',
            '- **Looks:** "Tall, hollow-cheeked, charcoal from collar to cuff, his posture that of someone who had spent decades being the cleverest person in dangerous rooms. He carried nothing." (Ch13)',
        ),
        (
            '- **Sound:** "A Belfast voice, unhurried and precise." (Ch13)',
            '- **Sound:** Belfast worn smooth by decades of being right in rooms that resented it. Brings a number, not an interpretation. (Ch13)',
        ),
        (
            '- **Sound:** Speaks "so softly that Lolly found herself leaning forward." (Ch14)',
            '- **Sound:** Speaks without competing with the room, so that she has to come to the sentence. The size of the claim does the volume. (Ch14)',
        ),
        (
            '"A respectable grey building in Coldharrow, between a Department of Minor Infrastructure and an',
            '"A respectable soot-softened stone building in Coldharrow, between a Department of Minor Infrastructure and an',
        ),
        (
            'Its graphite had left a grey smudge across the pad of her thumb."',
            'Its graphite had left a graphite smudge across the pad of her thumb."',
        ),
    ]
    for a, b in reps:
        if a in t:
            t = t.replace(a, b, 1)
    path.write_text(t, encoding="utf-8")


def main() -> None:
    if not LIVE.exists():
        raise SystemExit(f"missing {LIVE}")
    if OUT.exists():
        OUT.unlink()
    shutil.copy2(LIVE, OUT)
    doc = Document(str(OUT))
    n0 = len(doc.paragraphs)
    log = prose_pass(doc)
    doc.save(str(OUT))
    log.extend(xml_pass(OUT))
    verify(OUT)
    try:
        if LIVE.exists():
            LIVE.unlink()
        shutil.move(str(OUT), str(LIVE))
    except PermissionError:
        print("LOCKED: patched copy waiting at", OUT)
        for line in log:
            print(" ", line)
        return
    print("paragraphs", n0, "->", len(Document(str(LIVE)).paragraphs))
    for line in log:
        print(" ", line)
    print("wrote", LIVE, "mb", round(LIVE.stat().st_size / 1e6, 2))
    try:
        shutil.copy2(LIVE, WS)
        print("copied to workspace KINDLE file")
    except PermissionError:
        print("workspace KINDLE locked; series copy is live")
    update_guide(GUIDE)
    update_guide(GUIDE_WS)
    patch_character_guide(CHAR)
    if CHAR_SERIES.exists():
        patch_character_guide(CHAR_SERIES)
    print("guides updated")


if __name__ == "__main__":
    main()
