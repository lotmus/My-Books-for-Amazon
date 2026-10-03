# -*- coding: utf-8 -*-
from pathlib import Path
from docx import Document

p = Path(
    r"D:\My Books for Amazon"
    r"\Lolly Wren's Curious Science Adventures - SERIES"
    r"\Book 1 - Schrodingers Paperwork"
    r"\Schrodingers_Paperwork_BOOK_1_KINDLE_FINAL.docx"
)
d = Document(str(p))
out = Path(__file__).with_name("inspect_paras2.txt")
chunks = [
    (1024, 1028),
    (1060, 1070),
    (1238, 1245),
    (1316, 1325),
    (1428, 1438),
    (1624, 1630),
    (1774, 1795),
    (1970, 1985),
    (1986, 2010),
    (2130, 2145),
    (2340, 2352),
    (2574, 2585),
    (2728, 2745),
    (2800, 2815),
    (2828, 2850),
]
lines = []
for a, b in chunks:
    lines.append(f"\n===== {a}-{b} =====")
    for i in range(a, min(b + 1, len(d.paragraphs))):
        t = d.paragraphs[i].text.strip()
        if t:
            lines.append(f"{i} {t[:500]}")
out.write_text("\n".join(lines), encoding="utf-8")
print("wrote", out)
