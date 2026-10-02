# -*- coding: utf-8 -*-
from pathlib import Path
from docx import Document

p = Path(
    r"c:\Users\lomus\OneDrive\My Books for Amazon"
    r"\Lolly Wren's Curious Science Adventures - SERIES"
    r"\Book 1 - Schrodingers Paperwork"
    r"\Schrodingers_Paperwork_BOOK_1_KINDLE_FINAL.docx"
)
d = Document(str(p))
out = Path(__file__).with_name("final_check.txt")
needles = [
    "The ones from yesterday",
    "mantelpiece at Twenty-Two Elm Grove",
    "Three weeks after the incident",
    "18 lectures in four parts",
    "Physics Notes have also been rebuilt",
    "unlikely to be relevant",
    "Lolly’s new work",
    "feed-rate clock",
    "She did go downstairs",
    "hospital door, not from Pauli",
    "18 lectures in five parts",
    "What sticks: the few things",
    "say it to Mrs Chain",
]
blob = "\n".join(x.text for x in d.paragraphs)
lines = [f"paragraphs {len(d.paragraphs)}"]
for n in needles:
    lines.append(f"{n!r}: {n in blob}")
out.write_text("\n".join(lines), encoding="utf-8")
print("\n".join(lines))
