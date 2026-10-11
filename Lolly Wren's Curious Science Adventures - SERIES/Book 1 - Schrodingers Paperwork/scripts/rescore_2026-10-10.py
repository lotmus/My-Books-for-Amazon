import docx, copy, re, sys
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
SRC, DST = sys.argv[1], sys.argv[2]
d = docx.Document(SRC); P = d.paragraphs
RT = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink"
BMID = [max(int(b.get(qn('w:id'))) for b in d.element.body.iter(qn('w:bookmarkStart')))]
LOG = []
def find(prefix, start=0, contains=None):
    hits = [i for i, p in enumerate(P) if i >= start and p.text.startswith(prefix) and (contains is None or contains in p.text)]
    assert len(hits) == 1, (prefix, hits); return hits[0]
def base_rpr(p, skip_links=True):
    for el in p._p:
        if el.tag == qn('w:r') and ''.join(x.text or '' for x in el.iter(qn('w:t'))).strip():
            r = el.find(qn('w:rPr')); return copy.deepcopy(r) if r is not None else None
    for r in p._p.iter(qn('w:r')):
        rp = r.find(qn('w:rPr')); 
        if rp is not None:
            rp = copy.deepcopy(rp)
            for c in list(rp):
                if c.tag in (qn('w:rStyle'), qn('w:color'), qn('w:u')): rp.remove(c)
            return rp
    return None
def mkrun(text, rpr, link=False):
    r = OxmlElement('w:r')
    rp = copy.deepcopy(rpr) if rpr is not None else OxmlElement('w:rPr')
    if link:
        for c in list(rp):
            if c.tag in (qn('w:rStyle'), qn('w:color'), qn('w:u')): rp.remove(c)
        rs = OxmlElement('w:rStyle'); rs.set(qn('w:val'), 'Hyperlink'); rp.insert(0, rs)
    if len(rp): r.append(rp)
    t = OxmlElement('w:t'); t.set(qn('xml:space'), 'preserve'); t.text = text; r.append(t)
    return r
TOK = re.compile(r'\[([^\]]+)\]\(([^)\s]+)\)')
def setp(i, markup, bookmark=None, keep_bookmarks=True, rpr=None):
    p = P[i]; rpr = rpr if rpr is not None else base_rpr(p)
    oldbm = [copy.deepcopy(el) for el in p._p if el.tag in (qn('w:bookmarkStart'), qn('w:bookmarkEnd'))] if keep_bookmarks else []
    for el in list(p._p):
        if el.tag != qn('w:pPr'): p._p.remove(el)
    starts = [b for b in oldbm if b.tag == qn('w:bookmarkStart')]; ends = [b for b in oldbm if b.tag == qn('w:bookmarkEnd')]
    for b in starts: p._p.append(b)
    if bookmark:
        BMID[0] += 1; a = OxmlElement('w:bookmarkStart'); a.set(qn('w:id'), str(BMID[0])); a.set(qn('w:name'), bookmark); p._p.append(a)
    pos = 0
    for m in TOK.finditer(markup):
        if m.start() > pos: p._p.append(mkrun(markup[pos:m.start()], rpr))
        h = OxmlElement('w:hyperlink'); tgt = m.group(2)
        if tgt.startswith('#'): h.set(qn('w:anchor'), tgt[1:])
        else: h.set(qn('r:id'), d.part.relate_to(tgt, RT, is_external=True))
        h.set(qn('w:history'), '1'); h.append(mkrun(m.group(1), rpr, link=True)); p._p.append(h)
        pos = m.end()
    if pos < len(markup): p._p.append(mkrun(markup[pos:], rpr))
    if bookmark:
        b = OxmlElement('w:bookmarkEnd'); b.set(qn('w:id'), str(BMID[0])); p._p.append(b)
    for b in ends: p._p.append(b)
    LOG.append(('setp', i, P[i].text[:70]))
def rep(i, old, new, count=1):
    """replace text inside runs (single run match first, else across runs)"""
    p = P[i]
    for t in p._p.iter(qn('w:t')):
        if t.text and old in t.text:
            t.text = t.text.replace(old, new, count); LOG.append(('rep', i, old[:40], new[:40])); return
    # cross-run
    ts = [t for t in p._p.iter(qn('w:t'))]; full = ''.join(t.text or '' for t in ts); k = full.find(old)
    assert k >= 0, (i, old, full[:200])
    pos = 0; first = True; end = k + len(old)
    for t in ts:
        s, e = pos, pos + len(t.text or ''); pos = e
        if e <= k or s >= end: continue
        a, b = max(k, s) - s, min(end, e) - s
        t.text = t.text[:a] + (new if first else '') + t.text[b:]; first = False
    LOG.append(('rep*', i, old[:40], new[:40]))
def link_text(i, old_text, new_text=None, new_target=None):
    p = P[i]
    for h in p._p.iter(qn('w:hyperlink')):
        txt = ''.join(x.text or '' for x in h.iter(qn('w:t')))
        if txt == old_text:
            if new_text is not None:
                ts = list(h.iter(qn('w:t'))); ts[0].text = new_text
                for t in ts[1:]: t.text = ''
            if new_target:
                if new_target.startswith('#'):
                    h.attrib.pop(qn('r:id'), None); h.set(qn('w:anchor'), new_target[1:])
                else:
                    h.attrib.pop(qn('w:anchor'), None); h.set(qn('r:id'), d.part.relate_to(new_target, RT, is_external=True))
            LOG.append(('link', i, old_text[:30], new_text, new_target)); return
    raise AssertionError((i, old_text))

# 1 copyright page
rep(find('The physics explained in the lectures in the Appendix'), 'in the lectures in the Appendix', 'in the lessons at the back of the book')
# 2 Lolly's age (thirty-four in Chapter 1; no birthday in between)
rep(find('“That is exactly wrong, and exactly what I wanted you to say,”'), 'thirty-five years', 'thirty-four years')
# 3 Tengelman is a composite (A Word on Who Is Who says so)
i = find('Max Tengelman — cheerfully believes'); rep(i, 'Max Tegmark', 'a composite: Hugh Everett III’s many worlds and Max Tegmark’s mathematical universe')
# 4 companion site: no visible URL
i = find('A number of the experiments described in the lectures')
rep(i, 'described in the lectures have', 'described in the lessons have')
rep(i, 'viewable at ', 'viewable on ')
link_text(i, 'lotmus.github.io/lolly-wren-companion', 'the Lolly Wren companion site')
# 5 Lecture N -> Lesson N (the lessons are headed "Lesson")
i = find('If a demonstration is mentioned in a lecture’s Going-deeper block')
rep(i, 'in a lecture’s Going-deeper block, the companion page for that lecture', 'in a lesson’s Going-deeper block, the companion page for that lesson'); rep(i, 'Lecture 1’s', 'Lesson 1’s')
i = find('Lecture 1: Two Slits')
for a in ('Lecture 1:', 'Lecture 2:', 'Lecture 4:', 'Lecture 13:'): rep(i, a, a.replace('Lecture', 'Lesson'))
rep(i, 'the lecture still stands', 'the lesson still stands')
rep(find('On the wavelength of a proton (from Lecture 12)'), 'Lecture 12', 'Lesson 12')
rep(find('A heavy particle is not thereby a classical one.'), 'Lecture 6', 'Lesson 6')
rep(find('On the number nobody derived (from Lecture 18)'), 'Lecture 18', 'Lesson 18')
rep(find('On the vacuum that is not empty (from Lectures 6 and 14)'), 'Lectures 6 and 14', 'Lessons 6 and 14')
rep(find('On the number that is a bound, not a score (from Lecture 13)'), 'Lecture 13', 'Lesson 13')
rep(find('Not every good fact fit its assigned lecture.'), 'its assigned lecture', 'its assigned lesson')
setp(find('The standard door into no-cloning'), 'The standard door into no-cloning, syndromes, and thresholds. The grown-up cousin of [Lesson 8](#Lecture08NoCloning) and [Lesson 15](#Lecture15ErrorCorrection).')
setp(find('Holography as an argument people actually had.'), 'Holography as an argument people actually had. Pair with [Lesson 14b](#Lecture14bHolography).')
# 6 video entries: no visible URL, dates checked against YouTube
FIXD = {'NIk_0AW5hFU': ('25 Dec. 2025', '18 Dec. 2025'), '8DGgvE6hLAU': ('4 Dec. 2013', '5 Dec. 2013'), '8vMwzkOi0v4': ('Sept. 2012', '26 Sept. 2012'),
        '-IfmgyXs7z8': ('2 June 2016', '1 June 2016'), 'TQKELOE9eY4': ('YouTube, 2017', 'YouTube, 16 Sept. 2014'), 'z-syaCoqkZA': ('13 Jul. 2021', '13 July 2021')}
n = 0
for i, p in enumerate(P):
    m = re.search(r', youtube\.com/watch\?v=([\w-]+)\.', p.text)
    if m:
        v = m.group(1)
        if v in FIXD: rep(i, *FIXD[v])
        rep(i, m.group(0), '.'); n += 1
LOG.append(('videos de-URLed', n))
# 7 Further Reading: titles link to the book (Open Library) or paper (DOI); talk videos kept only from official channels, named in words
OL = 'https://openlibrary.org/works/'
FR = {
 'Al-Khalili, Jim.': ('Al-Khalili, Jim. [Quantum: A Guide for the Perplexed](https://openlibrary.org/isbn/9780297843054). Weidenfeld & Nicolson, 2003.', None),
 'Barbour, Julian': (f'Barbour, Julian. [The End of Time: The Next Revolution in Physics]({OL}OL8003307W). Oxford University Press, 1999.', ('Barbour on the same idea, at the Institute of Art and Ideas', 'https://www.youtube.com/watch?v=GoTeGW2csPk')),
 'Becker, Adam.': (f'Becker, Adam. [What Is Real?]({OL}OL20580450W) Basic Books, 2018.', None),
 'Bekenstein, Jacob D.': ('Bekenstein, Jacob D. “[Black Holes and Entropy](https://doi.org/10.1103/PhysRevD.7.2333).” Physical Review D, vol. 7, no. 8, 1973, pp. 2333–2346.', None),
 'Bell, J. S.': (f'Bell, J. S. [Speakable and Unspeakable in Quantum Mechanics]({OL}OL2936940W). 2nd ed., Cambridge University Press, 2004.', None),
 'Bohm, David.': ('Bohm, David. “A Suggested Interpretation of the Quantum Theory in Terms of ‘Hidden’ Variables, [I](https://doi.org/10.1103/PhysRev.85.166) & [II](https://doi.org/10.1103/PhysRev.85.180).” Physical Review, vol. 85, no. 2, 1952, pp. 166–193.', None),
 'Carroll, Sean.': (f'Carroll, Sean. [Something Deeply Hidden: Quantum Worlds and the Emergence of Spacetime]({OL}OL20128006W). Dutton, 2019.', ('Carroll’s book talk at Harvard', 'https://www.youtube.com/watch?v=XBQorAnnICs')),
 'Feynman, Richard P. The Character': (f'Feynman, Richard P. [The Character of Physical Law]({OL}OL271750W). MIT Press, 1967.', None),
 'Feynman, Richard P. QED': (f'Feynman, Richard P. [QED: The Strange Theory of Light and Matter]({OL}OL271752W). Princeton University Press, 1985.', None),
 'Gamow, George.': (f'Gamow, George. [Mr Tompkins in Wonderland]({OL}OL1136285W). Cambridge University Press, 1940.', ('an old educational film of Mr Tompkins', 'https://www.youtube.com/watch?v=WIKOFLRQuc8')),
 'Gell-Mann, Murray.': (f'Gell-Mann, Murray. [The Quark and the Jaguar: Adventures in the Simple and the Complex]({OL}OL3454689W). W. H. Freeman, 1994.', ('Gell-Mann on simplicity and complexity', 'https://www.youtube.com/watch?v=jKFZSzJMqpg')),
 'Greene, Brian.': (f'Greene, Brian. [The Elegant Universe: Superstrings, Hidden Dimensions, and the Quest for the Ultimate Theory]({OL}OL7920347W). W. W. Norton, 1999.', ('Greene’s TED talk on string theory', 'https://www.youtube.com/watch?v=YtdE662eY_M')),
 'Gribbin, John.': (f'Gribbin, John. [In Search of Schrödinger’s Cat: Quantum Physics and Reality]({OL}OL7918929W). Bantam Books, 1984.', ('Gribbin at the Royal Institution', 'https://www.youtube.com/watch?v=8V669ohFOkI')),
 'Griffiths, David J.': (f'Griffiths, David J., and Darrell F. Schroeter. [Introduction to Quantum Mechanics]({OL}OL34317637W). 3rd ed., Cambridge University Press, 2018.', ('Griffiths lecturing at IIT Bombay', 'https://www.youtube.com/watch?v=FY076-fEZ-s')),
 'Hawking, Stephen W.': (f'Hawking, Stephen W. [A Brief History of Time]({OL}OL1892617W). Bantam Books, 1988.', ('Hawking’s Reith Lecture on black holes, from BBC Radio 4', 'https://www.youtube.com/watch?v=EPNGZpp6Glo')),
 'Hawking, Stephen, ed.': (f'Hawking, Stephen, ed. [The Dreams That Stuff Is Made Of: The Most Astounding Papers of Quantum Physics—and How They Shook the Scientific World]({OL}OL16546694W). Running Press, 2011.', None),
 'Heisenberg, Werner.': (f'Heisenberg, Werner. [Physics and Philosophy: The Revolution in Modern Science]({OL}OL326412W). Harper & Row, 1958.', None),
 'McFadden, Johnjoe': (f'McFadden, Johnjoe, and Jim Al-Khalili. [Life on the Edge: The Coming of Age of Quantum Biology]({OL}OL19665622W). Crown, 2015.', ('Al-Khalili’s talk on quantum biology at Oxford', 'https://www.youtube.com/watch?v=kZY0lhgZC0c')),
 'Nielsen, Michael A.': (f'Nielsen, Michael A., and Isaac L. Chuang. [Quantum Computation and Quantum Information]({OL}OL13301633W). Cambridge University Press, 2010.', None),
 'Penrose, Roger.': (f'Penrose, Roger. [The Emperor’s New Mind]({OL}OL34018124W). Oxford University Press, 1989.', None),
 'Preskill, John.': ('Preskill, John. [Lecture notes on Quantum Computation](https://www.preskill.caltech.edu/ph219/). Caltech, various years. Freely available online.', None),
 'Rovelli, Carlo.': (f'Rovelli, Carlo. [Helgoland: Making Sense of the Quantum Revolution]({OL}OL24513723W). Riverhead Books, 2021.', ('Rovelli’s talk at Google', 'https://www.youtube.com/watch?v=gpWf2wyGQ0Q')),
 'Schrödinger, Erwin.': (f'Schrödinger, Erwin. [What Is Life? The Physical Aspect of the Living Cell]({OL}OL1210514W). Cambridge University Press, 1944.', ('a lecture on the book’s legacy from the Dublin Institute for Advanced Studies', 'https://www.youtube.com/watch?v=uGGfmZGgoAE')),
 'Susskind, Leonard.': (f'Susskind, Leonard. [The Black Hole War]({OL}OL5827350W). Little, Brown, 2008.', None),
 'Wallace, David.': (f'Wallace, David. [The Emergent Multiverse]({OL}OL17525007W). Oxford University Press, 2012.', None),
 'Wheeler, John Archibald': (f'Wheeler, John Archibald, with Kenneth Ford. [Geons, Black Holes, and Quantum Foam: A Life in Physics]({OL}OL2738201W). W. W. Norton & Company, 1998.', ('Wheeler on naming the black hole', 'https://www.youtube.com/watch?v=SsBACBVtPYs')),
 'Zurek, Wojciech H.': ('Zurek, Wojciech H. “[Quantum Darwinism](https://doi.org/10.1038/nphys1202).” Nature Physics, vol. 5, 2009, pp. 181–188. [Free preprint on arXiv](https://arxiv.org/abs/0903.5082).', None),
}
frh = find('Further Reading', contains=None) if False else [i for i, p in enumerate(P) if p.text == 'Further Reading' and p.style.name == 'Heading 1'][0]
who = [i for i, p in enumerate(P) if p.text == 'A Word on Who Is Who' and p.style.name == 'Heading 1'][0]
for key, (markup, video) in FR.items():
    hits = [i for i in range(frh, who) if P[i].text.startswith(key)]
    assert len(hits) == 1, (key, hits); i = hits[0]
    bm = 'frSchrodingerWhatIsLife' if key.startswith('Schrödinger') else ('frGribbinSchrodingersCat' if key.startswith('Gribbin') else None)
    setp(i, markup, bookmark=bm)
    if video:
        j = i + 1; assert P[j].text.strip(), (key, 'annotation')
        txt = P[j].text.rstrip()
        setp(j, txt + f' Watch: [{video[0]}]({video[1]}).')
# the arXiv URL line under Zurek goes (now a named link in the entry)
z = [i for i in range(frh, who) if P[i].text.strip() == 'https://arxiv.org/pdf/0903.5082'][0]
P[z]._p.getparent().remove(P[z]._p); LOG.append(('removed arXiv URL line', z))
P = d.paragraphs
# Al-Khalili annotation names Gribbin: link it to his entry
i = find('Friendly without being false.'); setp(i, 'Friendly without being false. A spare on-ramp if [Gribbin](#frGribbinSchrodingersCat) is already on the shelf.')
# in-text citation in Lesson 17
i = find('Going deeper (skip freely). In What Is Life? (1944)')
rep(i, 'In What Is Life? (1944), Schrödinger', 'In What Is Life? (1944), Schrödinger')
p = P[i]
# wrap "What Is Life?" in an internal link
ts = list(p._p.iter(qn('w:t'))); full = ''.join(t.text or '' for t in ts)
assert 'In What Is Life? (1944)' in full
for r in list(p._p.iter(qn('w:r'))):
    t = r.find(qn('w:t'))
    if t is not None and t.text and 'What Is Life?' in t.text:
        a, b = t.text.split('What Is Life?', 1)
        rp = r.find(qn('w:rPr'))
        before = mkrun(a, rp); h = OxmlElement('w:hyperlink'); h.set(qn('w:anchor'), 'frSchrodingerWhatIsLife'); h.set(qn('w:history'), '1'); h.append(mkrun('What Is Life?', rp, link=True)); after = mkrun(b, rp)
        r.addprevious(before); r.addprevious(h); r.addprevious(after); r.getparent().remove(r); LOG.append(('cite link', i)); break
else: raise AssertionError('What Is Life run')
# 8 If This Made You Curious: the sentence split over two paragraphs
i = find('Carlo Rovelli’s '); j = i + 1; assert P[j].text.startswith('Kenneth Ford’s The Quantum World')
rep(i, ' And if you want the same territory covered at considerably greater length and with far less bureaucracy, ', '')
setp(j, 'And if you want the same territory covered at considerably greater length and with far less bureaucracy, Kenneth Ford’s [The Quantum World](https://openlibrary.org/works/OL8063137W) goes much further into most of the questions this book only had time to dramatise.')
# 9 one review request only: the closing page keeps its farewell, drops the second ask
h = [i for i, p in enumerate(P) if p.text == 'If You Enjoyed This Book' and p.style.name == 'Heading 1'][0]
rep(h, 'If You Enjoyed This Book', 'Before You Go')
assert P[h + 1].text.startswith('Reviews are how small books find their readers.')
P[h + 1]._p.getparent().remove(P[h + 1]._p); P = d.paragraphs
toc = [i for i, p in enumerate(P[:80]) if p.text.strip() == 'If You Enjoyed This Book'][0]; rep(toc, 'If You Enjoyed This Book', 'Before You Go')
# 10 the cat link: TED-Ed instead of a hobby channel
link_text(find('“The whole trouble with my cat'), 'my cat', None, 'https://www.youtube.com/watch?v=UjaAxUO6-Uw')
# 11 Also-by: current ILWM subtitle
rep(find('In Love with Murder, Actually: A Comic Mystery'), 'A Comic Mystery', 'A Derailed Mystery in Provence')
# 12 tidy: glossary dash spacing, trailing spaces, empty List Paragraphs, runs of empty paragraphs
for i, p in enumerate(P):
    for t in p._p.iter(qn('w:t')):
        if t.text and '  —  ' in t.text: t.text = t.text.replace('  —  ', ' — ')
    ts = [t for t in p._p.iter(qn('w:t')) if t.text]
    if ts and ts[-1].text != ts[-1].text.rstrip() and p.text.strip(): ts[-1].text = ts[-1].text.rstrip()
def removable(p):
    x = p._p
    return not p.text.strip() and x.find('.//' + qn('w:br')) is None and x.find('.//' + qn('w:bookmarkStart')) is None and x.find('.//' + qn('w:drawing')) is None and x.find('.//' + qn('w:sectPr')) is None and not p.paragraph_format.page_break_before
rm = 0; prev_empty = False
for p in list(d.paragraphs):
    if not p.text.strip():
        if (prev_empty or p.style.name == 'List Paragraph') and removable(p): p._p.getparent().remove(p._p); rm += 1; continue
        prev_empty = True
    else: prev_empty = False
LOG.append(('empty paragraphs removed', rm))
# orphaned hyperlink rels
xml = d.element.xml
for rid, rel in list(d.part.rels.items()):
    if rel.reltype == RT and f'"{rid}"' not in xml: del d.part.rels[rid]
d.save(DST)
for l in LOG: print(*l)
# (second pass on the saved file)
d = docx.Document(DST); P = d.paragraphs
i = [k for k, p in enumerate(P) if p.text.startswith('And if you want more paperwork failing at physics')][0]
rep(i, 'And if you want more paperwork', 'If you want more paperwork')
d.save(DST); print(LOG[-1])
