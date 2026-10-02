# -*- coding: utf-8 -*-
from pathlib import Path
from lxml import etree
import zipfile
from docx import Document

LIVE = Path(
    r"c:\Users\lomus\OneDrive\My Books for Amazon"
    r"\Lolly Wren's Curious Science Adventures - SERIES"
    r"\Book 1 - Schrodingers Paperwork"
    r"\Schrodingers_Paperwork_BOOK_1_KINDLE_FINAL.docx"
)
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"

d = Document(str(LIVE))
print("paras", len(d.paragraphs))

needles = [
    "They reached in and fixed a pointer",
    "a moon that just hangs there being a moon is a pointer",
    "On a noticeboard in Sub-District 6",
    "Fainrose had called it tunnelling",
    "vulgar public explanation",
    "The arrows already have two parts",
    "No collapse. No privileged instant",
    "same unhelpful voice",
    "locked London garden square",
    "A nucleus at comparable speed",
    "doors into the Glossary",
    "A novel of the Ministry of Eventualities",
    "Book 1 of Lolly Wren",
    "each weighted by a small rotating arrow",
]
for n in needles:
    hits = [i for i, p in enumerate(d.paragraphs) if n in p.text]
    print(f"  {n[:50]!r} -> {hits}")

print("\n--- pointer ch3 ---")
print(next(p.text for p in d.paragraphs if "They reached in and fixed a pointer" in p.text))
print("\n--- pointer ch10 ---")
print(next(p.text for p in d.paragraphs if "moon that just hangs there being a moon" in p.text))
print("\n--- garden ---")
print(next(p.text for p in d.paragraphs if "Fainrose had called it tunnelling" in p.text))
print("\n--- htr ---")
print(next(p.text for p in d.paragraphs if "doors into the Glossary" in p.text))

with zipfile.ZipFile(LIVE) as z:
    root = etree.fromstring(z.read("word/document.xml"))
    relroot = etree.fromstring(z.read("word/_rels/document.xml.rels"))
relmap = {rel.get("Id"): rel.get("Target") for rel in relroot}
bms = {bm.get(f"{W}name") for bm in root.findall(f".//{W}bookmarkStart")}

needed = [
    "glSuperposition","glMeasurementBasis","Lecture03PrematureCollapse","glQuantumZeno",
    "Lecture05LawfulEvolution","glHawking","glEntanglement","glNocloning","glUncertainty",
    "Lecture10Decoherence","glMeasurementProblem","glPilot","glBellstheorem","glString",
    "glQuantumErrorCorrection","glLandauers","glNegentropy","glSchrdingers","glPointer",
    "glQuantumtunnelling","glDecoherence","glEinselection",
]
missing = [n for n in needed if n not in bms]
print("\nmissing bookmarks", missing)

print("\n=== retargeted novel links (first 2100 xml) ===")
n = 0
for i, p in enumerate(root.findall(f".//{W}p")):
    if i > 2100:
        break
    for h in p.findall(f".//{W}hyperlink"):
        a = h.get(f"{W}anchor")
        rid = h.get(f"{R}id")
        tgt = relmap.get(rid, "") if rid else ""
        vis = "".join(t.text or "" for t in h.findall(f".//{W}t"))
        if a in needed or (tgt and "youtube" in tgt.lower()):
            print(f"p{i} anchor={a} yt={bool(tgt and 'youtube' in tgt.lower())} vis={vis!r}")
            n += 1
print("count", n)

# straight quotes in new paras
print("\n=== straight quote scan in new text ===")
for p in d.paragraphs:
    t = p.text
    if any(s in t for s in (
        "doors into the Glossary", "nucleus at comparable speed", "locked London garden square",
        "arrows already have two parts", "No collapse. No privileged instant",
        "same unhelpful voice", "Fainrose had called it tunnelling",
        "A novel of the Ministry",
    )):
        if "'" in t or '"' in t:
            print("STRAIGHT", t[:200])
        else:
            print("curly-ok", t[:80])

print("\nbm start/end", len(root.findall(f".//{W}bookmarkStart")), len(root.findall(f".//{W}bookmarkEnd")))
