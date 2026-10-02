# -*- coding: utf-8 -*-
"""Continue recast-not-cut: pointer wraps, novel YouTube, unique arrow jobs, title, L12, garden."""
from __future__ import annotations

import shutil
import zipfile
from copy import deepcopy
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from lxml import etree

SERIES = Path(
    r"c:\Users\lomus\OneDrive\My Books for Amazon"
    r"\Lolly Wren's Curious Science Adventures - SERIES"
    r"\Book 1 - Schrodingers Paperwork"
)
LIVE = SERIES / "Schrodingers_Paperwork_BOOK_1_KINDLE_FINAL.docx"
OUT = SERIES / "bak" / "Schrodingers_Paperwork_BOOK_1_KINDLE_FINAL.docx.next"
WS = Path(
    r"c:\Users\lomus\OneDrive\My Books for Amazon"
    r"\Schrodingers_Paperwork"
    r"\Schrodingers_Paperwork_BOOK_1_KINDLE_FINAL.docx"
)
GUIDE = SERIES / "notes" / "MANUSCRIPT_GUIDE.md"
GUIDE_WS = Path(
    r"c:\Users\lomus\OneDrive\My Books for Amazon"
    r"\Schrodingers_Paperwork"
    r"\MANUSCRIPT_GUIDE.md"
)

LDQ = "\u201c"
RDQ = "\u201d"
LSQ = "\u2018"
RSQ = "\u2019"
EM = "\u2014"

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"


def set_text(p, new: str) -> None:
    runs = p.runs
    if not runs:
        p.add_run(new)
        return
    runs[0].text = new
    for r in runs[1:]:
        r.text = ""
    texts = p._p.findall(".//" + qn("w:t"))
    joined = "".join(t.text or "" for t in texts)
    if joined != new:
        for i, t in enumerate(texts):
            t.text = new if i == 0 else ""


def replace_once(p, old: str, new: str) -> bool:
    full = p.text
    if old not in full:
        return False
    for r in p.runs:
        if old in (r.text or ""):
            r.text = (r.text or "").replace(old, new, 1)
            return True
    set_text(p, full.replace(old, new, 1))
    return True


def insert_after(paragraph, text: str) -> Paragraph:
    new_el = deepcopy(paragraph._p)
    for child in list(new_el):
        if child.tag != qn("w:pPr"):
            new_el.remove(child)
    paragraph._p.addnext(new_el)
    new_p = Paragraph(new_el, paragraph._parent)
    new_p.add_run(text)
    return new_p


def must(cond, msg):
    if not cond:
        raise SystemExit(msg)


def find_para(doc, pred, msg):
    for p in doc.paragraphs:
        if pred(p.text):
            return p
    raise SystemExit(msg)


def make_hyperlink(anchor: str, text: str):
    hl = etree.Element(f"{W}hyperlink")
    hl.set(f"{W}anchor", anchor)
    hl.set(f"{W}history", "1")
    r = etree.SubElement(hl, f"{W}r")
    rpr = etree.SubElement(r, f"{W}rPr")
    rstyle = etree.SubElement(rpr, f"{W}rStyle")
    rstyle.set(f"{W}val", "Hyperlink")
    t = etree.SubElement(r, f"{W}t")
    if text[:1].isspace() or text[-1:].isspace():
        t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    t.text = text
    return hl


def make_run(text: str):
    r = etree.Element(f"{W}r")
    t = etree.SubElement(r, f"{W}t")
    if text[:1].isspace() or text[-1:].isspace():
        t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    t.text = text
    return r


def wipe_content(p):
    for child in list(p):
        if child.tag != f"{W}pPr":
            p.remove(child)


def wrap_phrase_whole_para(p, phrase: str, anchor: str) -> bool:
    full = "".join(t.text or "" for t in p.findall(f".//{W}t"))
    idx = full.find(phrase)
    if idx < 0:
        return False
    wipe_content(p)
    prefix, suffix = full[:idx], full[idx + len(phrase) :]
    if prefix:
        p.append(make_run(prefix))
    p.append(make_hyperlink(anchor, phrase))
    if suffix:
        p.append(make_run(suffix))
    return True


def rebuild_pointer_ch3(p) -> bool:
    full = "".join(t.text or "" for t in p.findall(f".//{W}t"))
    needle = "They reached in and fixed a pointer state before the environment"
    if needle not in full:
        return False
    left, right = full.split("Decoherence", 1)
    after_dec, rest = right.split("pointer state", 1)
    wipe_content(p)
    p.append(make_run(left))
    p.append(make_hyperlink("glDecoherence", "Decoherence"))
    p.append(make_run(after_dec))
    p.append(make_hyperlink("glPointer", "pointer state"))
    p.append(make_run(rest))
    return True


def rebuild_pointer_ch10(p) -> bool:
    full = "".join(t.text or "" for t in p.findall(f".//{W}t"))
    if "answerthat" not in full and "pointer state: an answer that" not in full:
        return False
    full = full.replace("an answerthat", "an answer that", 1)
    left, rest = full.split("(einselection)", 1)
    # left includes the opening paren before einselection
    if not left.endswith("("):
        # split on the linked word only
        left, rest = full.split("einselection", 1)
        mid_left, mid_right = "", rest
        # left is before einselection, rest starts with )
        wipe_content(p)
        p.append(make_run(left))
        p.append(make_hyperlink("glEinselection", "einselection"))
        before_ptr, after_ptr = rest.split("pointer state", 1)
        p.append(make_run(before_ptr))
        p.append(make_hyperlink("glPointer", "pointer state"))
        p.append(make_run(after_ptr))
        return True
    before_ptr, after_ptr = rest.split("pointer state", 1)
    wipe_content(p)
    p.append(make_run(left))
    p.append(make_hyperlink("glEinselection", "einselection"))
    p.append(make_run(")" + before_ptr if not before_ptr.startswith(")") else before_ptr))
    # wait, rest after (einselection) already includes ') and a moon...'
    return False


def rebuild_pointer_ch10_v2(p) -> bool:
    full = "".join(t.text or "" for t in p.findall(f".//{W}t"))
    full = full.replace("an answerthat", "an answer that", 1)
    if "einselection" not in full or "pointer state" not in full:
        return False
    a, b = full.split("einselection", 1)
    c, d = b.split("pointer state", 1)
    wipe_content(p)
    p.append(make_run(a))
    p.append(make_hyperlink("glEinselection", "einselection"))
    p.append(make_run(c))
    p.append(make_hyperlink("glPointer", "pointer state"))
    p.append(make_run(d))
    return True


def prose_pass(doc) -> list[str]:
    log = []

    title = find_para(
        doc,
        lambda t: t.strip() == "Lolly Wren" + RSQ + "s Curious Science Adventures",
        "title series line",
    )
    set_text(title, "A novel of the Ministry of Eventualities")
    log.append("title page: Ministry world-line (KDP series name kept in About the Series)")

    p904 = find_para(doc, lambda t: "vulgar public explanation" in t, "p904 hawking arrows")
    must(
        replace_once(
            p904,
            "is that a quantum field does not sit politely empty in curved spacetime; "
            "you sum every history the field could have and weight each by its own small rotating arrow, "
            "and what comes out at infinity is thermal. Mr Feynman built an entire, perfectly serious theory "
            "out of adding up arrows for paths no sensible person would take, and it kept being right, "
            "which he found nearly as irritating as I do.",
            "is that a quantum field does not sit politely empty in curved spacetime. "
            "Apply the same summing already on the table to every history the field could have, "
            "and what comes out at infinity is thermal. Mr Feynman made a working calculus of that "
            "arithmetic for paths no sensible person would take, and it kept being right, "
            "which he found nearly as irritating as I do.",
        ),
        "p904 replace failed",
    )
    log.append("arrows/Hawking: unique job = thermal vacuum, not a second tutorial")

    p1104 = find_para(doc, lambda t: "Those relations have a name" in t, "p1104 phase arrows")
    must(
        replace_once(
            p1104,
            "Phase. Picture it the way a better physicist than anyone in this building once did: "
            "every possible way your questionnaire could have gone carries its own small rotating arrow, "
            "and nothing becomes real until you add the arrows tip to tail and see what is left standing. "
            "Length tells you how likely an outcome is.",
            "Phase. The arrows already have two parts, and a vault is built to store only one of them. "
            "Length tells you how likely an outcome is.",
        ),
        "p1104 replace failed",
    )
    log.append("arrows/phase: unique job = length vs angle, vault cannot file phase")

    p1407 = find_para(
        doc, lambda t: "interesting question instead of the frightening one" in t, "p1407 eilstein"
    )
    must(
        replace_once(
            p1407,
            "He does not ask what happened; he asks what could have happened (every route), "
            "gives each one a small turning arrow, and adds them up to see where they point. "
            "It works to more decimal places than anything else I have seen tried, and I cannot tell you why.",
            "He keeps the summing already in play and refuses to wait for a moment when the universe "
            "makes up its mind. No collapse. No privileged instant. Only the total, which happens to work "
            "to more decimal places than anything else I have seen tried, and I cannot tell you why.",
        ),
        "p1407 replace failed",
    )
    log.append("arrows/Eilstein: unique job = calculation without a collapse moment")

    p2029 = find_para(
        doc, lambda t: "a physicist would eventually explain why" in t, "p2029 epilogue arrows"
    )
    set_text(
        p2029,
        "Somewhere else, a physicist would eventually explain why, in the same unhelpful voice Lolly"
        + RSQ
        + "s tutor had used: the form had taken every route it could take, the arithmetic had been "
        "trusted over the intuition, and nothing had been explained, only calculated, because that happened to work.",
    )
    log.append("arrows/epilogue: callback to the tutor, not a fifth tutorial")

    p1811 = find_para(
        doc, lambda t: "A barrier, she thought, was only ever a classical promise" in t, "garden 1811"
    )
    must(
        replace_once(
            p1811,
            "leaked through from the pavement.",
            "leaked through from the pavement. Fainrose had called it tunnelling, over the horizon curves, "
            "which had sounded like a burglary and had turned out to be a law.",
        ),
        "p1811 replace failed",
    )
    log.append("garden square: kept, now callbacks to Lesson 6 tunnelling")

    p2372 = find_para(
        doc, lambda t: "The usual tale says that a pair" in t and "virtual particles" in t,
        "L6 popular version",
    )
    must(
        replace_once(
            p2372,
            "applied to the horizon itself; it reproduces Hawking"
            + RSQ
            + "s temperature exactly, with no factory and no story about where the pair was born.",
            "applied to the horizon itself; it reproduces Hawking"
            + RSQ
            + "s temperature exactly, with no factory and no story about where the pair was born. "
            "Picture a locked London garden square: railings, a key, a view you are not entitled to. "
            "Classically the grass on the other side stays other people"
            + RSQ
            + "s grass. Quantum-mechanically a small amplitude leaks through anyway, "
            "which is the barrier-crossing this derivation applies to the horizon.",
        ),
        "p2372 replace failed",
    )
    log.append("Lesson 6: garden-square analogy for tunnelling (novel scene pays it off)")

    p2600 = find_para(
        doc,
        lambda t: t.startswith("where") and "matter wavelength" in t and "Bohm later" in t,
        "L12 going deeper body",
    )
    must(
        replace_once(
            p2600,
            "so it remains an interpretation rather than an experimentally established replacement.",
            "so it remains an interpretation rather than an experimentally established replacement. "
            "A nucleus at comparable speed has a still shorter wavelength (a proton is about 1836 times "
            "heavier than an electron); that is the precise sense in which chemistry may treat nuclei as "
            "nearly fixed points, while the Sun still fuses by tunnelling. The rest of that argument, "
            "including nuclear shells, lives in the Backlog, because it is true and this lecture is already doing enough.",
        ),
        "p2600 replace failed",
    )
    log.append("Lesson 12 Going deeper: nucleus wavelength clause, backlog keeps the full argument")

    htr = None
    for p in doc.paragraphs:
        if "remaining lectures have no toy on purpose" in p.text or (
            "If a demo is down" in p.text and "lecture still stands" in p.text
        ):
            htr = p
            break
        if p.text.startswith("Demos exist for Lectures 1, 2, 4 and 13"):
            htr = p
    must(htr is not None, "How to Read demos para")
    insert_after(
        htr,
        "Phrases marked in the novel are doors into the Glossary and the lectures, not exits to the open internet. "
        "The videos live in Further Reading, where a reader who wants them can find them on purpose.",
    )
    log.append("How to Read: novel marks -> glossary; YouTube stays in Further Reading")
    return log


NOVEL_YT = {
    "rId106": "glSuperposition",
    "rId107": "glMeasurementBasis",
    "rId108": "Lecture03PrematureCollapse",
    "rId109": "glQuantumZeno",
    "rId110": "Lecture05LawfulEvolution",
    "rId111": "glHawking",
    "rId112": "glEntanglement",
    "rId113": "glNocloning",
    "rId114": "glUncertainty",
    "rId115": "Lecture10Decoherence",
    "rId116": "glMeasurementProblem",
    "rId117": "glPilot",
    "rId118": "glBellstheorem",
    "rId119": "glString",
    "rId120": "glQuantumErrorCorrection",
    "rId121": "glLandauers",
    "rId122": "glNegentropy",
    "rId123": "glSchrdingers",
}


def xml_pass(docx_path: Path) -> list[str]:
    log = []
    buf = BytesIO(docx_path.read_bytes())
    with zipfile.ZipFile(buf, "r") as zin:
        infos = {i.filename: zin.read(i.filename) for i in zin.infolist()}
        namelist = zin.namelist()

    root = etree.fromstring(infos["word/document.xml"])
    relroot = etree.fromstring(infos["word/_rels/document.xml.rels"])
    relmap = {rel.get("Id"): rel.get("Target") for rel in relroot}

    paras = root.findall(f".//{W}p")

    ch3 = ch10 = epi = garden = False
    for p in paras:
        t = "".join(x.text or "" for x in p.findall(f".//{W}t"))
        if "They reached in and fixed a pointer" in t:
            ch3 = rebuild_pointer_ch3(p)
        elif "a moon that just hangs there being a moon is a pointer" in t:
            ch10 = rebuild_pointer_ch10_v2(p)
        elif t.startswith("On a noticeboard in Sub-District 6") and "pointer state" in t:
            epi = wrap_phrase_whole_para(p, "pointer state", "glPointer")
        elif "Fainrose had called it tunnelling" in t:
            garden = wrap_phrase_whole_para(p, "tunnelling", "glQuantumtunnelling")

    log.append(f"pointer wraps: ch3={ch3} ch10={ch10} epilogue={epi} garden={garden}")
    must(ch3 and ch10, "pointer wrap rebuild failed")

    retargeted = 0
    for p in paras:
        t = "".join(x.text or "" for x in p.findall(f".//{W}t"))
        for h in list(p.findall(f".//{W}hyperlink")):
            rid = h.get(f"{R}id")
            if rid in NOVEL_YT:
                tgt = relmap.get(rid, "")
                must("youtube" in (tgt or "").lower(), f"{rid} not youtube")
                if h.get(f"{W}anchor"):
                    continue
                h.attrib.pop(f"{R}id", None)
                h.set(f"{W}anchor", NOVEL_YT[rid])
                h.set(f"{W}history", "1")
                retargeted += 1
    log.append(f"novel YouTube -> glossary/lecture anchors: {retargeted}")
    must(retargeted == 18, f"expected 18 novel youtube retargets, got {retargeted}")

    unwrapped = 0
    for p in paras:
        t = "".join(x.text or "" for x in p.findall(f".//{W}t"))
        if "YouTube" in t or "youtube.com" in t.lower():
            continue
        for h in list(p.findall(f"{W}hyperlink")):
            rid = h.get(f"{R}id")
            tgt = relmap.get(rid, "") if rid else ""
            if rid and "youtube" in tgt.lower():
                parent = h.getparent()
                for child in list(h):
                    if child.tag == f"{W}r":
                        rpr = child.find(f"{W}rPr")
                        if rpr is not None:
                            for rs in list(rpr.findall(f"{W}rStyle")):
                                if rs.get(f"{W}val") == "Hyperlink":
                                    rpr.remove(rs)
                    parent.insert(list(parent).index(h), child)
                parent.remove(h)
                unwrapped += 1
    log.append(f"unwrapped YouTube-from-book-authors: {unwrapped}")
    must(unwrapped >= 10, f"expected book-author youtube unwraps, got {unwrapped}")

    infos["word/document.xml"] = etree.tostring(
        root, xml_declaration=True, encoding="UTF-8", standalone=True
    )

    out_buf = BytesIO()
    with zipfile.ZipFile(out_buf, "w", zipfile.ZIP_DEFLATED) as zout:
        with zipfile.ZipFile(docx_path, "r") as zin:
            for item in zin.infolist():
                data = infos.get(item.filename, zin.read(item.filename))
                zout.writestr(item, data)
    docx_path.write_bytes(out_buf.getvalue())
    return log


def verify(path: Path) -> None:
    d = Document(str(path))
    blob = "\n".join(p.text for p in d.paragraphs)
    checks = {
        "answer that": "an answer that survives being asked" in blob,
        "no answerthat": "answerthat" not in blob,
        "title ministry": any(
            p.text.strip() == "A novel of the Ministry of Eventualities" for p in d.paragraphs[:15]
        ),
        "kdp series kept": "Book 1 of Lolly Wren" in blob and "Curious Science Adventures" in blob,
        "hawking no rotating arrow": "vulgar public explanation" in blob
        and "weight each by its own small rotating arrow" not in blob,
        "phase two parts": "The arrows already have two parts" in blob,
        "eilstein no collapse": "No collapse. No privileged instant." in blob,
        "epilogue tutor": "same unhelpful voice Lolly" in blob and "tutor had used" in blob,
        "garden callback": "Fainrose had called it tunnelling" in blob,
        "L6 garden": "locked London garden square" in blob,
        "L12 nucleus": "A nucleus at comparable speed" in blob and "1836 times" in blob,
        "htr glossary doors": "doors into the Glossary" in blob,
        "tutor seed kept": "each weighted by a small rotating arrow" in blob,
        "backlog proton kept": "On the wavelength of a proton" in blob,
    }
    failed = [k for k, v in checks.items() if not v]
    if failed:
        raise SystemExit("VERIFY FAILED: " + ", ".join(failed))

    with zipfile.ZipFile(path) as z:
        root = etree.fromstring(z.read("word/document.xml"))
        relroot = etree.fromstring(z.read("word/_rels/document.xml.rels"))
        relmap = {rel.get("Id"): rel.get("Target") for rel in relroot}

    paras = root.findall(f".//{W}p")
    fr_i = None
    for i, p in enumerate(paras):
        t = "".join(x.text or "" for x in p.findall(f".//{W}t"))
        if t.strip() == "Further Reading" or t.startswith("Further Reading"):
            # first heading, not TOC
            if i > 200:
                fr_i = i
                break
    must(fr_i is not None, "Further Reading heading")

    novel_yt = 0
    for i, p in enumerate(paras):
        if i >= fr_i:
            break
        t = "".join(x.text or "" for x in p.findall(f".//{W}t"))
        if t.strip() in ("Glossary", "Backlog"):
            # glossary videos stay
            if t.strip() == "Glossary":
                break
        for h in p.findall(f".//{W}hyperlink"):
            rid = h.get(f"{R}id")
            tgt = relmap.get(rid, "") if rid else ""
            if tgt and "youtube" in tgt.lower():
                novel_yt += 1
                print("LEAK", i, tgt, t[:80])
    must(novel_yt == 0, f"youtube still in novel body: {novel_yt}")

    vis_ok = 0
    for p in paras:
        for h in p.findall(f".//{W}hyperlink"):
            if h.get(f"{W}anchor") == "glPointer":
                vis = "".join(x.text or "" for x in h.findall(f".//{W}t"))
                if vis in ("pointer state", "pointer states"):
                    vis_ok += 1
                else:
                    raise SystemExit(f"bad glPointer vis {vis!r}")
    must(vis_ok >= 3, f"glPointer wraps {vis_ok}")

    glued = 0
    for p in paras:
        for h in p.findall(f"{W}hyperlink"):
            vis = "".join(x.text or "" for x in h.findall(f".//{W}t"))
            if not vis:
                continue
            prev = h.getprevious()
            nxt = h.getnext()
            prev_txt = ""
            if prev is not None and prev.tag == f"{W}r":
                prev_txt = "".join(x.text or "" for x in prev.findall(f"{W}t"))
            next_txt = ""
            if nxt is not None and nxt.tag == f"{W}r":
                next_txt = "".join(x.text or "" for x in nxt.findall(f"{W}t"))
            if prev_txt and prev_txt[-1].isalpha() and vis[0].isalpha():
                glued += 1
            if next_txt and vis[-1].isalpha() and next_txt[0].isalpha():
                glued += 1
    must(glued == 0, f"still glued wraps: {glued}")
    print("VERIFY OK paras", len(d.paragraphs), "xml", len(paras))


GUIDE_ADDENDUM = """
## Recast continue-fix (28 Sep 2026)

Same recast-not-cut rule. Applied to the live Kindle file.

- Ch3/Ch10 `pointer state` glossary wraps repaired (they had split mid-word: `pointer s`+`tate`, `pointer st`+`ate: an answer`+`that`). Space restored: `an answer that survives`. Epilogue list-as-pointer wrapped cleanly to the same glossary door.
- In-novel YouTube hyperlinks (18 stealth links on chapter phrases) retargeted to Glossary/lecture bookmarks. How to Read now says marked phrases are glossary doors; videos stay in Further Reading, where they belong. Book-author names in the bibliography no longer jump to unrelated YouTube interviews.
- Feynman-arrow visits given unique jobs: the tutor still plants the method; Hawking uses the summing for a thermal vacuum; Fainrose uses length vs angle (phase the vault cannot file); Eilstein uses calculation without a collapse moment; the epilogue callbacks the tutor instead of teaching a fifth time.
- Title page series line recast to `A novel of the Ministry of Eventualities`. KDP series name kept in About the Series.
- Lesson 12 Going deeper now carries the nucleus-wavelength clause (proton ~1836× an electron) and points at the Backlog, which still holds the full argument (Born–Oppenheimer, solar tunnelling, nuclear shells).
- Garden square kept. Lesson 6 now uses a locked London square as the tunnelling analogy; the corridor scene callbacks Fainrose's word for it.

"""


def update_guide(path: Path) -> None:
    if not path.exists():
        return
    text = path.read_text(encoding="utf-8")
    if "Recast continue-fix (28 Sep 2026)" in text:
        return
    path.write_text(text.rstrip() + "\n" + GUIDE_ADDENDUM, encoding="utf-8")


def main() -> None:
    if not LIVE.exists():
        raise SystemExit(f"missing {LIVE}")
    if OUT.exists():
        OUT.unlink()
    shutil.copy2(LIVE, OUT)

    doc = Document(str(OUT))
    n0 = len(doc.paragraphs)
    log = prose_pass(doc)
    doc.save(str(OUT))
    log.extend(xml_pass(OUT))
    verify(OUT)

    try:
        if LIVE.exists():
            LIVE.unlink()
        shutil.move(str(OUT), str(LIVE))
    except PermissionError:
        print("LOCKED: patched copy waiting at", OUT)
        for line in log:
            print(" ", line)
        return

    print("paragraphs", n0, "->", len(Document(str(LIVE)).paragraphs))
    for line in log:
        print(" ", line)
    print("wrote", LIVE, "mb", round(LIVE.stat().st_size / 1e6, 2))
    if WS.exists():
        try:
            shutil.copy2(LIVE, WS)
            print("copied to workspace KINDLE file")
        except PermissionError:
            print("workspace KINDLE locked; series copy is live")
    update_guide(GUIDE)
    update_guide(GUIDE_WS)
    print("guides updated")


if __name__ == "__main__":
    main()
