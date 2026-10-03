# -*- coding: utf-8 -*-
from pathlib import Path
from docx import Document

LIVE = Path(r"D:\My Books for Amazon") / "Lolly Wren's Curious Science Adventures - SERIES" / "Book 1 - Schrodingers Paperwork" / "Schrodingers_Paperwork_BOOK_1_KINDLE_FINAL.docx"
OUT = LIVE.parent / "bak" / (LIVE.name + ".next")
src = OUT if OUT.exists() else LIVE
d = Document(str(src))
print("src", src.name, "paras", len(d.paragraphs))
for i, p in enumerate(d.paragraphs[1270:1295]):
    if "grey" in p.text.lower() or "gray" in p.text.lower() or "Prosser" in p.text or "apron" in p.text:
        print(f"{1270+i}|{p.text}")
