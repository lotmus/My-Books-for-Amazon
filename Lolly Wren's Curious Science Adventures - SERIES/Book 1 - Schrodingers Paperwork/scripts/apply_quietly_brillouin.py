# -*- coding: utf-8 -*-
"""Recast leftover 'quietly' dialogue tags; employ Brillouin in Ch18."""
from __future__ import annotations

import shutil
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

SERIES = Path(
    r"D:\My Books for Amazon"
    r"\Lolly Wren's Curious Science Adventures - SERIES"
    r"\Book 1 - Schrodingers Paperwork"
)
LIVE = SERIES / "Schrodingers_Paperwork_BOOK_1_KINDLE_FINAL.docx"
OUT = SERIES / "bak" / "Schrodingers_Paperwork_BOOK_1_KINDLE_FINAL.docx.next"
WS = Path(
    r"D:\My Books for Amazon"
    r"\Schrodingers_Paperwork"
    r"\Schrodingers_Paperwork_BOOK_1_KINDLE_FINAL.docx"
)
GUIDE = SERIES / "notes" / "MANUSCRIPT_GUIDE.md"
GUIDE_WS = Path(
    r"D:\My Books for Amazon"
    r"\Schrodingers_Paperwork"
    r"\MANUSCRIPT_GUIDE.md"
)

LDQ = "\u201c"
RDQ = "\u201d"
LSQ = "\u2018"
RSQ = "\u2019"
EM = "\u2014"


def set_text(p, new: str) -> None:
    runs = p.runs
    if not runs:
        p.add_run(new)
        return
    runs[0].text = new
    for r in runs[1:]:
        r.text = ""
    texts = p._p.findall(".//" + qn("w:t"))
    joined = "".join(t.text or "" for t in texts)
    if joined != new:
        for i, t in enumerate(texts):
            t.text = new if i == 0 else ""


def replace_once(p, old: str, new: str) -> bool:
    full = p.text
    if old not in full:
        return False
    for r in p.runs:
        if old in (r.text or ""):
            r.text = (r.text or "").replace(old, new, 1)
            return True
    set_text(p, full.replace(old, new, 1))
    return True


def must(cond, msg):
    if not cond:
        raise SystemExit(msg)


def find_para(doc, pred, msg):
    for p in doc.paragraphs:
        if pred(p.text):
            return p
    raise SystemExit(msg)


def req(p, old, new, label):
    must(replace_once(p, old, new), f"{label}: not found\n  {old[:90]!r}\n  in {p.text[:140]!r}")
    return label


def main() -> None:
    if not LIVE.exists():
        raise SystemExit(f"missing {LIVE}")
    if OUT.exists():
        OUT.unlink()
    shutil.copy2(LIVE, OUT)
    doc = Document(str(OUT))
    n0 = len(doc.paragraphs)
    log = []

    log.append(
        req(
            find_para(doc, lambda t: "Lolly agreed quietly" in t, "Lolly 216"),
            LDQ + "No," + RDQ + " Lolly agreed quietly. " + LDQ + "One oughtn" + RSQ + "t." + RDQ,
            LDQ + "No," + RDQ + " Lolly agreed, as if confirming a line she had already written. "
            + LDQ + "One oughtn" + RSQ + "t." + RDQ,
            "quietly/Lolly agree: notebook confirmation",
        )
    )
    log.append(
        req(
            find_para(doc, lambda t: "Lolly ventured quietly" in t, "Lolly 365"),
            "Lolly ventured quietly,",
            "Lolly ventured, without looking up from the lanes,",
            "quietly/Lolly venture: looking at the lanes",
        )
    )
    log.append(
        req(
            find_para(doc, lambda t: "Lolly noted quietly" in t, "Lolly 532"),
            LDQ + "It" + RSQ + "s been averaged," + RDQ + " Lolly noted quietly.",
            LDQ + "It" + RSQ + "s been averaged," + RDQ + " Lolly noted, hoping the notebook would contradict her.",
            "quietly/Lolly averaged: notebook hoped to contradict",
        )
    )
    log.append(
        req(
            find_para(doc, lambda t: "Bloody hell" in t and "managed, quietly" in t, "Bloody hell"),
            LDQ + "Bloody hell," + RDQ + " she managed, quietly.",
            LDQ + "Bloody hell," + RDQ + " she managed, without giving the room the satisfaction of a shout.",
            "quietly/Lolly bloody hell: no shout",
        )
    )
    log.append(
        req(
            find_para(doc, lambda t: "admitted, quietly" in t and "not domestic" in t, "Fainrose wrong"),
            LDQ + "I was wrong," + RDQ + " she admitted, quietly. " + LDQ + "That is not domestic." + RDQ,
            LDQ + "I was wrong," + RDQ + " she admitted, to the screen. " + LDQ + "That is not domestic." + RDQ,
            "quietly/Fainrose: to the screen",
        )
    )
    log.append(
        req(
            find_para(doc, lambda t: "Fainrose ventured, quietly" in t, "Fainrose decoherence"),
            "Fainrose ventured, quietly,",
            "Fainrose ventured, setting the word down carefully,",
            "quietly/Fainrose decoherence: word set down",
        )
    )
    log.append(
        req(
            find_para(doc, lambda t: "noted de Broccoli, quietly" in t, "de Broccoli"),
            LDQ + "Yes," + RDQ + " noted de Broccoli, quietly.",
            LDQ + "Yes," + RDQ + " noted de Broccoli, as if correcting a menu.",
            "quietly/de Broccoli: correcting a menu",
        )
    )
    log.append(
        req(
            find_para(doc, lambda t: "mentioned Max Tengelman quietly" in t, "Tengelman"),
            LDQ + "That" + RSQ + "s fair," + RDQ + " mentioned Max Tengelman quietly.",
            LDQ + "That" + RSQ + "s fair," + RDQ + " mentioned Max Tengelman, surprising himself by not talking over it.",
            "quietly/Tengelman: actually yields the floor",
        )
    )
    log.append(
        req(
            find_para(doc, lambda t: "Heisenburger noted, quietly" in t, "Heisenburger"),
            "Heisenburger noted, quietly,",
            "Heisenburger noted, with nothing underlined,",
            "quietly/Heisenburger: nothing underlined",
        )
    )
    log.append(
        req(
            find_para(doc, lambda t: "He noted it, quietly." in t, "Wittenberg number"),
            "He noted it, quietly.",
            "He let the number sit.",
            "quietly/von Wittenberg: number sits; claim does volume",
        )
    )

    p1942 = find_para(
        doc,
        lambda t: "negentropy, if you want the ugly word for it" in t and "Ministry salary" in t,
        "Schrottfinger negentropy",
    )
    log.append(
        req(
            p1942,
            "feeding on order "
            + EM
            + " negentropy, if you want the ugly word for it "
            + EM
            + " taking in structure",
            "feeding on order (negentropy, if you want the ugly word for it; a Frenchman named Brillouin later said I had meant information, a shape that stays a shape, and he was not wrong), taking in structure",
            "Brillouin/Ch18: Schrottfinger allows the Frenchman's reading; still a bill",
        )
    )

    doc.save(str(OUT))
    chk = Document(str(OUT))
    blob = "\n".join(p.text for p in chk.paragraphs)
    checks = {
        "lolly written": "confirming a line she had already written" in blob,
        "lanes": "without looking up from the lanes" in blob,
        "notebook contradict": "hoping the notebook would contradict her" in blob,
        "no shout": "without giving the room the satisfaction of a shout" in blob,
        "to the screen": "she admitted, to the screen" in blob,
        "word down": "setting the word down carefully" in blob,
        "menu": "as if correcting a menu" in blob,
        "tengelman yield": "surprising himself by not talking over it" in blob,
        "nothing underlined": "with nothing underlined" in blob,
        "number sit": "He let the number sit." in blob,
        "brillouin novel": "a Frenchman named Brillouin later said I had meant information" in blob,
        "lecture brillouin kept": "Léon Brillouin later read" in blob,
        "no agreed quietly": "agreed quietly" not in blob,
        "no noted quietly": "noted quietly" not in blob,
        "no ventured quietly": "ventured quietly" not in blob,
        "no admitted quietly": "admitted, quietly" not in blob,
        "institutional quietly kept": "quietly moving complexity" in blob,
        "injunction quietly kept": "renewed, quietly, every six months" in blob,
    }
    failed = [k for k, v in checks.items() if not v]
    if failed:
        raise SystemExit("VERIFY FAILED: " + ", ".join(failed))

    try:
        if LIVE.exists():
            LIVE.unlink()
        shutil.move(str(OUT), str(LIVE))
    except PermissionError:
        print("LOCKED: patched copy waiting at", OUT)
        for line in log:
            print(" ", line)
        return

    print("paragraphs", n0, "->", len(Document(str(LIVE)).paragraphs))
    for line in log:
        print(" ", line)
    print("wrote", LIVE, "mb", round(LIVE.stat().st_size / 1e6, 2))
    try:
        shutil.copy2(LIVE, WS)
        print("copied to workspace KINDLE file")
    except PermissionError:
        print("workspace KINDLE locked; series copy is live")

    addendum = """
## Quietly / Brillouin pass (28 Sep 2026)

Recast-not-cut. The leftover shared adverb after the voice-intro pass was `quietly` on dialogue tags: Lolly, Fainrose, de Broccoli, Tengelman, Heisenburger, and von Wittenberg were all "speaking quietly." Each tag now does that character's job instead (notebook confirmation; looking at the lanes; hoping the notebook will contradict her; no shout; to the screen; word set down; correcting a menu; Tengelman actually yielding the floor; nothing underlined; the number sits). Institutional/atmospheric quietly kept (complexity moved, email forgotten, injunction renewed, Bellboy's entrance, Venn reassigned).

Brillouin employed in Ch18: Schrottfinger allows the Frenchman's information-reading of negentropy (`a shape that stays a shape`) and does not take it as a loophole. Lesson 17 keeps the named lecture account.

"""
    for g in (GUIDE, GUIDE_WS):
        if g.exists():
            t = g.read_text(encoding="utf-8")
            if "Quietly / Brillouin pass (28 Sep 2026)" not in t:
                g.write_text(t.rstrip() + "\n" + addendum, encoding="utf-8")
    print("guides updated")


if __name__ == "__main__":
    main()
