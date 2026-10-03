# -*- coding: utf-8 -*-
from pathlib import Path
from docx import Document

p = Path(
    r"D:\My Books for Amazon"
    r"\Lolly Wren's Curious Science Adventures - SERIES"
    r"\Book 1 - Schrodingers Paperwork"
    r"\Schrodingers_Paperwork_BOOK_1_KINDLE_FINAL.docx"
)
d = Document(str(p))
out = Path(__file__).with_name("inspect_paras3.txt")
lines = []
for i in range(1245, 1256):
    t = d.paragraphs[i].text.strip()
    if t:
        lines.append(f"{i} {t[:400]}")
lines.append("--- banners ---")
for i, para in enumerate(d.paragraphs):
    t = para.text.strip()
    if t in {
        "What happens when you file the world (Lessons 6–10)",
        "What a state is, and what it means to ask (Lessons 1–5)",
        "What the experts still fight about (Lessons 11, 12 and 14)",
        "What they can still measure (Lesson 13)",
    } or t.startswith("What can be recovered") or t.startswith("What they can still") or (
        t.startswith("What ") and "Lesson" in t and len(t) < 90
    ):
        lines.append(f"{i} {t}")
lines.append("--- 2806 full ---")
lines.append(d.paragraphs[2806].text)
lines.append("--- 2736 full ---")
lines.append(d.paragraphs[2736].text)
lines.append("--- 2843 full ---")
lines.append(d.paragraphs[2843].text)
lines.append("--- 2005 full ---")
lines.append(d.paragraphs[2005].text)
lines.append("--- 112 full ---")
lines.append(d.paragraphs[112].text)
out.write_text("\n".join(lines), encoding="utf-8")
print("ok")
