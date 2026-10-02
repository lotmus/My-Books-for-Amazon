# -*- coding: utf-8 -*-
from pathlib import Path
from docx import Document
import re

LIVE = Path(
    r"c:\Users\lomus\OneDrive\My Books for Amazon"
    r"\Lolly Wren's Curious Science Adventures - SERIES"
    r"\Book 1 - Schrodingers Paperwork"
    r"\Schrodingers_Paperwork_BOOK_1_KINDLE_FINAL.docx"
)
d = Document(str(LIVE))
out = Path(__file__).with_name("voice_grey_scan.txt")
lines = []

VOICE_KEYS = (
    "unhurried", "precise", "flat and", "pitched low", "in a low", "quietly",
    "dry,", "dry.", "voice was", "voice matched", "register", "without raising",
    "not unkindly", "mildly", "faintly amused", "built for", "built, it seemed",
    "quietly essential", "quietly,", "in the same unhelpful",
)
CAST_START = None
for i, p in enumerate(d.paragraphs):
    if p.text.strip() == "Cast of Characters":
        CAST_START = i
        break

lines.append("=== CAST ===")
for i in range(CAST_START, CAST_START + 30):
    t = d.paragraphs[i].text.strip()
    if t:
        lines.append(f"{i}|{t[:400]}")
    if t.startswith("Nelson"):
        break

lines.append("\n=== VOICE / MANNER DESCRIPTORS (novel <2100) ===")
for i, p in enumerate(d.paragraphs[:2100]):
    t = p.text
    low = t.lower()
    if any(k in low for k in VOICE_KEYS) or "her voice" in low or "his voice" in low or "the voice" in low:
        if "Cast of Characters" in t or t.startswith("Lesson "):
            continue
        lines.append(f"{i}|{t[:420]}")

lines.append("\n=== DIALOGUE TAG VERBS sample (noted/remarked/observed/considered/stated) ===")
from collections import Counter
c = Counter()
pat = re.compile(
    r"\b(said|remarked|observed|considered|reflected|noted|stated|asked|replied|"
    r"announced|murmured|ventured|managed|finished|added|continued|explained|"
    r"insisted|declared|offered|muttered|whispered|snapped|dryly)\b",
    re.I,
)
# too heavy; just count in first 2100
for p in d.paragraphs[:2100]:
    for m in pat.findall(p.text):
        c[m.lower()] += 1
lines.append(str(c.most_common(30)))

lines.append("\n=== GREY ALL ===")
for i, p in enumerate(d.paragraphs):
    if re.search(r"\bgre[ya]\b", p.text, re.I):
        lines.append(f"{i}|{p.text[:280]}")

lines.append("\n=== UNRUH / ACCELERAT / THERMOMETER / VACUUM NOT EMPTY ===")
for i, p in enumerate(d.paragraphs):
    t = p.text
    if any(k in t.lower() for k in ("unruh", "accelerating thermometer", "legacy emptiness", "vacuum that is not empty")):
        lines.append(f"{i}|{t[:420]}")

lines.append("\n=== BRILLOUIN ===")
for i, p in enumerate(d.paragraphs):
    if "Brillouin" in p.text:
        lines.append(f"{i}|{p.text[:350]}")

# character first-speech samples
names = [
    "Lolly", "Mrs Chain", "Fainrose", "Beatrix", "Jago", "Gideon", "Priddy",
    "Venn", "Voller", "Pike", "Eilstein", "Bellboy", "Broccoli", "Heisenburger",
    "Tengelman", "Wittenberg", "Schrottfinger", "Pilbeam", "Prosser",
]
lines.append("\n=== FIRST INTRO LINES ===")
seen = set()
for i, p in enumerate(d.paragraphs[:2100]):
    t = p.text
    for n in names:
        if n in seen:
            continue
        if n == "Lolly" and t.startswith("Lolly Wren"):
            lines.append(f"{n} {i}|{t[:300]}")
            seen.add(n)
        elif n != "Lolly" and (t.startswith(n) or f"{n} " in t[:40]):
            if i < 140:  # skip TOC
                continue
            lines.append(f"{n} {i}|{t[:300]}")
            seen.add(n)

out.write_text("\n".join(lines), encoding="utf-8")
print("wrote", out, "lines", len(lines))
