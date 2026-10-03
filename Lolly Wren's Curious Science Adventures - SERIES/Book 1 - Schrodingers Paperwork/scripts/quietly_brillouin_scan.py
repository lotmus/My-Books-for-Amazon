# -*- coding: utf-8 -*-
from pathlib import Path
from docx import Document
import re

LIVE = Path(
    r"D:\My Books for Amazon"
    r"\Lolly Wren's Curious Science Adventures - SERIES"
    r"\Book 1 - Schrodingers Paperwork"
    r"\Schrodingers_Paperwork_BOOK_1_KINDLE_FINAL.docx"
)
d = Document(str(LIVE))
out = Path(__file__).with_name("quietly_brillouin_scan.txt")
lines = [f"paras {len(d.paragraphs)}"]

lines.append("\n=== quietly (first 2100) ===")
for i, p in enumerate(d.paragraphs[:2100]):
    if re.search(r"\bquietly\b", p.text, re.I):
        lines.append(f"{i}|{p.text[:380]}")

lines.append("\n=== Brillouin / Landauer / negentrop in novel ===")
for i, p in enumerate(d.paragraphs):
    t = p.text
    if any(k in t for k in ("Brillouin", "Landauer", "negentrop")):
        lines.append(f"{i}|{t[:400]}")

lines.append("\n=== leftover unhurried / pitched low / leaning ===")
for i, p in enumerate(d.paragraphs[:2100]):
    t = p.text.lower()
    if "unhurried" in t or "pitched low" in t or "leaning in" in t or "leaning forward" in t:
        lines.append(f"{i}|{p.text[:300]}")

lines.append("\n=== Schrottfinger / Ch17-18 Mrs Chain obituary window ===")
for i, p in enumerate(d.paragraphs):
    t = p.text
    if i < 1850 or i > 1980:
        continue
    if any(k in t.lower() for k in ("obituar", "negentrop", "what is life", "entropy", "schrott", "lid")):
        lines.append(f"{i}|{t[:350]}")

out.write_text("\n".join(lines), encoding="utf-8")
print("wrote", out, "n", len(lines))
