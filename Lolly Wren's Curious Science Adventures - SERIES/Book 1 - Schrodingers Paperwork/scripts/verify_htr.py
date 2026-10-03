# -*- coding: utf-8 -*-
from pathlib import Path
from docx import Document

p = Path(
    r"D:\My Books for Amazon"
    r"\Lolly Wren's Curious Science Adventures - SERIES"
    r"\Book 1 - Schrodingers Paperwork"
    r"\Schrodingers_Paperwork_BOOK_1_KINDLE_FINAL.docx"
)
d = Document(str(p))
out = Path(__file__).with_name("verify_htr.txt")
lines = []
for i in range(80, 102):
    para = d.paragraphs[i]
    runs = " | ".join(f"{'B' if r.bold else 'n'}:{(r.text or '')[:50]!r}" for r in para.runs[:6])
    lines.append(f"{i} [{para.style.name}] {para.text[:180]!r}")
    if runs:
        lines.append("   " + runs)
for name in ["Nelson", "Gideon, who had been standing"]:
    for i, para in enumerate(d.paragraphs):
        if name in para.text:
            lines.append(f"\nFOUND {i} {para.text[:250]!r}")
            break
# story lock run split
for i in (2735, 2805, 2842):
    para = d.paragraphs[i]
    lines.append(f"\nLOCK {i} nruns={len(para.runs)}")
    for r in para.runs:
        lines.append(f"  bold={r.bold} {r.text[:80]!r}")
out.write_text("\n".join(lines), encoding="utf-8")
print("ok")
