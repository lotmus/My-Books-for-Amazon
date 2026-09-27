# -*- coding: utf-8 -*-
import re
import sys
from pathlib import Path
from docx import Document

sys.stdout.reconfigure(encoding="utf-8")
out = Path(r"D:\My Books for Amazon\Schrodingers_Paperwork\_audit_appendix_checks2_out.txt")
lines = []

def P(s=""):
    lines.append(s)

doc = Document(r"D:\My Books for Amazon\Schrodingers_Paperwork\Schrodingers_Paperwork_BOOK_1.docx")

P("=== rarer / lid / lesson / simplifying ===")
for i, p in enumerate(doc.paragraphs):
    t = p.text
    if re.search(r"rarer than gravity|this lesson|There is no lid\.|simplifying.? accounts", t, re.I):
        P(f"[{i}] {t.strip()[:220]}")

P("\n=== TITLE MISMATCHES ===")
ch_headers = {}
lec_headers = {}
for i, p in enumerate(doc.paragraphs):
    t = p.text.strip()
    m = re.match(r"Chapter (\d+): Lecture(?: \(Part (One|Two) of Two\))?: (.+)$", t)
    if m:
        ch_headers[(m.group(1), m.group(2) or "")] = (i, m.group(3))
    m2 = re.match(r"Lecture (\d+) —(?: Part (One|Two):)? (.+)$", t)
    if m2:
        lec_headers[(m2.group(1), m2.group(2) or "")] = (i, m2.group(3))
    if t.startswith("→ Lecture for this chapter:"):
        P(f"LINK [{i}] {t}")

P("\nCH vs LEC:")
keys = sorted(set(list(ch_headers) + list(lec_headers)), key=lambda x: (int(x[0]), x[1]))
for k in keys:
    ch = ch_headers.get(k)
    lec = lec_headers.get(k)
    if ch and lec and ch[1] != lec[1]:
        P(f"  MISMATCH L{k}: CH=\"{ch[1]}\" vs LEC=\"{lec[1]}\"")
    elif ch and lec:
        P(f"  OK L{k}: {lec[1][:70]}")
    else:
        P(f"  PARTIAL L{k}: ch={ch} lec={lec}")

P("\n=== SECTION WORD COUNTS ===")
paras = [(i, p.text.strip()) for i, p in enumerate(doc.paragraphs)]
starts = [i for i, t in paras if t.startswith("Lecture ") and "—" in t and i >= 2164]
starts.append(2673)
for a, b in zip(starts, starts[1:]):
    header = paras[a][1]
    buckets = {"start": 0, "steps": 0, "chapter": 0, "wrong": 0, "deeper": 0, "breath": 0, "sticks": 0}
    mode = None
    for i in range(a, b):
        t = paras[i][1]
        if not t:
            continue
        if t.startswith("In one breath"):
            mode = "breath"
            buckets["breath"] += len(t.split())
            continue
        if t.startswith("Start here"):
            mode = "start"
            buckets["start"] += len(t.split())
            continue
        if t.startswith("The idea, step by step"):
            mode = "steps"
            continue
        if t.startswith("What this chapter"):
            mode = "chapter"
            buckets["chapter"] += len(t.split())
            continue
        if t.startswith("Where the popular"):
            mode = "wrong"
            buckets["wrong"] += len(t.split())
            continue
        if t.startswith("Going deeper"):
            mode = "deeper"
            buckets["deeper"] += len(t.split())
            continue
        if t.startswith("What sticks"):
            mode = "sticks"
            continue
        if t.startswith(("One-line", "Stop.", "Pretend", "You need", "When you walk", "Lecture ", "Chapter ", "←", "•")):
            if t.startswith("•") and mode == "sticks":
                buckets["sticks"] += len(t.split())
                continue
            mode = None
            continue
        if mode and mode in buckets:
            buckets[mode] += len(t.split())
    P(
        f"{header[:48]:48} br={buckets['breath']:3} st={buckets['start']:3} "
        f"steps={buckets['steps']:3} ch={buckets['chapter']:3} wr={buckets['wrong']:3} "
        f"dp={buckets['deeper']:3}"
    )

# Flag numbered step lines that are pure definition (is/are/means)
P("\n=== DRIEST STEP LINES (definitional openers) ===")
for i, p in enumerate(doc.paragraphs):
    if i < 2164 or i > 2672:
        continue
    t = p.text.strip()
    if re.match(r"^\d+\.\s+", t):
        # dry if starts with Noun is / X is a
        if re.match(r"^\d+\.\s+[A-Z][^.]{0,40}\s+(is|are|means|refers|describes)\b", t) and "you" not in t.lower() and "picture" not in t.lower() and "imagine" not in t.lower():
            if len(t) > 70:
                P(f"[{i}] {t[:140]}")

out.write_text("\n".join(lines) + "\n", encoding="utf-8")
print(f"Wrote {out} ({len(lines)} lines)")
