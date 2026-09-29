# -*- coding: utf-8 -*-
from pathlib import Path
from docx import Document
from docx.oxml.ns import qn

p = Path(
    r"c:\Users\lomus\OneDrive\My Books for Amazon"
    r"\Lolly Wren's Curious Science Adventures - SERIES"
    r"\Book 1 - Schrodingers Paperwork"
    r"\Schrodingers_Paperwork_BOOK_1_KINDLE_FINAL.docx"
)
d = Document(str(p))
out = Path(__file__).with_name("verify_recast.txt")
lines = []

def dump(i, n=1):
    for j in range(i, i + n):
        para = d.paragraphs[j]
        bolds = []
        for r in para.runs:
            bolds.append(f"{'B' if r.bold else 'n'}:{r.text[:40]!r}")
        lines.append(f"{j} [{para.style.name}] {para.text[:220]!r}")
        if bolds:
            lines.append("   " + " | ".join(bolds[:8]))

lines.append("=== How to Read layout ===")
for i, para in enumerate(d.paragraphs):
    if para.text.strip() == "How to Read This Book":
        dump(i, 20)
        break

lines.append("\n=== Shape of Course titles ===")
for i, para in enumerate(d.paragraphs):
    t = para.text.strip()
    if t.startswith("What happens when you file") or t.startswith("What the experts still") or t.startswith("What they can still") or t.startswith("What can be recovered") or t.startswith("What a state is"):
        dump(i, 2)

lines.append("\n=== story locks ===")
for i, para in enumerate(d.paragraphs):
    if para.text.startswith("What this chapter was actually showing you. The first Bell"):
        dump(i)
    if para.text.startswith("What this chapter was actually showing you. Lolly delivers"):
        dump(i)
    if para.text.startswith("What this chapter was actually showing you. The scene has already"):
        dump(i)

lines.append("\n=== sample inserts ===")
needles = ["feed-rate clock", "She did go downstairs", "car battery and a van", "hospital door", "one extra chair", "Miss Dorothy Kell", "The log climbed"]
for i, para in enumerate(d.paragraphs):
    for n in needles:
        if n in para.text:
            dump(i)
            break

lines.append("\n=== leftover What sticks in How to Read ===")
blob = "\n".join(p.text for p in d.paragraphs[:120])
for s in ["What sticks", "One-line summary", "Say it to Mrs Chain", "four parts", "five parts"]:
    lines.append(f"{s}: {s in blob}")

out.write_text("\n".join(lines), encoding="utf-8")
print("wrote", out)
