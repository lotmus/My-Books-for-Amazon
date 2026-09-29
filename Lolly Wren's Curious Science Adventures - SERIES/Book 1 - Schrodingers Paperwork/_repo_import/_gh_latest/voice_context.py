# -*- coding: utf-8 -*-
from pathlib import Path
from docx import Document

LIVE = Path(
    r"c:\Users\lomus\OneDrive\My Books for Amazon"
    r"\Lolly Wren's Curious Science Adventures - SERIES"
    r"\Book 1 - Schrodingers Paperwork"
    r"\Schrodingers_Paperwork_BOOK_1_KINDLE_FINAL.docx"
)
d = Document(str(LIVE))
out = Path(__file__).with_name("voice_context.txt")
lines = []

def dump(idxs, label):
    lines.append(f"\n===== {label} =====")
    for i in idxs:
        if 0 <= i < len(d.paragraphs):
            t = d.paragraphs[i].text
            if t.strip():
                lines.append(f"{i}|{t}")

# find Cast of Characters heading not TOC
for i, p in enumerate(d.paragraphs):
    if p.text.strip() == "Cast of Characters" and i > 80:
        dump(range(i, i+22), "CAST BODY")
        break

dump(range(198, 215), "LOLLY+CHAIN VOICE")
dump(range(400, 420), "JAGO? VOICE 410")
dump(range(800, 820), "FAINROSE VOICE")
dump(range(780, 795), "BEATRIX 788")
dump(range(750, 765), "PRIDDY")
dump(range(1148, 1162), "GIDEON 1157")
dump(range(1328, 1345), "EILSTEIN VOICE")
dump(range(1444, 1456), "DE BROCCOLI")
dump(range(1545, 1560), "BELLBOY/HEISEN")
dump(range(2368, 2380), "L6 going deeper")
dump(range(2660, 2720), "L14 area")
dump(range(2868, 2876), "BACKLOG vacuum")
dump(range(3048, 3058), "GLOSSARY Unruh")

# search Lesson 6/14 for Unruh, accelerating, thermometer
lines.append("\n===== Unruh-adjacent in lectures =====")
for i, p in enumerate(d.paragraphs):
    t = p.text
    if i < 2100:
        continue
    if any(k in t.lower() for k in ("unruh", "accelerat", "inertial observer", "warm bath", "thermometer")):
        lines.append(f"{i}|{t[:400]}")

# Venn intro
lines.append("\n===== Venn intro =====")
for i, p in enumerate(d.paragraphs[:900]):
    t = p.text
    if "Alistair Venn" in t or (t.startswith("The man") and "spiral" in t) or "Senior Review Officer" in t:
        lines.append(f"{i}|{t[:350]}")

# Bellboy voice description
lines.append("\n===== Bellboy voice phrases =====")
for i, p in enumerate(d.paragraphs):
    t = p.text
    if "Bellboy" in t and any(k in t.lower() for k in ("voice", "unhurried", "blunt", "accent", "belfast")):
        lines.append(f"{i}|{t[:350]}")

out.write_text("\n".join(lines), encoding="utf-8")
print("wrote", out)
