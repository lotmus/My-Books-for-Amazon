# -*- coding: utf-8 -*-
"""Follow-up: restore bold splits, Cast dashes, How to Read tail items."""
from __future__ import annotations

import shutil
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

LIVE = Path(
    r"c:\Users\lomus\OneDrive\My Books for Amazon"
    r"\Lolly Wren's Curious Science Adventures - SERIES"
    r"\Book 1 - Schrodingers Paperwork"
    r"\Schrodingers_Paperwork_BOOK_1_KINDLE_FINAL.docx"
)
OUT = LIVE.with_suffix(".docx.next")
WS = Path(
    r"c:\Users\lomus\OneDrive\My Books for Amazon"
    r"\Schrodingers_Paperwork"
    r"\Schrodingers_Paperwork_BOOK_1_KINDLE_FINAL.docx"
)
EM = "\u2014"


def wipe_content(p):
    for child in list(p._p):
        if child.tag != qn("w:pPr"):
            p._p.remove(child)


def write_runs(p, parts):
    """parts: list of (text, bold_or_None)."""
    wipe_content(p)
    for text, bold in parts:
        r = p.add_run(text)
        r.bold = bold


def insert_after(paragraph, text, style_name=None) -> Paragraph:
    new_el = deepcopy(paragraph._p)
    for child in list(new_el):
        if child.tag != qn("w:pPr"):
            new_el.remove(child)
    paragraph._p.addnext(new_el)
    new_p = Paragraph(new_el, paragraph._parent)
    new_p.add_run(text)
    return new_p


def main():
    doc = Document(str(LIVE))

    p84 = doc.paragraphs[84]
    assert p84.text.startswith("The lectures live together")
    write_runs(
        p84,
        [
            ("The lectures", True),
            (
                " live together in the Appendix as a real class, 18 lectures in five parts, one per chapter.",
                False,
            ),
        ],
    )

    # Rebuild list items 87-94 as clean visible runs
    expected = {
        87: "You need first: what this lecture assumes, named in one line.",
        88: "When you walk out, you should be able to: the skill the lecture is for.",
        89: "Start here: an everyday analogy, no jargon.",
        90: "The idea, step by step: the argument assembled in order.",
        91: "In one breath: the whole idea in a sentence or two.",
        92: "What this chapter was actually showing you: the story and the science locked together.",
        93: "Where the popular version goes wrong: the misconception you have probably been taught, corrected.",
        94: "Going deeper: the real machinery, the equations, the named experiments. Explicitly skippable.",
    }
    for i, text in expected.items():
        write_runs(doc.paragraphs[i], [(text, False)])

    # Fill empty list paras 95-96, then insert the stop line after 96
    write_runs(
        doc.paragraphs[95],
        [("What sticks: the few things worth keeping without notes.", False)],
    )
    write_runs(
        doc.paragraphs[96],
        [("One-line summary: the sentence to keep.", False)],
    )
    insert_after(
        doc.paragraphs[96],
        "Stop. Before you go on: say it to Mrs Chain. She will not be impressed by the word \u2018basis.\u2019 She will be impressed if the sentence is true.",
    )

    # Cast em-dash spacing
    for i, para in enumerate(doc.paragraphs):
        if para.text.startswith("Miss Dorothy Kell") and "unrecovered" in para.text:
            write_runs(
                para,
                [
                    (
                        "Miss Dorothy Kell "
                        + EM
                        + " wrote to the inquiry every week after the freeze. Her brother"
                        "\u2019s apprenticeship is among the unrecovered. Brought flowers to the list, and did not explain herself, which was correct.",
                        False,
                    )
                ],
            )
        if para.text.startswith("Nelson") and "terrier" in para.text:
            write_runs(
                para,
                [
                    (
                        "Nelson "
                        + EM
                        + " Mr Pilbeam"
                        "\u2019s terrier. Not permitted in the building. Attends anyway. Achieves, on balance, more clarity per sitting than the Board itself.",
                        False,
                    )
                ],
            )

    label = "What this chapter was actually showing you. "
    for i, para in enumerate(doc.paragraphs):
        t = para.text
        if t.startswith(label) and (
            "The first Bell run" in t
            or "Lolly delivers Mrs Chain a figure" in t
            or "The scene has already done the box" in t
        ):
            rest = t[len(label) :]
            write_runs(para, [(label, True), (rest, False)])

    if OUT.exists():
        OUT.unlink()
    doc.save(str(OUT))
    chk = Document(str(OUT))
    blob = "\n".join(p.text for p in chk.paragraphs[:110])
    assert "What sticks: the few things" in blob
    assert "One-line summary: the sentence to keep." in blob
    assert "say it to Mrs Chain" in blob
    p84t = chk.paragraphs[84]
    assert p84t.runs[0].bold is True
    assert p84t.runs[0].text == "The lectures"
    assert "five parts" in p84t.text
    # find locks
    split_ok = 0
    for para in chk.paragraphs:
        if not para.text.startswith(label):
            continue
        if not (
            "The first Bell run" in para.text
            or "Lolly delivers Mrs Chain a figure" in para.text
            or "The scene has already done the box" in para.text
        ):
            continue
        assert para.runs[0].bold is True, para.text[:80]
        assert para.runs[0].text == label, para.runs[0].text
        assert para.runs[1].bold in (False, None)
        split_ok += 1
    assert split_ok == 3, split_ok
    dorothy = next(p.text for p in chk.paragraphs if p.text.startswith("Miss Dorothy Kell"))
    assert "Kell " + EM + " wrote" in dorothy
    shutil.move(str(OUT), str(LIVE))
    try:
        shutil.copy2(LIVE, WS)
    except PermissionError:
        pass
    print("cleanup OK", "paras", len(chk.paragraphs))


if __name__ == "__main__":
    main()
