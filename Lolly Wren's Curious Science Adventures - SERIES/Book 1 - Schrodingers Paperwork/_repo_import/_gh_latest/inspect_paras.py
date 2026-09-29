# -*- coding: utf-8 -*-
from pathlib import Path
from docx import Document

p = Path(
    r"c:\Users\lomus\OneDrive\My Books for Amazon"
    r"\Lolly Wren's Curious Science Adventures - SERIES"
    r"\Book 1 - Schrodingers Paperwork"
    r"\Schrodingers_Paperwork_BOOK_1_KINDLE_FINAL.docx"
)
d = Document(str(p))
needles2 = [
    "three weeks after",
    "late August",
    "You need first",
    "About the Series",
    "If the apparatus shares",
    "What this chapter was actually showing you. Lolly delivers",
    "Start here.",
    "The lectures live together",
    "Lecture 1:",
    "interactive demonstrations",
    "Erich Schrottfinger",
    "Mrs Prosser",
    "Mr Pilbeam",
    "Albrecht Eilstein",
]
out = Path(__file__).with_name("inspect_paras.txt")
lines = []
for i, para in enumerate(d.paragraphs):
    t = para.text
    for n in needles2:
        if n in t:
            lines.append(f"{i} [{n}] {t[:350].replace(chr(10), ' | ')}")
            break
lines.append("=== CAST 117-145 ===")
for i in range(117, 146):
    t = d.paragraphs[i].text.strip()
    if t:
        lines.append(f"{i} {t[:200]}")
lines.append("=== HOW TO READ 80-115 ===")
for i in range(80, 116):
    t = d.paragraphs[i].text.strip()
    if t:
        lines.append(f"{i} {t[:250]}")
lines.append("=== APPENDIX BANNERS ===")
for i, para in enumerate(d.paragraphs):
    t = para.text.strip()
    if t.startswith("What happens when") or t.startswith("Part ") or "Lessons 6" in t and len(t) < 120:
        lines.append(f"{i} {t}")
lines.append("=== ABOUT SERIES / END ===")
for i in range(3140, len(d.paragraphs)):
    t = d.paragraphs[i].text.strip()
    if t:
        lines.append(f"{i} {t[:220]}")
out.write_text("\n".join(lines), encoding="utf-8")
print("wrote", out, "paras", len(d.paragraphs))
