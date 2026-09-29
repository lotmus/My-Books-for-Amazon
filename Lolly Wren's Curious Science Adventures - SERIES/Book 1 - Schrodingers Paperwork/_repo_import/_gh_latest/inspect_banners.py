# -*- coding: utf-8 -*-
from pathlib import Path
from docx import Document

p = Path(
    r"c:\Users\lomus\OneDrive\My Books for Amazon"
    r"\Lolly Wren's Curious Science Adventures - SERIES"
    r"\Book 1 - Schrodingers Paperwork"
    r"\Schrodingers_Paperwork_BOOK_1_KINDLE_FINAL.docx"
)
d = Document(str(p))
out = Path(__file__).with_name("inspect_banners.txt")
lines = []
for i in [2530, 2531, 2532, 2533, 2608, 2609, 2610, 2611, 2715, 2716, 2717, 2718, 2788, 2789, 2790, 2820, 2821, 2822]:
    if i < len(d.paragraphs):
        t = d.paragraphs[i].text
        style = d.paragraphs[i].style.name if d.paragraphs[i].style else ""
        lines.append(f"{i} [{style}] {t[:200]!r}")
out.write_text("\n".join(lines), encoding="utf-8")
print("ok")
