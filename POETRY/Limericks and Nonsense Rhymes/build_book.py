from pathlib import Path
import subprocess
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor


HERE = Path(__file__).resolve().parent
SOURCE = HERE / "Limericks_and_Other_Rhymes_for_People_Who_Should_Know_Better_v2.md"
OUTBASE = "Limericks_and_Other_Rhymes_for_People_Who_Should_Know_Better_REV02"
PYTHON = Path(sys.executable)
PANDOC = Path(r"C:\Users\lomus\AppData\Local\Pandoc\pandoc.exe")
BLUE = RGBColor(0x2A, 0x5C, 0xAA)


def make_reference_docx(path: Path) -> None:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(3.6)
    section.page_height = Inches(4.8)
    section.top_margin = section.bottom_margin = Inches(0.3)
    section.left_margin = section.right_margin = Inches(0.3)

    normal = doc.styles["Normal"]
    normal.font.name = "Georgia"
    normal.font.size = Pt(9.5)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.05

    title = doc.styles["Title"]
    title.font.name = "Georgia"
    title.font.size = Pt(20)
    title.font.bold = True
    title.font.color.rgb = BLUE
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for name, size, page_before in (("Heading 1", 15, True), ("Heading 2", 12, False)):
        style = doc.styles[name]
        style.font.name = "Georgia"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = BLUE
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style.paragraph_format.page_break_before = page_before
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(8)
        style.paragraph_format.space_after = Pt(5)

    doc.save(path)


def put_cover_on_first_page(path: Path) -> None:
    doc = Document(path)
    body = doc._body._body
    paragraphs = doc.paragraphs
    cover_index = next(
        (i for i, paragraph in enumerate(paragraphs)
         if paragraph._p.xpath(".//w:drawing")),
        None,
    )
    if cover_index is None:
        raise RuntimeError("The cover image was not found in the generated DOCX")

    cover = paragraphs[cover_index]._p
    caption = None
    if cover_index + 1 < len(paragraphs):
        candidate = paragraphs[cover_index + 1]
        if candidate.text.startswith("Cover:"):
            caption = candidate._p

    body.remove(cover)
    body.insert(0, cover)
    if caption is not None:
        body.remove(caption)
        body.insert(1, caption)
        title_index = 2
    else:
        title_index = 1
    doc.paragraphs[title_index].paragraph_format.page_break_before = True
    doc.save(path)


def main() -> None:
    if not SOURCE.exists():
        raise SystemExit(f"Missing manuscript: {SOURCE}")
    if not PANDOC.exists():
        raise SystemExit(f"Missing Pandoc: {PANDOC}")

    ref = HERE / "bak" / "reference_REV02.docx"
    ref.parent.mkdir(parents=True, exist_ok=True)
    make_reference_docx(ref)

    epub = HERE / f"{OUTBASE}.epub"
    docx = HERE / f"{OUTBASE}.docx"
    common = [str(PANDOC), str(SOURCE), "--from=markdown+hard_line_breaks", "--standalone"]
    metadata = [
        "--metadata=title=Limericks and Other Rhymes for People Who Should Know Better (Illustrated)",
        "--metadata=subtitle=Classic nonsense verse and modern limericks",
        "--metadata=author=Seamus O'Leary",
        "--metadata=lang=en-US",
    ]
    subprocess.run(common + metadata + ["--to=epub3", "--epub-title-page=false", "--toc", "--toc-depth=2", "--epub-cover-image=images/cover.png", "--output", str(epub)], cwd=HERE, check=True)
    subprocess.run(common + metadata + ["--to=docx", f"--reference-doc={ref}", "--output", str(docx)], cwd=HERE, check=True)
    put_cover_on_first_page(docx)
    print(f"Built {epub.name} ({epub.stat().st_size:,} bytes)")
    print(f"Built {docx.name} ({docx.stat().st_size:,} bytes)")


if __name__ == "__main__":
    main()
