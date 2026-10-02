# -*- coding: utf-8 -*-
import os as _os
HERE = _os.path.dirname(_os.path.abspath(__file__))
ROOT = _os.path.dirname(HERE)
"""Audit the finished Complete QED Course.docx for the scorecard."""
from collections import Counter
from docx import Document
from docx.oxml.ns import qn

path = _os.path.join(ROOT, "Complete QED Course.docx")
doc = Document(path)
lines = []

paras = doc.paragraphs
words = 0
h1 = []
styles = Counter()
for p in paras:
    t = p.text.strip()
    words += len(p.text.split())
    name = p.style.name if p.style is not None else ""
    styles[name] += 1
    if name == "Heading 1" and t:
        h1.append(t)

body = doc.element.body
links = body.findall(".//" + qn("w:hyperlink"))
drawings = body.findall(".//" + qn("w:drawing"))
omath = body.findall(".//{http://schemas.openxmlformats.org/officeDocument/2006/math}oMath")

blob = "\n".join(p.text for p in paras)
checks = [
    "Start from the objects already in hand",
    "__QEDCourseTOC_END__",
    "press F9",
    "precision QED",
    "(4π)³",
    "(4π)^3",
    "(eE)²/(4π)³",
    "(eE)² / (4π)³",
    "4π³",
    "α/(2π)",
    "Copyright",
    "Bibliography",
    "How to read this book",
    "one-loop QED",
    "Schwinger",
    "Uehling",
]
lines.append("PARAS %d TABLES %d WORDS %d" % (len(paras), len(doc.tables), words))
lines.append("HYPERLINKS %d DRAWINGS %d OMATH %d" % (len(links), len(drawings), len(omath)))
lines.append("H1 COUNT %d" % len(h1))
lines.append("--- H1 ---")
for t in h1:
    lines.append(t)
lines.append("--- CHECKS ---")
for c in checks:
    lines.append("%s  %s" % ("HAS" if c in blob else "NO ", c))
lines.append("--- STYLES ---")
for n, k in styles.most_common(12):
    lines.append("%s %d" % (n, k))
lines.append("--- OPENING ---")
n = 0
for p in paras:
    t = p.text.strip()
    if not t:
        continue
    lines.append("[%s] %s" % (p.style.name, t[:220]))
    n += 1
    if n >= 35:
        break
lines.append("--- L75 sample ---")
grab = False
count = 0
for p in paras:
    t = p.text.strip()
    if t.startswith("Lesson 75 "):
        grab = True
    if grab and t:
        lines.append(t[:240])
        count += 1
    if count >= 12:
        break
lines.append("--- L85 sample ---")
grab = False
count = 0
for p in paras:
    t = p.text.strip()
    if t.startswith("Lesson 85 "):
        grab = True
    if grab and t:
        lines.append(t[:240])
        count += 1
    if count >= 10:
        break
lines.append("--- L76 generator? ---")
grab = False
count = 0
for p in paras:
    t = p.text.strip()
    if t.startswith("Lesson 76 "):
        grab = True
    if grab and t:
        lines.append(t[:200])
        count += 1
    if count >= 8:
        break

out = _os.path.join(HERE, "_score_audit.txt")
open(out, "w", encoding="utf-8").write("\n".join(lines))
print("wrote", out, "h1", len(h1), "links", len(links), "words", words)
