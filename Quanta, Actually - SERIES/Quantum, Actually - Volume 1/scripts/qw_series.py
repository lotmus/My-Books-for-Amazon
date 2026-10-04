# Quanta, Actually series pass for The Quantum World.docx (Volume 1)
import sys, copy, re, docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import os
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
# Usage: python qw_series.py <source.docx> <output.docx>  (already applied 2026-10-02; not idempotent)
SRC, DST = sys.argv[1], sys.argv[2]
d = docx.Document(SRC); P = d.paragraphs
W = lambda t: qn('w:'+t)
def sn(p):
    try: return p.style.name
    except Exception: return '?'
def set_style(p, sid):
    pPr = p._p.get_or_add_pPr()
    ps = pPr.find(W('pStyle'))
    if ps is None:
        ps = OxmlElement('w:pStyle'); pPr.insert(0, ps)
    ps.set(W('val'), sid)
def set_only_text(p, text):
    ts = list(p._p.iter(W('t')))
    ts[0].text = text
    for t in ts[1:]: t.text = ''
def runfmt(r, bold=True, color=None, sz=None):
    r.bold = bold
    rPr = r._r.get_or_add_rPr()
    if color:
        c = rPr.find(W('color'))
        if c is None:
            c = OxmlElement('w:color'); rPr.append(c)
        c.set(W('val'), color)
    if sz:
        for tag in ('sz','szCs'):
            e = rPr.find(W(tag))
            if e is None: e = OxmlElement('w:'+tag); rPr.append(e)
            e.set(W('val'), str(sz))
def title_like(p):
    """Title -> Heading 1, keeping the Title look (rule under it, space after)."""
    pPr = p._p.get_or_add_pPr()
    if pPr.find(W('pBdr')) is None:
        b = OxmlElement('w:pBdr'); bo = OxmlElement('w:bottom')
        for k, v in (('val','single'),('sz','8'),('space','4'),('color','4F81BD')): bo.set(W(k), v)
        b.append(bo); pPr.append(b)
    if pPr.find(W('spacing')) is None:
        s = OxmlElement('w:spacing'); s.set(W('after'), '300'); pPr.append(s)
    set_style(p, 'Heading1')
    # pPr child order: pStyle must be first, which set_style guarantees
changes = []
W14='{http://schemas.microsoft.com/office/word/2010/wordml}'
def dc(el):
    e = copy.deepcopy(el)
    for k in (W14+'paraId', W14+'textId'):
        if k in e.attrib: del e.attrib[k]
    return e

def find(text, start=0, style=None):
    for i in range(start, len(P)):
        if P[i].text.strip() == text and (style is None or sn(P[i]) == style): return i
    raise SystemExit('not found: ' + text)

# 1. Title page (same layout as Physics, Actually)
t = find('The Quantum World', 0, 'Title')
new_title = dc(P[t]._p); P[t]._p.addprevious(new_title)
d2 = docx.text.paragraph.Paragraph(new_title, P[t]._parent); set_only_text(d2, 'Quanta, Actually')
set_style(P[t], 'Heading1'); changes.append('title page: Title "Quanta, Actually" + Heading 1 "The Quantum World"')
P = d.paragraphs
c = find('A Companion to the Physics, Actually Series')
set_only_text(P[c], 'Volume 1 in the Quanta, Actually Series')
for r in P[c].runs: runfmt(r, True, '4F81BD', 26)
a = c + 1; assert P[a].text.strip() == 'By Lothar J. Musiol'
set_only_text(P[a], 'Lothar J. Musiol')
for r in P[a].runs: runfmt(r, True, '4F81BD', 26)
changes.append('series line + author line')

# 2a. Keep title + copyright on one page: the two inserted lines (series
# Title, copyright series line) are paid for by three of the empty spacer
# paragraphs between the author line and "All rights reserved."
ar = find('All rights reserved.', 0)
sp = [k for k in range(ar - 1, 0, -1) if P[k].text == '' and not P[k]._p.xpath('.//w:br|.//w:drawing|.//w:bookmarkStart|.//w:sectPr')][:3]
assert len(sp) == 3
for k in sp: P[k]._p.getparent().remove(P[k]._p)
P = d.paragraphs
changes.append('removed 3 empty spacer paragraphs on title page')
# 2. Copyright page: series line after "© Lothar J. Musiol"
cp = find('© Lothar J. Musiol', 0)
sl = dc(P[cp]._p); P[cp]._p.addnext(sl)
slp = docx.text.paragraph.Paragraph(sl, P[cp]._parent); set_only_text(slp, 'Quanta, Actually series, Volume 1')
for r in slp.runs: runfmt(r, False, None, 18)
P = d.paragraphs
toc = find('Table of contents', 0, 'Heading 1'); set_only_text(P[toc], 'Table of Contents')

# 3. TOC entries "N — Title" -> "Chapter N — Title"
for i in range(toc, toc + 40):
    ts = list(P[i]._p.iter(W('t')))
    if ts and re.match(r'^\d+ — ', ts[0].text or ''):
        ts[0].text = 'Chapter ' + ts[0].text
    if P[i].text.strip() == 'About the author': last_toc = i; break

# 4. "Also in This Series" page before "How to read this book"
hr = find('How to read this book', 0, 'Heading 2')
anchor = P[hr]._p
def add_before(text, style=None, pagebreak=False):
    np_ = d.add_paragraph(text, style=style)
    if pagebreak: np_.paragraph_format.page_break_before = True
    anchor.addprevious(np_._p); return np_
h = add_before('Also in This Series', 'Heading 1', True)
add_before('Volume 1 — The Quantum World: From Quanta and Entanglement to Quantum Fields, Gravity, and the Future of Computing. This book: the survey, from the break with classical physics to quantum fields, QED, cryptography, and computing.')
add_before('Volume 2 — The Quantum Conversation: Phase, Light, and the Hidden Architecture of Electromagnetism. One stop on this map, collective electrodynamics, taken all the way: electromagnetism read outward from quantum phase and the potential.')
add_before('Volume 3 — Complete Quantum Electrodynamics Course: From Mathematical Foundations to One-Loop QED. The calculation course: Feynman rules, renormalization, the electron’s anomalous magnetic moment, and the Lamb shift, derived step by step.')
P = d.paragraphs

# 5. Headings: chapter/appendix/prologue/epilogue labels and Title-styled headings -> Heading 1
nH = 0
body_start = find('Also in This Series', 0, 'Heading 1')
for i, p in enumerate(P):
    if i <= body_start: continue
    tx = p.text.strip()
    m = re.match(r'^(CHAPTER|APPENDIX) (\d+)$', tx)
    if m:
        set_only_text(p, m.group(1).capitalize() + ' ' + m.group(2)); set_style(p, 'Heading1'); nH += 1; continue
    if tx in ('PROLOGUE', 'EPILOGUE'):
        set_only_text(p, tx.capitalize()); set_style(p, 'Heading1'); nH += 1; continue
    if sn(p) == 'Title':
        title_like(p); nH += 1; continue
    if sn(p) == 'Subtitle' and P[i-1].text.strip() == 'Prologue':
        set_style(p, 'Heading1'); nH += 1
changes.append('headings converted: %d' % nH)
P = d.paragraphs

# 6. Cross-references (Volume N)
for p in P:
    for r in p.runs:
        if r.text == ' — that is the subject of a full companion volume, The Quantum Conversation.':
            r.text = ' — that is the subject of The Quantum Conversation (Volume 2).'; changes.append('xref vol 2')
i = next(i for i, p in enumerate(P) if p.text.startswith('If you want to see how far this picture can be pushed'))
nx = dc(P[i]._p); P[i]._p.addnext(nx)
set_only_text(docx.text.paragraph.Paragraph(nx, P[i]._parent),
  'If you would rather calculate than picture it, the Complete Quantum Electrodynamics Course (Volume 3) derives this chapter’s results with pencil and paper: the Feynman rules, renormalization, and the electron’s magnetic moment.')
changes.append('xref vol 3')
P = d.paragraphs

# 7. Back matter "Also by"
i = find('The Quantum World is the companion book to the three-volume Physics, Actually series:')
proto = P[i]._p
qa = [ 'The Quanta, Actually series:',
       'Volume 1 — The Quantum World (this book)',
       'Volume 2 — The Quantum Conversation: Phase, Light, and the Hidden Architecture of Electromagnetism',
       'Volume 3 — Complete Quantum Electrodynamics Course: From Mathematical Foundations to One-Loop QED',
       'The Physics, Actually series:' ]
set_only_text(P[i], qa[0]); prev = P[i]._p
for txt in qa[1:]:
    e = dc(proto); prev.addnext(e); set_only_text(docx.text.paragraph.Paragraph(e, P[i]._parent), txt); prev = e
P = d.paragraphs
j = find('And the follow-up to this book:'); k = j + 1
assert P[k].text.startswith('The Quantum Conversation: Phase, Light')
set_only_text(P[j], 'Where this series goes next:')
set_only_text(P[k], 'The Quantum Conversation (Volume 2) — if this book’s appendix on collective electrodynamics left questions unanswered, Volume 2 goes further: a full-length case for deriving Maxwell’s equations, and the electromagnetic force itself, from the quantum behavior of matter. The Complete Quantum Electrodynamics Course (Volume 3) then does the full calculation.')
changes.append('also-by updated')

cp_ = d.core_properties
cp_.author = 'Lothar J. Musiol'; cp_.subject = 'Quanta, Actually, Volume 1'; cp_.keywords = 'Quanta, Actually; Volume 1'
from order import fix_order
nfix = sum(fix_order(p._p) for p in d.paragraphs)
changes.append('reordered %d pPr/rPr' % nfix)
d.save(DST); print('\n'.join(changes))
