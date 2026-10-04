# -*- coding: utf-8 -*-
"""Lesson 1 exists only as docx. Idempotent patch, applied to the standalone file and to
its section in the complete book: Prologue-4 bridge and first-pass boxes, exercise labels
(Core/Extension/Challenge), a try-first line before the solutions, the short mastery
check, and the Volume 2 header/subtitle."""
import os, sys
from docx import Document
from docx.text.paragraph import Paragraph
from docx.oxml.ns import qn
HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import build_lesson as bl

BRIDGE = ("Coming from Prologue 4",
          ["Prologue 4 described an arrow by two numbers (x, y), added arrows by adding parts, and turned and shrank "
           "them by multiplying. Section 1 gives that the standard names: the arrow is the complex number z = x + iy, "
           "its length squared is |z|² = z*z, and multiplying by e^(iφ) turns it by φ. The stopwatch arrow of "
           "Prologue 1 is e^(−iEt/ℏ). Nothing in Section 1 is new physics; it is the same arithmetic written compactly."])
FP_TITLE = "First pass: the result before the derivation"
FP_BODY = ["Complex numbers are the arrows of Prologues 1–4; vectors, inner products and matrices are the bookkeeping "
           "for many arrows at once. Every later lesson writes amplitudes as complex numbers, probabilities as "
           "|amplitude|², and physical quantities as Hermitian matrices.",
           "If Prologue 4 felt comfortable, Section 1 is a fast review; the new material starts with vector spaces in Section 2."]
CORE = {1, 2, 3, 4, 5, 6, 9, 10, 13}
TRY = "Attempt each exercise before reading its solution. If you are stuck, read only the first sentence of the solution, then try again."
CHECK = ("Before moving on, do the Core exercises without looking at the solutions and state this lesson's main result "
         "from memory. If both go smoothly, continue; if not, reread the sections the missed exercises cite.")


def _text(el):
    return "".join(t.text or "" for t in el.iter(qn("w:t")))


def _style(el):
    ps = el.find(qn("w:pPr") + "/" + qn("w:pStyle"))
    return ps.get(qn("w:val")) if ps is not None else ""


def patch_range(doc, els):
    """els: list of body elements of Lesson 1 (from its title to before Lesson 2)."""
    changed = 0
    texts = [_text(e) for e in els]
    if any(FP_TITLE in t for t in texts):
        return 0
    sec1 = next(e for e in els if e.tag == qn("w:p") and _text(e).strip() == "1 Complex numbers")
    b = bl.Builder(doc)
    b.box(BRIDGE[0], BRIDGE[1])
    b.para("")
    b.box(FP_TITLE, FP_BODY)
    b.para("")
    for el in b.elements:
        sec1.addprevious(el)
    changed += 1
    # exercise table: first table after the "Exercises" heading
    ex_h = next(e for e in els if e.tag == qn("w:p") and _text(e).strip() == "Exercises")
    tbl = ex_h.getnext()
    while tbl.tag != qn("w:tbl"):
        tbl = tbl.getnext()
    from docx.table import Table
    t = Table(tbl, doc._body)
    for row in t.rows[1:]:
        no = int(row.cells[0].text.strip())
        p = row.cells[1].paragraphs[0]
        txt = p.text
        if txt.startswith(("Core.", "Extension.", "Challenge.")):
            continue
        if txt.startswith("Challenge:"):
            label, rest = "Challenge.", txt[len("Challenge:"):].strip()
            rest = rest[:1].upper() + rest[1:]
        else:
            label, rest = ("Core." if no in CORE else "Extension."), txt
        for r in p.runs:
            r._r.getparent().remove(r._r)
        p.add_run(label).bold = True
        bl.add_rich_text(p, " " + rest)
    # try-first line after "Solutions"
    sol_h = next(e for e in els if e.tag == qn("w:p") and _text(e).strip() == "Solutions")
    q = doc.add_paragraph()
    q.add_run(TRY).italic = True
    sol_h.addnext(q._p)
    # mastery checklist -> short check
    m_h = next((e for e in els if e.tag == qn("w:p") and _text(e).strip() == "Mastery checklist"), None)
    if m_h is not None:
        for t_ in m_h.iter(qn("w:t")):
            t_.text = ""
        list(m_h.iter(qn("w:t")))[0].text = "Mastery check"
        nxt = m_h.getnext()
        first = True
        while nxt is not None and not (nxt.tag == qn("w:p") and _style(nxt).startswith("Heading")):
            after = nxt.getnext()
            if nxt.tag == qn("w:p"):
                if first:
                    para = Paragraph(nxt, doc._body)
                    pPr = nxt.find(qn("w:pPr"))
                    if pPr is not None:
                        nxt.remove(pPr)
                    for r in para.runs:
                        r._r.getparent().remove(r._r)
                    para.add_run(CHECK)
                    first = False
                else:
                    nxt.getparent().remove(nxt)
            nxt = after
    return changed


def lesson1_range(doc, standalone):
    body = list(doc.element.body)
    if standalone:
        return [e for e in body if e.tag != qn("w:sectPr")]
    start = next(i for i, e in enumerate(body) if e.tag == qn("w:p") and _style(e) == "Heading1"
                 and bl.is_lesson_heading(_text(e), 1))
    end = next(i for i, e in enumerate(body) if i > start and e.tag == qn("w:p") and _style(e) == "Heading1"
               and bl.is_lesson_heading(_text(e), 2))
    return body[start:end]


def main():
    ch = os.path.join(bl.CHAPTERS, "Lesson 01 Complex Numbers and Linear Algebra.docx")
    d = Document(ch)
    patch_range(d, lesson1_range(d, True))
    for p in d.sections[0].header.paragraphs:
        if "QED Course" in p.text or "Lesson" in p.text:
            for r in p.runs[1:]:
                r.text = ""
            p.runs[0].text = bl.CHAPTER_HEAD % ("Lesson", 1)
    for p in d.paragraphs[:3]:
        if p.style.name == "Subtitle" and p.text.startswith("Lesson 1 of"):
            for r in p.runs[1:]:
                r.text = ""
            p.runs[0].text = "Lesson 1 of %s, Volume %d: %s" % (bl.SERIES, bl.SERIES_VOLUME, bl.VOLUME_TITLE)
    d.save(ch)
    c = Document(bl.COMPLETE)
    n = patch_range(c, lesson1_range(c, False))
    c.save(bl.COMPLETE)
    print("lesson 1 patched", n)


if __name__ == "__main__":
    main()
