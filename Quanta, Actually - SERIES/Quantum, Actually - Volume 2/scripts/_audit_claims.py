# -*- coding: utf-8 -*-
import os as _os
HERE = _os.path.dirname(_os.path.abspath(__file__))
ROOT = _os.path.dirname(HERE)
"""Audit factual claims in the critical review against the complete course."""
from docx import Document
from docx.oxml.ns import qn

course = _os.path.join(ROOT, "Quantum, Actually - Volume 2.docx")
review = _os.path.join(ROOT, "notes", "Critical Review of the QED Course.docx")
out = []

d = Document(course)
paras = d.paragraphs
words = sum(len(p.text.split()) for p in paras)
out.append("words %d paras %d tables %d" % (words, len(paras), len(d.tables)))
sec = d.sections[0]
out.append("page_in %.2f x %.2f" % (sec.page_width.inches, sec.page_height.inches))
out.append("margins L%.3f R%.3f T%.3f B%.3f" % (
    sec.left_margin.inches, sec.right_margin.inches, sec.top_margin.inches, sec.bottom_margin.inches))
links = d.element.body.findall(".//" + qn("w:hyperlink"))
out.append("hyperlinks %d" % len(links))
h1 = sum(1 for p in paras if p.style and p.style.name == "Heading 1")
out.append("heading1 %d" % h1)

blob = "\n".join(p.text for p in paras)
for needle in [
    "Course edition 1.0",
    "precision QED",
    "__QEDCourseTOC_END__",
    "Update Field",
    "sky is blue",
    "sky's color",
    "dimensional regularization",
    "4−2ε",
    "4-2ε",
    "Copyright",
    "copyright",
    "Glossary",
    "glossary",
    "e = e",
    "Running Electric Charge",
    "Lesson 86 ",
    "ISBN",
]:
    out.append("HAS %-28s %s" % (needle, needle in blob))

# formula index neighborhood
idx = blob.find("Consolidated formula index")
out.append("--- INDEX WINDOW ---")
out.append(blob[idx:idx+2500] if idx >= 0 else "MISSING")

# lesson 68 window
i68 = blob.find("Lesson 68 ")
out.append("--- L68 WINDOW ---")
out.append(blob[i68:i68+700] if i68 >= 0 else "MISSING")

# capstone
ic = blob.find("Course capstone")
out.append("--- CAPSTONE ---")
out.append(blob[ic:ic+800] if ic >= 0 else "MISSING")

# review file
try:
    r = Document(review)
    out.append("REVIEW paras %d tables %d" % (len(r.paragraphs), len(r.tables)))
    rt = "\n".join(p.text for p in r.paragraphs)
    out.append("REVIEW words %d" % len(rt.split()))
    out.append("REVIEW has pagebreak marker via sect? sections %d" % len(r.sections))
except Exception as e:
    out.append("REVIEW OPEN FAIL %s" % e)

dest = _os.path.join(HERE, "_audit.txt")
open(dest, "w", encoding="utf-8").write("\n".join(out))
print("wrote", dest)
