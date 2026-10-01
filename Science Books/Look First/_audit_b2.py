"""One-shot audit for Book 2, A Trip Is Not a New Life."""
from __future__ import annotations

import hashlib
import os
import re
import zipfile
from collections import Counter

from docx import Document
from docx.oxml.ns import qn

ROOT = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(ROOT, "A Trip Is Not a New Life - Manuscript")
DOCX = os.path.join(SRC, "A Trip Is Not a New Life - Kindle.docx")
PARTS = [
    "00_Front_Matter.md",
    "01_Part_One_Dirt_Delay_Dates.md",
    "02_Part_Two_The_Moon_First.md",
    "03_Part_Three_Vehicle_Not_City.md",
    "04_Part_Four_Who_Stays.md",
    "05_Part_Five_The_Classroom.md",
    "06_Part_Six_The_Long_Ticket.md",
    "07_Part_Seven_Many_Clocks.md",
    "08_Part_Eight_The_Organ.md",
    "09_Part_Nine_Not_A_Straight_Line.md",
    "10_Part_Ten_The_Honest_Body.md",
    "11_Appendix.md",
]


def words(s: str) -> int:
    return len(re.findall(r"[A-Za-z0-9’']+", s))


def main() -> None:
    chunks = []
    for name in PARTS:
        raw = open(os.path.join(SRC, name), encoding="utf-8").read()
        chunks.append((name, raw))
    pop = [(n, r) for n, r in chunks if n not in ("00_Front_Matter.md", "11_Appendix.md")]
    print("=== WORD COUNTS ===")
    total = 0
    for name, raw in chunks:
        w = words(raw)
        total += w
        print(f"  {w:6d}  {name}")
    print(f"  {total:6d}  TOTAL")

    print("\n=== CHAPTERS ===")
    found = []
    thin = []
    for name, raw in pop:
        sections = re.split(r"(?=^##\s+)", raw, flags=re.M)
        for sec in sections:
            m = re.match(r"^##\s+(\d+)\.\s+(.+)$", sec, re.M)
            if not m:
                continue
            n = int(m.group(1))
            title = m.group(2).strip()
            body = re.sub(r"^Figure\s+\d+\..*$", " ", sec, flags=re.M)
            body = re.sub(r"!\[.*?\]\(.*?\)", " ", body)
            body = re.sub(r"^#+\s+.*$", " ", body, flags=re.M)
            wc = words(body)
            figs = re.findall(r"(?:^Figure\s+|!\[Figure\s+)(\d+)\.", sec, re.M)
            found.append((n, title, wc, figs, name))
            flag = ""
            if wc < 800:
                thin.append((n, wc, title))
                flag = " THIN"
            if figs != [str(n)]:
                flag += f" FIG{figs}"
            print(f"  Ch {n:02d} {wc:5d}  {title[:70]}{flag}")
    nums = [n for n, *_ in found]
    print("missing", [i for i in range(1, 47) if i not in nums])
    print("dupes", [k for k, v in Counter(nums).items() if v > 1])
    print("thin", thin)
    print("out of order", nums != list(range(1, 47)))

    print("\n=== CHAPTER POINTERS ===")
    bad_ptr = []
    for name, raw in chunks:
        if name == "11_Appendix.md":
            continue
        # split by chapter
        parts = re.split(r"(?=^##\s+\d+\.\s+)", raw, flags=re.M)
        cur = None
        for sec in parts:
            m = re.match(r"^##\s+(\d+)\.\s+", sec)
            if m:
                cur = int(m.group(1))
            for hit in re.finditer(r"Chapter\s+(\d+)", sec):
                n = int(hit.group(1))
                if n < 1 or n > 46:
                    bad_ptr.append((name, cur, n, sec[max(0, hit.start() - 40): hit.end() + 40].replace("\n", " ")))
    print("out of range", len(bad_ptr))
    for row in bad_ptr[:30]:
        print(" ", row)

    print("\n=== APPENDIX NOTES ===")
    app = chunks[-1][1]
    notes = re.findall(r"^(A\d+(?:\s*/\s*A\d+)?)\.\s+(.+)$", app, re.M)
    titles = [a for a, _ in notes]
    print("count", len(notes))
    c = Counter(titles)
    for k, v in c.items():
        if v > 1:
            print(" DUPLICATE LABEL", k, v)
    # second block starts at the second APPENDIX line
    blocks = app.split("APPENDIX — The Scientific Detail")
    print("appendix blocks", len(blocks) - (0 if app.startswith("APPENDIX") else 0), "splits", len(blocks))
    if len(blocks) >= 3:
        second = blocks[2]
        old = []
        for hit in re.finditer(r"Chapter\s+(\d+)", second):
            n = int(hit.group(1))
            if n < 27:
                snip = second[max(0, hit.start() - 30): hit.end() + 50].replace("\n", " ")
                old.append((n, snip))
        print("body-appendix Chapter N < 27:", len(old))
        for row in old[:25]:
            print(" ", row)

    print("\n=== PATTERN CHECKS ===")
    text = "\n".join(r for _, r in chunks)
    checks = [
        ("mojibake", r"Ã.|Â.|â€™|â€œ|â€”|\ufffd"),
        ("TODO", r"\b(TODO|FIXME|PLACEHOLDER|lorem ipsum)\b"),
        ("film", r"\b(Interstellar|Arrival|The Martian|Star Wars|Star Trek|Contact)\b"),
        ("fence", r"```"),
        ("html", r"<(?:div|span|br|img)\b"),
        ("Chapter 47+", r"Chapter\s+(?:4[7-9]|[5-9]\d)"),
        ("Figure 47+", r"Figure\s+(?:4[7-9]|[5-9]\d)"),
    ]
    for name, pat in checks:
        hits = list(re.finditer(pat, text))
        print(f"  {name}: {len(hits)}")
        for h in hits[:4]:
            print("   ", text[max(0, h.start() - 30): h.end() + 40].replace("\n", " "))

    # emphasis balance per file
    print("\n=== MARKUP BALANCE ===")
    for name, raw in chunks:
        stars = raw.count("**")
        # curly quotes
        if stars % 2:
            print("  odd **", name, stars)

    # doubled words
    print("\n=== DOUBLED WORDS (sample) ===")
    doubles = []
    for name, raw in pop:
        for m in re.finditer(r"\b([A-Za-z’']{4,})\s+\1\b", raw):
            doubles.append((name, m.group(0)))
    print("count", len(doubles))
    for row in doubles[:15]:
        print(" ", row)

    print("\n=== DOCX ===")
    d = Document(DOCX)
    sec = d.sections[0]
    print(f"  page {sec.page_width.inches:.2f} x {sec.page_height.inches:.2f} in  bytes {os.path.getsize(DOCX)}")
    h2 = [p.text for p in d.paragraphs if p.style.name == "Heading 2" and re.match(r"^\d+\.\s", p.text)]
    print("  chapter headings", len(h2), "first", h2[0] if h2 else None, "last", h2[-1] if h2 else None)
    seq = [int(t.split(".", 1)[0]) for t in h2]
    print("  chapter seq ok", seq == list(range(1, 47)))
    caps = [p.text for p in d.paragraphs if p.style.name == "Figure Caption"]
    cap_n = []
    for c in caps:
        m = re.match(r"Figure\s+(\d+)\.", c)
        cap_n.append(int(m.group(1)) if m else None)
    print("  captions", len(caps), "seq ok", cap_n == list(range(1, 47)))
    body = d.element.body
    bms = [b.get(qn("w:name")) for b in body.findall(".//" + qn("w:bookmarkStart"))]
    bms = [n for n in bms if n and n != "_GoBack"]
    anchors = [h.get(qn("w:anchor")) for h in body.findall(".//" + qn("w:hyperlink"))]
    missing = [a for a in anchors if a and a not in set(bms)]
    print("  bookmarks", len(bms), "links", len(anchors), "broken", len(missing))
    # image uniqueness via relationship ids and hashes of media
    with zipfile.ZipFile(DOCX) as z:
        media = [n for n in z.namelist() if n.startswith("word/media/")]
        hashes = []
        for n in media:
            hashes.append(hashlib.md5(z.read(n)).hexdigest())
    print("  media files", len(media), "unique hashes", len(set(hashes)))
    dup_h = [h for h, v in Counter(hashes).items() if v > 1]
    print("  duplicate media hashes", len(dup_h))
    # rId reuse
    ids = []
    paras = d.paragraphs
    for i, para in enumerate(paras):
        if para.style.name != "Figure Caption":
            continue
        blips = paras[i - 1]._p.findall(".//" + qn("a:blip"))
        rid = blips[0].get(qn("r:embed")) if blips else None
        ids.append((para.text.split(".", 1)[0], rid, bool(blips)))
    print("  captions missing drawing", [c for c, rid, ok in ids if not ok])
    print("  shared rIds", {k: v for k, v in Counter(rid for _, rid, _ in ids).items() if v > 1})
    h1s = d.styles["Heading 1"]
    print("  h1 font", h1s.font.name, "size EMU", h1s.font.size, "color", h1s.font.color.rgb)


if __name__ == "__main__":
    main()
