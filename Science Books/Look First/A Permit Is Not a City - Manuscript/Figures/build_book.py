# Generalized Kindle .docx builder for the "Look First" series Markdown manuscripts.
# usage: python build_book.py permit|body
# One figure per chapter; captions are taken from the manuscript's own ![Figure N. ...](figNN.png) lines.
# Photos: figs/figNN.jpg if present, else figs/figNN_slot.png placeholder.
import os, re, glob, sys
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = os.path.dirname(os.path.abspath(__file__))
AUTHOR = "Lothar J. Musiol"
BODY_FONT = "Georgia"
BOOKS = {
 "permit": dict(
    SRC=r"C:\Users\lomus\OneDrive\My Books for Amazon\Science Books\Look First\A Permit Is Not a City - Manuscript",
    TITLE="A Trip Is Not a Settlement", SUBTITLE="The Moon, Mars, and Why Leaving Does Not Clean the Earth",
    VOLUME="Look First, Volume 2.", FIGS=os.path.join(HERE, "figs_permit"),
    PHOTOS={13: "Credit: NASA/DSCOVR EPIC"},           # figure number -> credit line when the real photo is present
 ),
 "body": dict(
    SRC=r"C:\Users\lomus\OneDrive\My Books for Amazon\Science Books\Look First\The Body Keeps Its Own Clock - Manuscript",
    TITLE="A Longer Life Is Not a New Body", SUBTITLE="Healthspan, the Brain You Already Use, and Why Ten Thousand Years Is Not a Straight Line",
    VOLUME="Look First, Volume 3.", FIGS=os.path.join(HERE, "figs_body"),
    PHOTOS={1: "", 16: ""},
 ),
}

def fig_path(cfg, n):
    if n in cfg["PHOTOS"]:
        for ext in ("jpg", "jpeg", "png"):
            p = os.path.join(cfg["FIGS"], f"fig{n:02d}.{ext}")
            if os.path.exists(p): return p, True
        return os.path.join(cfg["FIGS"], f"fig{n:02d}_slot.png"), False
    return os.path.join(cfg["FIGS"], f"fig{n:02d}.png"), True

# ---------------- parse markdown into blocks ----------------
def parse_all(SRC):
    files = ["00_Front_Matter.md"] + sorted(os.path.basename(f) for f in glob.glob(os.path.join(SRC, "0[1-9]_*.md")) + glob.glob(os.path.join(SRC, "10_*.md"))) + ["11_Appendix.md"]
    blocks = []; chapter = None; seen = set()
    for fname in files:
        front = fname.startswith("00_"); in_appendix = fname.startswith("11_")
        lines = open(os.path.join(SRC, fname), encoding="utf-8").read().splitlines()
        i = 0
        while i < len(lines):
            line = lines[i].rstrip()
            if not line.strip(): i += 1; continue
            if front and line.startswith("# ") and not blocks:
                blocks.append(("title", line[2:].strip())); i += 1; continue
            if front and line.startswith("*") and line.endswith("*") and all(b[0] in ("title", "titleline") for b in blocks):
                blocks.append(("titleline", line.strip("*").strip())); i += 1; continue
            if line.strip() == "%%TOC%%": blocks.append(("toc",)); i += 1; continue
            if line.startswith("---"): i += 1; continue
            if line.startswith("# "):
                blocks.append(("h1", line[2:].strip(), "appendix" if in_appendix else "part")); chapter = None; i += 1; continue
            if line.startswith("## "):
                t = line[3:].strip(); m = re.match(r"(\d+)\.\s", t)
                if front: blocks.append(("h1", t, "front")); chapter = 0
                elif in_appendix: blocks.append(("h2", t, "appendix")); chapter = None
                elif m: chapter = int(m.group(1)); blocks.append(("h2", t, "chapter", chapter))
                else: blocks.append(("h2", t, "other"))
                i += 1; continue
            m = re.match(r"^!\[(.*)\]\((\S+)\)\s*$", line)
            if m:
                cap = m.group(1); mm = re.match(r"Figure\s+(\d+)\.\s*(.*)", cap)
                n = int(mm.group(1)) if mm else chapter
                if n is not None and n not in seen:
                    blocks.append(("fig", n, mm.group(2).strip() if mm else cap)); seen.add(n)
                i += 1; continue
            if line.startswith("|"):
                rows = []
                while i < len(lines) and lines[i].startswith("|"):
                    r = lines[i].strip()
                    if not re.match(r"^\|\s*-", r): rows.append([c.strip() for c in r.strip("|").split("|")])
                    i += 1
                blocks.append(("table", rows)); continue
            if re.match(r"^\s*[-*]\s+", line):
                items = []
                while i < len(lines) and re.match(r"^\s*[-*]\s+", lines[i]):
                    items.append(re.sub(r"^\s*[-*]\s+", "", lines[i]).strip()); i += 1
                blocks.append(("list", items)); continue
            para = [line]; i += 1
            while i < len(lines) and lines[i].strip() and not re.match(r"^(#|!\[|\||---|%%)", lines[i]):
                para.append(lines[i].rstrip("\n")); i += 1
            blocks.append(("p", para))
    return blocks

# ---------------- inline formatting ----------------
TOKEN = re.compile(
    r"(\*\*.+?\*\*"
    r"|\*[^*\n]+?\*"
    r"|\^\{[^}]*\}"
    r"|\^[^\s\^_*(){}]+"
    r"|(?<=[A-Za-zα-ωΑ-Ωħℓ★₀-₉∫Σ])_\{[^}]*\}"
    r"|(?<=[A-Za-zα-ωΑ-Ωħℓ★∫Σ])_[A-Za-z0-9μνΛ★]+)"
)
def add_runs(par, text, bold=False, italic=False, sup=False, sub=False, size=None):
    text = text.replace("M87*", "M87\uE000")
    pos = 0
    for m in TOKEN.finditer(text):
        if m.start() > pos: _run(par, text[pos:m.start()], bold, italic, sup, sub, size)
        tok = m.group(0)
        if tok.startswith("**"): add_runs(par, tok[2:-2], True, italic, sup, sub, size)
        elif tok.startswith("*"): add_runs(par, tok[1:-1], bold, True, sup, sub, size)
        elif tok.startswith("^"): add_runs(par, tok[2:-1] if tok.startswith("^{") else tok[1:], bold, italic, True, False, size)
        elif tok.startswith("_"): add_runs(par, tok[2:-1] if tok.startswith("_{") else tok[1:], bold, italic, False, True, size)
        pos = m.end()
    if pos < len(text): _run(par, text[pos:], bold, italic, sup, sub, size)
def _run(par, s, bold, italic, sup, sub, size):
    if not s: return
    r = par.add_run(s.replace("\uE000", "*")); r.bold = bold or None; r.italic = italic or None
    if sup: r.font.superscript = True
    if sub: r.font.subscript = True
    if size: r.font.size = Pt(size)

# ---------------- document helpers ----------------
_bm = [100]
def add_bookmark(par, name):
    _bm[0] += 1
    s = OxmlElement("w:bookmarkStart"); s.set(qn("w:id"), str(_bm[0])); s.set(qn("w:name"), name)
    e = OxmlElement("w:bookmarkEnd"); e.set(qn("w:id"), str(_bm[0]))
    par._p.insert(0, s); par._p.append(e)
def add_internal_link(par, text, anchor, bold=False):
    h = OxmlElement("w:hyperlink"); h.set(qn("w:anchor"), anchor); h.set(qn("w:history"), "1")
    r = OxmlElement("w:r"); rpr = OxmlElement("w:rPr")
    if bold: rpr.append(OxmlElement("w:b"))
    r.append(rpr); t = OxmlElement("w:t"); t.text = text; t.set(qn("xml:space"), "preserve"); r.append(t); h.append(r); par._p.append(h)
def page_break(doc):
    p = doc.add_paragraph(); p.add_run().add_break(WD_BREAK.PAGE); p.paragraph_format.space_after = Pt(0)
def set_style_font(style, name=BODY_FONT, size=11, bold=None, italic=None):
    style.font.name = name; style.font.size = Pt(size)
    if bold is not None: style.font.bold = bold
    if italic is not None: style.font.italic = italic
    rpr = style.element.get_or_add_rPr(); rf = rpr.find(qn("w:rFonts"))
    if rf is None: rf = OxmlElement("w:rFonts"); rpr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"): rf.set(qn(a), name)
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if rf.get(qn(a)) is not None: del rf.attrib[qn(a)]
def setup_styles(doc):
    from docx.enum.style import WD_STYLE_TYPE
    n = doc.styles["Normal"]; set_style_font(n, size=11)
    pf = n.paragraph_format; pf.first_line_indent = Inches(0.3); pf.space_after = Pt(0); pf.space_before = Pt(0); pf.line_spacing = 1.15
    h1 = doc.styles["Heading 1"]; set_style_font(h1, size=20, bold=True)
    h1.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER; h1.paragraph_format.space_before = Pt(60); h1.paragraph_format.space_after = Pt(24); h1.paragraph_format.first_line_indent = Inches(0)
    h2 = doc.styles["Heading 2"]; set_style_font(h2, size=15, bold=True)
    h2.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT; h2.paragraph_format.space_before = Pt(30); h2.paragraph_format.space_after = Pt(14); h2.paragraph_format.first_line_indent = Inches(0)
    for s in (h1, h2):
        s.paragraph_format.keep_with_next = True; s.font.color.rgb = None
        rpr = s.element.get_or_add_rPr(); c = rpr.find(qn("w:color"))
        if c is not None: rpr.remove(c)
    def mk(name, base="Normal"):
        try: st = doc.styles[name]
        except KeyError: st = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        st.base_style = doc.styles[base]; return st
    fp = mk("First Paragraph"); fp.paragraph_format.first_line_indent = Inches(0)
    ni = mk("No Indent"); ni.paragraph_format.first_line_indent = Inches(0); ni.paragraph_format.space_after = Pt(6)
    eq = mk("Equation"); eq.paragraph_format.first_line_indent = Inches(0); eq.paragraph_format.left_indent = Inches(0.3); eq.paragraph_format.space_before = Pt(6); eq.paragraph_format.space_after = Pt(6)
    cap = mk("Figure Caption"); set_style_font(cap, size=9.5, italic=True); cap.paragraph_format.first_line_indent = Inches(0); cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER; cap.paragraph_format.space_before = Pt(4); cap.paragraph_format.space_after = Pt(14)
    cr = mk("Figure Credit"); set_style_font(cr, size=8.5); cr.paragraph_format.first_line_indent = Inches(0); cr.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER; cr.paragraph_format.space_after = Pt(14)
    tp = mk("Title Page"); set_style_font(tp, size=28, bold=True); tp.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER; tp.paragraph_format.first_line_indent = Inches(0); tp.paragraph_format.space_before = Pt(140); tp.paragraph_format.space_after = Pt(18)
    ts = mk("Title Sub"); set_style_font(ts, size=14, italic=True); ts.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER; ts.paragraph_format.first_line_indent = Inches(0); ts.paragraph_format.space_after = Pt(60)
    ta = mk("Title Author"); set_style_font(ta, size=13); ta.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER; ta.paragraph_format.first_line_indent = Inches(0)
    cp = mk("Copyright"); set_style_font(cp, size=9.5); cp.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER; cp.paragraph_format.first_line_indent = Inches(0); cp.paragraph_format.space_after = Pt(8)
    t1 = mk("TOC Part"); set_style_font(t1, size=11, bold=True); t1.paragraph_format.first_line_indent = Inches(0); t1.paragraph_format.space_before = Pt(8); t1.paragraph_format.space_after = Pt(2)
    t2 = mk("TOC Chapter"); set_style_font(t2, size=10.5); t2.paragraph_format.first_line_indent = Inches(0); t2.paragraph_format.left_indent = Inches(0.3); t2.paragraph_format.space_after = Pt(1)
    li = mk("List Item"); li.paragraph_format.first_line_indent = Inches(-0.2); li.paragraph_format.left_indent = Inches(0.4); li.paragraph_format.space_after = Pt(3)

def add_figure(doc, cfg, n, caption):
    path, real = fig_path(cfg, n)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Inches(0); p.paragraph_format.space_before = Pt(10); p.paragraph_format.keep_with_next = True
    shape = p.add_run().add_picture(path, width=Inches(4.5))
    shape._inline.docPr.set("descr", f"Figure {n}. {caption}"[:120]); shape._inline.docPr.set("title", f"Figure {n}")
    c = doc.add_paragraph(style="Figure Caption"); r = c.add_run(f"Figure {n}. "); r.italic = False; r.bold = True
    add_runs(c, caption)
    credit = cfg["PHOTOS"].get(n, "")
    if credit and real: doc.add_paragraph(credit, style="Figure Credit")
    return not real

def add_table(doc, rows):
    if not rows: return
    ncol = max(len(r) for r in rows)
    t = doc.add_table(rows=len(rows), cols=ncol); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for ri, row in enumerate(rows):
        for ci in range(ncol):
            cell = t.cell(ri, ci); cell.text = ""; par = cell.paragraphs[0]
            par.paragraph_format.first_line_indent = Inches(0); par.paragraph_format.line_spacing = 1.0
            add_runs(par, row[ci] if ci < len(row) else "", bold=(ri == 0), size=9)
    doc.add_paragraph(style="No Indent")

def build(key):
    cfg = BOOKS[key]; SRC = cfg["SRC"]
    OUT = os.path.join(SRC, f"{cfg['TITLE']} - Kindle.docx")
    blocks = parse_all(SRC)
    doc = Document(); sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(6), Inches(9)
    for m in ("left_margin", "right_margin", "top_margin", "bottom_margin"): setattr(sec, m, Inches(0.75))
    setup_styles(doc)
    heads = [("part", b[1]) if b[0] == "h1" else ("chapter", b[1]) for b in blocks if b[0] in ("h1", "h2")]
    anchors = {}
    for k, (kind, t) in enumerate(heads, 1): anchors[(kind, t)] = f"bm_{k:03d}"
    doc.add_paragraph(cfg["TITLE"], style="Title Page"); doc.add_paragraph(cfg["SUBTITLE"], style="Title Sub"); doc.add_paragraph(AUTHOR, style="Title Author")
    page_break(doc)
    for line in (f"{cfg['TITLE']}: {cfg['SUBTITLE']}", f"Copyright © 2026 {AUTHOR}. All rights reserved.", cfg["VOLUME"], "Kindle edition.",
                 "No part of this book may be reproduced in any form without written permission from the author, except for brief quotations in reviews."):
        doc.add_paragraph(line, style="Copyright")
    page_break(doc)
    missing = []; prev = "start"; nfig = 0
    for b in blocks:
        kind = b[0]
        if kind in ("title", "titleline"): continue
        if kind == "toc":
            h = doc.add_paragraph("Contents", style="Heading 1"); add_bookmark(h, "toc")
            for hk, t in heads:
                p = doc.add_paragraph(style="TOC Part" if hk == "part" else "TOC Chapter"); add_internal_link(p, t, anchors[(hk, t)], bold=(hk == "part"))
            prev = "toc"; continue
        if kind == "h1":
            page_break(doc); h = doc.add_paragraph(b[1], style="Heading 1"); add_bookmark(h, anchors[("part", b[1])]); prev = "head"; continue
        if kind == "h2":
            if b[2] == "chapter": page_break(doc)
            h = doc.add_paragraph(b[1], style="Heading 2"); add_bookmark(h, anchors[("chapter", b[1])]); prev = "head"; continue
        if kind == "fig":
            nfig += 1
            if add_figure(doc, cfg, b[1], b[2]): missing.append(b[1])
            prev = "fig"; continue
        if kind == "table": add_table(doc, b[1]); prev = "table"; continue
        if kind == "list":
            for it in b[1]:
                p = doc.add_paragraph(style="List Item"); p.add_run("•  "); add_runs(p, it)
            prev = "list"; continue
        if kind == "p":
            if any(l.endswith("  ") for l in b[1]):
                p = doc.add_paragraph(style="First Paragraph" if prev != "p" else "Normal")
                for j, l in enumerate(b[1]):
                    add_runs(p, l.strip())
                    if j < len(b[1]) - 1: p.add_run().add_break(WD_BREAK.LINE)
                prev = "p"; continue
            text = " ".join(l.strip() for l in b[1])
            if re.match(r"^\*\*\(\d+\)\*\*", text): style = "Equation"
            elif text.startswith("**") or text.startswith("*(") or text.lower().startswith("*where"): style = "No Indent"
            else: style = "First Paragraph" if prev != "p" else "Normal"
            p = doc.add_paragraph(style=style); add_runs(p, text); prev = "p"
    doc.core_properties.title = cfg["TITLE"]; doc.core_properties.author = AUTHOR; doc.core_properties.subject = cfg["SUBTITLE"]
    doc.save(OUT)
    print("saved", OUT); print("figures:", nfig, "photo slots still placeholders:", missing)

if __name__ == "__main__":
    build(sys.argv[1])
