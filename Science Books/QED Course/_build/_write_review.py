# -*- coding: utf-8 -*-
"""Editorial review of the Complete QED Course, as a single Word file."""
import os

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor, Twips, Cm, Emu

OUT = os.path.join(
    os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
    "Critical Review of the QED Course.docx",
)
NAVY = RGBColor(0x1F, 0x4E, 0x78)
INK = RGBColor(0x1A, 0x1A, 0x1A)
MUTED = RGBColor(0x33, 0x33, 0x33)


def shade(cell, fill):
    tc = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc.append(shd)


def set_run_font(run, name, size, bold=False, color=INK, italic=False):
    run.bold = bold
    run.italic = italic
    run.font.name = name
    run.font.size = Pt(size)
    run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(a), name)


def style_paragraph(p, before=0, after=8, line=1.12, align="left"):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    if align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "justify":
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def add_text(p, text, size=11.5, bold=False, italic=False, color=INK, name="Cambria"):
    r = p.add_run(text)
    set_run_font(r, name, size, bold, color, italic)
    return r


def P(doc, text, size=11.5, before=0, after=8, italic=False, bold=False, align="left", color=INK):
    p = doc.add_paragraph()
    style_paragraph(p, before, after, align=align)
    add_text(p, text, size, bold, italic, color)
    return p


def H(doc, text, level):
    p = doc.add_paragraph()
    p.style = "Heading %d" % level
    # clear default runs from style application by setting text ourselves
    if p.runs:
        p.runs[0].text = text
        set_run_font(p.runs[0], "Calibri", 16 if level == 1 else 13, True, NAVY)
    else:
        add_text(p, text, 16 if level == 1 else 13, True, False, NAVY, "Calibri")
    style_paragraph(p, 16 if level == 1 else 12, 6, 1.05)
    return p


def bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        style_paragraph(p, 1, 3, 1.08)
        add_text(p, item, 11.5)


def numbers(doc, items):
    for i, item in enumerate(items, 1):
        p = doc.add_paragraph()
        style_paragraph(p, 2, 4, 1.08)
        add_text(p, "%d.  " % i, 11.5, True, False, NAVY, "Calibri")
        add_text(p, item, 11.5)


def table(doc, headers, rows, widths):
    tbl = doc.add_table(rows=1 + len(rows), cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    tbl_pr = tbl._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl_pr.append(layout)
    grid = tbl._tbl.tblGrid
    for i, w in enumerate(widths):
        grid.findall(qn("w:gridCol"))[i].set(qn("w:w"), str(w))
    for i, h in enumerate(headers):
        cell = tbl.rows[0].cells[i]
        cell.width = Twips(widths[i])
        shade(cell, "1F4E78")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        style_paragraph(p, 2, 2, 1.0)
        add_text(p, h, 10.5, True, False, RGBColor(0xFF, 0xFF, 0xFF), "Calibri")
    for r_i, row in enumerate(rows):
        for c_i, val in enumerate(row):
            cell = tbl.rows[r_i + 1].cells[c_i]
            cell.width = Twips(widths[c_i])
            if r_i % 2 == 1:
                shade(cell, "F4F7FA")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            style_paragraph(p, 2, 2, 1.0)
            add_text(p, val, 10.5, c_i == 1, False, INK, "Calibri")
    # margins
    for r_i, row in enumerate(tbl.rows):
        tr = row._tr
        trPr = tr.get_or_add_trPr()
        cant = OxmlElement("w:cantSplit")
        trPr.append(cant)
        if r_i == 0:
            hdr = OxmlElement("w:tblHeader")
            trPr.append(hdr)
        for cell in row.cells:
            tc = cell._tc.get_or_add_tcPr()
            mar = OxmlElement("w:tcMar")
            for side, w in (("top", "60"), ("bottom", "60"), ("start", "80"), ("end", "80")):
                el = OxmlElement("w:" + side)
                el.set(qn("w:w"), w)
                el.set(qn("w:type"), "dxa")
                mar.append(el)
            tc.append(mar)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    return tbl


def add_page_number(paragraph):
    run = paragraph.add_run()
    set_run_font(run, "Calibri", 9, False, MUTED)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    run._r.append(fld_begin)
    run2 = paragraph.add_run()
    set_run_font(run2, "Calibri", 9, False, MUTED)
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    run2._r.append(instr)
    run3 = paragraph.add_run()
    set_run_font(run3, "Calibri", 9, False, MUTED)
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run3._r.append(fld_end)


def build():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    sec.left_margin = Inches(1.0)
    sec.right_margin = Inches(1.0)
    sec.top_margin = Inches(0.9)
    sec.bottom_margin = Inches(0.85)
    sec.header_distance = Inches(0.4)
    sec.footer_distance = Inches(0.4)

    normal = doc.styles["Normal"]
    normal.font.name = "Cambria"
    normal.font.size = Pt(11.5)
    normal.font.color.rgb = INK
    pf = normal.paragraph_format
    pf.space_after = Pt(8)
    pf.line_spacing = 1.12

    for i, size in ((1, 16), (2, 13)):
        st = doc.styles["Heading %d" % i]
        st.font.name = "Calibri"
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = NAVY
        st.font.italic = False

    footer = sec.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_paragraph(fp, 0, 0, 1.0, "center")
    add_text(fp, "Critical review  ·  Complete Quantum Electrodynamics Course  ·  ", 9, False, False, MUTED, "Calibri")
    add_page_number(fp)

    header = sec.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    style_paragraph(hp, 0, 0, 1.0, "right")
    add_text(hp, "Editorial report  ·  30 September 2026", 9, False, False, MUTED, "Calibri")
    # bottom border on header paragraph
    pPr = hp._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "1F4E78")
    pBdr.append(bottom)
    pPr.append(pBdr)

    P(doc, "EDITORIAL REPORT", 11, 0, 2, True, True, "left", NAVY)
    P(doc, "A critical review of the Complete Quantum Electrodynamics Course", 22, 0, 4, False, True, "left", NAVY)
    P(doc, "What has to change before this manuscript is a book a reader can finish, trust, and buy.", 13, 0, 10, True, False, "left", MUTED)
    P(doc, "The object under review is the complete Word file, Course edition 1.0, checked against the sixty-seven standalone lesson files and their sources. The complete file contains about 208,000 words, 8,017 paragraphs, and 208 tables, on US Letter pages, and it contains zero pictures and zero hyperlinks. Lessons 1 through 61 in that file are the written course. Lessons 62 through 86 are one template with the topic name swapped in. Standalone files for Lessons 62 through 67 do exist, and they are real lessons, and they are not the chapters a reader finds in the complete book.")

    H(doc, "1. The judgment", 1)
    P(doc, "Read as a course of instruction through tree-level scattering, the sources are serious work. The sequence is a sequence. The factors of i and e are derived rather than announced. The massless Compton result and the two-body formula for dσ/dΩ agree with the standard answers. A reader who works the standalone lessons through Lesson 67 with a pencil will know how a QED cross section is built, and will have seen a cutoff and a mass renormalization. A reader who opens only the complete file stops at Lesson 61, on the threshold of vacuum polarization, and then reads templates.")
    P(doc, "Read as a book, the complete file is not ready. The subtitle promises precision QED. From Lesson 62 onward the file does not deliver a derivation. Earlier chapters cite later lessons by number for the Ward identity, the running charge, the anomalous magnetic moment, the Lamb shift, the path integral, and pair production in an external field. A buyer who believes the citation arrives at a page that says, in substance, define the essential objects and derive the central relation, and then never states a relation worth deriving. The same buyer never meets the written Lessons 62 through 67, because those files were not spliced into the complete book.")
    P(doc, "That gap decides the rest of the report. The manuscript’s strength is the path from complex numbers to a cross section. Its failure, as a published book, is that the complete file parts company with the written lessons at Lesson 62 and continues in a template that hides the difference. Overall, as the book a reader would be handed today, it is a 3 out of 10. The written lessons alone, judged as a course manuscript and not as the file that contains them, are a 7.")

    H(doc, "2. Scores", 1)
    P(doc, "Each score is out of ten. The third column is the change that would move the score, not a compliment placed beside the number.", before=0, after=6)
    doc.add_page_break()
    table(
        doc,
        ["Aspect", "Score", "The change that moves it"],
        [
            ["Curriculum of Lessons 1–67", "9", "Keep the order. Mark a short path through it."],
            ["Derivations through tree level", "8", "Protect the calculations. Separate them from the template chapters."],
            ["Honesty about deferred proofs", "8", "Land every deferral on a real page, or delete the deferral."],
            ["Physics of Lessons 1–67", "7", "Correct Lesson 60, then audit every “same phenomenon” sentence."],
            ["Exercises", "6", "Cut the restatements. Keep problems with a numerical or structural surprise."],
            ["Worked examples", "4", "Put the answer after the question, on a later page."],
            ["Voice and pleasure", "4", "Shorten the openings. Put one concrete picture or measured number in each lesson."],
            ["Equation typography", "3", "Stack fractions. Stop setting a slashed momentum as a combining character."],
            ["Figures", "1", "Draw the diagrams the algebra is about, starting with Lessons 44 and 49 and 53–64."],
            ["Headings and contents", "2", "Lesson titles only at Heading 1. Replace the unfilled Word field."],
            ["Glossary and index", "2", "Build one symbol list from the sixty-seven notation tables. Rebuild the formula index from the lessons."],
            ["Links, internal and outward", "1", "Make lesson citations clickable. Cite the papers and the data the numbers come from."],
            ["Front matter and book apparatus", "2", "Add a copyright page, a reader’s map, and a statement of who the book is for."],
            ["Closing of the argument", "1", "Splice the written Lessons 62–67 into the complete file, and remove 68–86 until they are real."],
            ["Readiness for Kindle", "2", "Static contents, real equations, figures, working links."],
            ["Readiness for print", "3", "Gutter for the length, page numbers that survive trim, a contents list a binder can use."],
            ["Readiness to sell", "2", "Sell only the book you have finished, under a title that matches its last real result."],
            ["The complete file, as a book", "3", "Do the ordered revision in the last section of this report."],
        ],
        [3200, 720, 5440],
    )
    P(doc, "The 9 is not politeness. Very few single-author manuscripts actually walk from the complex numbers to five spin-summed amplitudes without changing the rules halfway. That achievement is currently carrying a complete file in which twenty-five chapters, Lessons 62 through 86, are the same generated paragraph with a new noun.", before=4)

    H(doc, "3. The manuscript breaks its own promise", 1)
    P(doc, "The title page says Complete Quantum Electrodynamics Course. The subtitle says From mathematical foundations to precision QED. Edition 1.0 sits under that claim with no copyright line, no author, and no sentence that tells a reader where the book stops.")
    P(doc, "An audit of headings in the complete file settles where the real book ends. Lessons 1 through 61 have the working shape: objectives, a notation table, numbered sections, worked examples, exercises, solutions, a mastery checklist. From Lesson 62, Vacuum Polarization, through Lesson 86, every chapter has the same eight headings and no others: Learning objectives, Core statement, Development, Derivation workflow, Worked checkpoint, QED connection, Exercises, Solution guidance. Lesson 66’s core statement in that file is “d = 4 − 2ε in dimensional regularization.” The standalone Lesson 66, which the complete file does not contain, teaches a momentum cutoff Λ and never uses that phrase. Lesson 67’s source shows how a divergence in the electron’s self-energy can be absorbed into a bare mass. The complete file’s Lesson 67 does not.")
    P(doc, "The sources still point past themselves. Lesson 43 says that α/(2π) is Schwinger’s result for (g−2)/2 and that the derivation is Lesson 75. Lesson 39 postpones the Dirac g = 2 reduction to that same lesson. Lessons 47 and 51 postpone the Ward identity to Lesson 71. The written Lesson 62 postpones the running charge to Lesson 68. Lesson 3 postpones the hydrogen hierarchy to Lesson 76. Lesson 15 postpones the Lamb shift to Lesson 76 and the “one part in a billion” to Lesson 78. Lessons 2 and 11 postpone the path integral to Lesson 79. Lesson 57 postpones an external field to Lesson 85. None of those landing pages exist as lessons. In the complete file, the reader cannot even reach the written versions of Lessons 62 through 67, which are the last real steps before those citations.")
    P(doc, "The template is empty in a specific way. Lesson 68, Running Electric Charge, opens: “This lesson develops running electric charge as part of the path from elementary mathematics to quantum electrodynamics.” The objectives are “define the essential objects,” “derive the central relation,” and “recognize how the result enters later calculations.” The exercises ask the reader to derive that relation and to construct a simple example, without a relation on the page worth deriving. The same skeleton, with the noun replaced, runs through gauge fixing, the Ward identities, infrared divergences, soft photons, the anomalous magnetic moment, the Lamb shift, precision tests, path integrals, finite temperature, external fields, and the Standard Model.")
    P(doc, "A template is not a sketch a reader can study. It is a promise in the shape of a chapter. Because the earlier lessons are careful about what they have not proved, the fake landing pages do more damage than a missing chapter would. The book has trained the reader to trust a forward reference. The forward reference then betrays the training.")
    P(doc, "The capstone is one paragraph. It asks the reader to walk from a free Dirac field to a one-loop remark, and it stops before every result the subtitle advertised. The consolidated formula index is one line per lesson, and it describes the template’s slogans rather than the lessons that were written. It cites dimensional regularization. The written Lesson 66 uses a cutoff. The index gives the electron-photon vertex as −ieγμ. Lesson 44 derives +ieγμ and says the sign was checked. An index that contradicts both the chapter in the file and the chapter on disk is worse than no index, because a returning reader will trust the index.")
    P(doc, "Until Lessons 68 through 86 are written at the standard of Lesson 56, they should not be in the file a reader can open. Every sentence in Lessons 1 through 67 that names them has to be rewritten in the same pass. “Derived in Lesson 75” is a false statement about this manuscript. The honest sentence is: the number is Schwinger’s, this volume does not derive it, and the derivation is reserved for a later volume.")

    H(doc, "4. Logical development", 1)
    P(doc, "The written arc is the right arc for the book this author is actually able to finish. Complex numbers and linear algebra, calculus, relativity, four-vectors, Lagrangians, Hamiltonians, Maxwell, potentials and gauge, quantum mechanics through identical particles, the Klein-Gordon failure, the Dirac equation, gamma matrices, spinors, antiparticles, the current, classical fields, Noether, quantization of the scalar, the Dirac field, and the photon, local U(1), the QED Lagrangian, the vertex, the three propagators, the Feynman rules, spin sums, Mandelstam variables, then electron-muon scattering, Møller scattering, annihilation, Compton scattering, pair production, the cross section, decay rates, and a first reading of what those formulas say in a laboratory. Only then loops, power counting, a cutoff, and mass renormalization.")
    P(doc, "That order has a virtue textbooks often fake. A tool appears in the lesson before the lesson that spends it. The trace identities of Lesson 20 come back as instruments, not as a review chapter. The decision to quantize the photon in Coulomb gauge, and to confess that the covariant propagator is being used ahead of the Ward identity, is the right kind of confession. The reader is told which step is provisional. The failure is only that the provisional step is never discharged.")
    P(doc, "Two structural repairs would make the logic visible on a first open.")
    bullets(doc, [
        "Put a one-page map after the title. Three routes: the foundation (Lessons 1–37) for a reader who needs it, the QED spine (Lessons 38–60) for a reader who already has the foundation, and a clearly labeled stop at Lesson 67. State that precision QED is not in this volume.",
        "Inside each part, name the one result the part exists to reach. Part VIII exists to put five differential cross sections on one page. Say that in the part opener, in one sentence, before the hours.",
        "When a symbol changes meaning, say so in a single line the glossary can harvest. Lesson 67’s split between m and m₀ is the most important such line in the book, and it currently lives only inside that lesson.",
    ])
    P(doc, "The logic also slips in one place where the book thinks it is unifying. Lesson 60 treats the angular peaks of all five processes as one phenomenon, Rutherford’s, caused by a massless photon at vanishing momentum transfer. That unification is false for annihilation, Compton scattering, and pair production. Those peaks sit in a fermion propagator, which goes on shell when a photon and an electron are collinear. The photon-exchange peak of electron-muon scattering and of Møller scattering really is the long-range Coulomb singularity. A reader who absorbs Lesson 60’s moral will mis-classify the infrared problem for the rest of the subject. The repair is a short section that separates the three singularities and refuses the single name. Lesson 73, if it is ever written, can then inherit a distinction instead of a confusion.")

    H(doc, "5. Method", 1)
    P(doc, "The lesson shape is consistent, and consistency is a gift in a book this long: objectives, a notation table, a derivation, worked checks, a “where this goes next” table, a summary, exercises, solutions, a mastery checklist. A reader knows where she is. The shape is also how the book avoids teaching.")
    P(doc, "A worked example in this manuscript states the question and the answer in the same paragraph. There is no pause, no facing page, no “try this before you turn.” The example functions as a second telling of the section. That is useful for a reader who is lost. It is useless for a reader who is ready, and the book’s own objectives claim to be training that reader. Split every worked example. Leave the question in the lesson. Put the answer at the end of the part, in the same style as the solutions, with one line that says which mistake is the common one.")
    P(doc, "The exercises have the same defect one level down. A typical problem says: substitute Lesson 53’s amplitude into Lesson 58’s formula and recover the table. The solution is the substitution. Sixteen such problems train stamina, and stamina is not nothing, but they do not train judgment. After the fifth process, the reader can see the pattern of the assignment before she sees the physics. Keep the problems that can come out wrong in an interesting way: a sign from Fermi statistics, a pole that is not the pole you thought, a numerical value of α/(2π) compared with a measured anomaly, a kinematics threshold. Eight of those will teach more than sixteen restorations of a formula the lesson just displayed.")
    P(doc, "Solutions printed under the problems make the restoration even easier. For a self-study book, solutions belong in the book. They belong after the reader has had to leave the page: end of the part, or a ruled-off section that starts on a new page with the label Answer only after an attempt. A solution manual sold separately is a poor idea until the problems are strong enough that someone would want it.")
    P(doc, "The time estimates are part of the method, and they currently work against the book. Openings ask for four hours, or six, or nine to eleven. Taken together, a straight pass through the written lessons is a few hundred hours, and the book never says so in one place. Nor does it say which lessons a physicist who already knows special relativity and spin may skip, with a checkpoint exercise that proves the skip was safe. A course with one speed becomes a monument. A course with a marked path becomes a book people finish.")
    P(doc, "One habit is worth keeping and then spending more carefully. The book often stops and says what has not been proved. Lesson 62 refuses to pretend that transversality has cancelled the quadratic divergence. That refusal is the best pedagogical instinct in the manuscript. It becomes a tic when every lesson spends a worked example congratulating itself for the refusal. Use the refusal when a specific temptation is in the room. Cut it when the reader has already learned the moral.")

    H(doc, "6. Errors to remove before anyone else reads Lesson 60", 1)
    P(doc, "Lesson 60, in the section on Compton scattering, says that the low-energy limit of Compton scattering is the Thomson cross section, and that this classical formula is why the sky is blue, short wavelengths scattering more strongly than long ones from the electrons in the air. Thomson scattering from a free electron does not depend on wavelength. The sky is blue because sunlight scatters from molecules much smaller than a visible wavelength. That is Rayleigh scattering, and the cross section falls as the inverse fourth power of the wavelength. The electrons in those molecules are bound. The low-energy Compton formula does not describe them. The same mistake is repeated in the solution to Exercise 14, which offers “the classical explanation of the sky’s color” as a use of Compton scattering.")
    P(doc, "The repair is three sentences. State the Thomson limit and the formula σ_T = 8π r_e² / 3, with r_e = α/m. State that the cross section is independent of wavelength. State that the blue sky is Rayleigh scattering from molecules, and that this course’s Compton calculation is not that process. Leave the polarimeter. It is a true use of the process.")
    P(doc, "The second error is the one already named under logical development, and it needs a writer’s hand, not only a physicist’s. Worked example 3 says the poles in annihilation and pair production differ from electron-muon scattering “in bookkeeping rather than in physical origin,” and that every pole in the five formulas is a massless photon at vanishing momentum transfer. Rewrite the example so the photon propagator and the fermion propagator are different objects with different singularities. The summary, the connections table, and the line that calls the forward peak “Rutherford-like” for every t-channel process need the same pass. Rutherford belongs to the photon exchange. Collinear fermions belong to annihilation, Compton scattering, and pair production.")
    P(doc, "A third defect is smaller and still disqualifying in an index. The formula index’s regularization entry describes dimensional regularization. The lesson describes a cutoff. Rebuild that entry from Lesson 66’s own displayed result, the split between a logarithm of Λ and a finite piece that depends on Δ. Do this for every index line, from the lesson text outward. An index written from memory of a different draft will keep lying.")
    P(doc, "I am not claiming the tree-level algebra is rotten. The checks I made against the standard massless Compton square and against dσ/dΩ = |ℳ|² / (64π² s) hold. The danger is local and memorable: the sentences a non-specialist will repeat are the sentences that are wrong, because they are the sentences about the sky and about Rutherford.")

    H(doc, "7. Whether the book is good to read", 1)
    P(doc, "It is clear. It is not pleasurable. Clarity and pleasure are different achievements, and a book of this length needs both or the clarity never gets used.")
    P(doc, "The typical opening is one sentence of eighty or a hundred and twenty words, built from a colon, a dash, and a stack of clauses that restate the previous lesson’s achievement before they allow the new one. Then a time estimate. Then the objectives, which restate the opening. Then the notation table. The reader has been told what will happen three times before a derivative appears. The information is accurate. The rhythm is the rhythm of a specification.")
    P(doc, "Pleasure, in a calculation book, is not a joke in the margin. It is the moment a cancellation becomes visible, a number meets an experiment, or a picture makes a diagram obvious. This manuscript withholds those moments. History appears as a name attached to a formula, Rutherford in 1911, Schwinger’s 1948 number quoted and not derived. The laboratory appears in Lesson 60, which is the right chapter for it, and the most vivid claim on that page is the false one about the sky. Positron emission tomography is mentioned in a sentence that could have been a paragraph a reader remembers: two photons, nearly opposite, a line through a body. It is given the same weight as the time estimate.")
    P(doc, "The repair is a quota, not a rewrite of the voice into something cute. In each lesson, one paragraph of fewer than eighty words that contains a picture, a measured number, or a wrong idea the algebra is about to kill. In each part, one page that is only that. Cut the openings by half. The lesson’s job can be said in two sentences. The hours can move to the map at the front, once, where a reader can plan a month.")
    P(doc, "There is also a sameness of moral tone. The book is anxious that the reader will skip, assert, or treat an expectation as a proof. The anxiety is justified by the subject. Expressed in every lesson in the same sentence shape, it reads as mistrust. Trust the reader who has done Lesson 20. She does not need the method narrated again. She needs the next identity.")

    H(doc, "8. Formatting, and the object in the hand", 1)
    P(doc, "The page is US Letter, which is the right choice for an Amazon print edition aimed at the United States. The margins are about 0.90 inches on the left, 0.80 on the right, 0.80 on the top, and 0.75 on the bottom. For a pamphlet those are generous enough. For this manuscript they are not. At textbook density, 208,000 words plus tables and display equations is a book of roughly 550 to 700 pages. A bound book of that thickness needs a gutter. The inside margin has to grow with the page count, and 0.75 inches at the bottom leaves a page number with nowhere to live once the printer trims the sheet. Set the inside margin for the final page count after the empty lessons are removed, add a footer page number at least 0.5 inches from the trim, and mirror the margins.")
    P(doc, "The heading structure will defeat the contents list even after the field is updated. The file has 885 paragraphs styled as Heading 1. Lesson titles are Heading 1, and so are “How to use this lesson,” “1 Complex numbers,” “Exercises,” “Solutions,” and “Mastery checklist.” A contents list built from Heading 1 is a list of the entire book. Demote without exception: the lesson title is Heading 1, the part title is Heading 1, a numbered section is Heading 2, a subsection is Heading 3. “Exercises” and “Solutions” can be Heading 2. Nothing else may be Heading 1.")
    P(doc, "The table of contents currently in the file is a Word field that has not been populated. The visible sentence is an instruction to right-click and press F9. Directly under it, the body contains the marker __QEDCourseTOC_END__. A reader in Word who does not update fields sees an error message where a book’s contents should be. A reader on a Kindle, in Google Docs, or in a PDF made without field updates sees the same message. That string is a production bug. Delete it. Generate a static contents list from the repaired headings, with page numbers, and put it in the file as text. Keep a Word field only in the working copy, never in the file you upload.")
    P(doc, "The visual system is a generated report that learned one table. Two hundred and eight tables, most of them a blue header and alternating rows, carry notation, connections, exercises, and summaries. The system is readable in a single lesson and exhausting across sixty-seven. Use the blue table for notation and for the connections page. Set exercises as a numbered list, which is what they are. A list can be scanned. A one-row-per-problem table cannot, once the problem is a paragraph.")
    P(doc, "There is no copyright page, no information page, no statement of units and the metric signature that governs the whole book, and no “how to read this” that is about the book rather than about Lesson 1. Edition 1.0 on the title page is a software version number. A book needs a copyright year, a rights line, a printing history, and an ISBN before it is offered for sale. Those pages are short. Their absence is how a file remains a file.")

    H(doc, "9. Equations and the absence of pictures", 1)
    P(doc, "Display equations are Unicode sentences placed inside a single math run. Word will italicize them. It will not stack a fraction, and it will not build a superscript from the character ² in a way a screen reader or a Kindle can be trusted to keep. The slashed momentum is a letter followed by the combining solidus U+0338. That character is a frequent casualty of font substitution. On a phone, p-slash and p often become the same glyph, which in this subject is the difference between a propagator and a momentum.")
    P(doc, "Every display equation that is a fraction should be a stacked fraction. Every p-slash should be a momentum with a proper slash accent, or the notation the book is willing to defend in a glossary, used the same way every time. Subscripts written as _(μν) and as unicode ₀ should become one convention. This is production work, and it is the production work the book is about. A reader of a calculation course is staring at the equation for longer than she is staring at the prose. The prose can survive a long sentence. The equation cannot survive being a line of typewriter text.")
    P(doc, "There are no figures. Not a vertex, not a propagator, not a loop, not a Mandelstam diagram, not the two Compton graphs whose interference Lesson 56 spends eight hours on. The book asks the reader to hold the topology in her head while the trace is reduced. That is how errors of the Lesson 60 kind survive: the picture that would have shown a fermion line, rather than a photon line, was never drawn. Draw, at minimum, the electron-photon vertex, the fermion and photon propagators as line styles with a legend, both Compton diagrams, the electron-muon exchange, the two Møller diagrams with the relative minus sign, annihilation and pair production, the vacuum-polarization loop, the electron self-energy, and the vertex correction. Put each figure where the diagram is introduced, with a caption that names the momenta. A caption is a sentence. “Figure 12. Compton scattering, s channel and u channel” is enough if the lines are labeled.")
    P(doc, "Figures have a second job on Amazon. The Look Inside sample is how the book is judged in nine seconds. A page of Heading 1 complex-number prose, with an unfilled contents field above it, tells a buyer this is a notes dump. A page with one diagram and one stacked amplitude tells her it is a calculation book. Choose the sample pages on purpose: the contents map, one Dirac equation done properly, one Feynman diagram, one cross section.")

    H(doc, "10. Glossary, index, and the way back in", 1)
    P(doc, "There is no glossary. There are sixty-seven notation tables. They are the right local tool and the wrong global one. A reader who returns after a month to Lesson 54 cannot find, in one place, what t means, what the metric signature is, what m₀ means as opposed to m, or which sign convention governs the covariant derivative. She can search a Word file if she owns Word. A print reader cannot. A Kindle search across 200,000 words for the letter t is a punishment.")
    P(doc, "Build the glossary from the tables you already wrote. One entry per symbol, the lesson where it is defined, and a second line when the meaning changes. Ten pages will do. Put it before the formula index, not instead of a real index.")
    P(doc, "The formula index should be what its title says: the working formula, the conditions under which it holds, and the lesson. The massless Compton square is an index entry. “Diagram topology maps to algebraic factors” is not. Rebuild it after the empty lessons are gone, so the index does not advertise results the volume lacks. Add a short subject index for the words a reader will actually hunt: gauge, Ward identity, Thomson, Rayleigh, positron, cutoff, bare mass. Twenty pages of real index entries are worth more than a complete alphabetical index generated from every noun.")
    P(doc, "The mastery checklists can feed a different apparatus, a one-page “I can” list at the end of each part. Sixty-seven checklists in a row are a diary. Six part-level checklists are a course.")

    H(doc, "11. Links", 1)
    P(doc, "The complete file contains zero hyperlinks. In 8,017 paragraphs, nothing is a link: not a lesson citation, not a DOI, not a particle-data page, not a measured value of α, not the companion books that already sit beside this manuscript in the same catalog.")
    P(doc, "Internal links are the ones that change how the book is used. “Lesson 58, Section 5.2” appears hundreds of times. Each of those should be a heading link in the Word file and a working internal link in the Kindle file. A 600-page argument that cites itself in plain text is asking the reader to be her own index. The heading repair in the formatting section is what makes those links possible. You cannot link to “Section 5.2” if four different styles are pretending to be the title of the book.")
    P(doc, "Outward links are the ones that make the numbers checkable. The fine-structure constant, the electron anomaly, the Lamb interval of 1057.8 MHz, Schwinger’s 1948 paper, the Thomson cross section, the Rayleigh law: each of these should point at a stable source the reader can open. A precision claim with no citation is a mood. Cite CODATA or the Particle Data Group for the constants, and cite the paper for the historical result. Use a short further-reading paragraph at the end of each part, four items, each with a reason to read it. A dump of twenty textbooks helps no one. Peskin and Schroeder, or Schwartz, or Mandl and Shaw, named at the moment the course’s result matches theirs, would tell the reader where this book sits.")
    P(doc, "The F9 instruction is the anti-link. It tells the reader that navigation is her job, and that the job requires Microsoft Word. Delete it.")

    H(doc, "12. Selling it", 1)
    P(doc, "This file cannot be the product the title describes. If it is uploaded as the complete course in precision QED, the Look Inside will show either complex numbers or, if a buyer jumps ahead, a template chapter. Both outcomes produce returns. Returns, on a new listing, are how the book stops being shown.")
    P(doc, "Sell the book you have. A truthful title is closer to Quantum Electrodynamics to the Cross Section, with a subtitle that names the last real result: tree-level amplitudes, decay rates, and the first step of renormalization. The description should say, in the first three lines, who can start at Lesson 38, what formula she will be able to derive, and what the book does not contain. Precision QED, the Schwinger term, the Lamb shift, and the path integral are a second volume. Advertising them now spends the trust you will need when that volume exists.")
    P(doc, "The catalog you already have is the funnel this book is missing. A popular book on the quantum world and a conversation-style book on the same physics can send a reader here, and this book can send a reader back when she wants the story rather than the trace. There is no such sentence anywhere in the file. Write one, in the front matter and in the product description. Link the series. A textbook that pretends to have no relatives is harder to find and easier to abandon.")
    P(doc, "Price follows the object. A notes file of template chapters is not a $40 textbook, and a real 500-page calculation book is not a $2.99 impulse buy. After the empty lessons are removed, the equations are stacked, and a dozen figures exist, price it as a serious paperback and a serious ebook, and let the blurb carry the specificity: five processes, the formulas written out, the solutions in the book. Keywords should name the calculations, not the mood: Compton scattering, Dirac equation, Feynman rules, renormalization. Categories should be quantum field theory and mathematical physics, not a general science bin where this book will look unreadable.")
    P(doc, "Do not launch a workbook, a video course, or a solution manual from this draft. The exercises are not yet the kind of work anyone pays for twice. Fix the book. A later edition can add a short computational companion, a notebook that evaluates one cross section numerically, linked from Lesson 60. That companion would do more for the listing than another hundred pages of restated algebra, because it gives the Look Inside a picture and the reader a check.")
    P(doc, "Kindle and print are two products. The Word field, the combining characters, and the absent figures will fail on Kindle first. Produce the print PDF from a file whose contents are text and whose equations are real, then produce the ebook from the same source with internal links tested on a phone. Upload neither until you have read Lesson 60 and Lesson 68 on that phone.")

    H(doc, "13. What to do, in order", 1)
    P(doc, "The list is the revision. Each item is finished only when a reader who was not in the room can see it in the file.", before=0, after=6)
    numbers(doc, [
        "Decide which file is the book. Splice the written Lessons 62 through 67 into the complete file, or stop the complete file at Lesson 61 and say so. Delete Lessons 68 through 86 from anything a reader can open. Rewrite the capstone so it describes the chain that file actually completes.",
        "Search the remaining lessons for every mention of Lessons 68 through 86 and of Parts X and XI. Replace each citation with an honest sentence: derived here, or not in this volume. Lesson 43’s Schwinger sentence is the model of what must change.",
        "Correct Lesson 60. Separate Thomson from Rayleigh. Separate the photon pole from the collinear fermion pole. Correct Exercise 14’s solution and the connections table in the same sitting.",
        "Rebuild the formula index from the lesson text. The regularization entry has to describe the cutoff, because that is what Lesson 66 computes.",
        "Change the heading styles. Generate a static table of contents. Remove the F9 sentence and the marker __QEDCourseTOC_END__.",
        "Add the front matter: copyright, a one-page map with the three routes and an honest statement of the hours, the metric signature, and who should not start this book.",
        "Typeset the display equations as stacked mathematics. Fix the slashed momenta. Do Lessons 49 through 67 first, because that is where a wrong glyph changes the result.",
        "Draw the diagrams listed in the figures section. Caption them with the momenta. Place them in the lesson where the diagram is first used.",
        "Build a glossary from the notation tables, including the lesson of first use and the m versus m₀ distinction.",
        "Split worked examples so the answer is not in the question. Cut each lesson’s exercises to the ones that can be failed instructively. Move solutions to the end of the part.",
        "Add internal links for lesson and section citations, and a short further-reading note at the end of each part with real sources for α, the electron anomaly, and the historical papers you name.",
        "Set print margins for the final page count, with a gutter and a footer page number. Read a printed signature of Lessons 53 through 60, because that is the block a buyer will judge.",
        "Retitle the book to match its last real result. Write the Amazon description from that title. Link the series. Choose the Look Inside pages on purpose.",
        "Only after that, write Volume II, beginning with the running charge and the Schwinger calculation, at the standard of Lesson 56. Do not paste the template back in and call it a draft chapter.",
    ])

    H(doc, "14. What to leave alone", 1)
    P(doc, "The order of the written lessons. The decision to derive the vertex before using it. The refusal, in the real chapters, to pretend a symmetry has done a calculation it has not done. The notation table at the start of a lesson. The five processes chosen, which are the right five. The tone of respect for algebra, which is the book’s character and should survive the shortening of its sentences.")
    P(doc, "A revision that “makes it friendlier” by cutting the derivations would destroy the only thing a buyer cannot get from a shorter popular book. The popular books can be the door. This one has to be the room where the cross section is actually computed. Make that room shorter to walk across, honest about its back wall, and possible to see.")

    P(doc, "End of report.", 11, 12, 0, True, False, "left", MUTED)

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
