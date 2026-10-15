"""Shared docx editing helpers for the EE series (python-docx + lxml)."""
import re, copy
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Inches
W = lambda t: qn('w:'+t)
def sty(p):
    pp = p.find(W('pPr')); s = pp.find(W('pStyle')) if pp is not None else None
    return s.get(W('val')) if s is not None else ''
def ptext(p): return ''.join(t.text or '' for t in p.iter(W('t')))
def body_ps(doc): return [e for e in doc.element.body if e.tag == W('p')]
def re_sub_para(p, rx, fn):
    """Regex substitution over a paragraph's full text, keeping run formatting (text lands in the run where the match starts)."""
    ts = [t for t in p.iter(W('t'))]
    full = ''.join(t.text or '' for t in ts)
    ms = list(re.finditer(rx, full))
    if not ms: return 0
    spans = []; pos = 0
    for t in ts: n = len(t.text or ''); spans.append((pos, pos+n)); pos += n
    for m in reversed(ms):
        new = fn(m) if callable(fn) else m.expand(fn)
        if new == m.group(0): continue
        a, b = m.start(), m.end(); first = True
        for (s0, e0), t in zip(spans, ts):
            if e0 <= a or s0 >= b: continue
            txt = t.text or ''; lo, hi = max(a, s0) - s0, min(b, e0) - s0
            t.text = txt[:lo] + (new if first else '') + txt[hi:]; first = False
            t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
    return len(ms)
def mkrun(text, bold=False, italic=False, rpr=None):
    r = OxmlElement('w:r')
    if rpr is not None: r.append(copy.deepcopy(rpr))
    elif bold or italic:
        rp = OxmlElement('w:rPr')
        if bold: rp.append(OxmlElement('w:b')); rp.append(OxmlElement('w:bCs'))
        if italic: rp.append(OxmlElement('w:i')); rp.append(OxmlElement('w:iCs'))
        r.append(rp)
    t = OxmlElement('w:t'); t.text = text; t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve'); r.append(t); return r
def mkpara(runs, ppr=None, style=None, shade=None, spacing=None, num=None):
    p = OxmlElement('w:p'); pp = OxmlElement('w:pPr')
    if style: s = OxmlElement('w:pStyle'); s.set(W('val'), style); pp.append(s)
    if num is not None:
        n = OxmlElement('w:numPr'); il = OxmlElement('w:ilvl'); il.set(W('val'), '0'); ni = OxmlElement('w:numId'); ni.set(W('val'), str(num)); n.append(il); n.append(ni); pp.append(n)
    if shade: sh = OxmlElement('w:shd'); sh.set(W('fill'), shade); sh.set(W('val'), 'clear'); pp.append(sh)
    if spacing:
        sp = OxmlElement('w:spacing')
        for k, v in spacing.items(): sp.set(W(k), str(v))
        pp.append(sp)
    p.append(pp)
    for r in runs: p.append(r)
    return p
def insert_after(ref, els):
    for e in reversed(els): ref.addnext(e)
def insert_before(ref, els):
    for e in els: ref.addprevious(e)
def picture_para(doc, path, width_in, title, alt, caption):
    tmp = doc.add_paragraph(); tmp.add_run().add_picture(path, width=Inches(width_in))
    p = tmp._p
    for dp in p.iter(qn('wp:docPr')): dp.set('descr', alt); dp.set('title', title)
    pp = OxmlElement('w:pPr'); sp = OxmlElement('w:spacing'); sp.set(W('after'), '60'); sp.set(W('before'), '200'); pp.append(sp)
    jc = OxmlElement('w:jc'); jc.set(W('val'), 'center'); pp.append(jc); p.insert(0, pp)
    rp = OxmlElement('w:rPr')
    for tag in ('i', 'iCs'): rp.append(OxmlElement('w:' + tag))
    c = OxmlElement('w:color'); c.set(W('val'), '555555'); rp.append(c)
    for tag in ('sz', 'szCs'): e = OxmlElement('w:' + tag); e.set(W('val'), '19'); rp.append(e)
    cap = mkpara([mkrun(caption, rpr=rp)], spacing={'after': 200})
    jc2 = OxmlElement('w:jc'); jc2.set(W('val'), 'center'); cap.find(W('pPr')).append(jc2)
    p.getparent().remove(p)
    return [p, cap]
