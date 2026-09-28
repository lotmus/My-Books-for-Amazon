#!/usr/bin/env python3
"""Assemble the manuscript Markdown into a Word (.docx) deliverable.

Reads manuscript/How to Publish and Make Good Money.md and writes
"How to Publish and Make Good Money - Complete Manuscript.docx" next to
this script, using python-docx directly (no pandoc dependency).

Covers the Markdown conventions documented in this book's own CLAUDE.md:
Title/Part/Chapter/section headings (# through #####), key-takeaway/
worked-example/case-study blockquotes, GFM tables, checklists, bullet
lists, inline **bold**/*italic*/[links](url), figure images with italic
captions, and horizontal rules.
"""
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches

BOOK_DIR = Path(__file__).resolve().parent
MANUSCRIPT = BOOK_DIR / "manuscript" / "How to Publish and Make Good Money.md"
OUTPUT = BOOK_DIR / "How to Publish and Make Good Money - Complete Manuscript.docx"

HEADING_RE = re.compile(r"^(#{1,6})\s+(.*)$")
IMAGE_RE = re.compile(r"^!\[[^\]]*\]\(([^)]+\.(?:png|jpg|jpeg))\)$", re.IGNORECASE)
CHECKLIST_RE = re.compile(r"^- \[ \]\s+(.*)$")
BULLET_RE = re.compile(r"^-\s+(.*)$")
RULE_RE = re.compile(r"^-{3,}$")
INLINE_RE = re.compile(r"(\*\*\*.+?\*\*\*|\*\*.+?\*\*|\*.+?\*|\[[^\]]+\]\([^)]+\))")
LINK_RE = re.compile(r"^\[([^\]]+)\]\(([^)]+)\)$")

HYPERLINK_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink"


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    r_id = part.relate_to(url, HYPERLINK_REL, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    rpr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rpr.append(underline)
    run.append(rpr)
    text_el = OxmlElement("w:t")
    text_el.text = text
    run.append(text_el)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_inline_runs(paragraph, text):
    pos = 0
    for match in INLINE_RE.finditer(text):
        if match.start() > pos:
            paragraph.add_run(text[pos : match.start()])
        token = match.group(0)
        link_match = LINK_RE.match(token)
        if link_match:
            add_hyperlink(paragraph, link_match.group(1), link_match.group(2))
        elif token.startswith("***") and token.endswith("***"):
            run = paragraph.add_run(token[3:-3])
            run.bold = True
            run.italic = True
        elif token.startswith("**") and token.endswith("**"):
            paragraph.add_run(token[2:-2]).bold = True
        elif token.startswith("*") and token.endswith("*"):
            paragraph.add_run(token[1:-1]).italic = True
        pos = match.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def add_horizontal_rule(doc):
    paragraph = doc.add_paragraph()
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "999999")
    p_bdr.append(bottom)
    p_pr.append(p_bdr)


def parse_table(lines, start):
    rows = []
    i = start
    while i < len(lines) and lines[i].strip().startswith("|"):
        cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
        rows.append(cells)
        i += 1
    rows = [r for r in rows if not all(re.fullmatch(r":?-+:?", c) for c in r)]
    return rows, i


def build():
    lines = MANUSCRIPT.read_text(encoding="utf-8").split("\n")
    doc = Document()
    doc.core_properties.title = "How to Publish and Make Good Money"
    doc.core_properties.author = "Kevin Drew Peters"

    doc.styles["Heading 1"].paragraph_format.page_break_before = True
    doc.styles["Heading 2"].paragraph_format.page_break_before = True

    first_h1_seen = False
    i, n = 0, len(lines)
    while i < n:
        stripped = lines[i].strip()

        if not stripped:
            i += 1
            continue

        if stripped.startswith("|"):
            rows, i = parse_table(lines, i)
            if not rows:
                continue
            table = doc.add_table(rows=0, cols=len(rows[0]))
            table.style = "Table Grid"
            for row_index, row in enumerate(rows):
                cells = table.add_row().cells
                for col_index, cell_text in enumerate(row):
                    if col_index >= len(cells):
                        continue
                    cell_paragraph = cells[col_index].paragraphs[0]
                    add_inline_runs(cell_paragraph, cell_text)
                    if row_index == 0:
                        for run in cell_paragraph.runs:
                            run.bold = True
            continue

        heading_match = HEADING_RE.match(stripped)
        if heading_match:
            level = len(heading_match.group(1))
            heading_text = heading_match.group(2)
            if level == 1 and not first_h1_seen:
                first_h1_seen = True
                paragraph = doc.add_heading(level=0)
            else:
                paragraph = doc.add_heading(level=level)
            add_inline_runs(paragraph, heading_text)
            i += 1
            continue

        if RULE_RE.fullmatch(stripped):
            add_horizontal_rule(doc)
            i += 1
            continue

        image_match = IMAGE_RE.match(stripped)
        if image_match:
            image_path = (MANUSCRIPT.parent / image_match.group(1)).resolve()
            doc.add_picture(str(image_path), width=Inches(6))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            i += 1
            continue

        checklist_match = CHECKLIST_RE.match(stripped)
        if checklist_match:
            paragraph = doc.add_paragraph()
            paragraph.add_run("☐ ")
            add_inline_runs(paragraph, checklist_match.group(1))
            i += 1
            continue

        bullet_match = BULLET_RE.match(stripped)
        if bullet_match:
            paragraph = doc.add_paragraph(style="List Bullet")
            add_inline_runs(paragraph, bullet_match.group(1))
            i += 1
            continue

        if stripped.startswith(">"):
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.4)
            add_inline_runs(paragraph, stripped.lstrip(">").strip())
            i += 1
            continue

        paragraph = doc.add_paragraph()
        add_inline_runs(paragraph, stripped)
        i += 1

    doc.save(str(OUTPUT))
    return OUTPUT


if __name__ == "__main__":
    output_path = build()
    print(f"Wrote {output_path} ({output_path.stat().st_size:,} bytes)")
