# -*- coding: utf-8 -*-
"""Rebuild the Book 1 REV17 candidate DOCX from its paired release-candidate source.

manuscript_text_REV17_RELEASE_CANDIDATE.txt is the source (see the docstring at
its top for the markup). REV17 remains a candidate; REV14 remains the release master.
The candidate DOCX is copied to ..\\bak\\ (timestamped) before it is overwritten.
Usage:  python rebuild_kindle.py            (rebuilds the REV17 candidate in place)
        python rebuild_kindle.py out.docx   (writes a test build elsewhere)
"""
from __future__ import annotations

import re
import shutil
import sys
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
SRC = HERE / "manuscript_text_REV17_RELEASE_CANDIDATE.txt"
FIG = HERE / "figures"
OUT = ROOT / "FINAL_REV17_RELEASE_CANDIDATE_The_Murder_That_Hadnt_Happened_Yet_KINDLE_READY.docx"
NAVY = RGBColor(0x0C, 0x44, 0x7C)
PIC_WIDTH = Inches(5.2)

TOKEN = re.compile(r"\[([^\]]+)\]\((#[^)\s]+|https?://[^)\s]+)\)|(?<!\\)\*((?:\\\*|[^*])+?)(?<!\\)\*")
MARKS = re.compile(r"(?:\s\{#([^}]+)\})+$")
_BM = [0]


def unesc(s):
    return s.replace("\\*", "*")


def font(run, name="Georgia", size=11, bold=False, color=None, italic=False):
    run.font.size = Pt(size)
    run.bold = bold
    if italic:
        run.italic = True
    run.font.name = name
    rf = run._element.get_or_add_rPr().get_or_add_rFonts()
    for k in ("w:ascii", "w:hAnsi", "w:eastAsia"):
        rf.set(qn(k), name)
    if color is not None:
        run.font.color.rgb = color


def link(p, text, target, size=11):
    h = OxmlElement("w:hyperlink")
    if target.startswith("#"):
        h.set(qn("w:anchor"), target[1:])
    else:
        rid = p.part.relate_to(target, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
        h.set(qn("r:id"), rid)
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    rf = OxmlElement("w:rFonts"); rf.set(qn("w:ascii"), "Georgia"); rf.set(qn("w:hAnsi"), "Georgia")
    sz = OxmlElement("w:sz"); sz.set(qn("w:val"), str(round(size * 2)))
    col = OxmlElement("w:color"); col.set(qn("w:val"), "0563C1")
    u = OxmlElement("w:u"); u.set(qn("w:val"), "single")
    rpr.extend([rf, sz, col, u]); r.append(rpr)
    t = OxmlElement("w:t"); t.set(qn("xml:space"), "preserve"); t.text = unesc(text)
    r.append(t); h.append(r); p._p.append(h)


def inline(p, s, **kw):
    pos = 0
    for m in TOKEN.finditer(s):
        if m.start() > pos:
            font(p.add_run(unesc(s[pos:m.start()])), **kw)
        if m.group(1) is not None:
            link(p, m.group(1), m.group(2), size=kw.get("size", 11))
        else:
            font(p.add_run(unesc(m.group(3))), italic=True, **kw)
        pos = m.end()
    if pos < len(s):
        font(p.add_run(unesc(s[pos:])), **kw)


def bookmark(p, name):
    _BM[0] += 1
    a = OxmlElement("w:bookmarkStart"); a.set(qn("w:id"), str(_BM[0])); a.set(qn("w:name"), name)
    b = OxmlElement("w:bookmarkEnd"); b.set(qn("w:id"), str(_BM[0]))
    p._p.insert(0 if p._p.pPr is None else 1, a); p._p.append(b)


def fmt(p, sa=None, sb=None, align=None, kwn=None, pbb=None):
    pf = p.paragraph_format
    if sa is not None: pf.space_after = Pt(sa)
    if sb is not None: pf.space_before = Pt(sb)
    if kwn: pf.keep_with_next = True
    if pbb: pf.page_break_before = True
    if align is not None: p.alignment = align


def picture(doc, name, alt):
    doc.add_picture(str(FIG / name), width=PIC_WIDTH)
    p = doc.paragraphs[-1]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for dp in p._p.iter(qn("wp:docPr")):
        dp.set("descr", alt)
    return p


def build(out: Path):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.left_margin = sec.right_margin = Inches(0.85)
    sec.top_margin = sec.bottom_margin = Inches(0.8)
    section = ""
    section_course = False  # the course keeps short quoted lines tight (2 pt after)
    paras = [x for x in SRC.read_text(encoding="utf-8").split("\n\n") if x.strip()]
    C = WD_ALIGN_PARAGRAPH.CENTER
    for raw in paras:
        line = raw.strip("\n")
        marks = []
        m = MARKS.search(line)
        if m:
            marks = re.findall(r"\{#([^}]+)\}", m.group(0)); line = line[: m.start()]
        pbb = line.startswith("!! ")
        if pbb:
            line = line[3:]
        im = re.fullmatch(r"\[\[IMAGE (.+?) \| (.*)\]\]", line)
        head = dict(name="Amazon Ember", size=18, bold=True, color=NAVY)
        if im:
            p = picture(doc, im.group(1), im.group(2))
        elif line.startswith("#### "):
            p = doc.add_paragraph(); inline(p, line[5:], **{**head, "size": 26}); fmt(p, 10, 6, C)
        elif line.startswith("### "):
            p = doc.add_paragraph(); inline(p, line[4:], **head); fmt(p, 10, 16, C)
        elif line.startswith("## "):
            p = doc.add_paragraph(style="Heading 2"); inline(p, line[3:], **head); fmt(p, 10, 16, C, kwn=True)
        elif line.startswith("# "):
            section = line[2:]
            section_course = section_course or section == "COURSE"
            p = doc.add_paragraph(style="Heading 1"); inline(p, line[2:], **head); fmt(p, 10, 16, C, kwn=True, pbb=True)
        elif line.startswith("^ "):
            p = doc.add_paragraph(); inline(p, line[2:])
            fmt(p, 12, 12, C) if line[2:] == "\u2217" else fmt(p, 8, None, C)
        elif line.startswith("**") and line.endswith("**"):
            p = doc.add_paragraph(); inline(p, line[2:-2], bold=True); fmt(p, 2, 8, WD_ALIGN_PARAGRAPH.LEFT, kwn=True)
        elif line.startswith("- ") and section.startswith("Also by"):
            p = doc.add_paragraph(); inline(p, line[2:]); fmt(p, 1, None, WD_ALIGN_PARAGRAPH.LEFT)
        elif line.startswith("- "):
            p = doc.add_paragraph()
            if section == "Table of Contents":
                inline(p, line[2:], size=10)
                fmt(p, 0)
            else:
                inline(p, line[2:]); fmt(p, 2)
        else:
            short_quote = section_course and len(line) <= 90 and line[:1] in "\u201c\u2018\"'"
            p = doc.add_paragraph(); inline(p, line); fmt(p, 2 if short_quote else 8)
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        if pbb:
            p.paragraph_format.page_break_before = True
        for name in marks:
            bookmark(p, name)
    props = doc.core_properties
    props.title = "The Murder That Hadn't Happened Yet"
    props.subject = "A Relativistic Investigation Bureau Mystery"
    props.author = ""
    props.last_modified_by = ""
    props.keywords = "relativity, mystery, The Relativistic Investigation Bureau"
    props.language = "en-GB"
    props.comments = ""
    doc.save(out)
    print("Wrote", out, "paragraphs", len(doc.paragraphs))


if __name__ == "__main__":
    if len(sys.argv) > 1:
        build(Path(sys.argv[1]))
    else:
        if OUT.exists():
            bak = ROOT / "bak"; bak.mkdir(exist_ok=True)
            shutil.copy2(OUT, bak / f"{OUT.stem} - {datetime.now():%Y-%m-%d %H%M} before rebuild.docx")
        build(OUT)
