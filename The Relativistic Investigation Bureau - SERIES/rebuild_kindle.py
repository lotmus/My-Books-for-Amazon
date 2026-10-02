# -*- coding: utf-8 -*-
"""Rebuild KINDLE_READY.docx from manuscript_text.txt."""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement

ROOT = Path(__file__).resolve().parent
SRC = ROOT / "manuscript_text.txt"
OUT = ROOT / "FINAL_REV10_The_Murder_That_Hadnt_Happened_Yet_KINDLE_READY.docx"
NAVY = RGBColor(0x0C, 0x44, 0x7C)
URL_RE = re.compile(r"(https://[^\s]+)")

HEAD_START = (
    "Chapter ",
    "Lesson ",
    "PROLOGUE",
    "EPILOGUE",
    "COURSE",
    "WORKSHOP",
    "SUMMARY",
    "GLOSSARY",
    "BIBLIOGRAPHY",
    "ACKNOWLEDGEMENTS",
    "FURTHER READING",
    "How to Read",
    "Author's Note",
    "Copyright Stuff",
    "Reality Check",
    "Coming Next",
    "Table of Contents",
    "One Last Thing",
    "Syllabus Map",
    "About the Author",
    "Interlude",
)

FIGURES = [
    ("Drag the velocity yourself", ROOT / "coordinates.png"),
    ("Watch a particle trace its geodesic", ROOT / "formula.jpg"),
    ("Lesson 8\nThe Door", ROOT / "newDoor.png"),
    ("Lesson 11\nThe Hologram", ROOT / "Appendix_11_Holographic_Principle_REVISED.png"),
    ("The filing cabinet is an analogy", ROOT / "cabinet hole.png"),
]

LESSON_ARROW = re.compile(r"^-> Lesson for this chapter: (\d+) - .+$")
CHAPTER_ARROW = re.compile(r"^<- Chapter for this lesson: Chapter (\d+)\s*$")
EPILOGUE_ARROW = re.compile(r"^<- Chapter for this lesson: Epilogue\s*$")
SKIP_ARROW = re.compile(r"^-> If you skipped the course: Lesson (\d+) - .+$")
GLOSS_LINE = re.compile(r"^→ Glossary: (.+)$")
GLOSS_ENTRY = re.compile(r"^([A-Z][^.]{1,80})\. ")
_BOOKMARK_ID = 1


def set_run_font(run, name="Georgia", size=11, bold=False, color=None, heading=False):
    run.font.size = Pt(size)
    run.bold = bold
    run.font.name = name
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    prefer = "Amazon Ember" if heading else name
    rFonts.set(qn("w:ascii"), prefer)
    rFonts.set(qn("w:hAnsi"), prefer)
    rFonts.set(qn("w:eastAsia"), prefer)
    if color is not None:
        run.font.color.rgb = color


def add_hyperlink(paragraph, url, text):
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rPr.append(color)
    rPr.append(u)
    new_run.append(rPr)
    t = OxmlElement("w:t")
    t.text = text
    new_run.append(t)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def gloss_anchor(term: str) -> str:
    return "gloss_" + re.sub(r"[^A-Za-z0-9]+", "_", term).strip("_")


def is_heading(line: str, nxt: str, prev: str = "") -> bool:
    if not line or line.startswith("->") or line.startswith("<-") or line.startswith("→"):
        return False
    # Film-list lines and the typed contents are not a second set of chapter titles.
    if re.match(r"Lesson \d+ —", line) or re.match(r"Chapter \d+ - ", line):
        return False
    if re.fullmatch(r"(?:Chapter \d+|Lesson \d+|PROLOGUE)", prev):
        return True
    if line.startswith(HEAD_START):
        # "Chapter " and "Lesson " only open a section as a bare label ("Chapter 3"),
        # or as the in-lesson "Lesson checkpoint" subhead. A sentence that happens to
        # start with "Chapter 16's ..." is body text.
        if line.startswith("Chapter ") and not re.fullmatch(r"Chapter \d+", line):
            return False
        if line.startswith("Lesson ") and not (re.fullmatch(r"Lesson \d+", line) or line == "Lesson checkpoint"):
            return False
        return True
    if line in {
        "The Murder That Hadn't Happened Yet",
        "A Relativistic Investigation Bureau Mystery",
        "The Surprise in the Morning",
        "Sophie",
        "Spacetime for Investigators",
        "The Twin Paradox",
        "In Which the Universe Refuses to Behave",
        "The Relativistic Investigation Bureau Discovers That",
        "Solving One Mystery Is Merely an Invitation to Acquire Another",
        "The Wrong Kind of Multiverse",
    }:
        return True
    return False


def add_internal_link(paragraph, anchor, text):
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("w:anchor"), anchor)
    new_run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    rFonts = OxmlElement("w:rFonts")
    rFonts.set(qn("w:ascii"), "Georgia")
    rFonts.set(qn("w:hAnsi"), "Georgia")
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "22")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rPr.extend([rFonts, sz, color, u])
    new_run.append(rPr)
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    new_run.append(t)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def mark_bookmark(paragraph, name):
    global _BOOKMARK_ID
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(_BOOKMARK_ID))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(_BOOKMARK_ID))
    _BOOKMARK_ID += 1
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def add_plain(paragraph, text):
    if not text:
        return
    run = paragraph.add_run(text)
    set_run_font(run)


def add_body(doc, text: str, glossary: bool = False):
    p = doc.add_paragraph()
    stripped = text.strip()
    short_talk = len(stripped) <= 90 and stripped[:1] in "\"'“"
    p.paragraph_format.space_after = Pt(2 if short_talk else 8)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE

    lesson = LESSON_ARROW.match(stripped) or SKIP_ARROW.match(stripped)
    if lesson:
        add_internal_link(p, f"lesson_{lesson.group(1)}", stripped)
        return p
    chapter = CHAPTER_ARROW.match(stripped)
    if chapter:
        add_internal_link(p, f"chapter_{chapter.group(1)}", stripped)
        return p
    if EPILOGUE_ARROW.match(stripped):
        add_internal_link(p, "epilogue", stripped)
        return p
    gloss = GLOSS_LINE.match(stripped)
    if gloss:
        add_plain(p, "→ Glossary: ")
        terms = [part.strip() for part in gloss.group(1).split("·")]
        for i, term in enumerate(terms):
            if i:
                add_plain(p, " · ")
            add_internal_link(p, gloss_anchor(term), term)
        return p

    pos = 0
    for m in URL_RE.finditer(text):
        if m.start() > pos:
            add_plain(p, text[pos : m.start()])
        url = m.group(1).rstrip(").,;")
        add_hyperlink(p, url, url)
        pos = m.end()
    if pos == 0:
        add_plain(p, text)
    elif pos < len(text):
        add_plain(p, text[pos:])
    entry = GLOSS_ENTRY.match(stripped)
    if glossary and entry:
        mark_bookmark(p, gloss_anchor(entry.group(1)))
    return p


# Top-level sections. These get Word's Heading 1 style so Kindle and Word build
# navigation from them, and a page break before them. Subtitle lines under a
# label ("The Surprise in the Morning") stay plain headings so the navigation
# pane does not list every chapter twice. Short front-matter notes that follow
# How to Read get Heading 2 and no page break.
H1_EXACT = {"PROLOGUE", "EPILOGUE", "Interlude", "COURSE", "Syllabus Map", "WORKSHOP",
            "SUMMARY", "GLOSSARY", "About the Author", "FURTHER READING",
            "ACKNOWLEDGEMENTS", "BIBLIOGRAPHY", "One Last Thing", "How to Read This Book",
            "Table of Contents", "Coming Next in the Relativistic Investigation Bureau Series"}
H2_EXACT = {"Author's Note", "Copyright Stuff (The Serious Part)", "Reality Check (Sort Of)"}
SECTION_ANCHOR = {"PROLOGUE": "prologue", "How to Read This Book": "how_to_read",
                  "Interlude": "interlude", "Coming Next in the Relativistic Investigation Bureau Series": "coming_next",
                  "COURSE": "course", "GLOSSARY": "glossary", "BIBLIOGRAPHY": "bibliography"}


def toc_anchor(entry: str):
    m = re.match(r"Chapter (\d+) - ", entry)
    if m:
        return "chapter_" + m.group(1)
    for key, anchor in (("PROLOGUE", "prologue"), ("How to Read", "how_to_read"),
                        ("Interlude", "interlude"), ("EPILOGUE", "epilogue"),
                        ("Coming Next", "coming_next"), ("COURSE", "course")):
        if entry.startswith(key):
            return anchor
    return None


def heading_level(line: str, ahead: str):
    if re.fullmatch(r"(?:Chapter|Lesson) \d+", line) or line in H1_EXACT:
        if line == "WORKSHOP" and ahead.startswith("Do this after"):
            return 2  # the pointer on the syllabus page, not the workshop itself
        return 1
    if line in H2_EXACT:
        return 2
    return 0


def add_heading_line(doc, text: str, first=False, level=0):
    p = doc.add_paragraph()
    if level:
        p.style = doc.styles["Heading %d" % level]
        if level == 1:
            p.paragraph_format.page_break_before = True
        p.paragraph_format.keep_with_next = True
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6 if first else 16)
    p.paragraph_format.space_after = Pt(10)
    run = p.add_run(text)
    size = 26 if first else 18
    set_run_font(run, name="Amazon Ember", size=size, bold=True, color=NAVY, heading=True)
    return p


def maybe_figure(doc, joined_tail: str, used: set):
    for key, path in FIGURES:
        if key in used or not path.exists():
            continue
        if key in joined_tail:
            try:
                doc.add_picture(str(path), width=Inches(5.2))
                last = doc.paragraphs[-1]
                last.alignment = WD_ALIGN_PARAGRAPH.CENTER
                used.add(key)
            except Exception:
                pass


def main():
    lines = SRC.read_text(encoding="utf-8").splitlines()
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)

    used = set()
    i = 0
    first = True
    prev = ""
    in_glossary = False
    in_toc = False
    toc_count = 0
    seen = set()
    while i < len(lines):
        line = lines[i].rstrip()
        nxt = lines[i + 1].rstrip() if i + 1 < len(lines) else ""
        if not line:
            if in_toc and toc_count:
                in_toc = False
            i += 1
            continue
        if in_toc:
            anchor = toc_anchor(line)
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            if anchor:
                add_internal_link(p, anchor, line)
            else:
                add_plain(p, line)
            toc_count += 1
            prev = line
            i += 1
            continue
        if line == "GLOSSARY":
            in_glossary = True
        elif line in {"About the Author", "FURTHER READING", "ACKNOWLEDGEMENTS", "BIBLIOGRAPHY"}:
            in_glossary = False
        if is_heading(line, nxt, prev):
            ahead = lines[i + 2].strip() if i + 2 < len(lines) else ""
            level = 0 if first else heading_level(line, ahead)
            para = add_heading_line(doc, line, first=first, level=level)
            first = False
            if line == "Table of Contents":
                in_toc = True
            if line in SECTION_ANCHOR and SECTION_ANCHOR[line] not in seen:
                mark_bookmark(para, SECTION_ANCHOR[line])
                seen.add(SECTION_ANCHOR[line])
            if re.fullmatch(r"Chapter \d+", line):
                mark_bookmark(para, "chapter_" + line.split()[1])
            elif re.fullmatch(r"Lesson \d+", line):
                mark_bookmark(para, "lesson_" + line.split()[1])
            elif line == "EPILOGUE":
                mark_bookmark(para, "epilogue")
            tail = "\n".join(lines[i : i + 4])
            maybe_figure(doc, tail, used)
        else:
            add_body(doc, line, glossary=in_glossary)
            maybe_figure(doc, "\n".join(lines[max(0, i - 2) : i + 3]), used)
        prev = line
        i += 1

    props = doc.core_properties
    props.title = "The Murder That Hadn't Happened Yet"
    props.subject = "A Relativistic Investigation Bureau Mystery"
    props.author = ""  # the byline is Lothar's decision; never leave "python-docx" here
    props.last_modified_by = ""
    props.keywords = "relativity, mystery, The Relativistic Investigation Bureau"
    props.language = "en-GB"
    props.comments = ""
    doc.save(OUT)
    print("Wrote", OUT, "paragraphs", len(doc.paragraphs))


if __name__ == "__main__":
    main()
