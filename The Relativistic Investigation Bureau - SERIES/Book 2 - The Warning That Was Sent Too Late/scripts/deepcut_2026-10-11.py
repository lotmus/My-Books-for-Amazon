"""Deeper cut for Book 2 (11 Oct 2026, PT): Chapter 18's argument cut to its turns, reflective stretches in
Chapters 14 and 16 shortened; the clock-and-cone exposition moved to Lesson 6 with a link back to Chapter 18.
Usage: python deepcut_2026-10-11.py IN.docx OUT.docx
EDITS replace one exact phrase in one paragraph (must match exactly once). DELETES remove one whole paragraph by
its exact text ('TEXT|AFTER:PREV' = the paragraph with TEXT that directly follows the unique paragraph PREV).
Clues kept: Herbert's freehand order (sheet, door, voice, reading, refusal; sheet to train, nothing back; Sunday
POSSIBLE. NOT SHOWN; the empty circle), 'you cannot get a sender out of an order, only out of a mechanism',
Seventeen as a duration, Barbarian's letter and the hurry. Lesson exercises untouched."""
import sys, copy, docx
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

EDITS = [
 # ---- Chapter 14 (Monday evening): reflective trims ----
 ("She had sat through a second meeting after Marylebone. She had chaired it. Nobody in the room would have known that her daughter’s name had been on a railway board at lunchtime, because she had made sure nobody would, and because the minutes of a meeting about the refurbishment of a community hall are not the place for it. She had approved a quote for new radiators. She had declined a quote for a mural. She had driven home through the ordinary Monday traffic with her hands at ten to two on the wheel and had not once looked at her phone, because there was nothing on it she needed to read while moving.",
  "She had chaired a second meeting after Marylebone, about the refurbishment of a community hall, and nobody in the room would have known that her daughter’s name had been on a railway board at lunchtime. She had approved a quote for new radiators and declined one for a mural, on the grounds that a mural is a forecast. She had driven home with her hands at ten to two and had not once looked at her phone."),
 ("Tuppence took her coat off. He held out a hanger for it, one of his own, a good one with the curve right, and she put the coat on it and hung it on the back of the kitchen door, where her coat had hung every evening for three years, and stood for a moment with her hand on the shoulder of it.",
  "Tuppence took her coat off. He held out one of his own hangers for it, a good one with the curve right, and she hung it on the back of the kitchen door and stood for a moment with her hand on its shoulder."),
 ("Tuppence had chaired committees for twenty years. She knew what a minute was for. A minute recorded a decision, and the time of the decision, and who moved it, so that afterwards nobody could pretend the decision had been made at a different meeting by different people. She had spent a large part of her adult life making sure that the time on a decision was the time the decision was made, and not the time somebody found out about it, or the time somebody first wished it had gone the other way.",
  "Tuppence had chaired committees for twenty years. A minute recorded a decision and the time it was made: not the time somebody found out about it, and not the time somebody first wished it had gone the other way."),
 ("Julian turned the tap back on, quietly, and then off again, and came and sat down at the end of the table with the beech hanger, which he had picked up again on his way past the worktop. He turned it over in his hands. He did not offer an opinion.",
  "Julian came and sat at the end of the table with the beech hanger, turning it over in his hands."),
 ("Tuppence took her phone out of her bag. There was the message she had sent at twenty past one, School told. You’re staying put. Good girl. Mum., which had not been answered and had not needed to be, and nothing since from Derek, because he had said he would not ring Sophie unless Sophie rang him, and he had evidently decided that the same rule covered her.",
  "Tuppence took her phone out of her bag. Nothing from Derek. He had said he would not ring Sophie unless Sophie rang him, and had evidently decided that the same rule covered her."),
 # ---- Chapter 16: the lecture kept, the rumination shortened ----
 ("It was, he thought, almost the whole lecture. The light came across the sill without being asked, and reached him, and nothing of his had to go back the other way for that to happen. Seeing was not boarding. Being lit by a thing was not being on it.",
  "It was, he thought, almost the whole lecture. The light came across the sill without being asked, and nothing of his had to go back the other way. Seeing was not boarding."),
 ("It was the first thing he had said since Penny left. His own voice in the empty office sounded exactly as it always sounded, the voice for murders and for tea, and it occurred to him that this was the voice the carriage had used, that morning, through three soft knocks. The same register. The same flat calm. Something worn smooth, like a step.",
  "It was the first thing he had said since Penny left. His own voice in the empty office sounded as it always did, the voice for murders and for tea, and it occurred to him that this was the voice the carriage had used that morning, through three soft knocks. The same flat calm, worn smooth, like a step."),
 # ---- Chapter 18: the long argument cut to its turns; the exposition goes to Lesson 6 ----
 ("The room Barbarian took him into had always been the room it was, and Herbert did not look at it. He looked at the table. He could not have done otherwise. It had been cleared, which in that house was itself an event, and on it, in four rows of five, stood twenty clocks.",
  "Herbert did not look at the room. He looked at the table, which had been cleared, an event in that house, and on which, in four rows of five, stood twenty clocks."),
 ("They were not a collection in the ordinary sense. There was a kitchen timer", "There was a kitchen timer"),
 (" Every one that told the time told a different one. The station clock said twenty past nine. The chronometer said a quarter to five. The tomato said nothing,", " Every one that told the time told a different one. The tomato said nothing,"),
 ("“In steps. One.” He laid the pencil beside the drawing, parallel to the ruled line, exactly. “Between the door and the reading there is a quantity every good clock at rest in that office would agree on, whatever it said at the start: the difference between its readings. It is not a convention. It is geometry. Two—”",
  "“In steps. One.” He took a pencil from his inside pocket and laid it beside the drawing, parallel to the ruled line, exactly. “Two good clocks at rest side by side in that office agree on the count between the door and the reading, whatever each said at the start. That is not a convention. It is geometry. Two—”"),
 ("“Stop. Before two. How would they agree?” said Barbarian. “By counting. Each clock counts its own motion. A pendulum counts swings. Braunschweig counts caesium atoms, because a committee in Paris told it to. Every second there has ever been is one motion counted against another.” He spread his hands over the twenty clocks. “These do not measure time. They measure each other.”",
  "“Stop. Before two.” Barbarian spread his hands over the twenty clocks. “They agree because they count. A pendulum counts swings. Braunschweig counts caesium, because a committee in Paris told it to. These do not measure time. They measure each other.”"),
 ("“Two. Two good clocks together at the door and together at the reading agree on the count between,” said Herbert. “The world enforces that.”",
  "“Two. The world enforces that agreement, whatever you call it.”"),
 ("“I give you that the world is extraordinarily well behaved about which counts agree. What I do not give you is the wire.” Barbarian tapped the ruled line once with the end of a finger, the only time he touched the paper. “You have the events. You have the order. You have the counts, if anybody counted, which that office mostly did not. Where in any of that is the line?”",
  "“I give you the counts. I do not give you the wire.” Barbarian waved at the ruled line without touching it. “The line is where you put it. It is the best convenience in physics, so good that people forget they put it there. Then a sheet arrives before a door, and they look at the wire for the answer, and the wire says nothing. YESTERDAY. Not a date; a direction. The most honest railway in Europe. It has never once promised to arrive.”"),
 ("“No. Mr Gent should like to know who sent the sheet. That is a different wish.” Barbarian’s eyes went to the question mark beside SHEET. “And it is the wrong question.”",
  "Barbarian’s eyes went to the question mark beside SHEET."),
 ("“It is wrong in the form he asks it. Who sent it. Sent is a verb with a before inside it.", "“Mr Gent wants to know who sent it. That is the wrong question. Sent is a verb with a before inside it."),
 ("He reached out, and for the second time touched the drawing, and this time he turned it face-down.", "He reached out and, for the first time, touched the drawing, turning it face-down."),
 ("Herbert considered refusing, and found that he did not want to. He asked for a clean sheet, which took some time to find, because the papers in that house were where they had been put for reasons, and he took his ruler out of his inside pocket and then, after a moment’s thought, put it back. He drew the five events freehand, as small circles, without a line between them. He joined the ones that could be ordered with short arrows: sheet to door, door to voice, voice to reading, reading to refusal. He drew the train, out to one side, as a sixth circle, and an arrow from the sheet to the train, and nothing at all from the train back. He drew Sunday’s radiator as a seventh circle, lower down, and an arrow from it to the door, and wrote on the arrow, in his immaculate hand: POSSIBLE. NOT SHOWN. And off on its own, unconnected to anything, an eighth circle, empty, with a question mark inside it, and no arrow either way.",
  "Herbert considered refusing, and found that he did not want to. He took his ruler out of his inside pocket and, after a moment’s thought, put it back. On a clean sheet he drew the five events freehand, as small circles, and joined the ones that could be ordered with short arrows: sheet to door, door to voice, voice to reading, reading to refusal. The train went out to one side, with an arrow from the sheet to it and nothing at all coming back. Sunday’s radiator went lower down, with an arrow to the door on which he wrote, in his immaculate hand, POSSIBLE. NOT SHOWN. And off on its own, unconnected to anything, an empty circle with a question mark inside it."),
 ("Herbert left it. He stood, and put on his coat, and buttoned it, and looked at his watch, which said a quarter to four, because Braunschweig said so; and he found, to his own surprise, that he looked at it as one more clock on a table of twenty and did not mind.",
  "Herbert left it. He put on his coat and looked at his watch, which said a quarter to four because Braunschweig said so, and found, to his own surprise, that he looked at it as one more clock on a table of twenty and did not mind."),
 # Ch18: speaker after the cut; Ch12, Ch15: two more reflective stretches shortened
 ("“It is beautiful.” He did not look up. “And it has a ruler in it.”", "“It is beautiful,” said Barbarian, without looking up. “And it has a ruler in it.”"),
 ("Derek had known it would be. He had known it all the way back, on foot, with Tabitha’s hand in his. He had known it on the stairs, and on the landing, where Mrs Marsh rose from her step at their approach with the clipboard raised and Penny and Trevor stood behind her, not going in. He knew it, and then he came into his office and saw it, and found that knowing a thing and seeing it were still, after all this, two separate events.",
  "Derek had known it would be. He had known it all the way back, with Tabitha’s hand in his, and on the landing, where Mrs Marsh rose from her step with the clipboard raised and Penny and Trevor stood behind her, not going in. Then he came into his office and saw it, and knowing a thing and seeing it were still, after all this, two separate events."),
 ("She had taken out the itinerary, and her phone, and a biro from the Bureau that said nothing on the side, and the photograph she had taken back off Derek’s monitor at five. She stood it against the sugar bowl, face-up. Four of them and a dog, at half past one on Sunday morning, standing in a row in the office because she had told them to, and herself in the middle, looking furious. She had looked at it a good many times since Sunday. It was still the right way up.",
  "She had taken out the itinerary, her phone, a biro from the Bureau that said nothing on the side, and the photograph she had taken back off Derek’s monitor at five, and stood the photograph against the sugar bowl, face-up: four of them and a dog at half past one on Sunday morning, in a row because she had told them to, and herself in the middle, looking furious. It was still the right way up."),
]

DELETES = [
 # Ch18
 "Barbarian studied it for a long time. He did not touch it.",
 "Herbert took a pencil out of his inside pocket, because he found it easier to think with a pencil in his hand than with twenty clocks in his ear.",
 "Herbert regarded the line he had drawn.",
 "“It is where I put it.”",
 "“Yes.”|AFTER:“It is where I put it.”",
 "“It is a convenience.”",
 "“It is the best convenience in physics. It is so good that people forget they put it there.” Barbarian sat back. “Then a sheet arrives before a door, and they look at the wire for the answer, and the wire says nothing. YESTERDAY. Not a date; a direction. The most honest railway in Europe. It has never once promised to arrive. I should like to meet whoever designed it.”",
 "“So should Mr Gent.”",
 "“It is not a wrong question. It is an unanswerable one, at present.”",
]

# Lesson 6 gets the exposition cut from Chapter 18, with a link back to the chapter.
LESSON_AFTER = "Keep this: earlier opening left in the past, boarding refused while it was still future"
LESSON_PARTS = [
 ("Behind the case: what Herbert and Barbarian argue about over twenty clocks in ", None),
 ("Chapter 18", "chapter_18"),
 (". Every clock is a counter. A pendulum counts swings; since 1967 the SI second has been defined as 9,192,631,770 periods of the radiation from a caesium-133 atom, so every second ever recorded is one motion counted against another. Two good clocks at rest side by side in the office, set to different readings, still agree on the count between the door and the reading. That count, the proper time, is geometry, not convention. The ruled line on Herbert’s squared paper, the office’s world line, is a drawing aid: what the evidence supplies is events, their order, and counts where anybody counted. The light cone turns that into a partial order. Pairs that a signal could connect are ordered for every observer. Pairs that no signal could connect are neither before nor after. An event with no known place, like the sender of the morning sheet, cannot be ordered at all, which is why that field stays blank.", None),
]
def replace_in_par(p, old, new):
    runs = p.runs
    text = ''.join(r.text for r in runs)
    i = text.find(old)
    if i < 0: return False
    # locate runs covering [i, i+len(old))
    pos = 0; first = last = None; starts = []
    for k, r in enumerate(runs):
        starts.append(pos); pos += len(r.text)
    j = i + len(old)
    for k, r in enumerate(runs):
        s, e = starts[k], starts[k] + len(r.text)
        if first is None and e > i: first = k
        if s < j: last = k
    fr = runs[first]; off = i - starts[first]
    if first == last:
        fr.text = fr.text[:off] + new + fr.text[off + len(old):]
    else:
        lr = runs[last]; loff = j - starts[last]
        fr.text = fr.text[:off] + new
        lr.text = lr.text[loff:]
        for k in range(first + 1, last): runs[k].text = ''
    return True

def make_run(src_rpr, text, link=False):
    r = OxmlElement('w:r')
    if src_rpr is not None:
        rpr = copy.deepcopy(src_rpr)
        if link:
            for tag in ('w:color', 'w:u'):
                for e in rpr.findall(qn(tag)): rpr.remove(e)
            c = OxmlElement('w:color'); c.set(qn('w:val'), '0563C1'); rpr.append(c)
            u = OxmlElement('w:u'); u.set(qn('w:val'), 'single'); rpr.append(u)
        r.append(rpr)
    t = OxmlElement('w:t'); t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve'); t.text = text
    r.append(t); return r

def main(src, dst):
    d = docx.Document(src)
    for old, new in EDITS:
        hits = [p for p in d.paragraphs if old in ''.join(r.text for r in p.runs)]
        full = [p for p in d.paragraphs if old in p.text]
        assert len(hits) == 1 and len(full) == 1, (len(hits), len(full), old[:70])
        assert replace_in_par(hits[0], old, new)
    pars = d.paragraphs; doomed = []
    for spec in DELETES:
        if '|AFTER:' in spec:
            text, prev = spec.split('|AFTER:')
            idx = [i for i, p in enumerate(pars) if p.text == prev]
            assert len(idx) == 1, ('prev', prev[:60]); nxt = pars[idx[0] + 1]
            assert nxt.text == text, (nxt.text[:60], text); doomed.append(nxt)
        else:
            hits = [p for p in pars if p.text == spec]
            assert len(hits) == 1, (len(hits), spec[:70]); doomed.append(hits[0])
    for p in doomed:
        el = p._element; el.getparent().remove(el)
    # Lesson 6 insertion
    anchor = [p for p in d.paragraphs if p.text.startswith(LESSON_AFTER)]
    assert len(anchor) == 1
    a = anchor[0]._p
    newp = copy.deepcopy(a)
    rpr = None
    for r in newp.findall(qn('w:r')):
        if rpr is None and r.find(qn('w:rPr')) is not None: rpr = r.find(qn('w:rPr'))
    for child in list(newp):
        if child.tag != qn('w:pPr'): newp.remove(child)
    for text, target in LESSON_PARTS:
        if target:
            h = OxmlElement('w:hyperlink'); h.set(qn('w:anchor'), target)
            h.append(make_run(rpr, text, link=True)); newp.append(h)
        else:
            newp.append(make_run(rpr, text))
    a.addnext(newp)
    d.save(dst)
    print(len(EDITS), 'edits,', len(doomed), 'paragraphs removed, 1 lesson paragraph added')

if __name__ == '__main__':
    main(sys.argv[1], sys.argv[2])
