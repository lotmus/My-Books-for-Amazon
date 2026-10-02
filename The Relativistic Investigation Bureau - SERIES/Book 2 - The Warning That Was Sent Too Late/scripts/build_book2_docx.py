# -*- coding: utf-8 -*-
"""Build the Book 2 working DOCX from manuscript_text.txt.

Same layout as Book 1's rebuild_kindle.py (Georgia 11 body, Amazon Ember 18pt
bold #0C447C centred headings, chapter/lesson bookmarks, internal links on the
'-> Lesson for this chapter: N' lines). Imports Book 1's helpers read-only;
it does not modify Book 1's script or Kindle file.
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent.parent  # scripts\ -> Book 2 root
sys.path.insert(0, str(HERE.parent / "scripts"))  # rebuild_kindle.py lives in series scripts\
import rebuild_kindle as rk  # noqa: E402  (Book 1 helpers; main() is not called)

from docx import Document  # noqa: E402
from docx.shared import Inches  # noqa: E402

SRC = HERE / "manuscript_text.txt"
OUT = HERE / "The_Warning_That_Was_Sent_Too_Late_BOOK_2_DRAFT.docx"

# Book 2 arrows have no " - title" suffix.
rk.LESSON_ARROW = re.compile(r"^-> Lesson for this chapter: (\d+)(?: - .+)?$")

EXTRA_HEADINGS = {
    "The Warning That Was Sent Too Late",
    "A Relativistic Investigation Bureau Mystery",
    "Book Two",
    "In Which a File Arrives Before Its Author",
    "Tense for Investigators",
    "Wednesday, Like a Person",
}


def is_heading(line: str, prev: str) -> bool:
    if line in EXTRA_HEADINGS:
        return True
    if re.match(r"Lesson \d+ —", line):
        return False
    if re.fullmatch(r"(?:Chapter \d+|Lesson \d+|PROLOGUE|EPILOGUE)", prev):
        return True
    return bool(re.fullmatch(r"Chapter \d+|Lesson \d+", line)) or line in {
        "PROLOGUE", "EPILOGUE", "COURSE", "GLOSSARY", "FURTHER READING",
        "BIBLIOGRAPHY", "ACKNOWLEDGEMENTS",
    }


def main():
    lines = SRC.read_text(encoding="utf-8").splitlines()
    doc = Document()
    s = doc.sections[0]
    s.page_width, s.page_height = Inches(8.5), Inches(11)
    s.left_margin = s.right_margin = Inches(0.85)
    s.top_margin = s.bottom_margin = Inches(0.8)
    first, prev, in_gloss = True, "", False
    for raw in lines:
        line = raw.rstrip()
        if not line:
            continue
        if line == "GLOSSARY":
            in_gloss = True
        elif line in {"FURTHER READING", "ACKNOWLEDGEMENTS", "BIBLIOGRAPHY"}:
            in_gloss = False
        if is_heading(line, prev):
            para = rk.add_heading_line(doc, line, first=first)
            first = False
            if re.fullmatch(r"Chapter \d+", line):
                rk.mark_bookmark(para, "chapter_" + line.split()[1])
            elif re.fullmatch(r"Lesson \d+", line):
                rk.mark_bookmark(para, "lesson_" + line.split()[1])
            elif line == "EPILOGUE":
                rk.mark_bookmark(para, "epilogue")
        else:
            rk.add_body(doc, line, glossary=in_gloss)
        prev = line
    doc.save(OUT)
    print("Wrote", OUT, "paragraphs", len(doc.paragraphs))


if __name__ == "__main__":
    main()
