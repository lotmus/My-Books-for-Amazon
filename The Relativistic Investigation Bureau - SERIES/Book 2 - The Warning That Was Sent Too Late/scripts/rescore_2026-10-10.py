"""The Warning That Was Sent Too Late: 10 Oct 2026 typography and back-matter fixes.
Straight quotes to curly (Book 1 is curly), and the Also-by page leaves out this book (canonical rule).
Usage: python rescore_2026-10-10.py SRC DST"""
import docx, re, sys
from docx.oxml.ns import qn
SRC, DST = sys.argv[1:3]
d = docx.Document(SRC)
ELIS = re.compile(r"(em|tis|twas|til|cause|n|ere|ow|appen|im|er|alf)\b|\d\d?s\b|\d\d\b", re.I)
OPEN_BEFORE = set(' \t\n([{\u2014\u2013-/\u201c\u2018')
def curl(s):
    out = list(s)
    for k, c in enumerate(s):
        prev = s[k-1] if k else ' '
        nxt = s[k+1] if k + 1 < len(s) else ' '
        if c == '"':
            out[k] = '\u201c' if (prev in OPEN_BEFORE or k == 0) and not nxt.isspace() else '\u201d'
            if prev in OPEN_BEFORE and nxt.isspace(): out[k] = '\u201d' if k else '\u201c'
        elif c == "'":
            if prev.isalnum() and (nxt.isalnum()):
                out[k] = '\u2019'
            elif prev in OPEN_BEFORE or k == 0:
                out[k] = '\u2019' if ELIS.match(s[k+1:]) else '\u2018'
            else:
                out[k] = '\u2019'
    return ''.join(out)
def paras(doc):
    body = doc.element.body
    for p in body.iter(qn('w:p')): yield p
n = 0
for p in paras(d):
    ts = [t for t in p.iter(qn('w:t'))]
    full = ''.join(t.text or '' for t in ts)
    if '"' not in full and "'" not in full: continue
    new = curl(full); pos = 0
    for t in ts:
        L = len(t.text or ''); 
        if L: t.text = new[pos:pos+L]
        pos += L
    n += 1
# headers/footers too
for s in d.sections:
    for part in (s.header, s.footer):
        for p in part._element.iter(qn('w:p')):
            for t in p.iter(qn('w:t')):
                if t.text and ('"' in t.text or "'" in t.text): t.text = curl(t.text)
# Also-by: drop this book's own line
P = d.paragraphs
hits = [i for i, p in enumerate(P) if p.text.startswith('The Warning That Was Sent Too Late (The Relativistic Investigation Bureau, Book 2)')]
for i in hits:
    el = P[i]._p; el.getparent().remove(el)

# Bibliography and Further Reading: book titles as keyword links (Open Library), no visible URLs
from docx.oxml import OxmlElement
import copy
RT = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink"
OL = 'https://openlibrary.org/works/'
LINKS = [
 ('Carroll, Sean.', 'From Eternity to Here: The Quest for the Ultimate Theory of Time', 'OL5752308W'),
 ('Deutsch, David.', 'The Fabric of Reality', 'OL2635835W'),
 ('Gott, J. Richard', 'Time Travel in Einstein\u2019s Universe: The Physical Possibilities of Travel Through Time', 'OL8074339W'),
 ('Nahin, Paul J.', 'Time Machines: Time Travel in Physics, Metaphysics, and Science Fiction', 'OL1907388W'),
 ('Novikov, Igor D.', 'The River of Time', 'OL2734517W'),
 ('Price, Huw.', 'Time\u2019s Arrow and Archimedes\u2019 Point: New Directions for the Physics of Time', 'OL2939060W'),
 ('Thorne, Kip S.', 'Black Holes and Time Warps: Einstein\u2019s Outrageous Legacy', 'OL3902081W'),
 ('Kip S. Thorne, Black Holes', 'Black Holes and Time Warps: Einstein\u2019s Outrageous Legacy', 'OL3902081W'),
 ('Igor D. Novikov, The River', 'The River of Time', 'OL2734517W'),
 ('Huw Price, Time', 'Time\u2019s Arrow and Archimedes\u2019 Point', 'OL2939060W'),
]
def linkify(p, phrase, url):
    for r in p._p.iter(qn('w:r')):
        if r.getparent().tag == qn('w:hyperlink'): continue
        t = r.find(qn('w:t'))
        if t is None or not t.text or phrase not in t.text: continue
        a, b = t.text.split(phrase, 1); rpr = r.find(qn('w:rPr'))
        def mk(text, link=False):
            nr = OxmlElement('w:r'); rp = copy.deepcopy(rpr) if rpr is not None else OxmlElement('w:rPr')
            if link:
                c = OxmlElement('w:color'); c.set(qn('w:val'), '0563C1'); rp.append(c); u = OxmlElement('w:u'); u.set(qn('w:val'), 'single'); rp.append(u)
            if len(rp): nr.append(rp)
            nt = OxmlElement('w:t'); nt.set(qn('xml:space'), 'preserve'); nt.text = text; nr.append(nt); return nr
        h = OxmlElement('w:hyperlink'); h.set(qn('r:id'), d.part.relate_to(url, RT, is_external=True)); h.set(qn('w:history'), '1'); h.append(mk(phrase, True))
        if a: r.addprevious(mk(a))
        r.addprevious(h)
        if b: r.addprevious(mk(b))
        r.getparent().remove(r); return True
    raise AssertionError((phrase, p.text[:100]))
P = d.paragraphs
for start, title, key in LINKS:
    hit = [p for p in P if p.text.startswith(start)]; assert len(hit) == 1, (start, len(hit))
    linkify(hit[0], title, OL + key)

d.save(DST); print('paragraphs curled', n, 'also-by self lines removed', len(hits))
