import re, glob, docx
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

TITLE="The Will Is Not the Legacy, Actually"
SUB="Wills, the real alternatives, video messages, and the family list that keeps it all usable"
AUTHOR="Lothar J. Musiol"
md_files=sorted(glob.glob('/workspace/will/md/*.md'))

doc=Document()
sec=doc.sections[0]
sec.page_width=Inches(6); sec.page_height=Inches(9)
sec.left_margin=sec.right_margin=Inches(0.75); sec.top_margin=sec.bottom_margin=Inches(0.8)
st=doc.styles['Normal']; st.font.name='Georgia'; st.font.size=Pt(10.5)
st.element.rPr.rFonts.set(qn('w:eastAsia'),'Georgia')
st.paragraph_format.space_after=Pt(5); st.paragraph_format.line_spacing=1.15
for lvl,sz in [(1,20),(2,15),(3,12.5),(4,11)]:
    h=doc.styles[f'Heading {lvl}']; h.font.name='Georgia'; h.font.size=Pt(sz); h.font.bold=True
    h.font.color.rgb=RGBColor(0x1F,0x2A,0x44)
    rpr=h.element.get_or_add_rPr(); rf=rpr.find(qn('w:rFonts'))
    if rf is None: rf=OxmlElement('w:rFonts'); rpr.append(rf)
    for a in ('w:ascii','w:hAnsi','w:eastAsia','w:cs'): rf.set(qn(a),'Georgia')
    h.paragraph_format.space_before=Pt(18 if lvl<=2 else 12); h.paragraph_format.space_after=Pt(6)
    h.paragraph_format.keep_with_next=True

bm_id=[0]
def add_bookmark(p,name):
    s=OxmlElement('w:bookmarkStart'); s.set(qn('w:id'),str(bm_id[0])); s.set(qn('w:name'),name)
    e=OxmlElement('w:bookmarkEnd'); e.set(qn('w:id'),str(bm_id[0]))
    p._p.insert(1 if p._p.pPr is not None else 0,s); p._p.append(e); bm_id[0]+=1

def add_internal_link(p,text,anchor):
    h=OxmlElement('w:hyperlink'); h.set(qn('w:anchor'),anchor); h.set(qn('w:history'),'1')
    r=OxmlElement('w:r'); rpr=OxmlElement('w:rPr')
    c=OxmlElement('w:color'); c.set(qn('w:val'),'1F2A44'); rpr.append(c)
    r.append(rpr); t=OxmlElement('w:t'); t.text=text; t.set(qn('xml:space'),'preserve'); r.append(t); h.append(r); p._p.append(h)

INL=re.compile(r'(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)')
LINK=re.compile(r'\[([^\]]+)\]\(#([^)]+)\)')
LINKB=re.compile(r'(\*\*)?\[([^\]]+)\]\(#([^)]+)\)(\*\*)?')
HEADMAP={}; USED_LINKS=[]
def add_runs(p,text,italic=False,bold=False):
    # internal links [label](#Heading text) -> hyperlink to that heading's bookmark
    if LINK.search(text):
        pos=0
        for m in LINKB.finditer(text):
            if m.start()>pos: add_runs(p,text[pos:m.start()],italic,bold)
            target=m.group(3).strip()
            if target not in HEADMAP: raise SystemExit('unresolved link target: '+target)
            add_internal_link(p,m.group(2),HEADMAP[target]); USED_LINKS.append(HEADMAP[target])
            rpr=p._p[-1].find(qn('w:r')).find(qn('w:rPr'))
            u=OxmlElement('w:u'); u.set(qn('w:val'),'single'); rpr.append(u)
            if bold or (m.group(1) and m.group(4)): rpr.append(OxmlElement('w:b'))
            pos=m.end()
        if pos<len(text): add_runs(p,text[pos:],italic,bold)
        return
    for part in INL.split(text):
        if not part: continue
        if part.startswith('**') and part.endswith('**'):
            if LINK.search(part): add_runs(p,part[2:-2],italic,True); continue
            r=p.add_run(part[2:-2]); r.bold=True; r.italic=italic
        elif part.startswith('*') and part.endswith('*') and len(part)>2: r=p.add_run(part[1:-1]); r.italic=not italic; r.bold=bold
        elif part.startswith('`'): r=p.add_run(part[1:-1]); r.font.name='Consolas'; r.font.size=Pt(9.5)
        else: r=p.add_run(part); r.italic=italic; r.bold=bold

def page_break():
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

def centered(text,size,bold=False,italic=False,space=12):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run(text); r.font.size=Pt(size); r.bold=bold; r.italic=italic
    p.paragraph_format.space_after=Pt(space); return p

# Title page
for _ in range(6): doc.add_paragraph()
centered(TITLE,24,bold=True,space=14)
centered(SUB,13,italic=True,space=60)
centered(AUTHOR,15,space=6)
page_break()
# Copyright page
for _ in range(10): doc.add_paragraph()
cp=["Copyright © 2026 Lothar J. Musiol. All rights reserved.",
"No part of this book may be reproduced or transmitted in any form or by any means without written permission from the author, except for brief quotations in reviews and for the worksheets and checklists in the appendices, which readers may copy for their own personal use.",
"This book is an educational guide. It is not legal, tax, financial, or medical advice, and it does not create a lawyer-client relationship. Laws on wills, trusts, compulsory shares, taxes, digital assets, powers of attorney, and healthcare directives differ by country and by state, and they change. Legal figures naming California, the United States, Germany, or the European Union were checked against official sources in October 2026. Have a qualified lawyer or notary where you live review any document that transfers property, names a guardian, or grants a power of attorney before you rely on it.",
"All people and families in this book are fictional. Any resemblance to real persons is coincidental.",
"First edition, 2026."]
for t in cp:
    p=doc.add_paragraph(); r=p.add_run(t); r.font.size=Pt(8.5)
page_break()

# parse md -> blocks
blocks=[]
for f in md_files:
    lines=open(f,encoding='utf-8').read().split('\n'); i=0
    while i<len(lines):
        l=lines[i]
        if not l.strip(): i+=1; continue
        m=re.match(r'^(#{1,4}) (.*)',l)
        if m: blocks.append(('h',len(m.group(1)),m.group(2).strip())); i+=1; continue
        if l.startswith('|'):
            rows=[]
            while i<len(lines) and lines[i].startswith('|'):
                cells=[c.strip() for c in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r':?-+:?',c) for c in cells if c): rows.append(cells)
                i+=1
            blocks.append(('table',rows)); continue
        if l.startswith('>'):
            q=[]
            while i<len(lines) and lines[i].startswith('>'):
                q.append(lines[i][1:].strip()); i+=1
            paras=[]; cur=[]
            for x in q:
                if x: cur.append(x)
                elif cur: paras.append(' '.join(cur)); cur=[]
            if cur: paras.append(' '.join(cur))
            blocks.append(('quote',paras)); continue
        m=re.match(r'^(\s*)([-*]|\d+\.) (.*)',l)
        if m:
            items=[]
            while i<len(lines):
                m=re.match(r'^(\s*)([-*]|\d+\.) (.*)',lines[i])
                if not m: break
                items.append((m.group(2)[0].isdigit(), m.group(3))); i+=1
            blocks.append(('list',items)); continue
        para=[l.strip()]; i+=1
        while i<len(lines) and lines[i].strip() and not re.match(r'^(#{1,4} |\||>|\s*([-*]|\d+\.) )',lines[i]):
            para.append(lines[i].strip()); i+=1
        blocks.append(('p',' '.join(para)))

# headings & anchors
heads=[]; n=0
for b in blocks:
    if b[0]=='h':
        n+=1; heads.append((b[1],b[2],f'_sec{n:04d}'))
EXTRA=[(1,'About the Author','_secabout'),(1,'Also by Lothar J. Musiol','_secalso')]
for lvl,text,anc in heads+EXTRA:
    HEADMAP.setdefault(text,anc)

# TOC
tp=doc.add_paragraph(style='Heading 1'); tp.add_run('Contents')
for lvl,text,anc in heads+EXTRA:
    p=doc.add_paragraph(); p.paragraph_format.left_indent=Inches(0.22*(lvl-1))
    p.paragraph_format.space_after=Pt(1 if lvl>2 else 3)
    if lvl==1: p.paragraph_format.space_before=Pt(6)
    add_internal_link(p,text,anc)
    for r in p._p.iter(qn('w:r')):
        rpr=r.find(qn('w:rPr')); sz=OxmlElement('w:sz'); sz.set(qn('w:val'),str(int(2*(10 if lvl<=2 else 9)))); rpr.append(sz)
        if lvl==1: rpr.append(OxmlElement('w:b'))

hi=0; first_h1=True
for b in blocks:
    if b[0]=='h':
        lvl,text,anc=heads[hi]; hi+=1
        p=doc.add_paragraph(style=f'Heading {lvl}'); add_runs(p,text); add_bookmark(p,anc)
        if lvl==1 or (lvl==2 and text.startswith(('Chapter','Appendix'))): p.paragraph_format.page_break_before=True
    elif b[0]=='p':
        p=doc.add_paragraph(); t=b[1]
        if t.startswith('*') and t.endswith('*') and t.count('*')==2: add_runs(p,t[1:-1],italic=True)
        else: add_runs(p,t)
    elif b[0]=='quote':
        for t in b[1]:
            p=doc.add_paragraph(); p.paragraph_format.left_indent=Inches(0.3); p.paragraph_format.right_indent=Inches(0.2)
            add_runs(p,t)
    elif b[0]=='list':
        for num,t in b[1]:
            box=t.startswith('[ ] ')
            if box: t=t[4:]
            p=doc.add_paragraph(style='List Number' if num and not box else 'List Bullet')
            if box:
                p.style=doc.styles['Normal']; p.paragraph_format.left_indent=Inches(0.3); p.paragraph_format.first_line_indent=Inches(-0.22)
                p.add_run('☐  ')
            add_runs(p,t)
    elif b[0]=='table':
        rows=b[1]; ncol=max(len(r) for r in rows)
        t=doc.add_table(rows=len(rows),cols=ncol); t.style='Table Grid'
        for ri,row in enumerate(rows):
            for ci in range(ncol):
                cell=t.cell(ri,ci); cell.text=''; pp=cell.paragraphs[0]
                add_runs(pp,row[ci] if ci<len(row) else '',bold=(ri==0))
                for r in pp.runs: r.font.size=Pt(8.5)
                pp.paragraph_format.space_after=Pt(1)
        doc.add_paragraph()

# Review invitation
page_break()
for _ in range(3): doc.add_paragraph()
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run("If this book gave you something, a laugh, an idea, or a better question, I'd love to hear about it. A short review on Amazon, even a sentence or two, helps other curious readers find it, and I read every one."); r.italic=True
# About the author (Version B)
page_break()
p=doc.add_paragraph(style='Heading 1'); p.add_run('About the Author'); add_bookmark(p,'_secabout')
about=open('/workspace/will/about.md',encoding='utf-8').read()
vb=about.split('## Version B',1)[1].split('Plain-text copies',1)[0].strip()
for para in [x.strip() for x in vb.split('\n\n') if x.strip()]:
    doc.add_paragraph(para)
# Also by
page_break()
p=doc.add_paragraph(style='Heading 1'); p.add_run('Also by Lothar J. Musiol'); add_bookmark(p,'_secalso')
ab=open('/workspace/will/alsoby.md',encoding='utf-8').read()
sect=ab.split('## Also by Lothar J. Musiol',1)[1].split('## Also by Kevin Drew Peters',1)[0]
for line in sect.strip().split('\n'):
    line=line.strip()
    if not line: continue
    if line.startswith('**') and line.endswith('**'):
        p=doc.add_paragraph(); r=p.add_run(line.strip('*')); r.bold=True; p.paragraph_format.space_before=Pt(8)
    elif line.startswith('- '):
        p=doc.add_paragraph(style='List Bullet'); add_runs(p,line[2:])

# footer page numbers
fp=sec.footer.paragraphs[0]; fp.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=fp.add_run(); f1=OxmlElement('w:fldChar'); f1.set(qn('w:fldCharType'),'begin'); r._r.append(f1)
it=OxmlElement('w:instrText'); it.text='PAGE'; r._r.append(it)
f2=OxmlElement('w:fldChar'); f2.set(qn('w:fldCharType'),'end'); r._r.append(f2)
# keep w:rPr children in schema order (Word is strict about it)
ORDER=['rStyle','rFonts','b','bCs','i','iCs','caps','smallCaps','strike','dstrike','outline','shadow','emboss','imprint','noProof','snapToGrid','vanish','webHidden','color','spacing','w','kern','position','sz','szCs','highlight','u','effect','bdr','shd','fitText','vertAlign','rtl','cs','em','lang','eastAsianLayout','specVanish','oMath']
for rpr in doc.element.body.iter(qn('w:rPr')):
    kids=list(rpr)
    if all(k.tag.startswith('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}') for k in kids):
        kids.sort(key=lambda k: ORDER.index(k.tag.split('}')[1]) if k.tag.split('}')[1] in ORDER else 99)
        for k in kids: rpr.remove(k)
        for k in kids: rpr.append(k)
doc.core_properties.title=TITLE; doc.core_properties.subject=SUB; doc.core_properties.author=AUTHOR
out='/workspace/will/out/The_Will_Is_Not_the_Legacy_Actually_MASTER.docx'
doc.save(out); print(out, len(heads), 'headings;', len(USED_LINKS), 'in-text links resolved')
