# -*- coding: utf-8 -*-
"""Document-level finish: retitle, copyright, map, heading surgery,
static hyperlinked TOC, bookmarks, lesson links, gutter, bibliography.
Run after lessons have been spliced into Quantum, Actually - Volume 2.docx.
"""
import copy
import os
import re
import sys
import time

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.shared import Inches, Pt, RGBColor, Twips
from docx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import build_lesson as bl

COMPLETE = bl.COMPLETE
BOOKMARK_TOC = "QEDCourseTOC"

KEEP_H1 = re.compile(
    r"^(Lesson \d+[: ]|Part [IVX0]|Prologue \d+|Interlude[: ]|Table of Contents|Copyright|"
    r"How to read this book|Also in This Series|Glossary of symbols|Course capstone|"
    r"Consolidated formula index|Bibliography|Also by Lothar J\. Musiol$|"
    r"Complete Quantum Electrodynamics Course$|Volume \d+ - .*$)"
)
LESSON_H1 = re.compile(r"^Lesson (\d+)[: ]")

# Quantum, Actually series front matter (same layout as Physics, Actually):
# Title = series name, Heading 1 = book title, subtitle, series line, author.
SERIES = bl.SERIES
SERIES_LINE = "A Volume in the %s Series" % SERIES  # as in Math, Actually
BOOK_TITLE = bl.VOLUME_TITLE
VOLUME_H1 = "Volume %d - %s" % (bl.SERIES_VOLUME, bl.VOLUME_TITLE)  # Math, Actually: "Volume 1 - From Arithmetic to Calculus"
SUBTITLE = "From Mathematical Foundations to One-Loop QED"
SERIES_BLUE = RGBColor(0x4F, 0x81, 0xBD)
ALSO_IN_SERIES = [
    "Volume 1 — Questions and Answers from the Double Slit to the Superconducting Wire. "
    "The conceptual book, with no calculation required: interference, entanglement and Bell tests, "
    "quantum fields, and, in its Part Two, electromagnetism read outward from quantum phase and the "
    "potential, the view this course's Interlude after Lesson 41 summarizes.",
    "Volume 2 — A QED Course: From Mathematical Foundations to One-Loop QED. "
    "This book: the calculation course.",
]

# Back matter: the canonical "Also by Lothar J. Musiol" list (same in every book; see
# notes/ALSO_BY - canonical list.md at the repo root). Titles as on each master's title page.
ALSO_BY_HEADING = "Also by Lothar J. Musiol"
# Rebuilt 10 Oct 2026 from notes/ALSO_BY - canonical list.md. The book itself is left out; Kevin Drew
# Peters books never appear on a Lothar J. Musiol page.
ALSO_BY = [
    ('Physics, Actually', [
        'Physics, Actually, Volume 1: Motion, Forces, Time, and Relativity',
        'Physics, Actually, Volume 2: Gravity, Cosmology, and the Limits of Spacetime',
        'Physics, Actually, Volume 3: The Standard Model, Chaos, and the Edge of Knowledge',
        'Life, Actually: From the First Cell to the Edited Genome and the Search for Life Elsewhere',
    ]),
    ('More Actually Books', [
        'Gravitation, Actually: A Short Book on Gravity',
        'The Universe, Actually: What Humanity Currently Understands About Reality',
    ]),
    ('Math, Actually', [
        'Math, Actually, Volume 1: From Arithmetic to Calculus',
        'Math, Actually, Volume 2: From Multivariable Calculus to Set Theory & Logic',
        'Math, Actually, Volume 3: From Differential Equations to Abstract Algebra',
        'Math, Actually, Volume 4: From Category Theory to the Frontier',
    ]),
    ('Quantum, Actually', [
        'Quantum, Actually, Volume 1: Questions and Answers from the Double Slit to the Superconducting Wire',
    ]),
    ('Science Sparks', [
        'Science Sparks: Physics, Life, and Mathematics — The Same Few Rules, Told in Highlights',
    ]),
    ('Look First', [
        'Look First, Volume 1: The Universe Has No Now',
        'Look First, Volume 2: A Trip Is Not a New Life',
    ]),
    ('Electrical Engineering Series', [
        'Foundations of Electronics (Book 1)',
        'Circuits, Components, and Control (Book 2)',
        'Semiconductor Physics and Devices (Book 3)',
        'RF, Microwave, and Transceivers (Book 4)',
        'Communications, Wireless, and SDR (Book 5)',
        'Power and Energy (Book 6)',
        'Packaging, Layout, EMC, and Test (Book 7)',
    ]),
    ('History', [
        "The Dolphins' View of History",
    ]),
    ('Fiction', [
        "The Murder That Hadn't Happened Yet (The Relativistic Investigation Bureau, Book 1), Lothar J. Musiol writing as Spezala Genara Relavi (S.G.R.)",
        'The Warning That Was Sent Too Late (The Relativistic Investigation Bureau, Book 2), Lothar J. Musiol writing as Spezala Genara Relavi (S.G.R.)',
        "Schrödinger's Paperwork (Lolly Wren's Curious Science Adventures, Book 1), Lothar J. Musiol writing as Quentin Edmund Darling (Q.E.D.)",
        "The Permitted Options (Lolly Wren's Curious Science Adventures, Book 2), Lothar J. Musiol writing as Quentin Edmund Darling (Q.E.D.)",
        'Protocol Flamingo (The Invasion Storybooks, Book 1), Lothar J. Musiol writing as George Herbert Fontaine',
        'In Love with Murder, Actually: A Comic Mystery, Lothar J. Musiol writing as Romuald Livingston',
    ]),
    ('In German', [
        'Auswandern – Ab in die USA, Kanada, Australien oder Neuseeland?',
    ]),
]


def normalize_numbering(doc):
    """Heading numbering in the Physics, Actually style (idempotent):
    'Lesson 1 Title' -> 'Lesson 1: Title', 'Prologue 1 Title' -> 'Prologue 1: Title',
    'Part I Title' -> 'Part I — Title', 'Interlude Mead's View' -> 'Interlude: Mead's View'."""
    rules = [
        (re.compile(r"^(Lesson \d+) (?![:—])(.+)$"), r"\1: \2"),
        (re.compile(r"^(Prologue \d+) (?![:—])(.+)$"), r"\1: \2"),
        (re.compile(r"^(Part (?:[IVX]+|0)) (?![:—])(.+)$"), r"\1 — \2"),
        (re.compile(r"^Interlude (?![:—])(.+)$"), r"Interlude: \1"),
    ]
    n = 0
    for child in doc.element.body:
        if _p_style(child) != "Heading1":
            continue
        text = _text(child)
        for rx, sub in rules:
            if rx.match(text):
                ts = list(child.iter(qn("w:t")))
                ts[0].text = rx.sub(sub, text)
                ts[0].set(qn("xml:space"), "preserve")
                for t in ts[1:]:
                    t.text = ""
                n += 1
                break
    return n


def _series_runs(p, text):
    for r in list(p.runs):
        r._r.getparent().remove(r._r)
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(13)
    run.font.color.rgb = SERIES_BLUE
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER


def _set_text(p, text):
    runs = p.runs
    if runs:
        runs[0].text = text
        for r in runs[1:]:
            r.text = ""
    else:
        p.add_run(text)


SERIES_LINE_RE = re.compile(r"^(Volume \d+ in the .* Series|A Volume in the .* Series)$")


def series_title_page(doc):
    """Idempotent, Math, Actually layout: Title = series name, Heading 1 =
    'Volume N - Title', subtitle, 'A Volume in the ... Series', author
    (bold blue 13 pt); running head = full title."""
    paras = doc.paragraphs[:14]
    title = next((p for p in paras if p.style is not None and p.style.name == "Title"), None)
    if title is not None and title.text.strip() != SERIES:
        _set_text(title, SERIES)
    for p in paras:
        if p.style is not None and p.style.name == "Heading 1":
            if p.text.strip() in ("Copyright", "Table of Contents"):
                break
            if p.text.strip() != VOLUME_H1:
                _set_text(p, VOLUME_H1)
            break
    paras = doc.paragraphs[:14]
    have = False
    for p in paras:
        if SERIES_LINE_RE.match(p.text.strip()):
            have = True
            if p.text.strip() != SERIES_LINE:
                for r in list(p.runs):
                    r._r.getparent().remove(r._r)
                _series_runs(p, SERIES_LINE)
    sub = next((p for p in paras if p.style is not None and p.style.name == "Subtitle"
                and "mathematical foundations" in p.text.lower()), None)
    if not have and sub is not None:
        line = doc.add_paragraph()
        sub._p.addnext(line._p)
        _series_runs(line, SERIES_LINE)
    for p in doc.paragraphs[:14]:
        if p.text.strip() == bl.AUTHOR and p.style is not None and p.style.name == "Subtitle":
            p.style = doc.styles["Normal"]
            _series_runs(p, bl.AUTHOR)
    for sec in doc.sections:
        for hdr in (sec.header, sec.even_page_header):
            for p in hdr.paragraphs:
                if "QED Course" in p.text or "Actually" in p.text:
                    _set_text(p, bl.RUNNING_HEAD)
    doc.core_properties.title = bl.FULL_TITLE
    doc.core_properties.subject = "%s series, Volume %d" % (SERIES, bl.SERIES_VOLUME)
    doc.core_properties.keywords = ("quantum electrodynamics; QED; Feynman diagrams; Dirac equation; "
                                    "quantum field theory; renormalization; g-2; Lamb shift; exercises with solutions")
    doc.core_properties.author = bl.AUTHOR
LESSON_MENTION = re.compile(r"Lesson (\d+)")


def _is_body_start(text):
    t = text or ""
    return (
        LESSON_H1.match(t) is not None and LESSON_H1.match(t).group(1) == "1"
        or t.startswith("Part 0")
        or t.startswith("Prologue 1")
    )


def _text(child):
    if child.tag != qn("w:p"):
        return ""
    return "".join(t.text or "" for t in child.iter(qn("w:t")))


def _p_style(child):
    if child.tag != qn("w:p"):
        return None
    el = child.find(qn("w:pPr") + "/" + qn("w:pStyle"))
    if el is None:
        return None
    return el.get(qn("w:val"))


def _set_p_style(child, style_id):
    pPr = child.find(qn("w:pPr"))
    if pPr is None:
        pPr = OxmlElement("w:pPr")
        child.insert(0, pPr)
    el = pPr.find(qn("w:pStyle"))
    if el is None:
        el = OxmlElement("w:pStyle")
        pPr.insert(0, el)
    el.set(qn("w:val"), style_id)


def _bookmark_p(child, name, bid):
    for s in child.findall(qn("w:bookmarkStart")):
        if s.get(qn("w:name")) == name:
            return
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bid))
    start.set(qn("w:name"), name)
    child.insert(0, start)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bid))
    child.append(end)


def _hyperlink_run(paragraph, text, anchor):
    hl = OxmlElement("w:hyperlink")
    hl.set(qn("w:anchor"), anchor)
    r = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    st = OxmlElement("w:rStyle")
    st.set(qn("w:val"), "Hyperlink")
    rPr.append(st)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    rPr.append(color)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rPr.append(u)
    r.append(rPr)
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    r.append(t)
    hl.append(r)
    paragraph._p.append(hl)


def find_first_pagebreak(doc):
    body = doc.element.body
    for child in body:
        if child.tag == qn("w:p"):
            for br in child.iter(qn("w:br")):
                if br.get(qn("w:type")) == "page":
                    return child
    raise RuntimeError("no title-page break")


def remove_old_toc(doc):
    body = doc.element.body
    in_block = False
    to_remove = []
    for child in list(body):
        text = _text(child)
        style = _p_style(child)
        starts = child.findall(qn("w:bookmarkStart")) if child.tag == qn("w:p") else []
        if any(s.get(qn("w:name")) == BOOKMARK_TOC for s in starts) or (
            style == "Heading1" and text == "Table of Contents"
        ):
            in_block = True
        if in_block:
            to_remove.append(child)
        if in_block and (
            text == "__QEDCourseTOC_END__"
            or (style == "Heading1" and _is_body_start(text))
        ):
            if _is_body_start(text):
                to_remove.pop()  # keep Part 0 / Prologue 1 / Lesson 1
            break
    for node in to_remove:
        parent = node.getparent()
        if parent is not None:
            parent.remove(node)


def remove_old_front(doc):
    """Remove previously inserted copyright / how-to-read blocks if re-run."""
    body = doc.element.body
    markers = {"Copyright", "Also in This Series", "How to read this book"}
    in_block = False
    to_remove = []
    for child in list(body):
        text = _text(child)
        style = _p_style(child)
        if style == "Heading1" and text in markers:
            in_block = True
        if in_block:
            to_remove.append(child)
        if in_block and style == "Heading1" and text not in markers and text:
            to_remove.pop()
            break
        if in_block and style == "Heading1" and _is_body_start(text):
            to_remove.pop()
            break
    for node in to_remove:
        parent = node.getparent()
        if parent is not None:
            parent.remove(node)


def retitle(doc):
    for p in doc.paragraphs[:12]:
        if p.style and p.style.name == "Subtitle":
            if "precision QED" in p.text or "mathematical foundations" in p.text.lower():
                for r in p.runs:
                    r.text = ""
                if p.runs:
                    p.runs[0].text = SUBTITLE
                else:
                    p.add_run(SUBTITLE)
        if p.text.strip() == "Mathematical foundations for quantum states, operators, spinors, and fields":
            for r in p.runs:
                r.text = ""
            if p.runs:
                p.runs[0].text = (
                    "A calculation course from complex numbers to a=α/(2π), "
                    "the running of α, and the Lamb and Schwinger mechanisms"
                )
            else:
                p.add_run(
                    "A calculation course from complex numbers to a=α/(2π), "
                    "the running of α, and the Lamb and Schwinger mechanisms"
                )
        if p.text.strip() == "Course edition 1.0":
            for r in p.runs:
                r.text = ""
            if p.runs:
                p.runs[0].text = "First edition, 2026"
            else:
                p.add_run("First edition, 2026")


def set_gutter(doc):
    for sec in doc.sections:
        sec.left_margin = Inches(1.15)
        sec.right_margin = Inches(0.95)
        sec.top_margin = Inches(0.95)
        sec.bottom_margin = Inches(0.9)
        pgMar = sec._sectPr.find(qn("w:pgMar"))
        if pgMar is not None:
            # w:gutter is in twips, not EMU: Inches(0.3) is 274320 EMU, which as twips
            # is 190 in, leaving a negative text width (one letter per line).
            pgMar.set(qn("w:gutter"), str(Inches(0.3).twips))


def demote_headings(doc):
    n = 0
    for child in doc.element.body:
        if _p_style(child) != "Heading1":
            continue
        text = _text(child)
        if KEEP_H1.match(text or ""):
            continue
        _set_p_style(child, "Heading2")
        n += 1
    return n


def bookmark_lessons(doc):
    n = 0
    for child in doc.element.body:
        if _p_style(child) != "Heading1":
            continue
        m = LESSON_H1.match(_text(child) or "")
        if not m:
            continue
        num = int(m.group(1))
        _bookmark_p(child, "Lesson%d" % num, 1000 + num)
        n += 1
    for child in doc.element.body:
        if _p_style(child) != "Heading1":
            continue
        text = _text(child) or ""
        if text.startswith("Part "):
            token = text.split()[1] if len(text.split()) > 1 else ""
            if token.startswith("0"):
                slug = "Part0"
            else:
                slug = "Part" + re.sub(r"[^IVX]", "", token)
            _bookmark_p(child, slug or "Part", 2000 + hash(text) % 500)
        elif text == "Glossary of symbols":
            _bookmark_p(child, "Glossary", 3001)
        elif text == "Course capstone":
            _bookmark_p(child, "Capstone", 3002)
        elif text.startswith("Consolidated formula"):
            _bookmark_p(child, "FormulaIndex", 3003)
        elif text == "Bibliography":
            _bookmark_p(child, "Bibliography", 3004)
        elif text.startswith("Prologue "):
            pm = re.match(r"^Prologue (\d+)", text)
            if pm:
                _bookmark_p(child, "Prologue%s" % pm.group(1), 4000 + int(pm.group(1)))
        elif text.startswith("Interlude"):
            _bookmark_p(child, "Mead", 4010)
        elif text == "How to read this book":
            _bookmark_p(child, "HowToRead", 3005)
        elif text == "Copyright":
            _bookmark_p(child, "Copyright", 3006)
    return n


def collect_toc_entries(doc):
    entries = []
    for child in doc.element.body:
        if _p_style(child) != "Heading1":
            continue
        text = _text(child).strip()
        if not text or text in ("Table of Contents", "Copyright"):
            continue
        m = LESSON_H1.match(text)
        if m:
            entries.append(("Lesson%d" % int(m.group(1)), text, False))
            continue
        pm = re.match(r"^Prologue (\d+)", text)
        if pm:
            entries.append(("Prologue%s" % pm.group(1), text, False))
            continue
        if text.startswith("Interlude"):
            entries.append(("Mead", text, False))
            continue
        if text.startswith("Part "):
            token = text.split()[1] if len(text.split()) > 1 else ""
            if token.startswith("0"):
                slug = "Part0"
            else:
                slug = "Part" + re.sub(r"[^IVX]", "", token)
            entries.append((slug or "Part", text, True))
        elif text == "How to read this book":
            entries.append(("HowToRead", text, False))
        elif text == "Glossary of symbols":
            entries.append(("Glossary", text, False))
        elif text == "Course capstone":
            entries.append(("Capstone", text, False))
        elif text.startswith("Consolidated formula"):
            entries.append(("FormulaIndex", text, False))
        elif text == "Bibliography":
            entries.append(("Bibliography", text, False))
        elif text == ALSO_BY_HEADING:
            entries.append(("AlsoBy", text, False))
    return entries


def insert_static_toc(doc, entries):
    remove_old_toc(doc)
    # Place TOC after How to read this book if present, else after first page break
    body = doc.element.body
    anchor = None
    for child in body:
        if _p_style(child) == "Heading1" and _text(child) == "How to read this book":
            # find the page break after this section, or the Lesson 1 heading
            node = child
            last = child
            while node is not None:
                nxt = node.getnext()
                if nxt is not None and _p_style(nxt) == "Heading1" and _is_body_start(_text(nxt)):
                    anchor = last
                    break
                last = nxt if nxt is not None else last
                node = nxt
            break
    if anchor is None:
        anchor = find_first_pagebreak(doc)

    heading = doc.add_paragraph("Table of Contents", style="Heading 1")
    _bookmark_p(heading._p, BOOKMARK_TOC, 9001)
    intro = doc.add_paragraph(
        "Click any entry to jump. This table lists the course in reading order."
    )

    created = [heading._p, intro._p]
    for anchor_name, label, is_part in entries:
        p = doc.add_paragraph()
        if is_part:
            p.paragraph_format.space_before = Pt(8)
            run = p.add_run(label)
            run.bold = True
        else:
            p.paragraph_format.left_indent = Inches(0.2)
            _hyperlink_run(p, label, anchor_name)
        created.append(p._p)

    pb = doc.add_paragraph()
    pb.add_run().add_break(WD_BREAK.PAGE)
    created.append(pb._p)

    for el in created:
        anchor.addnext(el)
        anchor = el


def insert_front_matter(doc):
    remove_old_front(doc)
    anchor = find_first_pagebreak(doc)
    b = bl.Builder(doc)

    h = b.heading("Copyright", 1)
    _bookmark_p(h._p, "Copyright", 3006)
    b.para(SERIES)
    b.para("Volume %d: %s" % (bl.SERIES_VOLUME, BOOK_TITLE))
    b.para(SUBTITLE)
    b.para("Lothar J. Musiol")
    b.para("First edition, 2026")
    b.para(
        "Copyright © 2026 Lothar J. Musiol. All rights reserved. Independently published. "
        "No ISBN has been assigned to this file. Permission is granted to the purchaser "
        "to print one copy for personal study. Redistribution of the digital file is not granted."
    )
    b.para(
        "Units throughout: ℏ = c = 1, Heaviside–Lorentz, mostly-minus metric g = diag(1, −1, −1, −1). "
        "The electron charge in the vertex is +ieγ^μ as derived in Lesson 45. "
        "α = e²/4π ≈ 1/137.036 at vanishing momentum transfer."
    )
    b.para(
        "This book derives tree-level QED processes, one-loop mass and charge renormalization, "
        "the running of α for one lepton, the Ward identity, infrared cancellation, Dirac g = 2, "
        "Schwinger’s a = α/(2π), the leading Lamb mechanism, the Uehling −27 MHz piece, "
        "and the Schwinger pair-production exponent and prefactor. It does not compute two-loop g−2, "
        "a complete Lamb shift to kHz, α(M_Z) including quarks, weak decays, QCD, or gravity."
    )
    b.para("%s series" % SERIES)  # series line as in Physics, Actually
    b.pagebreak()

    h3 = b.heading("Also in This Series", 1)
    _bookmark_p(h3._p, "AlsoInSeries", 3007)
    for line in ALSO_IN_SERIES:
        b.para(line)
    b.pagebreak()

    h2 = b.heading("How to read this book", 1)
    _bookmark_p(h2._p, "HowToRead", 3005)
    b.para(
        "Work with pencil and paper. Attempt each worked example before reading its answer, "
        "and the exercises before the solutions. Solutions sit at the end of the lesson that posed them."
    )
    b.para(
        "Exercises carry three labels. Core exercises are the minimum before moving on. Extension exercises deepen the lesson "
        "and are worth doing on a second pass. Each lesson has one Challenge, which is optional. Before most major derivations "
        "a shaded First pass box states the result and why it matters, so that you know where the algebra is going."
    )
    b.para("The core track. If you want the shortest honest route from arrows to loop corrections, read Prologues 1–6 and these "
           "35 lessons, in order, doing the Core exercises: 1, 3, 4, 5, 7, 8, 10, 12, 18, 20, 21, 22, 23, 26, 28, 33, 36, 38, "
           "39, 42, 45, 46, 47, 48, 50, 52, 53, 54, 57, 59, 62, 63, 67, 68, 69. Skim Lessons 55 and 56 for the trace results "
           "used later; Lesson 76 (g − 2) is the natural finish. The other lessons fill in proofs, alternatives and applications; "
           "come back to them when a core lesson cites them.")
    b.image("Front_route.png", 5.8, "The core track (shaded) inside the full course.")
    b.para("Routes through the full course:")
    b.bullet(
        "Prologues 1–6 first, always: Feynman’s easy QED in this course’s words (photons, arrows, partial reflection, all paths, "
        "arrows as complex numbers, the three basic actions, loops and infinities), with no assumed background. "
        "Prologues 7–9 add Maxwell as the many-photon alternative, bras and kets with the Schrödinger equation, "
        "and the S-matrix with Feynman diagrams. Plan on 30–50 hours. Then Lesson 1."
    )
    b.bullet(
        "Lessons 1–38, the foundation, if complex numbers, Dirac matrices, or canonical quantization are not yet automatic. "
        "Lesson 9 is the bridge from Maxwell's waves to photons (optics, Huygens–Fresnel, the sum over light paths); "
        "skim its first three sections if undergraduate optics is familiar, but work Sections 4 to 6. "
        "Plan on 150–250 hours."
    )
    b.bullet(
        "After Lesson 41, read the Interlude on Mead’s view (A as the phase standard; E and B derived). Do not move it to the front. "
        "Volume 1, Part Two, develops that view at book length."
    )
    b.bullet(
        "Lessons 39–61, the QED spine, if the foundation is already in hand. Start at Lesson 39. "
        "The last tree-level result is Compton’s (1/4)Σ|ℳ|² = −2e⁴(s/u+u/s) and the conversion to a cross section. "
        "Plan on 80–120 hours."
    )
    b.bullet(
        "Lessons 62–87, one-loop QED and the path-integral language. Do not skip 54–58; later chapters use those amplitudes. "
        "Plan on 80–120 hours. The last derived numbers are a = α/(2π), the Lamb mechanism at ~10³ MHz, "
        "and Γ/V = (eE)²/(4π³) exp(−πm²/eE)."
    )
    b.para(
        "A reader who wants only the Feynman rules and tree processes can stop after Lesson 61. "
        "A reader who wants renormalization can stop after Lesson 70. "
        "Lessons 80–87 re-derive the same rules in the path integral and place QED inside the Standard Model; they are optional for computing a cross section."
    )
    b.para(
        "Do not start this book as a first course in calculus. Do start it as a first course in QED if linear algebra "
        "and ordinary quantum mechanics are willing to be rebuilt rather than assumed."
    )
    b.para(
        "This is Volume 2 of Quantum, Actually and the technical companion to Volume 1. "
        "Volume 1 is the conceptual map, and its Part Two "
        "reads electromagnetism outward from quantum phase and the potential. It is not required here; "
        "this course is where what it describes gets calculated."
    )
    b.pagebreak()

    for el in b.elements:
        anchor.addnext(el)
        anchor = el


def also_by_page(doc):
    """'Also by Lothar J. Musiol' as the last page, after the Bibliography (idempotent:
    an earlier copy, with its page break, is replaced)."""
    body = doc.element.body
    old = None
    for child in body:
        if _p_style(child) == "Heading1" and _text(child).strip() == ALSO_BY_HEADING:
            old = child
            break
    if old is not None:
        start = old
        prev = old.getprevious()
        if prev is not None and prev.tag == qn("w:p") and not _text(prev).strip() and any(
                br.get(qn("w:type")) == "page" for br in prev.iter(qn("w:br"))):
            start = prev
        cur = start
        while cur is not None and cur.tag != qn("w:sectPr"):
            if cur is not old and cur is not start and _p_style(cur) == "Heading1":
                break
            nxt = cur.getnext()
            body.remove(cur)
            cur = nxt
    b = bl.Builder(doc)
    b.pagebreak()
    h = b.heading(ALSO_BY_HEADING, 1)
    _bookmark_p(h._p, "AlsoBy", 3008)
    for group, titles in ALSO_BY:
        p = b.para()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        p.add_run(group).bold = True
        for title in titles:
            q = b.para()
            q.paragraph_format.left_indent = Inches(0.25)
            q.paragraph_format.first_line_indent = Inches(0)
            q.paragraph_format.space_before = Pt(0)
            q.paragraph_format.space_after = Pt(1)
            q.add_run(title)
    # b.para appends at the end of the body, which is where this page belongs


def link_lesson_mentions(doc):
    """Hyperlink 'Lesson N' only in short, single-run Normal paragraphs (Next-lesson lines)."""
    n = 0
    for p in doc.paragraphs:
        style = p.style.name if p.style is not None else ""
        if style != "Normal":
            continue
        if p._p.findall(qn("w:hyperlink")):
            continue
        runs = p.runs
        if len(runs) != 1:
            continue
        text = p.text
        if "Lesson " not in text or len(text) > 280:
            continue
        matches = list(LESSON_MENTION.finditer(text))
        if not matches:
            continue
        parts = []
        pos = 0
        for m in matches:
            num = int(m.group(1))
            if num < 1 or num > 87:
                continue
            if m.start() > pos:
                parts.append(("t", text[pos:m.start()]))
            parts.append(("h", m.group(0), "Lesson%d" % num))
            pos = m.end()
        if pos < len(text):
            parts.append(("t", text[pos:]))
        if not any(x[0] == "h" for x in parts):
            continue
        r = runs[0]
        r.text = ""
        r._r.getparent().remove(r._r)
        for item in parts:
            if item[0] == "t":
                p.add_run(item[1])
            else:
                _hyperlink_run(p, item[1], item[2])
                n += 1
    return n


def set_author(doc):
    """Author line under the subtitle on the title page, and the core property."""
    doc.core_properties.author = bl.AUTHOR
    paras = doc.paragraphs[:8]
    if any(p.text.strip() == bl.AUTHOR for p in paras):
        return
    for p in paras:
        if p.style is not None and p.style.name == "Subtitle":
            new = copy.deepcopy(p._p)
            p._p.addnext(new)
            for t in list(new.iter(qn("w:t")))[1:]:
                t.text = ""
            ts = list(new.iter(qn("w:t")))
            ts[0].text = bl.AUTHOR
            return


def main():
    doc = Document(COMPLETE)
    retitle(doc)
    set_author(doc)
    series_title_page(doc)
    renumbered = normalize_numbering(doc)
    set_gutter(doc)
    demoted = demote_headings(doc)
    insert_front_matter(doc)
    also_by_page(doc)
    # demote again in case front matter used H1 correctly (keep those)
    demoted += 0
    n_bm = bookmark_lessons(doc)
    entries = collect_toc_entries(doc)
    insert_static_toc(doc, entries)
    n_links = link_lesson_mentions(doc)
    # Do not ask Word for F9 field updates
    settings = doc.settings.element
    uf = settings.find(qn("w:updateFields"))
    if uf is not None:
        settings.remove(uf)
    last = None
    tmp = COMPLETE + ".tmp.docx"
    for attempt in range(8):
        try:
            doc.save(tmp)
            os.replace(tmp, COMPLETE)
            last = None
            break
        except OSError as e:
            last = e
            time.sleep(2 + attempt)
    if last is not None:
        raise last
    print(
        "finish_book: renumbered %d headings, demoted %d headings, bookmarks %d, TOC %d, links %d"
        % (renumbered, demoted, n_bm, len(entries), n_links)
    )


if __name__ == "__main__":
    main()
