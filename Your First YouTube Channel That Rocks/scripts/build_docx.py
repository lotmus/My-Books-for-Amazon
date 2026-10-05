"""Build Your First YouTube Channel That Rocks.docx in the same shape as the Kindle book."""
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor, Twips
from docx.text.paragraph import Paragraph

ROOT = Path(__file__).resolve().parent.parent  # scripts\ -> book root
FIG = ROOT / "figures"
ORDER = [
    (1, ROOT / "chapters" / "Name the Viewer.md"),
    (2, ROOT / "chapters" / "The Click Is the Whole Business.md"),
    (3, ROOT / "chapters" / "Nobody Can Tell What It Cost.md"),
    (4, ROOT / "chapters" / "Record It So They Stay.md"),
    (5, ROOT / "chapters" / "Eight Videos, Each With One Job.md"),
    (6, ROOT / "chapters" / "Six Weeks to a Working Channel.md"),
    (7, ROOT / "chapters" / "Read the Count.md"),
    (8, ROOT / "chapters" / "Grow Toward the Gate.md"),
    (9, ROOT / "chapters" / "The Gates and the Review.md"),
    (10, ROOT / "chapters" / "The Rules That Can Switch Off the Money.md"),
    (11, ROOT / "chapters" / "Ads, RPM, and the Shorts Pool.md"),
    (12, ROOT / "chapters" / "Money From the People Who Watch.md"),
    (13, ROOT / "chapters" / "Sponsors and Affiliate Links.md"),
    (14, ROOT / "chapters" / "Keep It Paying.md"),
    (15, ROOT / "chapters" / "The Channel Workbook.md"),
]
# Unnumbered front matter, rendered from Markdown like the chapters.
FRONT = [
    ROOT / "chapters" / "Start Here.md",
    ROOT / "chapters" / "Start This Week.md",
]
TITLE = "Your First YouTube Channel That Rocks"
SUBTITLE = "Grow Watch Time and Subscribers, Reach the Partner Program, and Earn From Ads, Fans, and Sponsors"
SERIES = "A working guide for creators who want a channel, not a hobby."
# Back matter: the canonical "Also by Lothar J. Musiol" list (same in every book; see
# notes/ALSO_BY - canonical list.md at the repo root). Titles as on each master's title page.
ALSO_BY_HEADING = "Also by Lothar J. Musiol"
ALSO_BY = [
    ('Physics, Actually', [
        'Physics, Actually, Volume 1: Motion, Forces, Time, and Relativity',
        'Physics, Actually, Volume 2: Gravity, Cosmology, and the Limits of Spacetime',
        'Physics, Actually, Volume 3: The Standard Model, Chaos, and the Edge of Knowledge',
        'Life, Actually: From the First Cell to the Edited Genome and the Search for Life Elsewhere',
    ]),
    ('Math, Actually', [
        'Math, Actually, Volume 1: From Arithmetic to Calculus',
        'Math, Actually, Volume 2: From Multivariable Calculus to Set Theory & Logic',
        'Math, Actually, Volume 3: From Differential Equations to Abstract Algebra',
        'Math, Actually, Volume 4: From Category Theory to the Frontier',
    ]),
    ('Quanta, Actually', [
        'Quanta, Actually, Volume 1: The Quantum World',
        'Quanta, Actually, Volume 2: The Quantum Conversation',
        'Quanta, Actually, Volume 3: Complete Quantum Electrodynamics Course',
    ]),
    ('Science Sparks', [
        'Science Sparks: Physics, Life, and Mathematics — The Same Few Rules, Told in Highlights',
    ]),
    ('Look First', [
        'Look First, Volume 1: The Universe Has No Now',
        'Look First, Volume 2: A Trip Is Not a New Life',
    ]),
    ('Electrical Engineering Series', [
        'Foundations of Electronics (Book 1)',
        'Circuits, Components, and Control (Book 2)',
        'Semiconductor Physics and Devices (Book 3)',
        'RF, Microwave, and Transceivers (Book 4)',
        'Communications, Wireless, and SDR (Book 5)',
        'Power and Energy (Book 6)',
        'Packaging, Layout, EMC, and Test (Book 7)',
    ]),
    ('History', [
        "The Dolphins' View of History",
    ]),
    ('Fiction', [
        "The Murder That Hadn't Happened Yet (The Relativistic Investigation Bureau, Book 1)",
        'The Warning That Was Sent Too Late (The Relativistic Investigation Bureau, Book 2)',
        'Schrödinger’s Paperwork (Lolly Wren’s Curious Science Adventures, Book 1)',
        'The Permitted Options (Lolly Wren’s Curious Science Adventures, Book 2)',
        'Protocol Flamingo (The Invasion Storybooks, Book 1), as George Herbert Fontaine',
    ]),
    ('How-To', [
        'Your First Book That Sells',
        'Your First YouTube Channel That Rocks',
    ]),
]
TOKEN = re.compile(r"(\*\*[^*]+?\*\*|\*[^*]+?\*|`[^`]+`|\[[^\]]+?\]\([^)]+?\))")


def shade(cell, fill):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def set_cell_width(cell, inches):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcW = OxmlElement("w:tcW")
    dxa = str(int(inches * 1440))
    tcW.set(qn("w:w"), dxa)
    tcW.set(qn("w:type"), "dxa")
    tcPr.append(tcW)


def add_hyperlink(paragraph, text, url):
    r_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    rStyle = OxmlElement("w:rStyle")
    rStyle.set(qn("w:val"), "Hyperlink")
    rPr.append(rStyle)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "22")
    rPr.append(sz)
    rFonts = OxmlElement("w:rFonts")
    rFonts.set(qn("w:ascii"), "Calibri")
    rFonts.set(qn("w:hAnsi"), "Calibri")
    rPr.append(rFonts)
    run.append(rPr)
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    run.append(t)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_inlines(paragraph, text):
    for part in TOKEN.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            inner = part[2:-2]
            if "](" in inner:
                before = len(paragraph.runs)
                add_inlines(paragraph, inner)
                for run in paragraph.runs[before:]:
                    run.bold = True
                continue
            run = paragraph.add_run(inner)
            run.bold = True
            run.font.name = "Calibri"
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Calibri"
        elif part.startswith("*") and part.endswith("*"):
            run = paragraph.add_run(part[1:-1])
            run.italic = True
            run.font.name = "Calibri"
        elif part.startswith("["):
            match = re.match(r"\[([^\]]+)\]\(([^)]+)\)", part)
            if match:
                add_hyperlink(paragraph, match.group(1), match.group(2))
            else:
                run = paragraph.add_run(part)
                run.font.name = "Calibri"
        else:
            run = paragraph.add_run(part)
            run.font.name = "Calibri"


def is_italic_block(text):
    return text.startswith("*") and text.endswith("*") and not text.startswith("**")


def add_picture(doc, rel, alt, base):
    src = (base / rel).resolve()
    png = src.with_suffix(".png")
    path = png if png.exists() else src
    if path.suffix.lower() not in {".png", ".jpg", ".jpeg"}:
        paragraph = doc.add_paragraph()
        run = paragraph.add_run(alt)
        run.italic = True
        return
    paragraph = doc.add_paragraph()
    paragraph.alignment = 1
    run = paragraph.add_run()
    run.add_picture(str(path), width=Inches(5.6))
    doc_pr = run._r.xpath(".//wp:docPr")
    if doc_pr:
        doc_pr[0].set("descr", alt)


def add_table(doc, rows):
    cols = max(len(row) for row in rows)
    widths = column_widths(cols)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.autofit = False
    table.allow_autofit = False
    table.style = "Table Grid"
    tbl = table._tbl
    tblPr = tbl.tblPr if tbl.tblPr is not None else OxmlElement("w:tblPr")
    tblW = OxmlElement("w:tblW")
    total = int(sum(widths) * 1440)
    tblW.set(qn("w:w"), str(total))
    tblW.set(qn("w:type"), "dxa")
    tblPr.append(tblW)
    for i, row in enumerate(rows):
        for j in range(cols):
            cell = table.rows[i].cells[j]
            set_cell_width(cell, widths[j])
            cell.text = ""
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(2)
            paragraph.paragraph_format.space_before = Pt(2)
            text = row[j] if j < len(row) else ""
            add_inlines(paragraph, text)
            for run in paragraph.runs:
                run.font.size = Pt(10.5)
                run.font.name = "Calibri"
                if i == 0:
                    run.bold = True
            if i == 0:
                shade(cell, "F3F1EA")
    doc.add_paragraph()


def column_widths(cols):
    if cols == 4:
        return [1.2, 0.9, 1.75, 1.75]
    if cols == 3:
        return [1.5, 1.5, 2.6]
    share = 5.6 / cols
    return [share] * cols


def parse_blocks(text):
    lines = text.replace("\r\n", "\n").split("\n")
    blocks = []
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line.strip():
            i += 1
            continue
        if line.strip() == "---":
            i += 1
            continue
        if line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                raw = lines[i].strip().strip("|")
                cells = [cell.strip() for cell in raw.split("|")]
                if not all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells):
                    rows.append(cells)
                i += 1
            blocks.append(("table", rows))
            continue
        if line.startswith("- "):
            items = []
            while i < len(lines) and lines[i].startswith("- "):
                items.append(lines[i][2:].strip())
                i += 1
            blocks.append(("bullets", items))
            continue
        if re.match(r"\d+\. ", line):
            items = []
            while i < len(lines) and re.match(r"\d+\. ", lines[i]):
                items.append(re.sub(r"^\d+\. ", "", lines[i]).strip())
                i += 1
            blocks.append(("numbers", items))
            continue
        if line.startswith("!["):
            match = re.match(r"!\[([^\]]*)\]\(([^)]+)\)", line.strip())
            if match:
                blocks.append(("image", match.group(2), match.group(1)))
                i += 1
                continue
        if line.startswith(">"):
            blocks.append(("callout", line.lstrip(">").strip()))
            i += 1
            continue
        paragraph = [line]
        i += 1
        while i < len(lines) and lines[i].strip() and not lines[i].startswith(("#", "|", "- ", "!", ">")) and not re.match(r"\d+\. ", lines[i]) and lines[i].strip() != "---":
            paragraph.append(lines[i].rstrip())
            i += 1
        blocks.append(("p", " ".join(paragraph).strip()))
    return blocks


def apply_page(section):
    section.page_width = Inches(7)
    section.page_height = Inches(10)
    section.top_margin = section.bottom_margin = Inches(0.7)
    section.left_margin = section.right_margin = Inches(0.7)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)


def set_running_head(section, text):
    apply_page(section)
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    header = section.header.paragraphs[0]
    header.clear()
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header.add_run(text)
    run.italic = True
    run.font.name = "Calibri"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
    footer = section.footer.paragraphs[0]
    footer.clear()
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    footer.add_run()._r.append(begin)
    footer.add_run()._r.append(instr)
    footer.add_run()._r.append(end)
    for item in footer.runs:
        item.font.name = "Calibri"
        item.font.size = Pt(9)


def add_body(doc, blocks, number, base):
    skipped_series = False
    for block in blocks:
        kind = block[0]
        if kind == "p":
            text = block[1]
            if text.startswith("# "):
                title = text[2:].strip()
                section = doc.add_section(WD_SECTION.NEW_PAGE)
                if number is None:
                    set_running_head(section, title)
                    doc.add_paragraph(title, "Heading 1")
                else:
                    set_running_head(section, f"{number}  {title}")
                    doc.add_paragraph(f"{number} {title}", "Heading 1")
                continue
            if text.startswith("## "):
                doc.add_paragraph(text[3:].strip(), "Heading 2")
                continue
            if not skipped_series and text.strip("*") == SERIES.strip("."):
                skipped_series = True
                continue
            if not skipped_series and "working guide for creators" in text:
                skipped_series = True
                continue
            paragraph = doc.add_paragraph()
            if is_italic_block(text):
                add_inlines(paragraph, text[1:-1])
                for run in paragraph.runs:
                    run.italic = True
            else:
                add_inlines(paragraph, text)
        elif kind == "bullets":
            for item in block[1]:
                paragraph = doc.add_paragraph(style="List Bullet")
                add_inlines(paragraph, item)
        elif kind == "numbers":
            num_id = new_restarted_num(doc)
            for item in block[1]:
                paragraph = doc.add_paragraph(style="List Number")
                if num_id:
                    force_num(paragraph, num_id)
                add_inlines(paragraph, item)
        elif kind == "callout":
            add_callout(doc, block[1])
        elif kind == "table":
            add_table(doc, block[1])
        elif kind == "image":
            add_picture(doc, block[1], block[2], base)


def add_callout(doc, text):
    table = doc.add_table(rows=1, cols=1)
    table.autofit = False
    table.allow_autofit = False
    cell = table.cell(0, 0)
    set_cell_width(cell, 5.6)
    shade(cell, "F3F1EA")
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.space_before = Pt(2)
    add_inlines(paragraph, text)
    for run in paragraph.runs:
        run.font.size = Pt(11)
        run.font.name = "Calibri"
    doc.add_paragraph()


def new_restarted_num(doc):
    numbering = doc.part.numbering_part._element
    nums = numbering.findall(qn("w:num"))
    if not nums:
        return None
    source = next((n for n in nums if n.get(qn("w:numId")) == "5"), nums[-1])
    abstract = source.find(qn("w:abstractNumId"))
    if abstract is None:
        return None
    new_id = max(int(n.get(qn("w:numId"))) for n in nums) + 1
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(new_id))
    abs_el = OxmlElement("w:abstractNumId")
    abs_el.set(qn("w:val"), abstract.get(qn("w:val")))
    num.append(abs_el)
    override = OxmlElement("w:lvlOverride")
    override.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:startOverride")
    start.set(qn("w:val"), "1")
    override.append(start)
    num.append(override)
    numbering.append(num)
    return str(new_id)


def force_num(paragraph, num_id):
    pPr = paragraph._p.get_or_add_pPr()
    existing = pPr.find(qn("w:numPr"))
    if existing is not None:
        pPr.remove(existing)
    numPr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    nid = OxmlElement("w:numId")
    nid.set(qn("w:val"), num_id)
    numPr.append(ilvl)
    numPr.append(nid)
    pPr.append(numPr)


def insert_contents(doc, anchor):
    """Fill the Contents page at `anchor` with internal links to every Heading 1 and Heading 2."""
    headings = [p for p in doc.paragraphs if p.style.name in ("Heading 1", "Heading 2") and p._p is not anchor._p]
    first = anchor
    contents = anchor
    contents.style = doc.styles["Heading 1"]
    contents.add_run("Contents")
    last = contents
    for i, heading in enumerate(headings):
        start = OxmlElement("w:bookmarkStart")
        start.set(qn("w:id"), str(i + 1))
        start.set(qn("w:name"), f"toc{i}")
        heading._p.insert(0, start)
        end = OxmlElement("w:bookmarkEnd")
        end.set(qn("w:id"), str(i + 1))
        heading._p.append(end)
        p_el = OxmlElement("w:p")
        last._p.addnext(p_el)
        last = Paragraph(p_el, first._parent)
        last.style = doc.styles["Normal"]
        if heading.style.name == "Heading 2":
            last.paragraph_format.left_indent = Inches(0.3)
        link = OxmlElement("w:hyperlink")
        link.set(qn("w:anchor"), f"toc{i}")
        run = OxmlElement("w:r")
        rPr = OxmlElement("w:rPr")
        rFonts = OxmlElement("w:rFonts")
        rFonts.set(qn("w:ascii"), "Calibri")
        rFonts.set(qn("w:hAnsi"), "Calibri")
        rPr.append(rFonts)
        sz = OxmlElement("w:sz")
        sz.set(qn("w:val"), "20" if heading.style.name == "Heading 2" else "22")
        rPr.append(sz)
        color = OxmlElement("w:color")
        color.set(qn("w:val"), "1F4E79")
        rPr.append(color)
        u = OxmlElement("w:u")
        u.set(qn("w:val"), "single")
        rPr.append(u)
        run.append(rPr)
        t = OxmlElement("w:t")
        t.text = heading.text
        run.append(t)
        link.append(run)
        p_el.append(link)


def main():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(7)
    section.page_height = Inches(10)
    section.top_margin = section.bottom_margin = Inches(0.7)
    section.left_margin = section.right_margin = Inches(0.7)
    for name in ["Normal", "Title", "Subtitle", "Heading 1", "Heading 2"]:
        doc.styles[name].font.name = "Calibri"
        doc.styles[name].font.color.rgb = RGBColor(0, 0, 0)
    doc.styles["Normal"].font.size = Pt(11)
    doc.styles["Normal"].paragraph_format.space_after = Pt(8)
    doc.styles["Normal"].paragraph_format.line_spacing = 1.15
    doc.styles["Title"].font.size = Pt(32)
    doc.styles["Heading 1"].font.size = Pt(21)
    doc.styles["Heading 1"].paragraph_format.page_break_before = False
    doc.styles["Heading 2"].font.size = Pt(16)
    doc.styles["Heading 2"].paragraph_format.space_before = Pt(16)
    doc.styles["Heading 2"].paragraph_format.space_after = Pt(8)
    section.different_first_page_header_footer = True
    set_running_head(section, "Start here")

    doc.add_paragraph("Your First YouTube\nChannel That Rocks", "Title")
    doc.add_paragraph(
        SUBTITLE,
        "Subtitle",
    )
    doc.add_paragraph(SERIES)
    author_line = doc.add_paragraph()
    author_line.add_run("Lothar J. Musiol").bold = True
    doc.add_paragraph("Updated October 2026")
    doc.add_paragraph("Copyright \u00a9 2026 Lothar J. Musiol. All rights reserved.")

    contents_section = doc.add_section(WD_SECTION.NEW_PAGE)
    set_running_head(contents_section, "Contents")
    contents_anchor = doc.add_paragraph()

    for path in FRONT:
        add_body(doc, parse_blocks(path.read_text(encoding="utf-8")), None, path.parent)

    for number, path in ORDER:
        add_body(doc, parse_blocks(path.read_text(encoding="utf-8")), number, path.parent)

    last = doc.add_section(WD_SECTION.NEW_PAGE)
    set_running_head(last, "The last word")
    doc.add_paragraph("The last word", "Heading 1")
    for text in CLOSE:
        paragraph = doc.add_paragraph()
        add_inlines(paragraph, text)

    gloss = doc.add_section(WD_SECTION.NEW_PAGE)
    set_running_head(gloss, "Glossary")
    doc.add_paragraph("Glossary", "Heading 1")
    glossary_intro = doc.add_paragraph()
    add_inlines(
        glossary_intro,
        "Short definitions for terms this book uses. A definition here is not a substitute for the platform page cited in the chapter.",
    )
    for term, definition, why_it_matters in GLOSSARY:
        paragraph = doc.add_paragraph()
        run = paragraph.add_run(term + ". ")
        run.bold = True
        run.font.name = "Calibri"
        add_inlines(paragraph, definition)
        why_run = paragraph.add_run(" Why it matters: ")
        why_run.italic = True
        why_run.font.name = "Calibri"
        add_inlines(paragraph, why_it_matters)

    sources = doc.add_section(WD_SECTION.NEW_PAGE)
    set_running_head(sources, "Sources")
    doc.add_paragraph("Official sources and updates", "Heading 1")
    for text in SOURCES:
        paragraph = doc.add_paragraph()
        if text.startswith("*") and text.endswith("*"):
            run = paragraph.add_run(text[1:-1])
            run.italic = True
            run.font.name = "Calibri"
        else:
            add_inlines(paragraph, text)

    also = doc.add_section(WD_SECTION.NEW_PAGE)
    set_running_head(also, "Also by")
    doc.add_paragraph(ALSO_BY_HEADING, "Heading 1")
    for group, titles in ALSO_BY:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(8)
        paragraph.paragraph_format.space_after = Pt(2)
        paragraph.paragraph_format.keep_with_next = True
        run = paragraph.add_run(group)
        run.bold = True
        run.font.name = "Calibri"
        for title in titles:
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.25)
            paragraph.paragraph_format.space_after = Pt(1)
            paragraph.add_run(title).font.name = "Calibri"

    insert_contents(doc, contents_anchor)
    doc.core_properties.title = TITLE
    doc.core_properties.author = "Lothar J. Musiol"
    doc.core_properties.subject = SUBTITLE
    doc.core_properties.keywords = "YouTube Partner Program; monetization; RPM; channel memberships; Super Thanks; sponsorships"
    out = ROOT / f"{TITLE}.docx"
    doc.save(out)
    words = sum(len(p.text.split()) for p in doc.paragraphs)
    print(out)
    print("Words", words)


CLOSE = [
    "The channel in this book does not start with money. It starts with one viewer and a question they would type, and it earns the right to be paid one finished minute at a time: a title that names the problem, an opening that keeps the promise, a method given away, a next video named inside the last one.",
    "The gates are arithmetic. Subscribers, plus hours or Shorts views, inside a window that rolls. Once you can work out how many views a day your channel needs at its own view duration, the gate stops being a mystery and becomes a distance, and chapter 4’s work on the opening becomes the cheapest way to shorten it.",
    "Past the gate, the money comes in five kinds, and each one pays for something different. Ads pay for reach. Memberships pay for loyalty. Supers pay for presence. Sponsors pay for a precise viewer. Affiliate commissions pay for recommendations people act on. A channel that keeps all five small and honest survives the next rule change. A channel that bets on one does not.",
    "The numbers will move again before the next edition of this book. Check the Earn tab, read the help page, write the date in the decision log. Then make the next video for the same person you named on the first page.",
]

SOURCES = [
    "These notes support the claims in the chapters. Claims carry one of four labels, explained in Start here: official rule, documented platform guidance, creator heuristic, or author recommendation. All dollar figures in worked examples are invented arithmetic. YouTube’s Partner Program, monetization, fan-funding, Shorts, and AI-disclosure pages, and the tool and other-platform pages, were checked in October 2026.",
    "Independent guide. Not affiliated with or endorsed by YouTube or any other platform named here. Not legal, tax, or financial advice. Check current official terms before acting.",
    "*Chapter 1. The money map, the rule that a channel outside the Partner Program does not share in ad revenue, the inactivity rule, clickable links, and posts are cited to YouTube’s own help pages in the chapter.*",
    "*Chapter 2. Title, description, and tag limits, captions, thumbnails, the title-and-thumbnail test, audience retention, recommendation signals, Shorts discovery, and series playlists are cited to YouTube’s help pages in the chapter. The eight jobs are the author’s production set; YouTube publishes no such quota. Title wording advice is labeled as creator heuristic.*",
    "*Chapter 3. Each tool claim was checked on that tool’s own license, terms, or product page in October 2026; claims that could not be checked were removed. Export settings are YouTube’s recommended upload encoding settings. That a higher-resolution upload gets a better playback encode is creator heuristic, and the chapter says so.*",
    "*Chapter 4. Video chapter rules, end screen timing, and the eight-minute mid-roll rule are cited to YouTube’s help pages. The speaking-rate range is a planning guess to replace with your own timed reading. That viewers forgive picture before sound is creator heuristic.*",
    "*Chapters 5 and 6 are the author’s growth path: eight videos with one job each and a six-week calendar. They state no ranking formula.*",
    "*Chapter 7. Qualified counts, impression counting, the 2%–10% click-through band, traffic source definitions, and the August 24, 2026 change to when a view is counted are cited to YouTube’s help pages. The worked example is invented arithmetic.*",
    "*Chapter 8. Thresholds, the live-stream rule, and the rule that paid campaign views do not count are cited to YouTube’s help pages and the August 2026 announcement. Subscribe wording, sub-for-sub effects, and collaborations are creator heuristics, labeled as tests.*",
    "*Chapter 9. Partner Program requirements, including the changes that take effect on 1 February 2027 and the 2027 activity rule, were checked in October 2026 against YouTube’s eligibility page, its monetization overview, its Changes to the YouTube Partner Program page, and its 10 August 2026 announcement.*",
    "*Chapter 10. Warnings, strikes, copyright removal, Content ID claims, limited ads and self-certification, paid promotion, AI use disclosure, made-for-kids features, the spam, fake engagement, and external links policies, and the reused, inauthentic-content, and AI-persona monetization rules are cited to YouTube’s help pages and blog in the chapter.*",
    "*Chapter 11. Ad formats, mid-rolls, CPM, playback-based CPM, RPM, the 55% and 45% shares, the Premium and Premium Lite pools, the Shorts Creator Pool, targeted Shorts ads, and the February 2027 Shorts floor are cited to YouTube’s help pages and blog. The worked example is invented arithmetic.*",
    "*Chapter 12. Fan-funding minimums, the 70% shares, US price points, banned perks, paused mode, the not-a-donation rule, and where Super Chat, Super Stickers, and Super Thanks are unavailable are cited to YouTube’s help pages. Advice on asking is labeled.*",
    "*Chapter 13. Creator Partnerships eligibility, the paid promotion setting, clickable links, and the Shopping affiliate program’s 500-subscriber threshold are cited to YouTube’s help pages and blog. The FTC pages are US guidance. The media kit, pricing method, and contract terms are author recommendations.*",
    "*Chapter 14. Other platforms were checked against the linked official page in October 2026. X’s Help Center page was read in the Internet Archive’s copy of 22 September 2026, because the live page refuses automated requests. The diagnostic order is the author’s recommendation.*",
    "Pages cited in the chapters, so you can open them:",
    "[YouTube Partner Program eligibility](https://support.google.com/youtube/answer/72851), [Changes to the YouTube Partner Program](https://support.google.com/youtube/answer/12843009), [2027 Partner Program announcement](https://blog.youtube/news-and-events/youtube-partner-program-updates-2027-new-opportunities-earn/), [monetization overview](https://support.google.com/youtube/answer/94522), [how to earn money on YouTube](https://support.google.com/youtube/answer/72857), [channel monetization policies](https://support.google.com/youtube/answer/1311392), [which links are clickable](https://support.google.com/youtube/answer/13748639), [cards](https://support.google.com/youtube/answer/6140493), [end screens](https://support.google.com/youtube/answer/6388789).",
    "[YouTube recommended upload encoding settings](https://support.google.com/youtube/answer/1722171), [YouTube Audio Library](https://support.google.com/youtube/answer/3376882), [supported caption files](https://support.google.com/youtube/answer/2734698), [automatic captions](https://support.google.com/youtube/answer/6373554).",
    "Tools: [DaVinci Resolve](https://www.blackmagicdesign.com/products/davinciresolve), [CapCut terms of service](https://www.capcut.com/clause/terms-of-service), [Pexels license](https://www.pexels.com/license/), [Pixabay license summary](https://pixabay.com/service/license-summary/), [Uppbeat pricing](https://uppbeat.io/pricing) and [user agreement](https://uppbeat.io/user-agreement), [OBS Studio](https://obsproject.com/). AI examples: [ElevenLabs pricing](https://elevenlabs.io/pricing), [Runway pricing](https://runway.com/pricing) and [terms of use](https://runway.com/terms-of-use).",
    "[Ad revenue analytics: RPM and CPM](https://support.google.com/youtube/answer/9314357), [partner earnings overview](https://support.google.com/youtube/answer/72902), [ad formats](https://support.google.com/youtube/answer/2467968), [mid-roll ads](https://support.google.com/youtube/answer/6175006), [Shorts monetization policies](https://support.google.com/youtube/answer/12504220), [advertiser-friendly content guidelines](https://support.google.com/youtube/answer/6162278).",
    "[Commerce Products monetization policies](https://support.google.com/youtube/answer/13195878), [channel memberships](https://support.google.com/youtube/answer/7636690), [membership levels and perks](https://support.google.com/youtube/answer/7544492), [US membership prices](https://support.google.com/youtube/answer/10119895), [Super Chat and Super Stickers](https://support.google.com/youtube/answer/7288782), [Super Chat eligibility](https://support.google.com/youtube/answer/9277801), [Super Thanks eligibility](https://support.google.com/youtube/answer/10879035), [Super Thanks tips](https://support.google.com/youtube/answer/13615971).",
    "[YouTube Creator Partnerships](https://support.google.com/youtube/answer/9385307), [YouTube Shopping affiliate program](https://support.google.com/youtube/answer/13376398), [Shopping expansion to 500 subscribers, March 2026](https://blog.youtube/creator-and-artist-stories/youtube-shopping-expansion-500-subscribers/).",
    "[Instagram, editing your profile](https://help.instagram.com/936495066470190/) and [Instagram link sticker](https://help.instagram.com/192168966243613). [TikTok, adding a website to your profile](https://support.tiktok.com/en/getting-started/setting-up-your-profile/adding-a-website-to-your-profile).",
    "[FTC Endorsement Guides, 16 CFR Part 255](https://www.ecfr.gov/current/title-16/chapter-I/subchapter-B/part-255) and [Disclosures 101 for Social Media Influencers](https://www.ftc.gov/business-guidance/resources/disclosures-101-social-media-influencers).",
    "[Video chapters](https://support.google.com/youtube/answer/9884579). [Impressions and watch time](https://support.google.com/youtube/answer/9314486), [impressions and click-through rate FAQ](https://support.google.com/youtube/answer/7628154), [reach reports](https://support.google.com/youtube/answer/9314355), [content performance](https://support.google.com/youtube/answer/12220281).",
    "[Community Guidelines strike basics](https://support.google.com/youtube/answer/2802032), [copyright strikes](https://support.google.com/youtube/answer/2814000), [copyright claims](https://support.google.com/youtube/answer/6013276), [paid promotion](https://support.google.com/youtube/answer/154235), [disclosing AI use](https://support.google.com/youtube/answer/14328491), [audience setting and made for kids](https://support.google.com/youtube/answer/9527654), [spam policy](https://support.google.com/youtube/answer/2801973), [fake engagement policy](https://support.google.com/youtube/answer/3399767), [external links policy](https://support.google.com/youtube/answer/9054257).",
]

GLOSSARY = [
    ("Eight jobs", "The only set of videos this book asks you to finish: who it is for, the problem in their words, the method, proof, the comparison, one objection, the deep dive, and which video to watch next. Chapters 2, 5, and 6 use this same set.", "Skipping one job is usually the missing step between videos nobody finishes and a path strangers follow."),
    ("Deep dive", "The long, complete video in the eight jobs, built to be watched for a long time. The video that fills the hour bar, and the one that can honestly pass eight minutes for mid-roll ads.", "It is the one video the hour bar depends on most, so its length and pacing earn more care than the other seven."),
    ("Expanded program", "In countries where YouTube has opened it, the earlier tier of the Partner Program: 500 subscribers, three public uploads in 90 days, and either 3,000 long-form hours in a year or 3 million Shorts views in 90 days. Fan funding, Shopping, and Creator Partnerships. Not a share of watch-page ads. Unchanged by the 2027 update.", "It can open memberships and Super Chat months before the ad gate does, which is real income a new channel can use earlier."),
    ("YouTube Partner Program", "The higher gate: 1,000 subscribers and either 4,000 long-form hours in 12 months or 10 million qualified Shorts views in 90 days. For channels that apply from 1 February 2027: 8,000 hours in 365 days or 20 million Shorts views in 90 days. This is the gate that adds watch-page ads, Shorts Feed ads, and YouTube Premium revenue. Channels already in the program keep their status.", "It is the gate every ad, Premium, and Shorts dollar in this book waits behind."),
    ("Qualified watch hours", "Public long-form viewing that YouTube counts toward the hour bars. Hours watched in the Shorts feed do not count. Private, unlisted, deleted, and ad-campaign views do not count.", "A channel can have real viewers and still fail the gate if those hours don't count."),
    ("Qualified Shorts views", "Public views of Shorts in the Shorts feed that YouTube counts toward the Shorts bars: 3 million in 90 days for the expanded program, or 10 million in 90 days for the Partner Program (20 million for channels applying from 1 February 2027). They do not fill the long-form hour bars.", "Shorts fame does not substitute for the long-form hour bar; it has its own separate bar."),
    ("Follow-on views", "Organic views from people who watch more of your videos after seeing a promoted one. The only part of a paid campaign that counts toward the Partner Program.", "It is the only way a paid ad campaign can actually help you reach the gate."),
    ("Advanced features", "The Studio status that makes an address in a long-form description, and in a long-form comment, clickable. Phone verification comes first. A Short’s description and comments stay unclickable after it is on. The Partner Program requires it.", "Without it, a link in a description or comment is just text nobody can tap."),
    ("AdSense", "AdSense for YouTube, the Google account YouTube uses to pay a channel it has accepted. Meeting a subscriber number does not open it. You apply, and YouTube reviews the channel.", "Meeting every number in this book still does not pay you until this account exists and is approved."),
    ("Modules", "The separate sets of terms you accept in Studio’s Earn tab once inside the program: Watch Page Monetization for long-form ads and Premium, Shorts Monetization for the Shorts feed, and the Commerce Product Module for fan funding.", "Accepting the Partner Program does not turn any income on by itself; each module has to be switched on separately."),
    ("CPM", "What advertisers paid per 1,000 ad impressions on your videos, before YouTube’s share. Ads and Premium only, monetized views only.", "It explains why RPM moved, even though it is not the number that lands in your account."),
    ("Playback-based CPM", "What advertisers paid per 1,000 playbacks that showed at least one ad. Often higher than CPM, because one playback can carry two ads.", "It is the number that best matches what one video actually earned per view shown an ad."),
    ("RPM", "Your total revenue after YouTube’s share, per 1,000 engaged views, including views with no ad. Includes ads, Premium, memberships, Super Chat, and Super Stickers. Usually lower than CPM.", "It is the one number that reflects everything a channel actually earned, which is why this book plans by it."),
    ("Mid-roll", "An ad break during a video. Available on monetized videos eight minutes or longer, placed automatically, by hand, or both. Breaks at natural pauses are more likely to serve an ad.", "Placing it at a natural pause instead of mid-sentence can change whether it serves an ad at all."),
    ("Limited ads", "The monetization status of a video that does not fully meet the advertiser-friendly guidelines: fewer advertisers, or none. Not a strike. You can request a human review.", "It silently caps a video's earning potential without ever touching the channel's standing."),
    ("Shorts Creator Pool", "The monthly pot of Shorts Feed ad revenue that YouTube shares out by each monetizing channel’s share of engaged Shorts views, after a share for music licensing. The channel keeps 45% of its allocation. From 1 February 2027, a month counts only if the channel had 10 million qualified Shorts views in the previous 90 days.", "A Short's payout depends on the whole pool and your share of it, not on a fixed rate per view."),
    ("YouTube Premium revenue", "A share of Premium and Premium Lite subscription fees, pooled and paid by members’ watch time and views: 55% for long-form, 45% for Shorts, of the pool allocated to creators.", "It pays even on videos with no ads at all, which is easy to miss when reading only the ads line."),
    ("Channel memberships", "Monthly payments from viewers for perks you choose, in up to six levels. The channel receives 70% after taxes and fees. Downloads of YouTube content, in-person one-to-one meetings, and random prizes are not allowed as perks.", "It is often the first stream available, well before the ad gate."),
    ("Super Chat and Super Stickers", "Paid, highlighted messages and animated stickers in the live chat of a live stream or premiere. The channel receives 70% after taxes and fees. YouTube says they are not donation tools.", "They only exist during a live stream, so they are a reason to go live, not a general-purpose stream."),
    ("Super Thanks", "A one-time paid animation and highlighted comment on a long-form video or a Short. Not available on claimed, unlisted, made-for-kids, or comments-off videos, or while a stream is live. The channel receives 70% after taxes and fees.", "A single claimed video loses access to it, which is one more cost of an unresolved Content ID claim."),
    ("Sponsor", "One company paying you for one video or segment. Select paid promotion in Studio, say so in the video, and say so next to the link.", "It is usually the first stream large enough to pay real money, and the one most dependent on a track record."),
    ("Creator Partnerships", "YouTube’s matching tool between brands and creators. Open to channels in the Partner Program, 18 or older, in available countries, with no active Community Guidelines strikes.", "It is YouTube's own path to sponsors, not a requirement to find one yourself."),
    ("Affiliate link", "A link that pays you a commission if the viewer buys a product you recommended. Say so next to the link. The FTC pages in the sources list are the US disclosure guidance this book points at.", "An undisclosed one is the kind of mistake that invites regulatory attention, not just a viewer complaint."),
    ("Shopping affiliate program", "YouTube’s own affiliate program for tagging other brands’ products in Shorts, long videos, and live streams. Open to partners with at least 500 subscribers in listed countries since March 2026.", "It opens far earlier than the ad gate, at only 500 subscribers."),
    ("Inauthentic content", "YouTube’s monetization name, since July 2025, for templated, repetitive, or mass-produced videos, including generic AI-made ones. It can keep a whole channel out of the Partner Program.", "It can block the whole channel from the Partner Program, not just one video."),
    ("Reused content", "Someone else’s material republished without significant original commentary, modification, or educational value. A monetization policy that applies to the whole channel.", "Like inauthentic content, it is a channel-wide risk, not a per-video one."),
    ("End screen", "An element in the last 5 to 20 seconds of a video at least 25 seconds long. One that opens a site outside YouTube requires the Partner Program.", "It is one of the few places a viewer can be pointed to the next video without leaving YouTube."),
    ("Info card", "A small panel attached to the video. One that opens a site outside YouTube also requires the Partner Program.", "It has the same gate as end screens for an outside link, which surprises creators who expect cards to be freely available."),
    ("Series playlist", "A playlist YouTube can feature as the next video while someone is watching one of yours. The account has to be verified, the videos have to be yours, and a video can sit in only one series playlist.", "It is how YouTube can keep recommending your own next video while someone is already watching one."),
    ("Short", "A vertical video, 1080 by 1920. An address in a Short’s description or comments is not clickable. A Short can point at a long video.", "Its unclickable description means a Short's main job is pointing viewers toward a long video, not toward a link."),
    ("H.264", "The video codec named in YouTube’s recommended upload settings for an MP4. The audio codec named beside it is AAC-LC.", "Matching it avoids a re-encode that can soften a sharp upload."),
    ("Closed captions", "A text track the viewer can turn on or off. The words can be searched. They are not burned into the picture.", "They are what makes a video's words searchable and accessible, which burned-in text never is."),
    ("Burned-in captions", "Words that are part of the picture. They stay on. They are a different choice from closed captions.", "They cannot be turned off, which matters for viewers who find on-screen text distracting."),
    ("Fader", "The volume slider in an editor. A music fader at about a tenth to a fifth of the way up is a position on that slider, not a measurement of loudness.", "Getting this wrong is a common reason a voice is hard to hear under music."),
    ("Text-to-speech", "A tool that reads a script aloud. Usable when the voice is not a clone of someone else. Cloning someone else’s voice is a consent question the tool does not answer for you.", "It is a usable part of the free production stack only as long as the voice is not a clone."),
    ("Stock license", "The terms on one clip or track. “Free” and “free to use commercially” are different sentences. Read the line on that file before you publish it, including on a second site.", "A library's current policy is not a permanent grant on a file downloaded under an older one."),
    ("Impression", "In Studio, a thumbnail shown on YouTube for more than one second with at least half of it visible. Thumbnails on other websites, in end screens, and in notifications are not counted.", "It is the denominator behind click-through rate, so a change in the number does not always mean a change in interest."),
    ("Impressions click-through rate", "How often a counted impression became a view. YouTube says half of channels and videos fall between 2% and 10%. It is a band, not a target.", "Chasing a number above the band usually means chasing clickbait, not better packaging."),
    ("Traffic source", "Studio’s label for how a viewer reached a video: search, suggested videos, browse features, playlists, external, and others. Views from links in video descriptions are filed under suggested videos.", "Knowing where viewers came from decides whether the fix is the title, the thumbnail, or neither."),
    ("Video chapters", "Timestamps in the description that split a video into named parts. The first is 00:00, there are at least three, and each part is at least ten seconds.", "They let a viewer who wants one part jump there instead of leaving the video."),
    ("Community Guidelines warning", "What a first violation typically gets. An optional policy training lets it expire 90 days after the training. A further violation of the same policy in that window can become a strike.", "Letting the training lapse without doing it leaves the channel exposed to a strike on the very next similar mistake."),
    ("Copyright strike", "The result of a valid legal removal request. The video comes down. It expires 90 days after it was applied once Copyright School is done, or it can be resolved by a retraction or a counter notification.", "Three of them in 90 days put the whole channel at risk of termination."),
    ("Content ID claim", "An automatic match against a rights holder’s file. It affects that video: monetize, track, or block. It is not a strike, unless a dispute without a valid reason leads to a removal request.", "It is not a strike, but disputing one without a real reason can become one."),
    ("Paid promotion setting", "The option in video details you select when a company paid you or gave you the product for a video. YouTube then shows a disclosure. Say it in the video as well.", "Switching it on is a legal disclosure, not just a Studio formality."),
    ("AI use", "The upload setting for realistic content that generative AI made or meaningfully altered, including AI-generated music that is the main focus of the video. Scripts, ideas, captions, thumbnails, and a clone of your own voice for voice-overs do not need it.", "Leaving it off on content that needed it risks a disclosure penalty, not just a label."),
    ("Made for kids", "The audience setting for videos directed to children, required under US law. It turns off comments, end screens, personalized ads, memberships, Supers, and more on those videos.", "It switches off several income streams at once, so setting it by habit rather than by fact costs real money."),
]


if __name__ == "__main__":
    main()
