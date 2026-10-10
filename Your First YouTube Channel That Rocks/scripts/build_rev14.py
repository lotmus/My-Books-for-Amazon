
from pathlib import Path
from copy import deepcopy
import json,re
from docx import Document
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[1]
SRC=ROOT/'Your First YouTube Channel_REV13_2026-10-07.docx'
OUT=ROOT/'Your First YouTube Channel_REV14_2026-10-07.docx'
d=Document(SRC);p=list(d.paragraphs);changes=[]
def change(i,text):
    old=p[i].text
    for c in list(p[i]._p):
        if c.tag not in (qn('w:pPr'),qn('w:bookmarkStart'),qn('w:bookmarkEnd')):
            p[i]._p.remove(c)
    p[i].add_run(text)
    changes.append(dict(paragraph=i,before=old,after=text))
def delete(i):
    if p[i]._p.getparent() is not None:
        changes.append(dict(paragraph=i,before=p[i].text,after='[removed]'))
        p[i]._p.getparent().remove(p[i]._p)
def link(par,label,anchor=None,url=None):
    h=OxmlElement('w:hyperlink')
    if anchor: h.set(qn('w:anchor'),anchor)
    if url:
        from docx.opc.constants import RELATIONSHIP_TYPE
        h.set(qn('r:id'),d.part.relate_to(url,RELATIONSHIP_TYPE.HYPERLINK,is_external=True))
    r=OxmlElement('w:r'); prop=OxmlElement('w:rPr')
    sty=OxmlElement('w:rStyle');sty.set(qn('w:val'),'Hyperlink');prop.append(sty);r.append(prop)
    t=OxmlElement('w:t');t.text=label;t.set(qn('xml:space'),'preserve');r.append(t);h.append(r);par._p.append(h)
def bookmark(i,name):
    b=OxmlElement('w:bookmarkStart'); b.set(qn('w:id'),'9010'); b.set(qn('w:name'),name)
    e=OxmlElement('w:bookmarkEnd'); e.set(qn('w:id'),'9010')
    p[i]._p.insert(1,b);p[i]._p.append(e)

def idx(prefix):
    ids=[i for i,par in enumerate(p) if par.text.startswith(prefix)]
    assert len(ids)==1,(prefix,ids);return ids[0]
def after(i,text):
    par=p[i].insert_paragraph_before(text);p[i]._p.addnext(par._p);return par
# Restore the two route-level facts without restoring the old policy tour.
i=idx('Tonight’s checklist.')
gate=p[i].insert_paragraph_before('Official rule: new applicants from 1 February 2027 need 1,000 subscribers plus either 8,000 qualified watch hours in 365 days or 20 million qualified Shorts views in 90 days; the dated requirements are in ')
link(gate,'chapter 9.','toc72')
short=p[i].insert_paragraph_before('For this book’s long-form route, a Short comes after its long video is published, names that video, and points to it through Studio’s Related video setting.')
change(idx('Continue with the calendar in chapter 6.'),'Continue with the calendar in chapter 6. Open a reference chapter when the current job needs it.')
change(idx('Finish this sentence before you design a banner:'),'Before the first video, finish this sentence: “This channel helps [specific person] who wants [specific result], one video at a time.” “People who like cooking” is not a person. “Adults who cook for one on a weeknight and are tired of throwing food away” is.')
filled=idx('The next video —')
change(filled,'The path video — “Price a bathroom repair: the path video and watch order.” Action: choose the video for the step you need, then follow the named order.')
# One disclaimer line beside the copyright; retain the fuller sources note.
i=idx('Copyright ©')
disclaimer=after(i,'Independent guide. Not affiliated with or endorsed by YouTube. Not legal, tax, or financial advice.')
for run in disclaimer.runs:run.font.size=Pt(9)
# The corrected REV13 has one Part III contents entry; enforce that in this saved revision.
contents_start=next(i for i,par in enumerate(d.paragraphs) if par.text=='Contents' and par.style.name=='Heading 1')
contents_end=next(i for i,par in enumerate(d.paragraphs) if i>contents_start and par.style.name=='Heading 1')
toc=d.paragraphs[contents_start:contents_end]
assert sum(par.text=='Part III Workbook' for par in toc)==1
assert sum(par.text=='Part III Workbook' and par.style.name=='Heading 1' for par in d.paragraphs)==1
# Match the claim-label typography.
run=gate.runs[0];rest=run.text[len('Official rule:'):];run.text='Official rule:';run.bold=True
new=deepcopy(run._r);new.find(qn('w:t')).text=rest
prop=new.find(qn('w:rPr'))
if prop is not None:
    for b in prop.findall(qn('w:b')):prop.remove(b)
run._r.addnext(new)
d.core_properties.comments='REV14: final route reminders, terminology and example consistency, single Part III contents entry, and front-matter disclaimer. 7 October 2026.'
d.save(OUT)
fresh=Document(OUT)
body='\n'.join(par.text for par in fresh.paragraphs)
assert 'open a shelf chapter' not in body.lower()
assert 'before you design a banner' not in body
assert 'The next video —' not in body
assert len(fresh.tables)==3 and len(fresh.inline_shapes)==1
anchors={b.get(qn('w:name')) for b in fresh.element.xpath('.//w:bookmarkStart')}
assert all(h.get(qn('w:anchor')) in anchors for h in fresh.element.xpath('.//w:hyperlink[@w:anchor]'))
assert sum(par.text.startswith('In this chapter:') for par in fresh.paragraphs)==15
first=next(i for i,par in enumerate(fresh.paragraphs) if par.text=='Who it is for —')
assert [par.text for par in fresh.paragraphs[first:first+8]]==['Who it is for —','The problem in their words —','The method, given away —','Proof —','Comparison —','One objection —','The deep dive —','The path video and watch order —']
qa=ROOT/'bak/rev14_qa';qa.mkdir(parents=True,exist_ok=True)
(qa/'changes.json').write_text(json.dumps(changes,ensure_ascii=False,indent=2),encoding='utf-8')
(qa/'paragraphs.txt').write_text('\n'.join(f'{i}\t{par.text}' for i,par in enumerate(fresh.paragraphs)),encoding='utf-8')
print('Saved',OUT,'Contents, eight jobs, and internal links verified')
