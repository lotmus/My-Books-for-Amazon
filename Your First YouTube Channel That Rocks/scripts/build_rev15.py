from pathlib import Path
import json,re
from docx import Document
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[1]
d=Document(ROOT/'Your First YouTube Channel_REV14_2026-10-07.docx'); changes=[]
def edit(prefix,text):
 ps=[p for p in d.paragraphs if p.text.startswith(prefix)];assert len(ps)==1,(prefix,len(ps))
 p=ps[0];changes.append({'before':p.text,'after':text})
 for c in list(p._p):
  if c.tag not in (qn('w:pPr'),qn('w:bookmarkStart'),qn('w:bookmarkEnd')):p._p.remove(c)
 p.add_run(text)
def remove(prefix):
 ps=[p for p in d.paragraphs if p.text.startswith(prefix)];assert len(ps)==1,prefix
 p=ps[0];assert not p._p.xpath('.//w:bookmarkStart');changes.append({'before':p.text,'after':'[removed]'})
 p._p.getparent().remove(p._p)
edit('Creator heuristic: words for desperation','Author recommendation: for sensitive subjects, describe the actual problem accurately rather than intensifying it for attention, and check the current advertiser-friendly guidelines before publishing.')
edit('Some words can limit a video', 'Official rule: advertiser suitability depends on the content and its context, including the title, thumbnail, description, and tags. Sensitive subjects can receive limited ads or none. YouTube publishes categories and examples; avoiding a few words does not establish suitability.')
edit('Most bad sound comes from','Author recommendation: check microphone distance and room noise before buying a different microphone.')
edit('Get the microphone close.','Move the microphone closer and make a short test. A phone across the room can pick up more room sound than clear speech. Try the same phone or a clip-on microphone nearer your mouth, then listen for breath noise and clothing rustle. Find a usable position before comparing prices.')
edit('Eight finished videos aimed at','Build eight for one viewer before you build forty for a calendar.')
edit('If the tool accepts references,','Author recommendation: make a short test with the motion the video needs before committing the whole video. Inspect it in motion, not only as a still. Reject visible distortions or changes that undermine the demonstration.')
edit('A clip that looks dense is rarely','Author recommendation: treat the generator as a camera, not as an editor. Generate only the shot you need, then edit it with your own footage and sound.')
for prefix in ['Lock the still before you move it.','Keep each generation short enough','Names change. Author recommendation:','Cut the approved takes as one sequence,']:remove(prefix)
edit('That one detailed prompt produces','That a detailed prompt guarantees a usable clip. Inspect a short test before building the video around generated footage.')
edit('Nobody records eleven clean minutes.','You do not have to record eleven clean minutes in one take. Record one section at a time. If you stumble, pause and restart the sentence. A clean restart is easier to cut than a correction mid-sentence.')
edit('You do not need lights.','Author recommendation: try available daylight before buying lights. Start by facing a window.')
edit('Do this before the banner.','Do this before the first video. Who the channel is for:')
# Remove only references that repeat the navigation without changing the next action.
for prefix in ['Use the evening checklist for packaging. This chapter adds','The evening checklist covers packaging. This section turns']:remove(prefix)
anchors={b.get(qn('w:name')) for b in d.element.xpath('.//w:bookmarkStart')}
assert all(h.get(qn('w:anchor')) in anchors for h in d.element.xpath('.//w:hyperlink[@w:anchor]'))
assert len(d.tables)==3 and len(d.inline_shapes)==1
out=ROOT/'Your First YouTube Channel_REV15_2026-10-07.docx';d.save(out)
qa=ROOT/'bak/rev15_qa';qa.mkdir(parents=True,exist_ok=True)
(qa/'changes.json').write_text(json.dumps(changes,ensure_ascii=False,indent=2),encoding='utf-8')
print('Saved',out,'Changes:',len(changes),'Internal links verified')
