
from pathlib import Path
from copy import deepcopy
import json,re
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
ROOT=Path(__file__).resolve().parents[1]
SRC=ROOT/'Your First YouTube Channel_REV10_2026-10-07.docx'
OUT=ROOT/'Your First YouTube Channel_REV11_2026-10-07.docx'
d=Document(SRC);p=list(d.paragraphs)
assert len(p)==1349,'Source revision changed'
changes=[]
def change(i,text):
    old=p[i].text
    for c in list(p[i]._p):
        if c.tag not in (qn('w:pPr'),qn('w:bookmarkStart'),qn('w:bookmarkEnd')):
            p[i]._p.remove(c)
    p[i].add_run(text)
    changes.append(dict(paragraph=i,before=old,after=text))
def delete(i):
    if p[i]._p.getparent() is not None:
        changes.append(dict(paragraph=i,before=p[i].text,after='[removed]'))
        p[i]._p.getparent().remove(p[i]._p)
def link(par,label,anchor=None,url=None):
    h=OxmlElement('w:hyperlink')
    if anchor: h.set(qn('w:anchor'),anchor)
    if url:
        from docx.opc.constants import RELATIONSHIP_TYPE
        h.set(qn('r:id'),d.part.relate_to(url,RELATIONSHIP_TYPE.HYPERLINK,is_external=True))
    r=OxmlElement('w:r'); prop=OxmlElement('w:rPr')
    sty=OxmlElement('w:rStyle');sty.set(qn('w:val'),'Hyperlink');prop.append(sty);r.append(prop)
    t=OxmlElement('w:t');t.text=label;t.set(qn('xml:space'),'preserve');r.append(t);h.append(r);par._p.append(h)
def bookmark(i,name):
    b=OxmlElement('w:bookmarkStart'); b.set(qn('w:id'),'9010'); b.set(qn('w:name'),name)
    e=OxmlElement('w:bookmarkEnd'); e.set(qn('w:id'),'9010')
    p[i]._p.insert(1,b);p[i]._p.append(e)

def after(i,text,style=None):
    new=p[i].insert_paragraph_before(text,style=style)
    p[i]._p.addnext(new._p)
    return new
def cross(i,text,ch):
    change(i,text)
    link(p[i],' See chapter '+str(ch)+'.',chapters[ch])
chapters={}
heading_indices={}
for i,par in enumerate(p):
    m=re.match(r'^(\d{1,2}) ',par.text)
    bs=par._p.xpath('./w:bookmarkStart')
    if m and par.style.name=='Heading 1' and bs:
        ch=int(m.group(1));chapters[ch]=bs[0].get(qn('w:name'));heading_indices[ch]=i
def table_at(i,rows,widths):
    t=d.add_table(rows=0,cols=len(widths));t.autofit=False
    for col,w in zip(t.columns,widths):col.width=Inches(w)
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        x=OxmlElement('w:'+edge);x.set(qn('w:val'),'single');x.set(qn('w:sz'),'4');x.set(qn('w:color'),'D9D9D9');borders.append(x)
    t._tbl.tblPr.append(borders)
    for ri,data in enumerate(rows):
        row=t.add_row();pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
        if ri==0:pr.append(OxmlElement('w:tblHeader'))
        for ci,(cell,text) in enumerate(zip(row.cells,data)):
            cell.width=Inches(widths[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text=text
            pr=cell._tc.get_or_add_tcPr()
            mar=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                x=OxmlElement('w:'+side);x.set(qn('w:w'),'90');x.set(qn('w:type'),'dxa');mar.append(x)
            pr.append(mar)
            sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E7EDF2' if ri==0 else ('F5F7F9' if ri%2==0 else 'FFFFFF'));pr.append(sh)
            for para in cell.paragraphs:
                para.paragraph_format.space_after=Pt(3);para.paragraph_format.space_before=Pt(2);para.paragraph_format.line_spacing=1.0;para.paragraph_format.keep_with_next=(ri==0)
                for run in para.runs:
                    run.font.name='Palatino Linotype';run.font.size=Pt(9);run.bold=ri==0;run.font.color.rgb=RGBColor(0,0,0)
    p[i]._p.addprevious(t._tbl)
    return t
# Short main contents. Move the problem shelf after it; put subsection links in each chapter.
toc_chapters=[27,35,45,56,66,74,96,104,114,123,136,144,151,160,169]
nav={}
for ch,start in enumerate(toc_chapters,1):
    end=toc_chapters[ch] if ch<15 else 180
    nav[ch]=[]
    for i in range(start+1,end):
        if re.match(r'^[IVX]+\. ',p[i].text):
            nav[ch].append(deepcopy(p[i]._p))
for ch,items in nav.items():
    layer_i=next(i for i in range(heading_indices[ch]+1,min(heading_indices[ch]+8,len(p))) if p[i].text.startswith('Layer:'))
    par=after(layer_i,'In this chapter: ')
    par.paragraph_format.space_before=Pt(6);par.paragraph_format.space_after=Pt(10)
    par.paragraph_format.line_spacing=1.0
    for j,node in enumerate(items):
        if j:par.add_run(' · ')
        hs=node.xpath('.//w:hyperlink')
        if hs:
            for h in hs:par._p.append(deepcopy(h))
        else:
            par.add_run(''.join(node.xpath('.//w:t/text()')))
    for r in par._p.xpath('.//w:r'):
        pr=r.find(qn('w:rPr'))
        if pr is None:pr=OxmlElement('w:rPr');r.insert(0,pr)
        sz=OxmlElement('w:sz');sz.set(qn('w:val'),'19');pr.append(sz)
keep={18,19,25,180,181,182,183,*toc_chapters,*range(82,96)}
last=p[183]._p
for i in range(82,96):
    last.addnext(p[i]._p);last=p[i]._p
for i in range(18,184):
    if i not in keep:delete(i)
for i in [19,25,*toc_chapters,180,181,182,183]:
    p[i].paragraph_format.space_after=Pt(3);p[i].paragraph_format.space_before=Pt(0);p[i].paragraph_format.line_spacing=1.0
    for r in p[i]._p.xpath('.//w:r'):
        pr=r.find(qn('w:rPr'))
        if pr is None:pr=OxmlElement('w:rPr');r.insert(0,pr)
        sz=OxmlElement('w:sz');sz.set(qn('w:val'),'22');pr.append(sz)
# Durable method; fewer absolute claims.
change(192,'Start this week assigns dates to this route; chapter 6 continues the calendar.')
change(243,'Use the live Earn tab and the official eligibility page for your country before planning around a gate. The dated summary below explains the change that shapes this edition; chapter 9 is the detailed reference.')
change(245,'Official rule, checked 7 October 2026: the ad-and-Premium gate currently requires 1,000 subscribers plus either 4,000 qualified watch hours in the previous 12 months or 10 million qualified Shorts views in 90 days. From 1 February 2027, new entrants need 8,000 qualified hours in 365 days or 20 million qualified Shorts views in 90 days, still with 1,000 subscribers. Existing YPP creators keep their status; the earlier fan-funding tier is unchanged. Chapter 9 separates entry, activity, and payment requirements.')
link(p[245],' Current eligibility.',url='https://support.google.com/youtube/answer/72851?hl=en')
link(p[245],' Published 2027 change.',url='https://support.google.com/youtube/answer/12843009?hl=en')
change(272,'“A second subject is a second channel” is this book’s operating rule, not YouTube’s. Use it when the new subject serves a different viewer. A how-to and a vlog of the same job may belong together; a broader channel can work if the same viewer wants both. Test that overlap before splitting the work. A new channel must qualify separately.')
change(302,'Record the viewer and set-up decisions on the first two workbook sheets. The money map identifies which reference chapter to open when an income stream becomes relevant.')
change(314,'“The click is the transaction” is a useful shorthand for this chapter’s packaging job, not a published ranking rule. A click begins the viewing decision; retention, satisfaction, and continuation determine whether the promise held.')
# One short gate pointer per chapter, not repeated numerical lectures.
gate_note='Thresholds change. Use the live Earn tab and the dated gate reference in chapter 9.'
note_i=[i for i,par in enumerate(p) if par.text.startswith('Thresholds change. Before planning around')]
seen=set()
for i in note_i:
    owner=max((ch for ch,idx in heading_indices.items() if idx<i),default=0,key=lambda ch:heading_indices[ch])
    if owner==9:delete(i)
    elif owner in seen:delete(i)
    else:
        cross(i,gate_note,9);seen.add(owner)
# Map carries dependencies, not another threshold schedule.
t=d.tables[0]
dependencies=[
'Outside programs set their own terms. Clickable long-form links need advanced features. YouTube Shopping has its own YPP and country requirements.',
'A direct deal needs no YPP acceptance. Creator Partnerships has YPP, age, country, and channel-standing requirements.',
'YPP acceptance; the earlier tier where offered; age and country eligibility; Commerce Product Module. Chapter 9 gives the gates.',
'The higher ad-and-Premium gate and Watch Page Module. See chapter 9 for the dated entry requirements.',
'The higher gate and Shorts Module. The monthly payment floor is separate from entry; chapter 11, §V explains it.',
'Watch Page Module for long videos; Shorts Module for Shorts. The Shorts payment floor is in chapter 11, §V.'
]
for ri,dep in enumerate(dependencies,1):
    cell=t.cell(ri,2);cell.text=dep
    label=t.cell(ri,0).text.split('\n')[0]
    ch=13 if ri<=2 else (12 if ri==3 else 11)
    cell=t.cell(ri,0);cell.text=label
    cell.paragraphs[0].add_run('\n')
    link(cell.paragraphs[0],'Chapter '+str(ch),chapters[ch])
t.cell(4,3).text='No ad share before acceptance. Meeting a gate starts a review.'
t.cell(6,3).text='Paid from subscription pools. Shorts payment conditions differ from long-form.'
# Skepticism pass: delete hacks and fragile numerical production advice.
change(409,'Mixing is a listening test (author recommendation). Start with the music muted, make the speech clear, then bring the music up only while every word stays intelligible. Check on headphones and on the phone speaker. An editor’s fader position is not a loudness measurement. Use sound effects only where they explain something the viewer can see.')
change(428,'If the tool accepts references, reuse the approved set for each shot. Make a short test with the face, motion, and camera move the video actually needs. Keep the take only if the identity and geometry hold; price and free-tier status do not tell you that.')
change(431,'Keep each generation short enough to inspect. Compare the takes in motion, not as single frames, and record which failures make a shot unusable. A successful test is more useful than a promised number of seconds or regenerations.')
change(433,'Cut the approved takes as one sequence, match their color, and add the sound after the motion works. An upscale cannot repair an unusable take.')
change(1301,'Fader. The volume slider in an editor. Its position is not a loudness measurement; listen to the speech with the music added. ')
link(p[1301],'See chapter 3.',chapters[3])
delete(697)
change(695,'Author recommendation: tie a subscription request to a useful next video. “Next week is the tile line, the one most quotes leave out. Subscribe if you want it.” Use the sentence only if that video is planned.')
delete(696)
change(716,'Author recommendation: test a collaboration when both channels serve the same viewer. Record whether those visitors stay or continue; arrival from a related channel does not establish that they will.')
change(712,'Bought engagement is a policy problem, not a shortcut to a reliable audience. YouTube’s spam policy names coordinated sub-for-sub schemes as engagement manipulation; the fake-engagement policy covers artificial views, likes, comments, and other metrics.')
# Give the eighth video a distinct, useful job.
change(552,'The path video. A short index to the seven teaching videos: where a beginner should start, which video solves each question, and what to watch next. It earns its slot when the series needs a watch order that a list of titles does not explain. Show the choices and name the destinations; do not reteach the method. Each teaching video still ends with its own next-video pointer. Those pointers and this index are different jobs.')
change(1174,'The eight jobs, as a checklist: seven teaching videos and one path video. The path video explains the watch order; each teaching video also names its next destination.')
# Name the workbook's routing entry consistently.
for i,par in enumerate(p):
    if i>1136 and ('which video to watch next' in par.text or 'and which video to watch next' in par.text):
        for node in par._p.xpath('.//w:t'):
            if node.text:node.text=node.text.replace('which video to watch next','the path video and watch order')
# Chapter 8: replace the second copy of the full route with an application hand-off.
change(681,'Use your own average view duration from chapter 7. Planning estimate: remaining qualified hours × 60 ÷ average minutes watched = additional views needed, if those new views qualify. Divide by the days available for a daily pace. The estimate ignores older hours falling out of the rolling window, so record that limitation beside it.')
change(717,'Record the route, the remaining distance, and one test on the week sheet. Review the test on the date you wrote down.')
change(719,'When Earn shows eligibility, move to chapter 9 for the application and review. Before then, continue the existing production route and record the qualified counts weekly.')
for i in range(720,729):delete(i)
change(730,'That Shorts views fill the long-form hour bar. YouTube counts them toward a separate view gate; watch time in the Shorts feed does not fill the hour gate.')
change(732,'That subscribers alone establish a working audience. They satisfy one eligibility requirement; returning viewing is a different question.')
change(734,'That a growth service or a posting streak substitutes for qualified viewing. Choose the route and test whether viewers continue through the videos.')
# Definitive gate table, interpretation visibly separated.
gate_rows=[
['Application date','Subscribers','Qualified long-form hours','Or qualified Shorts views'],
['Current gate before 1 February 2027','1,000','4,000 in the previous 12 months','10 million in 90 days'],
['From 1 February 2027, new entrants','1,000','8,000 in the previous 365 days','20 million in 90 days']
]
table_at(759,gate_rows,[1.6,.65,1.65,1.7]);delete(759);delete(760)
change(764,'Official rule, checked 7 October 2026: YouTube’s Changes page gives the effective date and says existing YPP creators keep their status. The table gives entry requirements, not the separate Shorts payment floor or ongoing activity test.')
link(p[764],' Published change.',url='https://support.google.com/youtube/answer/12843009?hl=en')
link(p[764],' Current gate.',url='https://support.google.com/youtube/answer/72851?hl=en')
change(766,'Author reading — application timing. Based on the stated 1 February effective date, this guide treats a completed application submitted before then under the threshold displayed at submission. This is an interpretation, not an express promise from YouTube about an application still under review on that date. Confirm the applicable threshold in Earn when you apply and save the dated notice.')
change(793,'The earlier tier and the ad-and-Premium tier unlock different features. Meeting either gate begins a review; acceptance still requires the modules and payment setup.')
# Diagnosis table and a prominent upload-checklist link, while keeping detailed sections.
change(805,'Use the table to identify the notice, then read that section and the linked official page. Record the affected video, the action, and any deadline in the rights log. A video-level problem and a channel-level problem call for different responses.')
after(804,'Before your first upload: ').add_run('complete the checklist in §XI.')
new=p[804]._p.getnext()
from docx.text.paragraph import Paragraph
link(Paragraph(new,p[804]._parent),' Open the upload checklist.','rev10UploadChecklist')
rows=[
['Notice','What it affects','Clock','Read'],
['Guidelines warning','A policy violation; training may be available.','Eligible warnings expire 90 days after completed training.','§II · Strike basics'],
['Guidelines strike','Posting privileges and channel standing.','First posting pause: 1 week; second: 2 weeks. Each strike remains 90 days.','§II · Strike basics'],
['Copyright strike','The video is removed; channel standing is affected.','90 days if Copyright School is completed; three active strikes risk termination.','§III · Copyright strikes'],
['Content ID claim','That video can be blocked, tracked, or monetized for the claimant.','Read the notice for any dispute or appeal deadline. No strike by itself.','§IV · Copyright claims'],
['Limited or no ads','Earnings on that video.','Use the review instructions shown in Studio; this is not a strike.','§V · Advertiser guidance']
]
diag=table_at(808,rows,[1.1,1.55,1.85,1.1])
urls=['https://support.google.com/youtube/answer/2802032?hl=en','https://support.google.com/youtube/answer/2802032?hl=en','https://support.google.com/youtube/answer/2814000?hl=en','https://support.google.com/youtube/answer/6013276?hl=en','https://support.google.com/youtube/answer/6162278?hl=en']
for ri,url in enumerate(urls,1):
    cell=diag.cell(ri,3);text=cell.text;cell.text=''
    link(cell.paragraphs[0],text,url=url)
for i in range(808,813):delete(i)
change(814,'Disclosures are checks you make before publication: §§VI–VII cover paid promotion and realistic AI use; §IX covers the audience setting. Section VIII covers whole-channel monetization policies. The diagnostic table is a summary; the notice and current official page control the action.')
change(829,'Copyright and Content ID are different. A copyright strike follows a legal removal request; a Content ID claim is a match that affects the video. If a dispute leads to a valid removal request, the result can be a strike. Read the license and keep its record before publishing.')
change(887,'[ ] If Studio asks for ad self-certification, the answers describe the actual video, opening, title, and thumbnail.')
p[883].paragraph_format.space_before=Pt(18)
p[883].paragraph_format.keep_with_next=True
# Synchronize the chapter 14 navigation label with its new section heading after editing.
# Keep the chapter summaries where they add a distinction; cut repeated gating instructions.
change(946,'Author recommendation: if your channel is below the monthly Shorts payment floor, plan income from the streams it can actually use. Shorts may still introduce viewers to long videos; verify that transition in Studio rather than assuming it.')
change(967,'Fan funding depends on viewers choosing to pay for the channel, not just on view volume. The earlier YPP tier can make it available before ad revenue in supported countries; chapter 9 gives the entry requirements.')
change(1006,'That fan funding requires the full ad gate. Where the earlier YPP tier is available, its features can open sooner. Check the country and feature requirements in chapter 9.')
change(1049,'Say the disclosure in the video before the recommendation. If they gave you the product, select paid promotion. Advanced features govern clickable long-form description links; YouTube Shopping product tags follow the separate requirements in §VI.')
# Six streams throughout.
change(1092,'A change can affect entry to a program, payment from one stream, or ongoing activity. Chapter 9 separates those conditions. The income record should show which part of the channel depends on each.')
stream_text=[
'Watch-page ads pay for advertising on eligible long videos.',
'Shorts revenue pays for eligible viewing in the Shorts feed.',
'Premium pays from subscription revenue allocated to viewing.',
'Fan funding includes memberships and Supers; it depends on viewers choosing to pay.',
'Sponsors pay for the audience and the work agreed in a deal.',
'Affiliate commissions pay for purchases attributed to a disclosed recommendation.'
]
for i,text in zip(range(1093,1099),stream_text):change(i,text)
change(1263,'The money map has six streams: watch-page ads, Shorts revenue, Premium, fan funding, sponsors, and affiliate commissions. They pay under different terms. Keep separate records so you can see what a policy change or a quiet month affected.')
# Chapter 14: a durable comparison, with a directory rather than another dated floor list.
change(1100,'IV. Compare another platform before cross-posting')
change(1101,'Use this comparison before choosing a second platform. It is an author recommendation, not a payout forecast. Copy the answers and the official page’s date into the decision log.')
change(1102,'Asset rights. Does every footage, music, voice, and image license allow this site and this use? A YouTube-specific grant does not answer that.')
change(1103,'Eligibility. What counts: followers, qualified viewing, activity, country, age, account type, application, or invitation? Record the current requirement and its window.')
change(1104,'Qualifying content. Which formats, lengths, topics, and originality rules apply? Check the finished video, not just the account.')
change(1105,'Payment. What creates revenue: ads, subscriptions, rewards, or another mechanism? What is excluded, and what account, fees, tax information, or payment minimum is required?')
change(1106,'Audience. Is the same viewer already there, and do they watch this kind of video? A second logo is not evidence of an audience.')
change(1107,'Additional work. Count the recut, caption, upload, moderation, and record-keeping time. Reuse one finished video as a bounded test before opening another production schedule.')
change(1108,'Current source. Start with the official program pages below. Use the terms for your country and the account’s own eligibility notice. Save the date and the relevant page; no follower-floor table in this book replaces that check.')
change(1109,'Official program directory: ')
directory=[
('TikTok Creator Rewards (US terms)','https://www.tiktok.com/legal/page/global/creator-rewards-program-us/en'),
('Facebook Content Monetization','https://www.facebook.com/business/help/1049081556813520'),
('Snapchat Monetization Program','https://help.snapchat.com/hc/en-us/articles/14669003687444-About-Snapchat-s-Monetization-Program'),
('X Original Content Rewards','https://help.x.com/en/using-x/original-content-rewards'),
('Rumble Creator Program','https://rumble.support/help/rumble-creator-program')
]
for j,(label,url) in enumerate(directory):
    if j:p[1109].add_run(' · ')
    link(p[1109],label,url=url)
delete(1110);delete(1111)
change(1113,'A second YouTube channel qualifies separately. It does not inherit subscribers, qualified viewing, or acceptance. Use chapter 9 for its applicable entry requirements.')
change(1117,'Activity is an ongoing eligibility condition, separate from the entry gate. Chapter 9, §VI holds the current test and the 2027 change; check that section before planning a long break.')
change(1119,'Author recommendations for a sustainable pace:')
change(1131,'That one large stream is automatically safer than several small ones. Compare the obligations and dependencies of each; diversification can also add work.')
change(1132,'That a platform program is permanent. Section IV records the current terms and the date, so a change can be compared with the decision you made.')
change(1134,'Sources: paid-campaign viewing, activity, and breaks are documented in YouTube’s linked pages in chapters 2, 8, and 9. Section IV is the author’s comparison framework; its directory links to official program pages rather than quoting payout floors. Links checked 7 October 2026; Facebook requires sign-in for the linked help page.')
change(1330,'Chapter 14. The diagnostic order and cross-platform comparison are author recommendations. The directory supplies official program pages, not a frozen list of thresholds. Entry and activity rules are centralized in chapter 9.')
change(1324,'Chapter 8. Qualified viewing, live-stream conditions, paid campaigns, and engagement manipulation are cited to YouTube’s official pages. The distance calculation is an estimate; subscription wording and collaboration choices are author recommendations.')
# Clear stale numeric PAGE caches in footers while retaining functional page fields.
for rel in d.part.rels.values():
    if rel.reltype.endswith('/footer'):
        root=rel.target_part.element
        if not root.xpath('.//w:fldChar | .//w:fldSimple | .//w:instrText') and not ''.join(root.xpath('.//w:t/text()')).strip():
            continue
        for para in root.xpath('.//w:p'):
            for child in list(para):
                if child.tag!=qn('w:pPr'):para.remove(child)
            field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),' PAGE ')
            r=OxmlElement('w:r');t=OxmlElement('w:t');t.text='';r.append(t);field.append(r);para.append(field)
# Headings and inline labels: typography supplies the distinction without reference boxes.
for name in ['Title','Subtitle','Heading 1','Heading 2','Heading 3']:
    if name in d.styles:d.styles[name].font.color.rgb=RGBColor(0,0,0)
for par in d.paragraphs:
    if par.style.name in ['Title','Subtitle','Heading 1','Heading 2','Heading 3']:
        for run in par.runs:run.font.color.rgb=RGBColor(0,0,0)
    # Bold only a leading claim label, preserving the following prose and links.
    for label in ['Official rule','Documented platform guidance','Creator heuristic','Author recommendation','Author reading']:
        if par.runs and par.runs[0].text.startswith(label):
            run=par.runs[0];rest=run.text[len(label):];run.text=label;run.bold=True
            new=deepcopy(run._r);new.find(qn('w:t')).text=rest
            prop=new.find(qn('w:rPr'))
            if prop is not None:
                for bold in prop.findall(qn('w:b')):prop.remove(bold)
            run._r.addnext(new)
            break
# Consistent table body styling after content edits, including links.
for table in d.tables:
    for ri,row in enumerate(table.rows):
        for cell in row.cells:
            for para in cell.paragraphs:
                para.paragraph_format.space_after=Pt(3);para.paragraph_format.space_before=Pt(2);para.paragraph_format.line_spacing=1.0
                for node in para._p.xpath('.//w:r'):
                    prop=node.find(qn('w:rPr'))
                    if prop is None:prop=OxmlElement('w:rPr');node.insert(0,prop)
                    for tag in ['w:sz','w:szCs']:
                        for old in prop.findall(qn(tag)):prop.remove(old)
                        size=OxmlElement(tag);size.set(qn('w:val'),'18');prop.append(size)
                    fonts=OxmlElement('w:rFonts');fonts.set(qn('w:ascii'),'Palatino Linotype');fonts.set(qn('w:hAnsi'),'Palatino Linotype');prop.append(fonts)
                    if ri==0:
                        bold=OxmlElement('w:b');prop.append(bold)
for par in d.paragraphs:
    if par.text.startswith('In this chapter:'):
        for node in par._p.xpath('.//w:t'):
            if node.text:node.text=node.text.replace('IV. Other platforms pay for the same video','IV. Compare another platform before cross-posting')
d.core_properties.comments='REV11: consistency, evidence, repetition, navigation, and production cleanup against the supplied review, 7 October 2026.'
d.save(OUT)
fresh=Document(OUT)
anchors={b.get(qn('w:name')) for b in fresh.element.xpath('.//w:bookmarkStart')}
missing=[h.get(qn('w:anchor')) for h in fresh.element.xpath('.//w:hyperlink[@w:anchor]') if h.get(qn('w:anchor')) not in anchors]
assert not missing,missing
body='\n'.join(par.text for par in fresh.paragraphs)
for banned in ['?sub_confirmation=1','two or three regenerations','reliably fail','five kinds','Open chapter.','A channel that meets the current numbers and applies before that date applies under them.']:
    assert banned not in body,banned
assert len(fresh.tables)==3
source=Document(SRC)
oldwords=sum(len(par.text.split()) for par in source.paragraphs)+sum(len(cell.text.split()) for tab in source.tables for row in tab.rows for cell in row.cells)
newwords=sum(len(par.text.split()) for par in fresh.paragraphs)+sum(len(cell.text.split()) for tab in fresh.tables for row in tab.rows for cell in row.cells)
qa=ROOT/'bak/rev11_qa';qa.mkdir(parents=True,exist_ok=True)
(qa/'changes.json').write_text(json.dumps(changes,ensure_ascii=False,indent=2),encoding='utf-8')
(qa/'paragraphs.txt').write_text('\n'.join(f'{i}\t{par.text}' for i,par in enumerate(fresh.paragraphs)),encoding='utf-8')
print('Saved',OUT,'Changes',len(changes),'Words',oldwords,'->',newwords)
