
from pathlib import Path
from copy import deepcopy
import re,json
from docx import Document
from docx.shared import Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[1]
SRC=ROOT/'Your First YouTube Channel_REV12_2026-10-07.docx'
OUT=ROOT/'Your First YouTube Channel_REV13_2026-10-07.docx'
d=Document(SRC);p=list(d.paragraphs);changes=[]
assert len(p)==1215,'Source revision changed'
deleted_anchors={}
def change(i,text):
    old=p[i].text
    for c in list(p[i]._p):
        if c.tag not in (qn('w:pPr'),qn('w:bookmarkStart'),qn('w:bookmarkEnd')):
            p[i]._p.remove(c)
    p[i].add_run(text)
    changes.append(dict(paragraph=i,before=old,after=text))
def delete(i):
    if p[i]._p.getparent() is not None:
        changes.append(dict(paragraph=i,before=p[i].text,after='[removed]'))
        p[i]._p.getparent().remove(p[i]._p)
def link(par,label,anchor=None,url=None):
    h=OxmlElement('w:hyperlink')
    if anchor: h.set(qn('w:anchor'),anchor)
    if url:
        from docx.opc.constants import RELATIONSHIP_TYPE
        h.set(qn('r:id'),d.part.relate_to(url,RELATIONSHIP_TYPE.HYPERLINK,is_external=True))
    r=OxmlElement('w:r'); prop=OxmlElement('w:rPr')
    sty=OxmlElement('w:rStyle');sty.set(qn('w:val'),'Hyperlink');prop.append(sty);r.append(prop)
    t=OxmlElement('w:t');t.text=label;t.set(qn('xml:space'),'preserve');r.append(t);h.append(r);par._p.append(h)
def bookmark(i,name):
    b=OxmlElement('w:bookmarkStart'); b.set(qn('w:id'),'9010'); b.set(qn('w:name'),name)
    e=OxmlElement('w:bookmarkEnd'); e.set(qn('w:id'),'9010')
    p[i]._p.insert(1,b);p[i]._p.append(e)

base_delete=delete
def delete(i):
    for b in p[i]._p.xpath('./w:bookmarkStart'):deleted_anchors[b.get(qn('w:name'))]='toc0'
    base_delete(i)
def replace_start(prefix,text):
    ids=[i for i,par in enumerate(p) if par.text.startswith(prefix)]
    assert len(ids)==1,(prefix,ids)
    change(ids[0],text);return ids[0]
def cut_start(prefix):
    ids=[i for i,par in enumerate(p) if par.text.startswith(prefix)]
    assert len(ids)==1,(prefix,ids)
    delete(ids[0])
def add_before(i,text,style=None):
    return p[i].insert_paragraph_before(text,style)
def add_after(i,text,style=None):
    par=add_before(i,text,style);p[i]._p.addnext(par._p);return par
def anchor(par,name,id):
    b=OxmlElement('w:bookmarkStart');b.set(qn('w:id'),str(id));b.set(qn('w:name'),name)
    e=OxmlElement('w:bookmarkEnd');e.set(qn('w:id'),str(id))
    par._p.insert(1,b);par._p.append(e)
def part_before(i,title,intro,name,id):
    h=add_before(i,title,'Heading 1');h.paragraph_format.page_break_before=True
    h.paragraph_format.keep_with_next=True
    anchor(h,name,id)
    body=add_before(i,intro)
    body.paragraph_format.keep_with_next=True
    p[i].paragraph_format.page_break_before=False
    return h
# The three parts are the physical hierarchy.
part_before(119,'Part I Build the Channel',
'Start this week and chapters 1–8 are the working guide: make the connected videos, read what viewers did, and change one break. Use the calendar at your own pace.',
'rev13PartI',9131)
part_before(604,'Part II When Money and Rules Matter',
'This is the reference manual. You do not have to learn it all while making video one. Use a chapter when you need a gate, a notice, or a payment rule. The exception is chapter 10, §XI: complete that checklist before every upload, starting with the first.',
'rev13PartII',9132)
part_before(1001 if p[1001].text=='15 The Channel Workbook' else next(i for i,par in enumerate(p) if par.text=='15 The Channel Workbook' and par.style.name=='Heading 1'),
'Part III Workbook',
'Keep the records here. Fill the sheet for the decision you are making, and keep the old answers when you change one.',
'rev13PartIII',9133)
# Part labels in the short main contents.
for i,title,an in [(20,'Part I Build the Channel','rev13PartI'),(29,'Part II When Money and Rules Matter','rev13PartII'),(35,'Part III Workbook','rev13PartIII')]:
    par=add_before(i,'')
    link(par,title,an)
    par.paragraph_format.space_before=Pt(8);par.paragraph_format.space_after=Pt(3)
    for run in par._p.xpath('.//w:r'):
        prop=run.find(qn('w:rPr'))
        if prop is None:prop=OxmlElement('w:rPr');run.insert(0,prop)
        prop.append(OxmlElement('w:b'))
change(40,'Find the page you need')
change(41,'Chapters 1–8 build the channel. Chapters 9–14 are the money and policy reference; chapter 15 holds the records. Choose the problem below, or follow Start this week.')
delete(43)
# Start here: remove about half the entrance, keeping the route and claim labels.
change(56,'This is for a first channel, or a channel whose videos have not yet found the same viewer twice. The test is a stranger who finishes one video and starts the next. A relative who says it is lovely does not count. They are being polite. A working channel is built from viewing people choose to continue; monetization then follows published eligibility and payment rules.')
change(58,'Tonight, make the first decision on paper. The rest of this page is the route; Start this week gives it dates.')
change(59,'Name one viewer and the result they want. Write twenty titles in the phrase they would type. Circle “who it is for” and the deep dive. Set up the channel with its audience setting answered truthfully. No banner this week.')
change(60,'Publish the first video. Open on the finished result and give the method away, including the hard step. Use a thumbnail readable on a phone. Name the next video at the end. Watch the finished file on a phone before you publish it; if you would have left, it is not done.')
change(61,'Publish the deep dive next. Ask three strangers who match the viewer sentence where they would have stopped. Write their answers down. Change the thumbnail or the opening, not both. Then finish the other six jobs in chapter 5.')
for i in range(63,68):delete(i)
change(68,'Before the first upload, complete the policy checklist. ')
link(p[68],'Chapter 10, §XI.','rev10UploadChecklist')
change(69,'Tonight’s checklist. Mark it on paper.')
change(75,'Paper. Next video named.')
for i in [79,80,81]:delete(i)
# One navigation convention: layer, workbook, chapter contents.
change(82,'Where to look')
change(83,'The problem links before this page find the count, policy, or income question you need. Part I is the work. Part II is the reference. Part III is the record. Each chapter’s layer and workbook lines say when to use it and which sheet to fill.')
for i in range(84,96):delete(i)
change(96,'The workbook is the spine')
change(97,'Keep the sheets in a notebook or a file you can revise. The viewer sentence, video job, weekly counts, rights, money, and decisions belong in one dated record. Chapter 15 supplies the blanks.')
for i in range(98,105):delete(i)
for i in range(112,118):delete(i)
# Precise openings, with the retained sharp chapter titles.
change(134,'A channel is a place a particular viewer recognizes, uses, and returns to. Name that viewer before choosing the tools or the income stream.')
change(143,'Author recommendation: put a second subject on a second channel when it serves a different viewer. A how-to and a vlog of the same job may belong together; a broader channel can work if the same viewer wants both. Test that overlap before splitting the work.')
change(172,'A post can ask which question the viewer wants answered next. Use it between videos, not in place of them. The monetized-channel activity rules are in the reference section.')
change(185,'A good video cannot help until someone chooses it. The title and thumbnail let the right viewer recognize the job. This chapter covers the form limits, caption options, and packaging tests that help the promise hold.')
delete(186)
change(221,'For the long-form watch-hour route, a Short that sends no one to a long video does not move that hour counter. It may still attract subscribers or contribute to the separate Shorts route.')
change(218,'The audience-retention report shows the percentage still watching after thirty seconds and compares the opening with recent videos of similar length. It usually takes one to two days to appear. Documented platform guidance: a high intro percentage can mean the opening matched the title and thumbnail. Use the comparison to inspect the opening, rather than as a target percentage.')
change(234,'Use the evening checklist for packaging. This chapter adds upload limits, captions, Studio testing, and discovery reports.') if p[234].text.startswith('For the packaging checklist') else None
# Licensing leads chapter 3; examples are explicitly subordinate and dated.
change(252,'Finish with tools you can use and output you have the right to publish. A price tag answers neither question by itself.')
change(253,'Start with the five checks below. Record the answers before choosing the application.')
# Move the stack and its graphic below the license framework.
last=p[265]._p
for i in [254,255,256]:
    last.addnext(p[i]._p);last=p[i]._p
intro=add_before(266,'Dated examples follow, retained from the October 2026 checks. Use the five questions to decide whether the current version and terms fit your video.')
intro.runs[0].italic=True;intro.runs[0].font.size=Pt(11)
for i,par in enumerate(p):
    if par.text.startswith('What matters most:'):delete(i)
for i in [259,260,261,262,263]:
    par=p[i];text=par.text
    split=text.find('. ')
    if split>0:
        change(i,text[:split+1]);par.runs[0].bold=True;par.add_run(text[split+1:])
# Named examples use a smaller example style, preserving their actual links.
example_prefixes=['DaVinci Resolve has','CapCut runs','Pexels:','Pixabay:','Uppbeat.','YouTube Studio’s own subtitle editor','OBS Studio is','Voice example:','Footage example:']
for par in d.paragraphs:
    if any(par.text.startswith(prefix) for prefix in example_prefixes):
        par.paragraph_format.left_indent=Pt(8)
        for run in par.runs:run.font.size=Pt(11)
        for h in par._p.xpath('.//w:hyperlink/w:r'):
            prop=h.find(qn('w:rPr'))
            if prop is None:prop=OxmlElement('w:rPr');h.insert(0,prop)
            size=OxmlElement('w:sz');size.set(qn('w:val'),'22');prop.append(size)
# Cut duplicate gate warnings and separate the necessary claim labels from defensive recaps.
for i,par in enumerate(p):
    if par.text.startswith('Thresholds change. Use the live Earn tab'):delete(i)
    if par.text.startswith('Shelf job:') or par.text.startswith('Workbook job:'):delete(i)
# Chapters 7–14: keep the distinct facts and tests; remove second summaries of the route.
change(487,'Read whether strangers find the videos, stay, and continue. The week sheet records those decisions; the reference section covers eligibility and revenue.')
delete(542)
change(531,'When YouTube changes a metric you track, record the date in the decision log. Note that break in the series before comparing the numbers.')
change(556,'The gate is arithmetic. Use qualified counts, their rolling windows, and your own view duration to estimate the distance.')
delete(557)
change(565,'Author recommendation: choose a route you can sustain. A steady long-form audience and a large Shorts audience require different production work.')
change(571,'Ask for a subscription once, with a reason tied to the next useful video.')
delete(572)
delete(574)
change(576,'Ten minutes watched adds the same time whether it is in one video or across several. Connect the videos so a viewer who needs another step can find it.')
delete(587) # the next paragraph carries the official fake-engagement rule
change(595,'When Earn shows eligibility, use the application instructions in Part II. Until then, keep the weekly qualified counts beside the production record.')
delete(606);delete(657)
# Activity: ordinary compliance and the restoration window are separate tests.
change(654,'Official rule, checked 7 October 2026. Until the 2027 update, YouTube may turn off monetization after six months without an upload or post. From 1 February 2027, the ordinary active-channel test is met by any one of: 1,000 qualified watch hours in 365 days; 1 million qualified Shorts views in 90 days; or two long-form videos or five Shorts uploaded every 90 days.')
link(p[654],' Current eligibility.',url='https://support.google.com/youtube/answer/72851?hl=en')
link(p[654],' Updated activity test.',url='https://support.google.com/youtube/answer/12843009?hl=en')
restore=add_after(654,'Restoration is a different test. If YouTube places the channel in the extended 90-day restoration window, its published restoration conditions list the qualified watch hours or qualified Shorts views above. Uploads alone satisfy the ordinary test, but are not listed as a way to restore status in that window.')
link(restore,' Restore active status.',url='https://support.google.com/youtube/answer/12843009?hl=en')
# Specific copy cuts in the reference chapters, without removing the policy table or checklist.
for prefix in [
'In short: A limited-ads label',
'In short: CPM is',
'In short: Fan funding',
'In short: A sponsor',
'In short: When the counts stall'
]:
    hits=[i for i,par in enumerate(p) if par.text.startswith(prefix)]
    for i in hits:delete(i)
# Focus the reference introductions on their job rather than re-stating the whole model.
for old,new in [
('YouTube Studio is generous with numbers.','Studio offers more numbers than a weekly decision needs. Read the reports below in order: arrival, retention, continuation.'),
('A monetized channel is not finished.','Use this chapter when the counts stall, a stream changes, or the production pace stops being sustainable.'),
('A channel that pays leaves a trail:','Keep the old answers when you revise a sheet. That record is how you can tell what changed and whether the change helped.')
]:
    hits=[i for i,par in enumerate(p) if par.text.startswith(old)]
    assert len(hits)==1,(old,hits)
    change(hits[0],new)
# The last word brings the reader back to the stranger rather than another gate lecture.
end_i=next(i for i,par in enumerate(p) if par.text=='The last word' and par.style.name=='Heading 1')
assert p[end_i+1].text.startswith('The channel in this book')
change(end_i+1,'A working channel can feel small before it pays. A stranger asks a question that shows they tried the method. Another comes back for the next step. Someone tells you which explanation finally made sense. Those are reasons to keep making the work; record what they teach you without turning them into an income forecast.')
change(end_i+2,'There is room to enjoy this. Try a clearer demonstration, a better camera angle, a way of explaining that sounds like you. Keep the viewer steady while you discover the style. One change gives you something to learn from, not just another restriction.')
change(end_i+3,'The reference pages will be there when you need a rule or a payment condition. The records will show what the work cost and what it earned. Neither has to occupy the whole evening.')
change(end_i+4,'Make the next useful video for the person you named at the beginning. If a stranger uses it and comes back, notice that too.')
# Minimize return-link scaffolding: retain chapter navigation and only the necessary upload link.
for par in d.paragraphs:
    hs=par._p.xpath('.//w:hyperlink[@w:anchor="tocEvening"]')
    for h in hs:
        # Keep the calendar's single return link, remove repeated add-ons.
        if not par.text.startswith('The evening page in Start here'):
            h.getparent().remove(h)
# Synchronize chapter navigation with section headings, retaining destinations.
headings={}
for par in d.paragraphs:
    if par.style.name.startswith('Heading'):
        for b in par._p.xpath('./w:bookmarkStart'):headings[b.get(qn('w:name'))]=par.text
for par in d.paragraphs:
    if par.text.startswith('In this chapter:'):
        for h in par._p.xpath('.//w:hyperlink[@w:anchor]'):
            dest=h.get(qn('w:anchor'))
            if dest in headings:
                ts=h.xpath('.//w:t')
                if ts:
                    ts[0].text=headings[dest]
                    for t in ts[1:]:t.text=''
# Anchor repair for removed orientation headings: send those links to Start here.
remaining={b.get(qn('w:name')) for b in d.element.xpath('.//w:bookmarkStart')}
start_anchor=next(b.get(qn('w:name')) for b in p[55]._p.xpath('./w:bookmarkStart'))
for h in d.element.xpath('.//w:hyperlink[@w:anchor]'):
    dest=h.get(qn('w:anchor'))
    if dest not in remaining:
        assert dest in deleted_anchors,(dest,'unexpected missing destination')
        h.set(qn('w:anchor'),start_anchor)
# Bold leading claim labels after rewritten paragraphs.
for par in d.paragraphs:
    for label in ['Official rule','Documented platform guidance','Creator heuristic','Author recommendation']:
        if par.runs and par.runs[0].text.startswith(label):
            run=par.runs[0]
            if run.text==label:break
            rest=run.text[len(label):];run.text=label;run.bold=True
            new=deepcopy(run._r);new.find(qn('w:t')).text=rest
            prop=new.find(qn('w:rPr'))
            if prop is not None:
                for b in prop.findall(qn('w:b')):prop.remove(b)
            run._r.addnext(new);break
d.core_properties.comments='REV13: three-part hierarchy, compressed entrance and reference prose, license-first production, precise activity-restoration distinction, and a human ending. 7 October 2026.'
d.save(OUT)
fresh=Document(OUT);source=Document(SRC)
def wc(doc):return sum(len(par.text.split()) for par in doc.paragraphs)+sum(len(cell.text.split()) for t in doc.tables for row in t.rows for cell in row.cells)
def entrance(doc):
 ps=doc.paragraphs;start=next(i for i,par in enumerate(ps) if par.text=='Start here' and par.style.name=='Heading 1');end=next(i for i,par in enumerate(ps) if i>start and par.style.name=='Heading 1');return sum(len(par.text.split()) for par in ps[start:end])
anchors={b.get(qn('w:name')) for b in fresh.element.xpath('.//w:bookmarkStart')}
missing=[h.get(qn('w:anchor')) for h in fresh.element.xpath('.//w:hyperlink[@w:anchor]') if h.get(qn('w:anchor')) not in anchors]
assert not missing,missing
text='\n'.join(par.text for par in fresh.paragraphs)
for bad in ['YouTube pays for finished minutes','Nobody watches a video because it is good','A Short that never points at a long video ends there','Shelf job:','Thresholds change. Use the live Earn tab']:
    assert bad not in text,bad
assert len(fresh.tables)==3 and len(fresh.inline_shapes)==1
assert sum(par.text.startswith('In this chapter:') for par in fresh.paragraphs)==15
assert wc(fresh)<wc(source)
parts=[(i,par.text) for i,par in enumerate(fresh.paragraphs) if par.style.name=='Heading 1' and par.text.startswith('Part ')]
assert [t for _,t in parts]==['Part I Build the Channel','Part II When Money and Rules Matter','Part III Workbook'],parts
for (i,_),expected in zip(parts,['Start this week','9 The Gates and the Review','15 The Channel Workbook']):
    assert fresh.paragraphs[i+2].text==expected,(i,expected)
qa=ROOT/'bak/rev13_qa';qa.mkdir(parents=True,exist_ok=True)
(qa/'changes.json').write_text(json.dumps(changes,ensure_ascii=False,indent=2),encoding='utf-8')
(qa/'paragraphs.txt').write_text('\n'.join(f'{i}\t{par.text}' for i,par in enumerate(fresh.paragraphs)),encoding='utf-8')
report={'words_before':wc(source),'words_after':wc(fresh),'entrance_before':entrance(source),'entrance_after':entrance(fresh),'tables':len(fresh.tables),'images':len(fresh.inline_shapes),'internal_links':'all resolve'}
(qa/'content_checks.json').write_text(json.dumps(report,indent=2),encoding='utf-8')
print('Saved',OUT);print(json.dumps(report))
