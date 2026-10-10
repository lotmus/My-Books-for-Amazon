
from pathlib import Path
from copy import deepcopy
import re,json
from docx import Document
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[1]
SRC=ROOT/'Your First YouTube Channel_REV11_2026-10-07.docx'
OUT=ROOT/'Your First YouTube Channel_REV12_2026-10-07.docx'
d=Document(SRC);p=list(d.paragraphs);changes=[]
assert len(p)==1213,'Source revision changed'
def change(i,text):
    old=p[i].text
    for c in list(p[i]._p):
        if c.tag not in (qn('w:pPr'),qn('w:bookmarkStart'),qn('w:bookmarkEnd')):
            p[i]._p.remove(c)
    p[i].add_run(text)
    changes.append(dict(paragraph=i,before=old,after=text))
def delete(i):
    if p[i]._p.getparent() is not None:
        changes.append(dict(paragraph=i,before=p[i].text,after='[removed]'))
        p[i]._p.getparent().remove(p[i]._p)
def link(par,label,anchor=None,url=None):
    h=OxmlElement('w:hyperlink')
    if anchor: h.set(qn('w:anchor'),anchor)
    if url:
        from docx.opc.constants import RELATIONSHIP_TYPE
        h.set(qn('r:id'),d.part.relate_to(url,RELATIONSHIP_TYPE.HYPERLINK,is_external=True))
    r=OxmlElement('w:r'); prop=OxmlElement('w:rPr')
    sty=OxmlElement('w:rStyle');sty.set(qn('w:val'),'Hyperlink');prop.append(sty);r.append(prop)
    t=OxmlElement('w:t');t.text=label;t.set(qn('xml:space'),'preserve');r.append(t);h.append(r);par._p.append(h)
def bookmark(i,name):
    b=OxmlElement('w:bookmarkStart'); b.set(qn('w:id'),'9010'); b.set(qn('w:name'),name)
    e=OxmlElement('w:bookmarkEnd'); e.set(qn('w:id'),'9010')
    p[i]._p.insert(1,b);p[i]._p.append(e)

change(1040,'The problem in their words —')
change(1046,'The path video and watch order —')
change(1132,'Eight jobs. Seven teaching videos—who it is for, the problem in their words, the method, proof, comparison, one objection, and the deep dive—plus a path video explaining where to start and what to watch next. ')
link(p[1132],'See chapter 5.','toc41')
change(143,'This book’s operating rule is to put a second subject on a second channel when it serves a different viewer. That is an author recommendation, not a YouTube requirement. A how-to and a vlog of the same job may belong together; a broader channel can work if the same viewer wants both. Test that overlap before splitting the work. A new channel must qualify separately.')
change(186,'“The click is the transaction” is shorthand for this chapter’s packaging job, not a published ranking rule. A click begins the viewing decision; retention, satisfaction, and continuation determine whether the promise held.')
change(982,'Author recommendation: use the following practices to keep a sustainable pace.')
p[982].runs[0].text='Author recommendation:'
p[982].runs[0].bold=True
p[982].add_run(' use the following practices to keep a sustainable pace.')
change(91,'Before upload one and before monetization.')
p[91].runs[0].bold=True
p[91].add_run(' Chapter 10 covers rules that apply from the first upload; chapter 9 covers the gates and the review.')
# Keep one prose statement and one graphic; shorten its extraction/accessibility description.
for sh in d.inline_shapes:
    sh._inline.docPr.set('descr','Diagram of the four production jobs: voice, music, editing, and footage. Generated material is an optional side door.')
    sh._inline.docPr.set('title','Figure 3.1 Production tools by job')
# Remove source notes for material absent from the current chapters, preserving live links.
orphan='The speaking-rate range is a planning guess to replace with your own timed reading. '
for par in d.paragraphs:
    for node in par._p.xpath('.//w:t'):
        if node.text:node.text=node.text.replace(orphan,'')
# Source prose can span runs: also handle the complete paragraph text if needed.
for i in [409,1185]:
    if orphan.strip() in p[i].text:
        change(i,p[i].text.replace(orphan,''))
# Chapter 13 has no worked numerical example.
for node in p[936]._p.xpath('.//w:t'):
    if node.text:node.text=node.text.replace('The worked example is invented arithmetic. ','')
if 'The worked example is invented arithmetic.' in p[936].text:
    change(936,p[936].text.replace('The worked example is invented arithmetic. ',''))
# Canonical URLs verified 7 October: runwayml.com redirects to runway.com.
for rel in d.part.rels.values():
    if rel.is_external and str(rel.target_ref).startswith('https://runwayml.com/'):
        rel._target=str(rel.target_ref).replace('https://runwayml.com/','https://runway.com/',1)
# Keep the six-stream record compatible with the map and the before-YPP exceptions.
change(1064,'Watch-page ad revenue:')
new=p[1064].insert_paragraph_before('Shorts revenue:');p[1064]._p.addnext(new._p)
nextp=p[1064].insert_paragraph_before('Premium revenue:');new._p.addnext(nextp._p)
change(1071,'If a video took ten hours and the month’s revenue was $20, write both down. The hours show what the work cost. Before YPP acceptance, YouTube revenue lines are zero; direct sponsors and outside affiliate programs can pay earlier if their terms are met. Record any such income separately.')
d.core_properties.comments='REV12: final copy pass against the supplied review, 7 October 2026. Preserves the REV11 structure.'
d.save(OUT)
fresh=Document(OUT)
text='\n'.join(par.text for par in fresh.paragraphs)
for term in ['See chapter 5.,','speaking-rate range','A second subject is a second channel is','The click is the transaction is','Author recommendations','?sub_confirmation=1']:
    assert term not in text,term
assert 'The problem in their words —' in text
assert 'The worked example is invented arithmetic.' not in next(par.text for par in fresh.paragraphs if par.text.startswith('Sources: Creator Partnerships'))
assert 'A banner reads' not in fresh.element.xml
anchors={b.get(qn('w:name')) for b in fresh.element.xpath('.//w:bookmarkStart')}
assert all(h.get(qn('w:anchor')) in anchors for h in fresh.element.xpath('.//w:hyperlink[@w:anchor]'))
assert len(fresh.tables)==3 and len(fresh.inline_shapes)==1
assert sum(par.text.startswith('In this chapter:') for par in fresh.paragraphs)==15
first=fresh.paragraphs.index(next(par for par in fresh.paragraphs if par.text=='Who it is for —')) if False else next(i for i,par in enumerate(fresh.paragraphs) if par.text=='Who it is for —')
assert [par.text for par in fresh.paragraphs[first:first+8]]==['Who it is for —','The problem in their words —','The method, given away —','Proof —','Comparison —','One objection —','The deep dive —','The path video and watch order —']
qa=ROOT/'bak/rev12_qa';qa.mkdir(parents=True,exist_ok=True)
(qa/'copy_changes.json').write_text(json.dumps(changes,ensure_ascii=False,indent=2),encoding='utf-8')
(qa/'paragraphs.txt').write_text('\n'.join(f'{i}\t{par.text}' for i,par in enumerate(fresh.paragraphs)),encoding='utf-8')
words=sum(len(par.text.split()) for par in fresh.paragraphs)+sum(len(c.text.split()) for t in fresh.tables for row in t.rows for c in row.cells)
print('Saved',OUT,'Words',words,'Tables',len(fresh.tables),'Internal destinations verified')
