from pathlib import Path
from copy import deepcopy
import json,re,zipfile
from docx import Document
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[1];d=Document(ROOT/'Your First YouTube Channel_REV16_2026-10-07.docx');changes=[]
def find(prefix):
 ps=[p for p in d.paragraphs if p.text.startswith(prefix)];assert len(ps)==1,(prefix,len(ps));return ps[0]
def replace(p,text):
 changes.append({'before':p.text,'after':text})
 for c in list(p._p):
  if c.tag not in (qn('w:pPr'),qn('w:bookmarkStart'),qn('w:bookmarkEnd')):p._p.remove(c)
 p.add_run(text)
def label(p,text):
 n=p.insert_paragraph_before(text);n.runs[0].bold=True;n.paragraph_format.keep_with_next=True;return n
def link(p,text,anchor):
 h=OxmlElement('w:hyperlink');h.set(qn('w:anchor'),anchor);r=OxmlElement('w:r');prop=OxmlElement('w:rPr');s=OxmlElement('w:rStyle');s.set(qn('w:val'),'Hyperlink');prop.append(s);r.append(prop);t=OxmlElement('w:t');t.text=text;t.set(qn('xml:space'),'preserve');r.append(t);h.append(r);p._p.append(h)
replace(find('Tonight’s checklist.'),'Tonight’s checklist. Mark it on paper; complete the upload checks before publishing.')
# Move all pre-publish decisions together; keep every existing action.
p=find('Paper. One viewer sentence.');label(p,'MAKE')
for p in list(d.paragraphs):
 if p.text.startswith('Paper. '):replace(p,'[ ] '+p.text[len('Paper. '):])
checks=[find('[ ] Audience setting'),find('[ ] Paid promotion checked'),find('[ ] Altered or synthetic content checked')]
last=find('[ ] Watched on a phone.')
for p in checks:last._p.addnext(p._p);last=p
label(checks[0],'BEFORE YOU PRESS PUBLISH')
right=checks[-1].insert_paragraph_before('[ ] Rights checked for every outside clip, image, voice, and track; license and date recorded in the rights log.')
checks[-1]._p.addnext(right._p)
label(find('[ ] Three strangers asked'),'AFTER PUBLISHING')
# Fix XML whitespace rather than flatten hyperlinks.
spaces=0
for part in [d.part]+[s.header.part for s in d.sections]+[s.footer.part for s in d.sections]:
 for t in part.element.xpath('.//w:t'):
  if t.text and (t.text[:1].isspace() or t.text[-1:].isspace()):t.set(qn('xml:space'),'preserve');spaces+=1
for p in d.paragraphs:
 for t in p._p.xpath('.//w:t'):
  if t.text:
   t.text=re.sub(r'((?:Official rule|Author recommendation|Creator heuristic|Documented platform guidance):)(?=\S)',r'\1 ',t.text)
   if t.text[:1].isspace() or t.text[-1:].isspace():t.set(qn('xml:space'),'preserve')
# Give section destinations their own anchors and direct problem links.
section_names={'II. Community Guidelines:':'rev17Guidelines','III. Copyright strikes:':'rev17Copyright','IV. Content ID claims:':'rev17Claim','V. Limited ads:':'rev17Limited'}
for k,name in section_names.items():
 p=find(k);b=OxmlElement('w:bookmarkStart');b.set(qn('w:name'),name);b.set(qn('w:id'),str(9100+list(section_names).index(k)));e=OxmlElement('w:bookmarkEnd');e.set(qn('w:id'),b.get(qn('w:id')));p._p.insert(1,b);p._p.append(e)
p=find('Layer: Set up before monetization — except');n=p.insert_paragraph_before();n.paragraph_format.keep_with_next=True
for i,(text,a) in enumerate([('CONTENT ID CLAIM → IV','rev17Claim'),('GUIDELINES STRIKE → II','rev17Guidelines'),('COPYRIGHT STRIKE → III','rev17Copyright'),('LIMITED ADS → V','rev17Limited'),('ABOUT TO UPLOAD → XI','rev10UploadChecklist')]):
 if i:n.add_run('\n')
 link(n,text,a)
sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'F0F0F0');n._p.get_or_add_pPr().append(sh)
# Advice and rules stay distinct.
p=find('A second channel is for a second viewer,');replace(p,'Author recommendation: split for a second viewer; a second format alone is not a reason to start over. Official rule: each channel must qualify for YPP separately. Chapter 14, §V helps you decide whether the split earns its extra work.')
# Retain specialist definitions and the signature eight jobs; cut the mini-book entries.
ps=d.paragraphs;a=next(i for i,p in enumerate(ps) if p.text=='Glossary' and p.style.name=='Heading 1');b=next(i for i,p in enumerate(ps) if i>a and p.text=='Official sources and updates' and p.style.name=='Heading 1')
removed=[]
for p in ps[a+1:b]:
 if p.text.startswith(('Fader.','Stock license.','Burned-in captions.','Text-to-speech.','Deep dive.','Sponsor.','Affiliate link.')):
  if not p._p.xpath('.//w:bookmarkStart'):
   removed.append(p.text);p._p.getparent().remove(p._p)
replace(find('Use the separate Printable Workbook'),'The edition package includes a separate eight-page Printable Workbook in PDF and Word formats. Print a fresh sheet when you need one, or copy these blanks into a notebook. Keep the first answers so you can see what changed. The filled example stays here.')
anchors={b.get(qn('w:name')) for b in d.element.xpath('.//w:bookmarkStart')};assert all(h.get(qn('w:anchor')) in anchors for h in d.element.xpath('.//w:hyperlink[@w:anchor]'))
for p in d.paragraphs:
 if p.text.startswith('[ ] '):p.paragraph_format.space_after=Pt(4)
assert not any(p.text.startswith('Paper.') for p in d.paragraphs)
out=ROOT/'Your First YouTube Channel_REV17_2026-10-07.docx';d.save(out)
qa=ROOT/'bak/rev17_qa';qa.mkdir(parents=True,exist_ok=True);(qa/'changes.json').write_text(json.dumps({'changes':changes,'glossary_removed':removed,'whitespace_nodes':spaces},ensure_ascii=False,indent=2),encoding='utf-8')
print('Saved REV17; glossary entries removed:',len(removed),'whitespace nodes:',spaces,'all anchors resolve.')

