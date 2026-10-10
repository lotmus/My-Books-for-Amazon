from pathlib import Path
from docx import Document
from docx.shared import Inches,Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[1];d=Document(ROOT/'Your First YouTube Channel_REV15_2026-10-07.docx')
def find(s):
 p=[p for p in d.paragraphs if p.text.startswith(s)];assert len(p)==1,(s,len(p));return p[0]
def replace(p,t):
 for c in list(p._p):
  if c.tag not in (qn('w:pPr'),qn('w:bookmarkStart'),qn('w:bookmarkEnd')):p._p.remove(c)
 p.add_run(t)
p=find('For this book');replace(p,'Author recommendation: for this book’s long-form route, a Short comes after its long video is published, names that video, and points to it through Studio’s Related video setting.')
p.runs[0].text=p.runs[0].text.removeprefix('Author recommendation:');p.runs[0].text='Author recommendation:'+p.runs[0].text
# Insert readable ticks, keeping the evening list’s existing style.
p=find('Paper. Next video named.')
for t in ['[ ] Audience setting matches who this video is directed to; made for kids is a required audience decision, not a growth tactic.','[ ] Paid promotion checked: if a sponsor or commercial relationship requires disclosure, select the setting and say the disclosure in the video.','[ ] Altered or synthetic content checked: disclose realistic generated or meaningfully altered material that could be mistaken for real.']:
 n=p.insert_paragraph_before(t);n.style=p.style
replace(find('Example videos, not a model to copy.'),'Author recommendation: watch a video that does the job yours will do, and write down its first sentence. Then write your own.')
p=find('Country availability. The expanded program,')
replace(p,'Country availability. Check the Partner Program list, the expanded-program list, and each feature’s own eligibility page. Availability in one list does not establish availability in the others.')
n=p.insert_paragraph_before('Your country is not on the list');n.runs[0].bold=True
n=p.insert_paragraph_before('Official rule: the earlier tier and each feature have their own country requirements. If the expanded tier is unavailable but YPP is offered where you live, use the main YPP entry route; fan funding and Shopping still depend on their separate requirements. If YPP itself is unavailable, the hour count does not override that restriction.')
n=p.insert_paragraph_before('Author recommendation: keep making useful videos and recording qualified hours where Studio shows them. The work and the count still matter, but rolling-window hours can expire before a program becomes available. Sponsors and outside affiliate programs do not require YPP entry; check their own availability and disclose the relationship. Those are the doors to test while a YouTube feature is closed.')
# A shaded box made from paragraphs avoids adding a navigation table.
for n in list(d.paragraphs):
 if n.text.startswith(('Your country is not on the list','Official rule: the earlier tier','Author recommendation: keep making useful videos')):
  prop=n._p.get_or_add_pPr();shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'F0F0F0');prop.append(shade)
replace(find('Copy these into a notebook.'),'Use the separate Printable Workbook supplied with this edition, or copy these sheets into a notebook. Keep the first answers so you can see what changed. The filled example stays here. Paste the words only after replacing every bracket and deleting anything that is not yet true.')
# Keep the architecture; tighten repeated mechanics and make the two models visible.
replace(find('Long-form means public videos whose hours accumulate.'),'Keep the long-form and Shorts counters separate; chapter 7 explains what each counts. Choose the route before planning the pace.')
replace(find('End screens. YouTube'), 'End screens. Use the timing and settings in chapter 4, §VI. Point at the next video in watch order, by name.')
replace(find('Watch Page Monetization Module'), 'Ads and Premium on long-form and live videos — in Studio, the Watch Page Monetization Module;')
replace(find('Shorts Monetization Module'), 'Shorts-feed ads and Premium — in Studio, the Shorts Monetization Module;')
replace(find('Commerce Product Module —'), 'Fan funding: memberships, Super Chat, Super Stickers, and Super Thanks — in Studio, the Commerce Product Module.')
for par in d.paragraphs:
 if 'files carrying C2PA content credentials' in par.text:
  replace(par,par.text.replace('files carrying C2PA content credentials','files carrying digital records of how they were created or altered (C2PA content credentials)'))
for prefix,text in [('I. Give every video one job','THE EIGHT JOBS · Watch order\nWho it is for → Problem → Method → Proof\nComparison → Objection → Deep dive → Path video\nAuthor recommendation: film who it is for and the deep dive first. The path video names where to start; every teaching video names the next step.'),('I. When the counts stall, find the earliest break','FIND THE EARLIEST BREAK\nViewer → Click → Stay → Next video → Gate → Money\nAuthor recommendation: inspect from the left. Write one change on the decision log at the first break you can support with evidence.')]:
 par=find(prefix);n=par.insert_paragraph_before(text);n.runs[0].font.size=Pt(11)
 sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'F0F0F0');n._p.get_or_add_pPr().append(sh)
for par in d.paragraphs:
 if par.text.startswith('Author recommendation: for this book’s long-form route'):
  text=par.text;replace(par,'');par.runs[0].text='Author recommendation:';par.runs[0].bold=True;par.add_run(text[len('Author recommendation:'):])

anchors={b.get(qn('w:name')) for b in d.element.xpath('.//w:bookmarkStart')};assert all(h.get(qn('w:anchor')) in anchors for h in d.element.xpath('.//w:hyperlink[@w:anchor]'))
d.save(ROOT/'Your First YouTube Channel_REV16_2026-10-07.docx')
# Seven sheets with writing space, plus a separate page for twenty titles.
w=Document();sec=w.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11);sec.top_margin=sec.bottom_margin=Inches(.65)
w.styles['Normal'].font.name='Calibri';w.styles['Normal'].font.size=Pt(10);w.styles['Normal'].paragraph_format.space_after=Pt(4)
ps=d.paragraphs
heads=['I. Viewer sheet','II. Set-up sheet','III. Video-job sheet','IV. Week sheet','V. Money sheet','VI. Decision log','VII. Rights log','VIII. One filled example']
indices=[next(i for i,p in enumerate(ps) if p.text==h) for h in heads]
for j in range(7):
 if j:w.add_page_break()
 w.add_heading('Your First YouTube Channel',0);w.add_paragraph('Printable Workbook • 7 October 2026 • Copy a fresh sheet when you need it.')
 w.add_heading(heads[j],1)
 for p in ps[indices[j]+1:indices[j+1]]:
  t=p.text.strip()
  if not t:continue
  if t.startswith(('Twenty video titles','A useful line sounds','The eight jobs,','The path video explains','The eight jobs are','Deep dive published','Video chapters','Each video names','Series playlist')):continue
  if j==2 and t.endswith('—'):continue
  w.add_paragraph(t)
  if not t.startswith(('If a line does not','Keep each stream','One line for','Copy before')):
   w.add_paragraph('_'*85)
 if j==0:
  w.add_page_break();w.add_heading('Viewer sheet • Twenty titles',1);w.add_paragraph('Write each as a question your viewer would type. Circle the first video and the deep dive.')
  for k in range(1,21):w.add_paragraph(f'{k:2}. '+ '_'*80)
w.save(ROOT/'Your First YouTube Channel_Printable Workbook_2026-10-07.docx')
(ROOT/'notes/Week_One_Reader_Test.md').write_text('''# Independent reader test
Status: prepared; not yet carried out.

Give a first-channel reader the evening page, Start this week, the book and printable sheets. Ask them to follow the evening page and week one independently. Do not explain the method or point to the policy checklist. They may open the book when the instructions send them there.

Record whether they can name one viewer, circle two titles, publish one video, and locate chapter 10, §XI before uploading. Record where they pause, what they expected, and the exact wording they misunderstood. Ask them to show the audience, paid-promotion and altered-or-synthetic decisions, and the sheet they will use for their first change. Do not score views or growth as a usability result.

Afterward ask: What will you do next? Where did you get stuck? What did you have to guess? Fix observed stuck points before adding explanations elsewhere. This test requires an independent human reader; manuscript checks do not replace it.
''',encoding='utf-8')
print('Saved REV16, printable workbook, and independent-reader test brief. Links verified.')
