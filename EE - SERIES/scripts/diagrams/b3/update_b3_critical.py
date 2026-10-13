"""Apply the 13 Oct 2026 critical-reader corrections to the Book 3 master."""
from pathlib import Path
import sys
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root = Path(__file__).resolve().parents[3]
docx_path = root / "Semiconductor_Physics_and_Devices_Book3.docx"
figure_path = Path(__file__).resolve().parent / "fig" / "b3_bands.png"
doc = Document(docx_path)

def replace_once(needle, replacement):
    hits = [p for p in doc.paragraphs if needle in p.text]
    if not hits:
        if not replacement:
            return
        if any(replacement and replacement in p.text for p in doc.paragraphs):
            return
        if "p-type" in needle and any("p-type material:" in p.text for p in doc.paragraphs):
            return
        if "thermal conductivity is the lowest in the list" in needle and any("thermal conductivity is the lowest in this comparison" in p.text for p in doc.paragraphs):
            return
    if len(hits) != 1:
        raise RuntimeError(f"Expected one paragraph containing {needle!r}; found {len(hits)}")
    p = hits[0]
    updated = p.text.replace(needle, replacement, 1)
    if p.runs:
        p.runs[0].text = updated
        for run in p.runs[1:]:
            run.text = ""
    else:
        p.add_run(updated)

replace_once(
    "2,290 to 3,450 W/m·K, the highest of any material.",
    "2,290 to 3,450 W/m·K, the highest in this comparison for bulk semiconductor crystals.",
)
replace_once(
    "No usable p-type doping exists, so there is no pn junction, no conventional MOSFET body, and no bipolar device.",
    "Its central limitation is the lack of a mature, controllable, device-quality p-type beta-Ga2O3 layer. Reports of hole conductivity exist, but reliable p-type material remains difficult; conventional homojunction pn devices and p-channel structures therefore remain research challenges. Heterojunctions that use a different p-type semiconductor are a separate possibility.",
)
replace_once(
    "Its central limitation is the lack of a mature, controllable, device-quality p-type beta-Ga2O3 layer. Reports of hole conductivity exist, but reliable p-type material remains difficult; conventional homojunction pn devices and p-channel structures therefore remain research challenges. Heterojunctions that use a different p-type semiconductor are a separate possibility.",
    "Its key limitation is p-type material: reports of hole conductivity exist, but reliable, controllable, device-quality layers remain difficult to make. Conventional homojunction pn devices and p-channel structures therefore remain research challenges; heterojunctions using a different p-type semiconductor are a separate possibility.",
)
replace_once("It has two deep weaknesses. ", "")
replace_once(
    "And its thermal conductivity is the lowest in the list, so heat must leave through the top surface, through thinned wafers, or through a bonded layer of a better conductor.",
    "Its thermal conductivity is the lowest in this comparison, so heat must leave through the top surface, a thinned wafer, or a bonded layer of a better conductor.",
)
replace_once(
    "Diamond has the opposite profile: the best thermal conductivity of any material,",
    "Diamond has the opposite profile: among crystalline semiconductors, it has one of the highest thermal conductivities,",
)
replace_once(
    "Ga2O3 has no usable p-type doping, so there is no pn junction, no inversion-channel MOSFET body, and no bipolar device;",
    "Reports of hole conductivity in Ga2O3 exist, but mature, controllable, device-quality p-type layers remain a challenge; conventional homojunction pn devices and p-channel structures are therefore research problems;",
)
replace_once(
    "Reports of hole conductivity in Ga2O3 exist, but mature, controllable, device-quality p-type layers remain a challenge; conventional homojunction pn devices and p-channel structures are therefore research problems; designers use Schottky barriers, depletion-mode and fin transistors, or junctions with other p-type materials.",
    "Reports of hole conductivity in Ga2O3 exist, but mature, controllable, device-quality p-type layers remain a challenge. Conventional homojunction pn devices and p-channel structures are therefore research problems. Designers can instead use Schottky barriers, depletion-mode and fin transistors, or heterojunctions with other p-type materials.",
)
replace_once(
    "Figure 1.1 — Energy bands with the gaps drawn to scale. A metal's bands overlap; silicon's 1.12 eV gap is narrow enough for a few electrons to cross at room temperature; an insulator's gap of about 9 eV is not.",
    "Figure 1.1 — Energy bands with the gaps drawn to scale. A metal's bands overlap; silicon's 1.12 eV gap is narrow enough for a few electrons to cross at room temperature; an insulator's gap of about 9 eV is not.",
)

# Replace the embedded band diagram while retaining its position and caption.
paras = list(doc.element.body.iter(qn("w:p")))
caption = next((p for p in paras if any(t.text and t.text.startswith("Figure 1.1") for t in p.iter(qn("w:t")))), None)
if caption is None:
    raise RuntimeError("Figure 1.1 caption not found")
pos = paras.index(caption) - 1
picture = None
while pos >= 0:
    p = paras[pos]
    blips = list(p.iter(qn("a:blip")))
    if blips:
        picture = (p, blips[0])
        break
    pos -= 1
if picture is None:
    raise RuntimeError("Figure 1.1 image not found")
rid = picture[1].get(qn("r:embed"))
part = doc.part.rels[rid].target_part
part._blob = figure_path.read_bytes()

# Add the primary study to the existing Part IX source paragraph.
source_needle = 'Woo et al., "From Wide to Ultrawide-Bandgap Semiconductors for High Power and High Frequency Electronic Devices" (Journal of Physics: Materials, 2024);'
park = 'Park et al., "P-type conductivity in oxygen-rich β-Ga2O3 thin films" (Materials Science in Semiconductor Processing, 2025), https://doi.org/10.1016/j.mssp.2025.110003'
source_paras = [p for p in doc.paragraphs if source_needle in p.text]
if len(source_paras) != 1:
    raise RuntimeError(f"Expected one Part IX source paragraph; found {len(source_paras)}")
if park not in source_paras[0].text:
    replace_once(source_needle, source_needle + " " + park + ";")

# Keep the linked-source list readable and avoid duplicating its linked titles in
# the main citation run. Earlier runs of this updater copied p.text, which also
# includes hyperlink labels, back into that run.
source_p = next(p for p in doc.paragraphs if "Woo et al." in p.text)
first_run = source_p._p.find(qn("w:r"))
first_text = first_run.find(qn("w:t"))
if first_text is None:
    raise RuntimeError("Part IX source text run not found")
tail_marker = "for OLED triplet harvesting and TADF."
if tail_marker in first_text.text:
    first_text.text = first_text.text[:first_text.text.index(tail_marker) + len(tail_marker)]
while first_text.text.count(park) > 1:
    second = first_text.text.find(park, first_text.text.find(park) + len(park))
    start = second - 2 if first_text.text[max(0, second - 2):second] == "; " else second
    first_text.text = first_text.text[:start] + first_text.text[second + len(park):]

children = list(source_p._p)
links = [c for c in children if c.tag == qn("w:hyperlink")]
if links:
    # Add a label once, immediately before the first clickable title.
    before = links[0]
    prev = before.getprevious()
    if prev is None or prev.tag != qn("w:r") or "Further links:" not in "".join(t.text or "" for t in prev.iter(qn("w:t"))):
        label_run = OxmlElement("w:r")
        label_text = OxmlElement("w:t"); label_text.text = " Further links: "
        label_text.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        label_run.append(label_text); before.addprevious(label_run)
    for link in links[:-1]:
        nxt = link.getnext()
        if nxt is not None and nxt.tag == qn("w:r"):
            txt = nxt.find(qn("w:t"))
            if txt is None:
                txt = OxmlElement("w:t"); nxt.append(txt)
            txt.text = "; "; txt.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")

doc.save(docx_path)
print(f"Updated {docx_path}")
