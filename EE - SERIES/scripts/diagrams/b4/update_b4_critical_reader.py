"""Apply the 13 Oct 2026 critical-reader repairs to the Book 4 master."""
from pathlib import Path
from docx import Document

root = Path(__file__).resolve().parents[3]
master = root / "RF_Microwave_and_Transceivers_Book4.docx"
doc = Document(master)

changes = [
    (
        "The frequency did, and that is the whole subject of this book.",
        "The frequency did, and that is the whole subject of this book. The capacitor keeps its name on the packet; the circuit has other opinions.",
    ),
    (
        "A coil, a variable capacitor, a germanium diode, and a long wire will pull a radio station out of the air with no battery at all. A few more parts make a Morse transmitter that can reach across a continent on a watt when propagation and the antenna cooperate, and a twenty-dollar USB dongle can receive signals from aircraft and weather satellites.",
        "A coil, a variable capacitor, a germanium diode, and a long wire will pull a radio station out of the air with no battery at all. The station supplies both the signal and the power, if it is generous enough. A few more parts make a one- to five-watt Morse transmitter; a continent-spanning contact is possible when propagation and the antenna cooperate. A twenty-dollar USB dongle receives aircraft and weather satellites. Building these receivers and transmitters teaches the ideas of this book faster than any amount of reading.",
    ),
    (
        "so a 150 V device leaves the margin a resonant load that drifts with the receiver's position needs.",
        "leaving 64.5 V of nominal headroom with a 150 V-rated device.",
    ),
]

for old, new in changes:
    hits = [p for p in doc.paragraphs if old in p.text]
    if len(hits) == 1:
        p = hits[0]
        for run in p.runs:
            if old in run.text:
                run.text = run.text.replace(old, new, 1)
                break
        else:
            raise RuntimeError(f"Target spans multiple runs; refusing to flatten formatting: {old[:70]}")
    elif not any(new in p.text for p in doc.paragraphs):
        raise RuntimeError(f"Expected one paragraph containing {old!r}; found {len(hits)}")

doc.save(master)
print(f"Updated {master}")
