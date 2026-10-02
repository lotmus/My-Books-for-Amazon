# -*- coding: utf-8 -*-
"""Nine-to-seven pass for the docx-only masters (Books 1-3), 2026-10-01.
Rewrites cross-book references run by run, updates the series table and the
'nine'/'eight' wording. Usage: python fix_docx_books.py <book_no> <in.docx> <out.docx>"""
import sys, copy
from docx import Document
from docx.oxml.ns import qn
from xref import xref
ROWS = [("Foundations", "Physics, mathematics, and circuit theory"),
        ("Circuits, Components, and Control", "Analog and digital, op-amps, passives, the classical loop"),
        ("Semiconductor Physics and Devices", "Diodes, BJTs, MOSFETs, FinFET, compound semiconductors"),
        ("RF, Microwave, and Transceivers", "Lines, matching, antennas, PAs, receivers, PLLs, radar"),
        ("Communications, Wireless, and SDR", "Modulation, band plans, DSP and software radio"),
        ("Power and Energy", "Regulators, magnetics, the grid, the rack"),
        ("Packaging, Layout, EMC, and Test", "Boards, packages, cooling, emissions, the accredited lab, robots")]
WORDS = [("first of nine", "first of seven"), ("second of nine", "second of seven"), ("third of nine", "third of seven"),
         ("Each of the nine books", "Each of the seven books"), ("other eight books", "other six books"), ("the filters that Book 4 and Book 4 assume", "the filters that Book 4 assumes")]
def set_cell(cell, text):
    ps = cell.paragraphs
    for p in ps[1:]: p._p.getparent().remove(p._p)
    rs = ps[0].runs
    rs[0].text = text
    for r in rs[1:]: r.text = ""
def main(n, inp, outp):
    n = int(n); doc = Document(inp); bad = []; changed = 0
    if any("RF, Microwave, and Transceivers" in c.text for t in doc.tables for r in t.rows for c in r.cells) or \
       any("of seven" in p.text for p in doc.paragraphs):
        sys.exit("already in seven-book numbering; xref must never run twice")
    def paras():
        for p in doc.paragraphs: yield p
        for t in doc.tables:
            for r in t.rows:
                for c in r.cells:
                    for p in c.paragraphs: yield p
    for p in paras():
        full = p.text; want = xref(full, n, internal=False)
        for a, b in WORDS: want = want.replace(a, b)
        if want == full: continue
        for r in p.runs:
            t = xref(r.text, n, internal=False)
            for a, b in WORDS: t = t.replace(a, b)
            r.text = t
        changed += 1
        if p.text != want: bad.append((full, want, p.text))
    for t in doc.tables:
        if len(t.rows) == 10 and t.rows[0].cells[0].text == "Book" and "Packaging" in t.rows[9].cells[0].text:
            for i, (title, covers) in enumerate(ROWS, 1):
                set_cell(t.rows[i].cells[0], f"{i}. {title}" + (" (this book)" if i == n else ""))
                set_cell(t.rows[i].cells[1], covers)
            for r in list(t.rows)[8:]: r._tr.getparent().remove(r._tr)
            print("series table updated")
    doc.save(outp); print("paragraphs changed:", changed, "run-split problems:", len(bad))
    for b in bad: print("  BAD:", b[0][:200], "\n   WANT:", b[1][:200], "\n   GOT:", b[2][:200])
if __name__ == "__main__":
    main(*sys.argv[1:])
