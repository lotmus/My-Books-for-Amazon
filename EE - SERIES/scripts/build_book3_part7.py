# -*- coding: utf-8 -*-
"""Book 3 extension (2026-10-01): Part VII (Chapters 20-22).

Input: Book 3 as rebuilt by build_book3_parts4_6.py (Parts I-VI, Ch 1-19).
Inserts Part VII (Ch 20 fabrication, Ch 21 piezoelectric films and acoustic
resonators, Ch 22 compact models and the PDK) before Appendix A in the house
style, adds answers (Appendix B) and constants (Appendix A), extends
Appendix C and Appendix E, and updates the front matter. DOCX only.
Usage: python build_book3_part7.py <in.docx> <out.docx>
"""
import copy, sys
from pathlib import Path
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC = Path(__file__).resolve().parents[1] / "chapters" / "Book3" / "part7" / "src"
CHS = [f"ch{n}" for n in range(20, 23)]
CONSTS = [
    "Boron diffusion in silicon, hand values: D = D0 exp(−Ea/kT), D0 = 1.0 cm²/s, Ea = 3.5 eV; diffusion length ≈ 2√(Dt).",
    "Lithography: half-pitch R = k1λ/NA (k1 ≥ 0.25 for one exposure); DOF ≈ k2λ/NA². ArF 193 nm, EUV 13.5 nm.",
    "Acoustic: AlN longitudinal velocity ≈ 11,350 m/s, kt² ≈ 6.5%, εr ≈ 9.5, TCF ≈ −25 ppm/K. BAW f ≈ v/(2t); kt² ≈ (π²/4)(fp − fs)/fp.",
    "Statistics: one-sided Q(3) = 1.35×10⁻³, Q(6) = 9.9×10⁻¹⁰; ±3σ holds 99.73%.",
]
SOURCES = [
    "Part VII: Plummer, Deal, and Griffin, Silicon VLSI Technology, for Chapter 20. Lakin, \"Thin Film Resonator Technology\" (IEEE UFFC, 2005); Akiyama et al., \"Enhancement of Piezoelectric Response in Scandium Aluminum Nitride Alloy Thin Films\" (Advanced Materials, 2009); Hashimoto, RF Bulk Acoustic Wave Filters for Communications, for Chapter 21. Chauhan et al., FinFET Modeling for IC Simulation and Design: Using the BSIM-CMG Standard; Gummel and Poon, \"An Integral Charge Control Model of Bipolar Transistors\" (Bell System Technical Journal, 1970), for Chapter 22.",
]
APPC_ADD = "Part VII adds the fab steps that set every device's limits (Chapter 20), the piezoelectric thin films and acoustic resonators behind the filters that Book 4 and Book 6 assume (Chapter 21), and the compact models and process design kits through which all of these devices reach a simulator (Chapter 22). Book 9 surveys the 2026 MEMS products built on the same films."

def strip_runs(p):
    for c in list(p):
        if c.tag in (qn("w:r"), qn("w:hyperlink"), qn("w:bookmarkStart"), qn("w:bookmarkEnd")):
            p.remove(c)
    for a in list(p.attrib):
        if a.endswith("paraId") or a.endswith("textId"):
            del p.attrib[a]
    return p

def make_p(tpl, text):
    p = strip_runs(copy.deepcopy(tpl))
    r0 = next((r for r in tpl.findall(qn("w:r")) if r.find(qn("w:t")) is not None), None)
    r = OxmlElement("w:r")
    if r0 is not None and r0.find(qn("w:rPr")) is not None:
        r.append(copy.deepcopy(r0.find(qn("w:rPr"))))
    t = OxmlElement("w:t"); t.set(qn("xml:space"), "preserve"); t.text = text
    r.append(t); p.append(r)
    return p

def set_text(par, new):
    for r in par.runs[1:]:
        r._r.getparent().remove(r._r)
    par.runs[0].text = new

def main(inp, outp):
    doc = Document(inp)
    P = doc.paragraphs
    H1 = lambda p: p.style is not None and p.style.name == "Heading 1"
    idx = lambda pred: next(i for i, p in enumerate(P) if pred(p))
    if any(H1(p) and p.text.startswith("Chapter 20.") for p in P):
        sys.exit("already rebuilt")
    idx(lambda p: H1(p) and p.text.startswith("Chapter 19."))
    i5 = idx(lambda p: H1(p) and p.text.startswith("Chapter 5."))
    T = {"H1": P[i5]._p, "P": P[i5 + 1]._p, "H2": P[i5 + 2]._p, "EQ": P[i5 + 4]._p}
    iwe = idx(lambda p: p.text == "Worked Example 5.1")
    T["WEP"], T["WES"] = P[iwe + 1]._p, P[iwe + 2]._p
    T["KEY"] = P[idx(lambda p: p.text.startswith("Key idea. Forward current"))]._p
    T["PR"] = P[idx(lambda p: p.text.startswith("1. Same diode."))]._p
    ipart = idx(lambda p: H1(p) and p.text.startswith("Part II."))
    T["PART"], T["PARTP"] = P[ipart]._p, P[ipart + 1]._p
    iab = idx(lambda p: p.style.name == "Heading 2" and p.text == "Chapter 11" and P.index(p) > iwe)
    T["AH"], T["AP"] = P[iab]._p, P[iab + 1]._p
    iA = idx(lambda p: H1(p) and p.text.startswith("Appendix A."))
    T["CONST"] = P[iA + 1]._p
    T = {k: copy.deepcopy(v) for k, v in T.items()}
    anchor = P[iA]._p
    answers, words = {}, {}
    for ch in CHS:
        prn = 0; answers[ch] = []
        for line in (SRC / f"{ch}.txt").read_text(encoding="utf-8").splitlines():
            if not line.strip():
                continue
            tag, _, b = line.partition(":"); b = b.strip()
            if tag == "ANS":
                answers[ch].append(b); continue
            if tag == "PR":
                prn += 1; b = f"{prn}. {b}"
            if tag not in T:
                raise ValueError(f"{ch}: unknown tag {tag}")
            anchor.addprevious(make_p(T[tag], b))
        assert prn == len(answers[ch]), ch
    # Appendix A constants
    P = doc.paragraphs
    iB = idx(lambda p: H1(p) and p.text.startswith("Appendix B."))
    last = P[iB - 1]._p
    for c in CONSTS:
        p = make_p(T["CONST"], c); last.addnext(p); last = p
    # Appendix B answers
    P = doc.paragraphs
    iC = idx(lambda p: H1(p) and p.text.startswith("Appendix C."))
    a = P[iC]._p
    for ch in CHS:
        a.addprevious(make_p(T["AH"], f"Chapter {int(ch[2:])}"))
        for k, s in enumerate(answers[ch], 1):
            a.addprevious(make_p(T["AP"], f"{k}. {s}"))
    # Appendix C: add a Part VII paragraph after the first body paragraph
    P = doc.paragraphs
    iC = idx(lambda p: H1(p) and p.text.startswith("Appendix C."))
    first = P[iC + 1]
    set_text(first, first.text.replace("and the junction as a source and detector of light and as a high-voltage switch (Part VI).",
        "the junction as a source and detector of light and as a high-voltage switch (Part VI), and how devices are made, modeled, and extended to piezoelectric films (Part VII)."))
    first._p.addnext(make_p(copy.deepcopy(first._p), APPC_ADD))
    # Appendix E sources
    P = doc.paragraphs
    iE = idx(lambda p: H1(p) and p.text.startswith("Appendix E."))
    set_text(P[iE], "Appendix E. Sources for Parts II to VII")
    last = P[-1]._p; stpl = copy.deepcopy(P[iE + 1]._p)
    for s in SOURCES:
        p = make_p(stpl, s); last.addnext(p); last = p
    # front matter
    P = doc.paragraphs
    hits = 0
    for p in P[:14]:
        t = p.text
        if t.startswith("Parts I–VI  —"):
            set_text(p, "Parts I–VII  —  From the Crystal to the Chip"); hits += 1
        elif t.startswith("Copyright © 2026") and "Parts I through VI." in t:
            set_text(p, t.replace("Parts I through VI.", "Parts I through VII.")); hits += 1
        elif t.startswith("Part II takes up the bipolar transistor"):
            set_text(p, t.replace("into a high-voltage switch.", "into a high-voltage switch. Part VII follows a transistor through the fab, builds the piezoelectric thin films and acoustic resonators behind every phone's filters, and explains the compact models and process design kits that carry all of these devices into a simulator.")); hits += 1
    print("front-matter edits:", hits)
    doc.save(outp); print("saved", outp)

if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
